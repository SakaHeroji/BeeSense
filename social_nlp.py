#!/usr/bin/env python3
"""Sentence clustering and evidence-preserving artifacts, independent of Streamlit."""

from __future__ import annotations

import argparse
import importlib.metadata
import json
import re
import shutil
import tempfile
import time
import unicodedata
from collections import Counter
from copy import deepcopy
from datetime import date, datetime, timezone
from hashlib import file_digest, sha256
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent
VERSION = "2.2.0"
ENCODER_CACHE_VERSION = "2.0.0"  # Coherence changes do not change SBERT encoding.
SCHEMA_VERSION = "beesense-sentence-clusters:2"
LEGACY_SCHEMA_VERSION = "beesense-sentence-clusters:1"
EXPLANATION_SCHEMA = "beesense-topic-explanations:1"
REQUIRED_COLUMNS = {
    "platform",
    "jenis",
    "teks",
    "url",
    "judul_post",
    "tanggal_post_wib",
    "akurasi_tanggal_post",
    "tanggal_konten_wib",
    "akurasi_tanggal_konten",
    "teks_dipotong",
}
INDONESIAN_STOPWORDS = set(
    """ada adalah agar akan aku anda apa apakah atau bagi bahwa banget bang baru banyak
belum bisa bro buat bukan cara dan dari dengan dia di dong gak ga gue gw hal hanya harus ini itu jadi juga
kak ka kk kalau kalian kami kamu karena kan ke ketika kok lagi lah lebih lo lu maka masih mau memang
mereka min mungkin nya oleh pada paling para pernah pun saja sangat saya sejak sudah sih sis suatu
supaya tak tapi telah tentang tersebut terus tiap tolong untuk udah udh ya yang yg dgn krn tdk sy dr
utk dpt tp si aja bgt btw emang gimana bagaimana nih deh wkwk wkwkwk hahaha haha hehehe gan kakak pak bu
mas mbak gitu begitu kayak seperti soal dong kah nya mending apakah""".split()
)


def file_hash(path):
    with Path(path).open("rb") as stream:
        return file_digest(stream, "sha256").hexdigest()


def json_write(path, value):
    Path(path).write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )


class SocialNLPClusterer:
    """Encode with SBERT, fit UMAP + DBSCAN and optionally merge to a topic count.

    Encoder injection is for tests. Training uses an immutable public SBERT
    model through SentenceTransformer, with L2 normalization on local CPU.
    Dates/titles/platforms are provenance, never embedding features.
    """

    def __init__(self, config=None, *, encoder=None, progress=None):
        base = json.loads((ROOT / "social_nlp.json").read_text())
        if isinstance(config, (str, Path)):
            config = json.loads(Path(config).read_text())
        if config:
            if not isinstance(config, dict) or set(config) - set(base):
                raise ValueError("Unknown NLP configuration fields")
            for key, value in config.items():
                if isinstance(base[key], dict):
                    if not isinstance(value, dict) or set(value) - set(base[key]):
                        raise ValueError("Unknown settings in " + key)
                    base[key].update(deepcopy(value))
                else:
                    base[key] = deepcopy(value)
        self.config = base
        self.encoder = encoder
        self.progress = progress or (lambda event: None)
        self._validate_config()
        self.model_info = {"test_encoder": encoder is not None}
        self.sentences = pd.DataFrame()
        self.occurrences = pd.DataFrame()
        self.skipped = pd.DataFrame()
        self.comparisons = []
        self.topics = []
        self._started = time.monotonic()

    def _validate_config(self):
        c = self.config
        for key in (
            "seed",
            "min_words",
            "min_characters",
            "max_segment_characters",
            "max_unique_sentences",
        ):
            if type(c[key]) is not int or c[key] < (0 if key == "seed" else 1):
                raise ValueError("Invalid " + key)
        if c["max_segment_characters"] < c["min_characters"]:
            raise ValueError("Segment limit smaller than minimum text length")
        if not c["include_kinds"] or not set(c["include_kinds"]) <= {
            "post",
            "video",
            "comment",
            "reply",
        }:
            raise ValueError("Invalid include_kinds")
        m = c["model"]
        if not re.fullmatch(r"[\w.-]+/[\w.-]+", m["repo"]) or not re.fullmatch(
            r"[a-f0-9]{40}", m["revision"]
        ):
            raise ValueError(
                "Model requires a repository and immutable 40-character revision"
            )
        if m["backend"] != "sentence-transformers":
            raise ValueError("Embedding backend must be sentence-transformers (SBERT)")
        if (
            not 16 <= m["max_tokens"] <= 512
            or not 0 <= m["stride"] < m["max_tokens"] - 2
        ):
            raise ValueError("Invalid tokenizer window/stride")
        if not 1 <= m["batch_size"] <= 256 or not 1 <= m["threads"] <= 32:
            raise ValueError("Invalid inference batch size/threads")
        u = c["umap"]
        if (
            type(u["n_neighbors"]) is not int
            or not 2 <= u["n_neighbors"] <= 100
            or type(u["n_components"]) is not int
            or not 2 <= u["n_components"] <= 50
            or type(u["min_dist"]) not in (int, float)
            or not np.isfinite(u["min_dist"])
            or not 0 <= u["min_dist"] <= 1
        ):
            raise ValueError("Invalid UMAP dimensions/neighbors/min_dist")
        d = c["dbscan"]
        if (
            not isinstance(d["eps"], list)
            or not 1 <= len(d["eps"]) <= 30
            or any(
                type(x) not in (int, float) or not np.isfinite(x) or not 0 < x <= 10
                for x in d["eps"]
            )
        ):
            raise ValueError(
                "DBSCAN eps must be positive finite UMAP Euclidean distances up to 10"
            )
        if (
            not isinstance(d["min_samples"], list)
            or not 1 <= len(d["min_samples"]) <= 10
            or any(type(x) is not int or not 2 <= x <= 10000 for x in d["min_samples"])
        ):
            raise ValueError("DBSCAN min_samples must contain integers in 2..10000")
        if (
            type(d["neighborhood_batch_size"]) is not int
            or not 1 <= d["neighborhood_batch_size"] <= 2048
            or type(d["max_neighbor_edges"]) is not int
            or not 1 <= d["max_neighbor_edges"] <= 50000000
        ):
            raise ValueError("Invalid DBSCAN neighborhood memory bounds")
        e = c["evaluation"]
        if (
            not 0 < e["min_coverage"] <= 1
            or not 0 < e["max_largest_cluster"] < 1
            or e["sample_size"] < 10
            or not 0 <= e["score_tolerance"] <= 0.05
        ):
            raise ValueError("Invalid evaluation settings")
        g = c["graph"]
        if (
            not 2 <= g["max_nodes"] <= 500
            or not 1 <= g["neighbors"] <= 10
            or not 0 <= g["min_similarity"] <= 1
        ):
            raise ValueError("Invalid graph bounds")
        self._validate_explanation_options(c["explanation"])
        target = c["target_topics"]
        if target is not None and (type(target) is not int or not 2 <= target <= 500):
            raise ValueError("target_topics must be null or an integer in 2..500")
        coherence = c["coherence"]
        if (
            coherence["metric"] != "c_v"
            or type(coherence["topn"]) is not int
            or not 2 <= coherence["topn"] <= 12
            or type(coherence["window_size"]) is not int
            or not 2 <= coherence["window_size"] <= 1000
            or type(coherence["min_score"]) not in (int, float)
            or not np.isfinite(coherence["min_score"])
            or not 0 <= coherence["min_score"] <= 1
        ):
            raise ValueError("Invalid coherence settings")

    def _event(self, stage, **values):
        self.progress({"stage": stage, **values})

    @staticmethod
    def sentence_spans(text, max_characters=1200):
        """Rule-based segmentation preserving exact offsets in the original text."""
        if not isinstance(text, str):
            return []
        protected = list(text)
        expressions = [
            r"https?://\S+",
            r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}",
            r"(?<=\d)\.(?=\d)",
            r"\b(?:dr|prof|dll|dsb|dst|mr|mrs|ms|no|e\.g|i\.e|s\.kom|s\.pd|s\.t|m\.sc)\.",
        ]
        for expression in expressions:
            for match in re.finditer(expression, text, re.I):
                protected_end = match.end()
                if expression.startswith("https?"):
                    protected_end -= len(match[0]) - len(match[0].rstrip(".,!?;:)]}"))
                for index in range(match.start(), protected_end):
                    if protected[index] in ".!?":
                        protected[index] = "\ufff0"
        masked = "".join(protected)
        boundaries = (
            [0]
            + [m.end() for m in re.finditer(r"[.!?]+[\"')\]]*(?:\s+|$)|\n+", masked)]
            + [len(text)]
        )
        result = []
        for start, end in zip(boundaries, boundaries[1:]):
            while start < end and text[start].isspace():
                start += 1
            while end > start and text[end - 1].isspace():
                end -= 1
            windowed = end - start > max_characters
            while end > start:
                stop = min(end, start + max_characters)
                if stop < end:
                    cut = text.rfind(" ", start + max_characters // 2, stop)
                    if cut > start:
                        stop = cut
                result.append(
                    (start, stop, "length_window" if windowed else "sentence")
                )
                start = stop
                while start < end and text[start].isspace():
                    start += 1
        return result

    @staticmethod
    def _clean(text):
        text = unicodedata.normalize("NFC", text)
        text = re.sub(r"https?://\S+", " ", text)
        return re.sub(r"\s+", " ", text).strip()

    def prepare_excel(self, path=None):
        for attribute in ("_term_matrix", "_term_vocabulary", "_coherence_reference"):
            self.__dict__.pop(attribute, None)
        from openpyxl import load_workbook

        path = Path(path or self.config["input_file"])
        if not path.is_absolute():
            path = ROOT / path
        self.input_path = path.resolve()
        self.input_sha256 = file_hash(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if "Data" not in workbook.sheetnames:
                raise ValueError("Input requires the SocialCrawler Data sheet")
            iterator = iter(workbook["Data"].values)
            headers = next(iterator, ())
            if not REQUIRED_COLUMNS <= set(headers):
                raise ValueError(
                    "Missing SocialCrawler columns: "
                    + ", ".join(sorted(REQUIRED_COLUMNS - set(headers)))
                )
            sources = []
            for row_number, values in enumerate(iterator, 2):
                row = dict(zip(headers, values))
                row["source_row"] = row_number
                sources.append(row)
        finally:
            workbook.close()
        self.source_data = pd.DataFrame(sources)
        for column in ("tanggal_post_wib", "tanggal_konten_wib"):
            self.source_data[column] = self.source_data[column].map(
                lambda value: value.isoformat() if hasattr(value, "isoformat") else None
            )
        unique = {}
        occurrences, skipped = [], []
        eligible = 0
        for row in sources:
            if row["jenis"] not in self.config["include_kinds"]:
                continue
            eligible += 1
            if row["teks_dipotong"]:
                raise ValueError(
                    f"Source row {row['source_row']} has truncated text; export complete input before training"
                )
            text = row["teks"]
            if not isinstance(text, str) or not text.strip():
                skipped.append(
                    {
                        "source_row": row["source_row"],
                        "sentence": "",
                        "reason": "empty_text",
                    }
                )
                continue
            for start, end, segmentation in self.sentence_spans(
                text, self.config["max_segment_characters"]
            ):
                original = text[start:end]
                cleaned = self._clean(original)
                words = re.findall(r"[^\W\d_]+(?:[\w'-]*[\w])?", cleaned, re.UNICODE)
                if (
                    len(cleaned) < self.config["min_characters"]
                    or len(words) < self.config["min_words"]
                ):
                    skipped.append(
                        {
                            "source_row": row["source_row"],
                            "sentence": original,
                            "reason": "too_short_or_nonlinguistic",
                        }
                    )
                    continue
                sid = sha256(cleaned.encode()).hexdigest()[:24]
                if sid in unique and unique[sid]["embedding_text"] != cleaned:
                    raise ValueError("Sentence ID hash collision")
                unique.setdefault(
                    sid,
                    {
                        "sentence_id": sid,
                        "sentence": original,
                        "embedding_text": cleaned,
                    },
                )
                if len(unique) > self.config["max_unique_sentences"]:
                    raise ValueError(
                        "Sentence budget exceeded; no silent sampling is performed"
                    )
                record = {
                    "occurrence_id": sha256(
                        f"{self.input_sha256}:{row['source_row']}:{start}:{end}".encode()
                    ).hexdigest()[:24],
                    "sentence_id": sid,
                    "source_row": row["source_row"],
                    "platform": str(row["platform"]),
                    "kind": str(row["jenis"]),
                    "post_title": str(row["judul_post"] or ""),
                    "sentence": original,
                    "span_start": start,
                    "span_end": end,
                    "segmentation": segmentation,
                    "url": str(row["url"] or ""),
                }
                for source, target in [
                    ("tanggal_post_wib", "post_date"),
                    ("tanggal_konten_wib", "content_date"),
                ]:
                    value = row[source]
                    record[target] = (
                        value.isoformat() if hasattr(value, "isoformat") else None
                    )
                record["post_date_accuracy"] = row["akurasi_tanggal_post"]
                record["content_date_accuracy"] = row["akurasi_tanggal_konten"]
                occurrences.append(record)
        self.sentences = pd.DataFrame(unique.values())
        self.occurrences = pd.DataFrame(occurrences)
        self.skipped = pd.DataFrame(
            skipped, columns=["source_row", "sentence", "reason"]
        )
        self.preparation = {
            "input_rows": len(sources),
            "eligible_source_rows": eligible,
            "sentence_occurrences": len(occurrences),
            "unique_sentences": len(unique),
            "duplicate_occurrences": len(occurrences) - len(unique),
            "skipped_segments": len(skipped),
            "included_kinds": self.config["include_kinds"],
            "source_timezone": "Asia/Jakarta",
        }
        if len(unique) < 12:
            raise ValueError("At least 12 informative unique sentences are required")
        self._event("prepared", **self.preparation)
        return self

    def _encode_sbert(self, texts):
        import torch
        from huggingface_hub import snapshot_download
        from sentence_transformers import SentenceTransformer

        model = self.config["model"]
        snapshot = Path(
            snapshot_download(
                repo_id=model["repo"],
                revision=model["revision"],
                cache_dir=str(self._cache_dir() / "models"),
                token=False,
                allow_patterns=[
                    "*.json",
                    "vocab.txt",
                    "1_Pooling/config.json",
                    "model.safetensors",
                ],
                ignore_patterns=["onnx/*", "openvino/*"],
            )
        )
        weights = snapshot / "model.safetensors"
        if not weights.is_file():
            raise ValueError("Pinned SBERT revision must provide safetensors weights")
        torch.set_num_threads(model["threads"])
        torch.manual_seed(self.config["seed"])
        encoder = SentenceTransformer(
            str(snapshot),
            device="cpu",
            local_files_only=True,
            trust_remote_code=False,
            model_kwargs={"use_safetensors": True},
        )
        encoder.max_seq_length = model["max_tokens"]
        encoder.eval()
        self.model_info = dict(
            family="SBERT",
            repo=model["repo"],
            revision=model["revision"],
            weights_sha256=file_hash(weights),
            tokenizer_sha256=file_hash(snapshot / "tokenizer.json"),
            backend="sentence-transformers-pytorch-cpu",
            pooling="SBERT modules; mean of overflow windows then L2",
            test_encoder=False,
            max_tokens=model["max_tokens"],
            stride=model["stride"],
            fine_tuned_on_source=False,
        )
        vectors, self.window_counts = [], []
        for offset in range(0, len(texts), model["batch_size"]):
            batch = texts[offset : offset + model["batch_size"]]
            tokens = encoder.tokenizer(
                batch,
                truncation=True,
                max_length=model["max_tokens"],
                stride=model["stride"],
                return_overflowing_tokens=True,
                padding=True,
                return_tensors="pt",
            )
            owners = tokens.pop("overflow_to_sample_mapping").numpy()
            pooled = [[] for _ in batch]
            with torch.inference_mode():
                for index in range(0, len(owners), model["batch_size"]):
                    features = {
                        key: value[index : index + model["batch_size"]]
                        for key, value in tokens.items()
                    }
                    embeddings = encoder(features)["sentence_embedding"].cpu().numpy()
                    for owner, vector in zip(
                        owners[index : index + model["batch_size"]], embeddings
                    ):
                        pooled[int(owner)].append(vector)
            self.window_counts.extend(len(group) for group in pooled)
            vectors.extend(np.mean(group, axis=0) for group in pooled)
            if offset % (model["batch_size"] * 10) == 0:
                self._event(
                    "embedding",
                    complete=offset + len(batch),
                    total=len(texts),
                    encoder="SBERT",
                )
        return np.asarray(vectors, dtype=np.float32)

    def _cache_dir(self):
        path = Path(self.config["cache_dir"])
        path = path if path.is_absolute() else ROOT / path
        path.mkdir(parents=True, exist_ok=True)
        return path

    def _reuse_embeddings(self, directory):
        """Reuse immutable SBERT features only after verifying data/model identity."""
        data = self.load_run(directory)
        summary = data["summary"]
        if summary["source_sha256"] != self.input_sha256:
            raise ValueError("Embedding source workbook differs from the current input")
        if summary["config"]["model"] != self.config["model"]:
            raise ValueError("Embedding model configuration differs from the saved run")
        for column in ("sentence_id", "embedding_text"):
            if data["sentences"][column].tolist() != self.sentences[column].tolist():
                raise ValueError("Prepared sentences differ from the embedding source")
        vectors = data["embeddings"]
        if (
            vectors.ndim != 2
            or not np.isfinite(vectors).all()
            or not np.allclose(np.linalg.norm(vectors, axis=1), 1, atol=1e-5)
        ):
            raise ValueError("Saved SBERT embeddings must be finite and L2 normalized")
        self.embeddings = vectors.copy()
        self.window_counts = data["sentences"]["embedding_windows"].astype(int).tolist()
        self.model_info = deepcopy(summary["model"])
        self.model_info.update(
            reused_from=str(directory),
            reused_manifest_sha256=data["artifact_manifest_sha256"],
        )
        self._event("embeddings_reused", **self.model_info)
        return self

    def encode(self, *, reuse_run=None):
        if reuse_run is not None:
            if self.encoder is not None:
                raise ValueError(
                    "Cannot combine an injected encoder and saved embeddings"
                )
            return self._reuse_embeddings(reuse_run)
        texts = self.sentences["embedding_text"].tolist()
        key = sha256(
            json.dumps(
                {
                    "model": self.config["model"],
                    "texts": texts,
                    "version": ENCODER_CACHE_VERSION,
                },
                ensure_ascii=False,
                sort_keys=True,
            ).encode()
        ).hexdigest()
        cached = self._cache_dir() / (key + ".npy")
        metadata = cached.with_suffix(".json")
        cache_hit = False
        if self.encoder is None and cached.exists() and metadata.exists():
            meta = json.loads(metadata.read_text())
            if file_hash(cached) == meta["sha256"]:
                vectors = np.load(cached, allow_pickle=False)
                self.model_info = meta["model_info"]
                self.window_counts = meta["window_counts"]
                cache_hit = True
        if not cache_hit:
            if self.encoder is not None:
                vectors = np.asarray(self.encoder(texts), dtype=np.float32)
                self.window_counts = [1] * len(texts)
            else:
                vectors = self._encode_sbert(texts)
        vectors = np.asarray(vectors, dtype=np.float32)
        if (
            vectors.ndim != 2
            or vectors.shape[0] != len(texts)
            or not np.isfinite(vectors).all()
        ):
            raise ValueError("Invalid encoder matrix")
        norms = np.linalg.norm(vectors, axis=1, keepdims=True)
        if np.any(norms < 1e-8):
            raise ValueError("Zero encoder vector")
        self.embeddings = (
            vectors
            if cache_hit and np.allclose(norms, 1, atol=1e-5)
            else vectors / norms
        )
        if self.encoder is None and not cache_hit:
            temporary = cached.with_suffix(".tmp.npy")
            np.save(temporary, self.embeddings, allow_pickle=False)
            temporary.replace(cached)
            json_write(
                metadata,
                {
                    "sha256": file_hash(cached),
                    "model_info": self.model_info,
                    "window_counts": self.window_counts,
                },
            )
        self.model_info.update(
            cache_hit=cache_hit,
            dimensions=self.embeddings.shape[1],
            long_sentence_windows=sum(n > 1 for n in self.window_counts),
        )
        self._event(
            "encoded",
            sentences=len(texts),
            dimensions=self.embeddings.shape[1],
            cache_hit=cache_hit,
        )
        return self

    @staticmethod
    def stability(first, second):
        from sklearn.metrics import adjusted_rand_score

        first, second = np.asarray(first), np.asarray(second)
        union = (first >= 0) | (second >= 0)
        if union.sum() < 8:
            return 0.0
        n = len(first)
        a = np.where(first >= 0, first, n + np.arange(n))[union]
        b = np.where(second >= 0, second, n + np.arange(n))[union]
        return float(adjusted_rand_score(a, b))

    @staticmethod
    def select_candidate(comparisons, tolerance):
        eligible = [row for row in comparisons if row["eligible"]]
        if not eligible:
            raise ValueError("No candidate passed the quality/coverage constraints")
        highest = max(row["selection_score"] for row in eligible)
        close = [
            row for row in eligible if row["selection_score"] >= highest - tolerance
        ]
        selected = max(
            close,
            key=lambda row: (
                row["silhouette_cosine"],
                row["stability_ari"],
                row["coverage"],
            ),
        )
        policy = {
            "highest_composite_score": highest,
            "score_tolerance": tolerance,
            "near_tie_candidates": [row["candidate"] for row in close],
            "tie_break": "prefer larger cosine silhouette in original embedding space, then stability and coverage",
            "interpretation": "heuristic operational tolerance, not statistical significance or labeled accuracy",
        }
        return selected, policy

    def _evaluate(self, name, labels, alternate, parameters, seconds):
        from sklearn.metrics import silhouette_score

        labels = np.asarray(labels)
        assigned = labels >= 0
        counts = Counter(labels[assigned].tolist())
        coverage = float(assigned.mean())
        largest = max(counts.values(), default=0) / max(int(assigned.sum()), 1)
        stability = self.stability(labels, alternate)
        sample = self.evaluation_indices
        selected = sample[labels[sample] >= 0]
        silhouette = None
        if 2 <= len(set(labels[selected])) < len(selected):
            if hasattr(self, "evaluation_distances"):
                mask = labels[sample] >= 0
                distances = self.evaluation_distances[np.ix_(mask, mask)]
                silhouette = float(
                    silhouette_score(distances, labels[selected], metric="precomputed")
                )
            else:
                silhouette = float(
                    silhouette_score(
                        self.embeddings[selected], labels[selected], metric="cosine"
                    )
                )
        valid = (
            len(counts) >= 2
            and silhouette is not None
            and coverage >= self.config["evaluation"]["min_coverage"]
            and largest <= self.config["evaluation"]["max_largest_cluster"]
        )
        score = (
            (
                0.6 * max(silhouette or 0, 0) * coverage
                + 0.25 * max(stability, 0)
                + 0.15 * coverage
            )
            if valid
            else None
        )
        result = {
            "candidate": name,
            "parameters": parameters,
            "clusters": len(counts),
            "coverage": coverage,
            "noise_fraction": 1 - coverage,
            "largest_cluster_fraction": largest,
            "silhouette_cosine": silhouette,
            "stability_ari": stability,
            "selection_score": score,
            "eligible": valid,
            "seconds": seconds,
            "silhouette_sample_assigned": len(selected),
            "stability_union_points": int(((labels >= 0) | (alternate >= 0)).sum()),
        }
        self.comparisons.append(result)
        self._event("candidate", **result)
        return result

    @staticmethod
    def _filter_radius(graph, eps):
        """Preserve explicit zero distances, including identical sentence vectors."""
        from scipy.sparse import csr_matrix

        keep = graph.data <= eps
        cumulative = np.concatenate(([0], np.cumsum(keep, dtype=np.int64)))
        return csr_matrix(
            (graph.data[keep], graph.indices[keep], cumulative[graph.indptr]),
            shape=graph.shape,
        )

    def _reduce_embeddings(self):
        from umap import UMAP

        u = self.config["umap"]
        n = len(self.embeddings)
        parameters = dict(
            n_neighbors=min(u["n_neighbors"], n - 1),
            n_components=min(u["n_components"], n - 2),
            min_dist=u["min_dist"],
            metric="cosine",
            n_jobs=1,
        )
        seeds = [self.config["seed"], self.config["seed"] + 1]
        identity = dict(
            embeddings_sha256=sha256(self.embeddings.tobytes()).hexdigest(),
            parameters=parameters,
            seeds=seeds,
            umap_version=importlib.metadata.version("umap-learn"),
        )
        key = sha256(json.dumps(identity, sort_keys=True).encode()).hexdigest()
        cached = self._cache_dir() / ("umap_" + key + ".npz")
        metadata = cached.with_suffix(".json")
        cache_hit = False
        if cached.exists() and metadata.exists():
            info = json.loads(metadata.read_text())
            if file_hash(cached) == info["sha256"]:
                with np.load(cached, allow_pickle=False) as archive:
                    reductions = [archive["seed_primary"], archive["seed_alternate"]]
                cache_hit = all(
                    x.shape == (n, parameters["n_components"]) and np.isfinite(x).all()
                    for x in reductions
                )
        if not cache_hit:
            reductions = []
            for seed in seeds:
                self._event(
                    "reducing",
                    algorithm="UMAP",
                    seed=seed,
                    dimensions=parameters["n_components"],
                )
                reductions.append(
                    np.asarray(
                        UMAP(**parameters, random_state=seed).fit_transform(
                            self.embeddings
                        ),
                        dtype=np.float32,
                    )
                )
            temporary = cached.with_suffix(".tmp.npz")
            np.savez_compressed(
                temporary, seed_primary=reductions[0], seed_alternate=reductions[1]
            )
            temporary.replace(cached)
            json_write(metadata, dict(sha256=file_hash(cached), identity=identity))
        self.reduced_embeddings = reductions
        self.reduction_info = dict(
            algorithm="UMAP",
            parameters=parameters,
            seeds=seeds,
            cache_hit=cache_hit,
            clustering_metric="euclidean",
            original_dimensions=self.embeddings.shape[1],
        )
        self._event("reduced", **self.reduction_info)
        return reductions

    def _radius_graph(self, features=None, metric="cosine"):
        from scipy.sparse import vstack
        from sklearn.neighbors import NearestNeighbors

        options = self.config["dbscan"]
        features = self.embeddings if features is None else features
        radius = max(options["eps"])
        neighbors = NearestNeighbors(
            radius=radius, metric=metric, algorithm="brute", n_jobs=1
        ).fit(features)
        batches, edges = [], 0
        for start in range(0, len(features), options["neighborhood_batch_size"]):
            block = neighbors.radius_neighbors_graph(
                features[start : start + options["neighborhood_batch_size"]],
                radius=radius,
                mode="distance",
                sort_results=True,
            )
            edges += block.nnz
            if edges > options["max_neighbor_edges"]:
                raise ValueError(
                    "DBSCAN radius graph exceeds max_neighbor_edges; review eps or the explicit memory budget"
                )
            batches.append(block.astype(np.float32))
            self._event(
                "neighborhoods",
                complete=min(start + block.shape[0], len(features)),
                total=len(features),
                edges=edges,
                radius=radius,
                metric=metric,
            )
        return vstack(batches, format="csr")

    def _merge_topics(self, labels):
        """Merge entire DBSCAN groups using sentence-weighted Ward linkage.

        Minimize the increase in within-topic squared error in original SBERT
        space. Noise stays -1; source groups never split. Ties follow source IDs.
        """
        labels = np.asarray(labels, dtype=int)
        ids = np.unique(labels[labels >= 0])
        target = self.config["target_topics"]
        if target is None:
            return labels.copy(), None
        if len(ids) < target:
            raise ValueError("DBSCAN found fewer groups than target_topics")
        if len(ids) > 2048:
            raise ValueError(
                "Topic reduction exceeds 2048 source groups; review the DBSCAN grid to bound pairwise memory"
            )
        assigned = labels >= 0
        positions = np.searchsorted(ids, labels[assigned])
        sums = np.zeros((len(ids), self.embeddings.shape[1]), dtype=np.float64)
        np.add.at(sums, positions, self.embeddings[assigned])
        sizes = np.bincount(positions, minlength=len(ids)).astype(np.float64)
        means = sums / sizes[:, None]
        norms = (means**2).sum(axis=1)
        distances = np.maximum(norms[:, None] + norms[None, :] - 2 * means @ means.T, 0)
        costs = (
            distances
            * sizes[:, None]
            * sizes[None, :]
            / (sizes[:, None] + sizes[None, :])
        )
        np.fill_diagonal(costs, np.inf)
        active = np.ones(len(ids), dtype=bool)
        members = [[int(cid)] for cid in ids]
        merges = []
        while active.sum() > target:
            a, b = np.unravel_index(np.argmin(costs), costs.shape)
            merges.append(
                {
                    "left": members[a].copy(),
                    "right": members[b].copy(),
                    "squared_error_increase": float(costs[a, b]),
                }
            )
            sums[a] += sums[b]
            sizes[a] += sizes[b]
            members[a] = sorted(members[a] + members[b])
            active[b] = False
            means[a] = sums[a] / sizes[a]
            distances = ((means - means[a]) ** 2).sum(axis=1)
            costs[a, :] = distances * sizes[a] * sizes / (sizes[a] + sizes)
            costs[:, a] = costs[a, :]
            costs[~active, :] = np.inf
            costs[:, ~active] = np.inf
            costs[a, a] = np.inf
        mapping = {
            cid: final_id
            for final_id, index in enumerate(np.flatnonzero(active))
            for cid in members[index]
        }
        reduced = np.full_like(labels, -1)
        reduced[assigned] = np.array([mapping[int(cid)] for cid in ids])[positions]
        return reduced, {
            "method": "sentence_weighted_ward",
            "embedding_space": "original_normalized_SBERT",
            "target_topics": target,
            "original_topics": len(ids),
            "final_topics": int(active.sum()),
            "noise_preserved": True,
            "mapping": {str(cid): int(final) for cid, final in mapping.items()},
            "merges": merges,
        }

    def compare(self):
        """Select DBSCAN parameters on UMAP; evaluate in original SBERT geometry."""
        from sklearn.cluster import DBSCAN
        from sklearn.metrics.pairwise import cosine_distances
        from sklearn.neighbors import sort_graph_by_row_values
        from threadpoolctl import threadpool_limits

        n = len(self.sentences)
        rng = np.random.default_rng(self.config["seed"])
        self.evaluation_indices = np.sort(
            rng.choice(
                n, min(n, self.config["evaluation"]["sample_size"]), replace=False
            )
        )
        self.comparisons, fitted = [], {}
        with threadpool_limits(limits=self.config["model"]["threads"]):
            self.evaluation_distances = cosine_distances(
                self.embeddings[self.evaluation_indices]
            )
            np.fill_diagonal(self.evaluation_distances, 0)
            reductions = self._reduce_embeddings()
            graphs = [
                self._radius_graph(features, metric="euclidean")
                for features in reductions
            ]
            for eps in self.config["dbscan"]["eps"]:
                filtered = [
                    sort_graph_by_row_values(
                        self._filter_radius(graph, eps),
                        copy=False,
                        warn_when_not_sorted=False,
                    )
                    for graph in graphs
                ]
                for samples in self.config["dbscan"]["min_samples"]:
                    if samples > n:
                        continue
                    started = time.monotonic()
                    models = [
                        DBSCAN(
                            eps=eps, min_samples=samples, metric="precomputed", n_jobs=1
                        ).fit(graph)
                        for graph in filtered
                    ]
                    name = f"sbert_umap_dbscan_eps{eps:g}_min{samples}"
                    parameters = dict(
                        algorithm="DBSCAN",
                        eps=eps,
                        min_samples=samples,
                        metric="euclidean",
                        fit_space="UMAP_of_normalized_Sbert_embeddings",
                        umap=self.reduction_info["parameters"],
                        umap_seeds=self.reduction_info["seeds"],
                        neighborhood="sparse_radius_graph",
                    )
                    raw_labels = models[0].labels_.copy()
                    labels, alternate = raw_labels, models[1].labels_
                    raw_counts = [len(set(x) - {-1}) for x in (labels, alternate)]
                    target = self.config["target_topics"]
                    topic_count_eligible = target is None or min(raw_counts) >= target
                    reduction = None
                    if target is not None and topic_count_eligible:
                        labels, reduction = self._merge_topics(labels)
                        alternate, _ = self._merge_topics(alternate)
                        parameters["topic_reduction"] = reduction["method"]
                        parameters["target_topics"] = target
                    result = self._evaluate(
                        name,
                        labels,
                        alternate,
                        parameters,
                        time.monotonic() - started,
                    )
                    result["dbscan_clusters"] = raw_counts[0]
                    result["alternate_dbscan_clusters"] = raw_counts[1]
                    result["topic_count_eligible"] = topic_count_eligible
                    result["eligible"] &= topic_count_eligible
                    result["geometry_eligible"] = result["eligible"]
                    result["coherence_cv"] = None
                    result["coherence_status"] = "not_evaluated_geometry_constraints"
                    core = np.zeros(n, dtype=bool)
                    core[models[0].core_sample_indices_] = True
                    fitted[name] = (labels, core, raw_labels, reduction)
        # Evaluate in composite-score order. Once a passing score is known,
        # only its tolerance band can win. Lower scores need no expensive
        # coherence calculation; this preserves the exact selection rule.
        best_passing = None
        tolerance = self.config["evaluation"]["score_tolerance"]
        for result in sorted(
            (row for row in self.comparisons if row["geometry_eligible"]),
            key=lambda row: row["selection_score"],
            reverse=True,
        ):
            if (
                best_passing is not None
                and result["selection_score"] < best_passing - tolerance
            ):
                result["eligible"] = False
                result["coherence_status"] = "not_evaluated_below_passing_score_band"
                continue
            coherence = self.score_coherence(
                self._ctfidf_terms(fitted[result["candidate"]][0])
            )
            result["coherence_cv"] = coherence["macro_mean"]
            result["coherence_topics"] = coherence["topics_evaluated"]
            result["eligible"] = coherence["target_met"]
            result["coherence_status"] = "passed" if result["eligible"] else "failed"
            if result["eligible"] and best_passing is None:
                best_passing = result["selection_score"]
            self._event("candidate_coherence", **result)
        self.selected, self.selection_policy = self.select_candidate(
            self.comparisons, self.config["evaluation"]["score_tolerance"]
        )
        self.selection_policy.update(
            candidate_family="SBERT + UMAP + DBSCAN + optional centroid topic merging",
            target_topics=self.config["target_topics"],
            stability_test="final topic partitions on two seeded UMAP projections, after merging when configured",
            coherence_gate=self.config["coherence"],
            coherence_search="descending composite score; all candidates in the best passing score tolerance band are evaluated; lower scores are explicitly marked unevaluated",
        )
        labels, core, raw_labels, self.topic_reduction = fitted[
            self.selected["candidate"]
        ]
        self.sentences["cluster_id"] = np.asarray(labels, dtype=int)
        self.sentences["dbscan_cluster_id"] = raw_labels
        self.sentences["dbscan_core_point"] = core
        self.sentences["membership_strength"] = np.nan
        self.sentences["embedding_windows"] = self.window_counts
        self.occurrences = self.occurrences.merge(
            self.sentences[["sentence_id", "cluster_id", "dbscan_cluster_id"]],
            on="sentence_id",
            validate="many_to_one",
        )
        self._event("selected", **self.selected)
        return self

    @staticmethod
    def _topic_vectorizer():
        from sklearn.feature_extraction.text import ENGLISH_STOP_WORDS, CountVectorizer

        return CountVectorizer(
            ngram_range=(1, 2),
            min_df=2,
            max_df=0.95,
            max_features=25000,
            stop_words=sorted(INDONESIAN_STOPWORDS | set(ENGLISH_STOP_WORDS)),
            token_pattern=r"(?u)\b[^\W\d_]\w+\b",
        )

    def _ctfidf_terms(self, labels):
        """Use the same class TF-IDF vocabulary for selection and publication."""
        from scipy.sparse import csr_matrix
        from sklearn.preprocessing import normalize

        labels = np.asarray(labels)
        ids = sorted(set(labels) - {-1})
        if not hasattr(self, "_term_matrix"):
            vectorizer = self._topic_vectorizer()
            self._term_matrix = vectorizer.fit_transform(
                self.sentences["embedding_text"]
            )
            self._term_vocabulary = np.asarray(vectorizer.get_feature_names_out())
        mapping = {cid: i for i, cid in enumerate(ids)}
        assigned_indices = np.flatnonzero(labels >= 0)
        group = csr_matrix(
            (
                np.ones(len(assigned_indices)),
                ([mapping[labels[i]] for i in assigned_indices], assigned_indices),
            ),
            shape=(len(ids), len(labels)),
        )
        counts = group @ self._term_matrix
        average_length = float(np.asarray(counts.sum(axis=1)).mean())
        idf = np.log1p(
            average_length / np.maximum(np.asarray(counts.sum(axis=0)).ravel(), 1)
        )
        ctfidf = normalize(counts, norm="l1", axis=1).multiply(idf).tocsr()
        result = {}
        for cid in ids:
            scores = ctfidf[mapping[cid]].toarray().ravel()
            ranked = np.argsort(scores)[::-1]
            # A short topic must not be padded with zero-weight unrelated words.
            result[int(cid)] = self._term_vocabulary[
                ranked[scores[ranked] > 0][:12]
            ].tolist()
        return result

    def score_coherence(self, topic_terms):
        """Macro c_v on ALL non-noise topics and a fixed full-corpus reference.

        Unigrams and bigrams use exactly the c-TF-IDF analyzer. A bigram is one
        token, not split after ranking. Duplicate sentence occurrences are not
        repeated; reference texts include noise. Invalid topics invalidate the
        aggregate instead of being silently discarded.
        """
        from gensim.corpora import Dictionary
        from gensim.models import CoherenceModel

        settings = self.config["coherence"]
        if not hasattr(self, "_coherence_reference"):
            analyzer = self._topic_vectorizer().build_analyzer()
            texts = [analyzer(text) for text in self.sentences["embedding_text"]]
            self._coherence_reference = (
                texts,
                Dictionary(texts),
                sha256(
                    json.dumps(
                        texts, ensure_ascii=False, separators=(",", ":")
                    ).encode()
                ).hexdigest(),
            )
        texts, dictionary, corpus_sha256 = self._coherence_reference
        rows = [
            {
                "cluster_id": int(cid),
                "terms": list(terms[: settings["topn"]]),
                "score": None,
            }
            for cid, terms in sorted(topic_terms.items())
            if cid >= 0
        ]
        valid = [
            row
            for row in rows
            if len(set(row["terms"])) >= 2
            and all(term in dictionary.token2id for term in row["terms"])
        ]
        if valid:
            model = CoherenceModel(
                topics=[row["terms"] for row in valid],
                texts=texts,
                dictionary=dictionary,
                coherence="c_v",
                topn=settings["topn"],
                window_size=settings["window_size"],
                processes=1,
            )
            for row, value in zip(valid, model.get_coherence_per_topic()):
                if np.isfinite(value):
                    row["score"] = float(value)
        evaluated = sum(row["score"] is not None for row in rows)
        mean = (
            float(np.mean([row["score"] for row in rows]))
            if rows and evaluated == len(rows)
            else None
        )
        return {
            **settings,
            "macro_mean": mean,
            "target_met": mean is not None and mean >= settings["min_score"],
            "topics_total": len(rows),
            "topics_evaluated": evaluated,
            "topics_below_target": sum(
                row["score"] is None or row["score"] < settings["min_score"]
                for row in rows
            ),
            "reference_documents": len(texts),
            "reference_sha256": corpus_sha256,
            "tokenization": "c-TF-IDF CountVectorizer analyzer: stopword-filtered unigrams followed by bigrams; phrases remain single tokens",
            "reference_scope": "all unique embedding_text sentences, including DBSCAN noise; no cluster concatenation",
            "aggregation": "unweighted mean of all non-noise topics; any unscorable topic invalidates the mean",
            "interpretation": "in-sample lexical coherence of topic terms, not accuracy or held-out generalization",
            "gensim_version": importlib.metadata.version("gensim"),
            "per_topic": rows,
        }

    @classmethod
    def evaluate_run_coherence(cls, directory, settings=None):
        data = cls.load_run(directory)
        evaluator = cls({"coherence": settings} if settings else None)
        evaluator.sentences = data["sentences"]
        result = evaluator.score_coherence(
            {topic["cluster_id"]: topic["terms"] for topic in data["topics"]}
        )
        result.update(
            source_run=str(Path(directory)),
            source_manifest_sha256=data["artifact_manifest_sha256"],
            source_sha256=data["summary"]["source_sha256"],
        )
        return result

    def describe_topics(self):
        labels = self.sentences["cluster_id"].to_numpy()
        topic_terms = self._ctfidf_terms(labels)
        self.coherence = self.score_coherence(topic_terms)
        if not self.coherence["target_met"]:
            raise ValueError("Selected topics failed the coherence c_v threshold")
        if hasattr(self, "selected"):
            self.selected["coherence_cv"] = self.coherence["macro_mean"]
        scores_by_id = {
            row["cluster_id"]: row["score"] for row in self.coherence["per_topic"]
        }
        ids = sorted(topic_terms)
        topics, centroids = [], []
        occurrence_counts = self.occurrences.groupby("cluster_id").size().to_dict()
        for cid in ids:
            indices = np.flatnonzero(labels == cid)
            centroid = self.embeddings[indices].mean(axis=0)
            centroid /= max(np.linalg.norm(centroid), 1e-8)
            centroids.append(centroid)
            similarities = self.embeddings[indices] @ centroid
            self.sentences.loc[self.sentences.index[indices], "centroid_similarity"] = (
                similarities
            )
            top = topic_terms[cid]
            chosen = []
            for term in top:
                if not any(
                    set(term.split()) <= set(other.split())
                    or set(other.split()) <= set(term.split())
                    for other in chosen
                ):
                    chosen.append(term)
                if len(chosen) == 3:
                    break
            representative = indices[np.argsort(similarities)[::-1][:5]]
            source = self.occurrences[self.occurrences["cluster_id"] == cid]
            topics.append(
                {
                    "cluster_id": int(cid),
                    "label": " / ".join(chosen) or f"Topik {cid}",
                    "terms": top,
                    "coherence_cv": scores_by_id[cid],
                    "unique_sentences": len(indices),
                    "occurrences": int(occurrence_counts.get(cid, 0)),
                    "platform_counts": {
                        str(k): int(v)
                        for k, v in source["platform"].value_counts().items()
                    },
                    "mean_centroid_similarity": float(similarities.mean()),
                    "representative_ids": self.sentences.iloc[representative][
                        "sentence_id"
                    ].tolist(),
                }
            )
        self.topics = topics
        self.centroids = np.asarray(centroids, dtype=np.float32)
        self.graph = self.build_topic_graph()
        return self

    def build_topic_graph(self):
        if not self.topics:
            return {"nodes": [], "edges": []}
        similarities = self.centroids @ self.centroids.T
        nodes = [
            {
                "id": str(t["cluster_id"]),
                **t,
                "kind": "topic",
                "inference": "unsupervised_cluster",
            }
            for t in self.topics
        ]
        edges = {}
        g = self.config["graph"]
        for i, topic in enumerate(self.topics):
            neighbors = [j for j in np.argsort(similarities[i])[::-1] if j != i][
                : g["neighbors"]
            ]
            for j in neighbors:
                score = float(similarities[i, j])
                if score < g["min_similarity"]:
                    continue
                a, b = sorted([topic["cluster_id"], self.topics[j]["cluster_id"]])
                edges[a, b] = {
                    "source": str(a),
                    "target": str(b),
                    "similarity": score,
                    "relation": "semantic_similarity",
                    "confidence": "INFERRED",
                    "score_kind": "cosine_not_probability",
                }
        return {
            "nodes": nodes,
            "edges": list(edges.values()),
            "meaning": "Similarity of mean normalized sentence embeddings; no social-user relation inferred",
        }

    @staticmethod
    def _validate_explanation_options(options):
        expected = json.loads((ROOT / "social_nlp.json").read_text())["explanation"]
        if set(options) != set(expected):
            raise ValueError("Invalid explanation option fields")
        if options["date_basis"] not in {"post_date", "content_date"}:
            raise ValueError("Explanation date_basis must be post_date or content_date")
        if type(options["include_estimated_dates"]) is not bool:
            raise ValueError("include_estimated_dates must be boolean")
        for key in [
            "window_months",
            "min_periods",
            "min_monthly_comments",
            "min_topic_comments",
            "min_window_topic_comments",
            "min_roots",
            "min_term_comments",
            "max_evidence",
        ]:
            if type(options[key]) is not int or options[key] < 1:
                raise ValueError("Invalid explanation " + key)
        if (
            not 1 <= options["window_months"] <= 12
            or not 3 <= options["min_periods"] <= 60
            or not 1 <= options["max_evidence"] <= 10
        ):
            raise ValueError("Explanation window/evidence limits exceeded")
        for key in ["dominance_warning_threshold", "kendall_threshold", "mmr_lambda"]:
            if (
                isinstance(options[key], bool)
                or not isinstance(options[key], (int, float))
                or not np.isfinite(options[key])
                or not 0 < options[key] <= 1
            ):
                raise ValueError("Invalid explanation " + key)
        for key in [
            "change_threshold_pp",
            "slope_threshold_pp",
            "platform_shift_threshold_pp",
        ]:
            if (
                isinstance(options[key], bool)
                or not isinstance(options[key], (int, float))
                or not np.isfinite(options[key])
                or not 0 < options[key] <= 100
            ):
                raise ValueError("Invalid explanation " + key)

    @staticmethod
    def parent_post_key(row):
        """Identify the parent thread, not the individual comment URL."""
        platform = str(row.get("platform", ""))
        parsed = urlparse(str(row.get("url", "")))
        host = (parsed.hostname or "").lower()
        if platform == "youtube" and (
            host == "youtube.com" or host.endswith(".youtube.com")
        ):
            video = parse_qs(parsed.query).get("v", [])
            if video and video[0]:
                return "youtube:" + video[0], "native_url"
        if platform == "youtube" and host == "youtu.be" and parsed.path.strip("/"):
            return "youtube:" + parsed.path.strip("/").split("/")[0], "native_url"
        if platform == "reddit" and (
            host == "reddit.com" or host.endswith(".reddit.com")
        ):
            found = re.search(r"/comments/([A-Za-z0-9]+)/", parsed.path + "/")
            if found:
                return "reddit:" + found[1], "native_url"
        context = f"{platform}|{row.get('post_title', '')}|{row.get('post_date', '')}"
        return "context:" + sha256(context.encode()).hexdigest()[
            :20
        ], "context_fallback"

    def _explanation_evidence(self, source, terms, options):
        """Extract real sentences using relevance plus MMR diversity."""
        if source.empty:
            return []
        sentence_indices = {sid: i for i, sid in enumerate(self.sentences.sentence_id)}
        similarities = self.sentences.set_index("sentence_id").get(
            "centroid_similarity", pd.Series(dtype=float)
        )
        candidates = (
            source.sort_values("source_row").drop_duplicates("sentence_id").copy()
        )
        patterns = [
            re.compile(r"(?<!\w)" + re.escape(term.casefold()) + r"(?!\w)")
            for term in terms
        ]
        weights = np.array([1 / (i + 1) for i in range(len(terms))])
        candidates["matched_terms"] = [
            [
                term
                for term, pattern in zip(terms, patterns)
                if pattern.search(self._clean(text).casefold())
            ]
            for text in candidates.sentence
        ]
        candidates["lexical_score"] = [
            sum(weight for term, weight in zip(terms, weights) if term in matched)
            / max(weights.sum(), 1)
            for matched in candidates.matched_terms
        ]
        candidates["centrality"] = (
            candidates.sentence_id.map(similarities).fillna(0).clip(0, 1)
        )
        candidates["length_score"] = [
            min(len(re.findall(r"\w+", text)) / 12, 1) for text in candidates.sentence
        ]
        candidates["relevance"] = (
            0.45 * candidates.centrality
            + 0.4 * candidates.lexical_score
            + 0.15 * candidates.length_score
        )
        candidates = (
            candidates.sort_values(
                ["relevance", "sentence_id"], ascending=[False, True]
            )
            .head(200)
            .reset_index(drop=True)
        )
        # If informative matches exist, exclude generic acknowledgements first.
        informative = (candidates.lexical_score > 0) & (candidates.length_score >= 0.5)
        if informative.sum() >= options["max_evidence"]:
            candidates = candidates[informative].reset_index(drop=True)
        vectors = getattr(self, "embeddings", None)
        if vectors is not None:
            if len(vectors) != len(self.sentences):
                raise ValueError("Invalid embeddings for topic explanation")
            indices = [sentence_indices[sid] for sid in candidates.sentence_id]
            matrix = np.asarray(vectors[indices], dtype=np.float32)
            matrix = matrix / np.maximum(
                np.linalg.norm(matrix, axis=1, keepdims=True), 1e-8
            )
        else:
            matrix = None
        selected, used_roots, used_comments = [], set(), set()
        for _ in range(min(options["max_evidence"], len(candidates))):
            remaining = [i for i in range(len(candidates)) if i not in selected]
            diverse = [
                i
                for i in remaining
                if candidates.iloc[i]._root not in used_roots
                and candidates.iloc[i]._comment not in used_comments
            ]
            if diverse:
                remaining = diverse

            def score(i):
                if not selected:
                    redundancy = 0.0
                elif matrix is not None:
                    redundancy = float(
                        np.maximum(matrix[selected] @ matrix[i], 0).max()
                    )
                else:
                    tokens = set(
                        self._clean(candidates.iloc[i].sentence).casefold().split()
                    )
                    redundancy = max(
                        len(
                            tokens
                            & set(
                                self._clean(candidates.iloc[j].sentence)
                                .casefold()
                                .split()
                            )
                        )
                        / max(
                            len(
                                tokens
                                | set(
                                    self._clean(candidates.iloc[j].sentence)
                                    .casefold()
                                    .split()
                                )
                            ),
                            1,
                        )
                        for j in selected
                    )
                return (
                    options["mmr_lambda"] * float(candidates.iloc[i].relevance)
                    - (1 - options["mmr_lambda"]) * redundancy
                )

            chosen = max(remaining, key=lambda i: (score(i), -i))
            selected.append(chosen)
            used_roots.add(candidates.iloc[chosen]._root)
            used_comments.add(candidates.iloc[chosen]._comment)
        evidence = []
        for i in selected:
            row = candidates.iloc[i]
            record = {
                key: row.get(key)
                for key in [
                    "sentence_id",
                    "source_row",
                    "sentence",
                    "platform",
                    "post_title",
                    "url",
                    "post_date",
                    "content_date",
                    "post_date_accuracy",
                    "content_date_accuracy",
                ]
            }
            record.update(
                source_row=int(row.source_row),
                parent_post_key=row._root,
                matched_terms=row.matched_terms,
                centroid_similarity=float(row.centrality),
                relevance_score=float(row.relevance),
                confidence="EXTRACTED",
                selection_method="MMR_cosine"
                if matrix is not None
                else "MMR_token_jaccard",
            )
            evidence.append(record)
        return evidence

    @staticmethod
    def _explanation_trend(series, cohort, topic_comments, options):
        from scipy.stats import kendalltau, theilslopes

        valid = [row for row in series if row["eligible_for_trend"]]
        slope, tau = None, None
        if len(valid) >= 3:
            x = np.array(
                [pd.Period(row["period"], freq="M").ordinal for row in valid],
                dtype=float,
            )
            y = np.array([row["share"] * 100 for row in valid])
            slope = float(theilslopes(y, x).slope)
            raw_tau = float(kendalltau(x, y).statistic)
            tau = raw_tau if np.isfinite(raw_tau) else 0.0
        complete = [row for row in series if row["full_calendar_month"]]
        width = options["window_months"]
        comparison = None
        cautions, blockers = [], []
        if len(complete) >= width * 2:
            before, after = complete[-2 * width : -width], complete[-width:]

            def aggregate(rows):
                denominator = sum(row["all_comments"] for row in rows)
                numerator = sum(row["topic_comments"] for row in rows)
                platforms = {}
                for platform in ["youtube", "reddit"]:
                    total = sum(
                        row["platforms"][platform]["all_comments"] for row in rows
                    )
                    hits = sum(
                        row["platforms"][platform]["topic_comments"] for row in rows
                    )
                    platforms[platform] = {
                        "all_comments": total,
                        "topic_comments": hits,
                        "share": hits / total if total else None,
                    }
                members = topic_comments[
                    topic_comments._month.isin([row["period"] for row in rows])
                ]
                return {
                    "periods": [row["period"] for row in rows],
                    "all_comments": denominator,
                    "topic_comments": numerator,
                    "share": numerator / denominator if denominator else None,
                    "eligible": all(row["eligible_for_trend"] for row in rows),
                    "roots": int(members._root.nunique()),
                    "platforms": platforms,
                }

            previous, recent = aggregate(before), aggregate(after)
            delta = (
                (recent["share"] - previous["share"]) * 100
                if previous["share"] is not None and recent["share"] is not None
                else None
            )
            active = [
                p
                for p in ["youtube", "reddit"]
                if previous["platforms"][p]["all_comments"]
                or recent["platforms"][p]["all_comments"]
            ]
            adjusted = None
            mix_shift = None
            if previous["all_comments"] and recent["all_comments"]:
                mix_shift = max(
                    abs(
                        recent["platforms"][p]["all_comments"] / recent["all_comments"]
                        - previous["platforms"][p]["all_comments"]
                        / previous["all_comments"]
                    )
                    * 100
                    for p in active
                )
                if all(
                    min(
                        previous["platforms"][p]["all_comments"],
                        recent["platforms"][p]["all_comments"],
                    )
                    >= options["min_monthly_comments"]
                    for p in active
                ):
                    adjusted = sum(
                        previous["platforms"][p]["all_comments"]
                        / previous["all_comments"]
                        * (
                            recent["platforms"][p]["share"]
                            - previous["platforms"][p]["share"]
                        )
                        * 100
                        for p in active
                    )
            comparison = {
                "previous": previous,
                "recent": recent,
                "delta_pp": delta,
                "platform_standardized_delta_pp": adjusted,
                "platform_mix_shift_pp": mix_shift,
                "standardization": "previous-window platform weights; both windows require platform support",
            }
            if not previous["eligible"] or not recent["eligible"]:
                blockers.append(
                    "Jendela pembanding memuat bulan dengan terlalu sedikit komentar atau tanpa pengamatan."
                )
            if any(
                window["topic_comments"] and window["roots"] < options["min_roots"]
                for window in [previous, recent]
            ):
                blockers.append(
                    "Salah satu jendela pembanding bergantung pada terlalu sedikit video/post induk."
                )
            if (
                max(previous["topic_comments"], recent["topic_comments"])
                < options["min_window_topic_comments"]
            ):
                blockers.append(
                    "Penyebutan topik pada jendela pembanding masih terlalu sedikit."
                )
            if (
                mix_shift is not None
                and mix_shift >= options["platform_shift_threshold_pp"]
            ):
                cautions.append(
                    "Komposisi YouTube–Reddit berubah cukup besar; perubahan porsi dapat sensitif terhadap sumber yang terkumpul."
                )
            if adjusted is None and len(active) > 1:
                cautions.append(
                    "Perbandingan dengan komposisi platform tetap belum dapat dihitung karena dukungan salah satu platform kurang."
                )
        else:
            blockers.append(
                "Belum tersedia dua jendela bulan kalender lengkap untuk dibandingkan."
            )
        supported_comments = topic_comments[
            topic_comments._month.isin([row["period"] for row in valid])
        ]
        roots = supported_comments._root.value_counts()
        dominance = (
            float(roots.iloc[0] / len(supported_comments)) if len(roots) else 0.0
        )
        if len(valid) < options["min_periods"]:
            blockers.append(
                "Jumlah bulan dengan dukungan pengamatan belum cukup untuk merangkum arah."
            )
        if len(supported_comments) < options["min_topic_comments"]:
            blockers.append("Jumlah komentar bertopik masih terlalu sedikit.")
        if len(roots) < options["min_roots"]:
            blockers.append("Topik berasal dari terlalu sedikit video/post induk.")
        if dominance > options["dominance_warning_threshold"]:
            cautions.append(
                "Sebagian besar komentar topik berasal dari satu video/post; pola ini dapat mencerminkan thread tersebut."
            )
        direction = "inconclusive"
        delta = comparison["delta_pp"] if comparison else None
        spread = (
            (max(row["share"] for row in valid) - min(row["share"] for row in valid))
            * 100
            if valid
            else None
        )
        if not blockers and delta is not None and slope is not None:
            change, gradient, rank = (
                options["change_threshold_pp"],
                options["slope_threshold_pp"],
                options["kendall_threshold"],
            )
            if delta >= change:
                direction = (
                    "rising" if slope >= gradient and tau >= rank else "recent_increase"
                )
            elif delta <= -change:
                direction = (
                    "falling"
                    if slope <= -gradient and tau <= -rank
                    else "recent_decrease"
                )
            else:
                direction = "fluctuating" if spread >= max(10, 5 * change) else "stable"
            adjusted = comparison["platform_standardized_delta_pp"]
            if (
                adjusted is not None
                and abs(adjusted) >= change
                and abs(delta) >= change
                and adjusted * delta < 0
            ):
                direction = "fluctuating"
                cautions.append(
                    "Arah berubah ketika komposisi platform dibuat tetap; satu label naik/turun akan menyesatkan."
                )
        labels = {
            "rising": "Indikasi meningkat",
            "falling": "Indikasi menurun",
            "recent_increase": "Menguat belakangan",
            "recent_decrease": "Melemah belakangan",
            "fluctuating": "Fluktuatif",
            "stable": "Relatif stabil belakangan",
            "inconclusive": "Arah belum cukup jelas",
        }
        peak = (
            max(valid, key=lambda row: (row["share"], row["topic_comments"]))
            if valid
            else None
        )
        return {
            "direction": direction,
            "label": labels[direction],
            "support": "terbatas"
            if blockers or cautions
            else "memadai_untuk_deskripsi_snapshot",
            "slope_pp_per_month": slope,
            "kendall_tau": tau,
            "eligible_months": len(valid),
            "share_range_pp": spread,
            "comparison": comparison,
            "peak_share_period": peak["period"] if peak else None,
            "roots": len(roots),
            "support_comments": len(supported_comments),
            "dominant_root_share": dominance,
            "dominant_root_key": str(roots.index[0]) if len(roots) else None,
            "blockers": blockers,
            "cautions": cautions,
            "confidence": "INFERRED",
            "uncertainty": "descriptive support rules, not calibrated confidence, significance, causality or forecast",
        }

    @staticmethod
    def _explanation_term_changes(source, terms, comparison, options):
        if comparison is None:
            return []
        windows = [comparison["previous"]["periods"], comparison["recent"]["periods"]]
        sets = []
        for periods in windows:
            subset = source[source._month.isin(periods)]
            documents = (
                subset.groupby("_comment")
                .sentence.agg(lambda values: " ".join(values))
                .str.casefold()
            )
            sets.append(documents)
        old, recent = sets
        if not len(old) or not len(recent):
            return []
        changes = []
        for term in terms:
            pattern = re.compile(r"(?<!\w)" + re.escape(term.casefold()) + r"(?!\w)")
            a, b = (
                int(old.str.contains(pattern).sum()),
                int(recent.str.contains(pattern).sum()),
            )
            if a + b < options["min_term_comments"]:
                continue
            log_odds = float(
                np.log((b + 0.5) / (len(recent) - b + 0.5))
                - np.log((a + 0.5) / (len(old) - a + 0.5))
            )
            delta = (b / len(recent) - a / len(old)) * 100
            if abs(delta) >= options["change_threshold_pp"]:
                changes.append(
                    {
                        "term": term,
                        "previous_comments": a,
                        "recent_comments": b,
                        "previous_total": len(old),
                        "recent_total": len(recent),
                        "delta_pp": delta,
                        "smoothed_log_odds_delta": log_odds,
                        "direction": "more_frequent" if delta > 0 else "less_frequent",
                        "meaning": "term presence among this topic's comments; not proof of a newly born topic",
                    }
                )
        return sorted(
            changes,
            key=lambda item: (-abs(item["smoothed_log_odds_delta"]), item["term"]),
        )[:5]

    def explain_topics(
        self, occurrences=None, *, date_start=None, date_end=None, options=None
    ):
        """Evidence-grounded topic descriptions and descriptive, source-aware trends.

        A comment is counted once per topic/month. Calendar-edge months and low
        exposure are excluded from trend fitting; source sampling remains partial.
        """
        from babel.dates import format_date
        from babel.numbers import format_decimal, format_percent

        opts = {**self.config["explanation"], **(options or {})}
        self._validate_explanation_options(opts)
        if not self.topics or not {"sentence_id", "cluster_id"} <= set(self.sentences):
            raise ValueError("Describe fitted topics before explaining them")
        frame = (self.occurrences if occurrences is None else occurrences).copy()
        required = {
            "source_row",
            "sentence_id",
            "sentence",
            "cluster_id",
            "platform",
            "post_title",
            "url",
            "post_date",
            "content_date",
            "post_date_accuracy",
            "content_date_accuracy",
        }
        if not required <= set(frame):
            raise ValueError("Explanation requires sentence occurrence provenance")
        if not self.sentences.sentence_id.is_unique:
            raise ValueError("Unique sentence IDs required for explanations")
        vectors = getattr(self, "embeddings", None)
        if vectors is not None and (
            np.asarray(vectors).ndim != 2
            or len(vectors) != len(self.sentences)
            or not np.isfinite(vectors).all()
        ):
            raise ValueError("Invalid explanation embeddings")
        known = self.sentences.set_index("sentence_id").cluster_id
        if (
            not frame.sentence_id.isin(known.index).all()
            or not (frame.sentence_id.map(known) == frame.cluster_id).all()
        ):
            raise ValueError("Occurrence/topic membership mismatch")
        if not set(frame.platform) <= {"youtube", "reddit"}:
            raise ValueError("Unsupported explanation platform")
        frame["_comment"] = frame.platform + ":" + frame.source_row.astype(str)
        root_keys = [
            self.parent_post_key(row) for row in frame.to_dict(orient="records")
        ]
        frame["_root"] = [item[0] for item in root_keys]
        frame["_root_method"] = [item[1] for item in root_keys]
        basis = opts["date_basis"]
        accuracy = basis + "_accuracy"
        frame["_date"] = pd.to_datetime(frame[basis], errors="coerce", format="mixed")
        if (
            frame.groupby("_comment")._date.nunique(dropna=False).gt(1).any()
            or frame.groupby("_comment")._root.nunique().gt(1).any()
        ):
            raise ValueError("Conflicting dates/parent posts for one comment")
        comments = frame.drop_duplicates("_comment")
        undated = int(comments._date.isna().sum())
        estimated = int((comments[accuracy] == "perkiraan").sum())
        dated = frame[frame._date.notna()].copy()
        if not opts["include_estimated_dates"]:
            dated = dated[dated[accuracy] != "perkiraan"]
        start = (
            pd.Timestamp(date_start).normalize()
            if date_start is not None
            else (dated._date.min().normalize() if len(dated) else None)
        )
        end = (
            pd.Timestamp(date_end).normalize()
            if date_end is not None
            else (dated._date.max().normalize() if len(dated) else None)
        )
        if (
            (start is not None and pd.isna(start))
            or (end is not None and pd.isna(end))
            or (start is not None and end is not None and start > end)
        ):
            raise ValueError("Invalid explanation date window")
        if start is not None:
            dated = dated[dated._date >= start]
        if end is not None:
            dated = dated[dated._date < end + pd.Timedelta(days=1)]
        dated["_month"] = dated._date.dt.to_period("M").astype(str)
        cohort = dated.drop_duplicates("_comment")
        if (
            start is not None
            and end is not None
            and end.to_period("M").ordinal - start.to_period("M").ordinal >= 120
        ):
            raise ValueError("Explanation is limited to 120 calendar months")
        periods = (
            list(pd.period_range(start, end, freq="M"))
            if start is not None and end is not None
            else []
        )
        if len(periods) > 120:
            raise ValueError("Explanation is limited to 120 calendar months")
        full = {
            str(period): bool(
                period.start_time.normalize() >= start
                and period.end_time.normalize() <= end
            )
            for period in periods
        }
        # Descriptions follow the caller's filtered rows; the date window and
        # calendar completeness further restrict trend fitting, not vocabulary.
        explanations = []
        for topic in sorted(self.topics, key=lambda item: item["cluster_id"]):
            cid = int(topic["cluster_id"])
            if cid < 0:
                raise ValueError("Noise cannot be explained as one coherent topic")
            source = frame[frame.cluster_id == cid].copy()
            if source.empty:
                continue
            temporal = dated[dated.cluster_id == cid]
            topic_comments = temporal.drop_duplicates("_comment")
            series = []
            for period in periods:
                key = str(period)
                all_rows = cohort[cohort._month == key]
                hits = topic_comments[topic_comments._month == key]
                count, total = len(hits), len(all_rows)
                platform_rows = {}
                for platform in ["youtube", "reddit"]:
                    n = int((all_rows.platform == platform).sum())
                    k = int((hits.platform == platform).sum())
                    platform_rows[platform] = {
                        "all_comments": n,
                        "topic_comments": k,
                        "share": k / n if n else None,
                    }
                series.append(
                    {
                        "period": key,
                        "full_calendar_month": full[key],
                        "all_comments": total,
                        "topic_comments": count,
                        "share": count / total if total else None,
                        "eligible_for_trend": full[key]
                        and total >= opts["min_monthly_comments"],
                        "platforms": platform_rows,
                    }
                )
            trend = self._explanation_trend(series, cohort, topic_comments, opts)
            if undated:
                trend["cautions"].append(
                    "Komentar tanpa tanggal yang dapat dibaca tidak masuk tren waktu."
                )
            if estimated and opts["include_estimated_dates"]:
                trend["cautions"].append(
                    "Sebagian tanggal yang dipakai adalah perkiraan; batas bulan dapat tidak pasti."
                )
            fallback = int(
                (
                    source.drop_duplicates("_comment")._root_method
                    == "context_fallback"
                ).sum()
            )
            if fallback:
                trend["cautions"].append(
                    "Sebagian identitas post induk memakai konteks judul/tanggal, bukan ID URL native."
                )
            if trend["cautions"]:
                trend["support"] = "terbatas"
            terms = [str(term) for term in topic.get("terms", [])[:12]]
            focus = []
            for term in terms:
                if not any(
                    set(term.split()) <= set(other.split())
                    or set(other.split()) <= set(term.split())
                    for other in focus
                ):
                    focus.append(term)
                if len(focus) == 3:
                    break
            evidence = self._explanation_evidence(source, terms, opts)
            description = (
                "Istilah pembeda pada kelompok ini: "
                + (", ".join(focus) or topic["label"])
                + "."
            )
            if evidence:
                quote = evidence[0]["sentence"]
                excerpt = (
                    quote if len(quote) <= 260 else quote[:260].rsplit(" ", 1)[0] + "…"
                )
                description += " Cuplikan penjelas dari percakapan: “" + excerpt + "”"
            else:
                description += " Tidak ada kalimat untuk topik ini pada filter aktif."
            comparison = trend["comparison"]
            narrative = trend["label"] + ". "
            if comparison and comparison["delta_pp"] is not None:
                previous, recent = comparison["previous"], comparison["recent"]

                def month_label(values):
                    return (
                        format_date(
                            pd.Period(values[0], freq="M").start_time,
                            "MMM y",
                            locale="id_ID",
                        )
                        + "–"
                        + format_date(
                            pd.Period(values[-1], freq="M").start_time,
                            "MMM y",
                            locale="id_ID",
                        )
                    )

                narrative += f"Porsi komentar bertopik pada {month_label(previous['periods'])} adalah {format_percent(previous['share'], format='0.0%', locale='id_ID')} ({previous['topic_comments']}/{previous['all_comments']}); pada {month_label(recent['periods'])} menjadi {format_percent(recent['share'], format='0.0%', locale='id_ID')} ({recent['topic_comments']}/{recent['all_comments']}). Perubahannya {format_decimal(comparison['delta_pp'], format='+0.0;-0.0', locale='id_ID')} poin persentase. "
            narrative += "Ini deskripsi komentar yang terkumpul pada filter, bukan kesimpulan tren seluruh platform."
            comments_all_dates = source.drop_duplicates("_comment")
            explanations.append(
                {
                    "cluster_id": cid,
                    "label": topic["label"],
                    "description": description,
                    "focus_terms": focus,
                    "terms_scope": "full-model c-TF-IDF vocabulary",
                    "occurrences": len(source),
                    "unique_sentences": int(source.sentence_id.nunique()),
                    "comments": len(comments_all_dates),
                    "dated_comments": len(topic_comments),
                    "platform_comments": {
                        p: int((comments_all_dates.platform == p).sum())
                        for p in ["youtube", "reddit"]
                    },
                    "trend": trend,
                    "trend_explanation": narrative,
                    "monthly": series,
                    "term_changes": self._explanation_term_changes(
                        temporal, terms, comparison, opts
                    ),
                    "evidence": evidence,
                    "description_method": "c-TF-IDF terms + extractive MMR; no abstractive claims",
                    "inference": "descriptive_cluster_explanation",
                }
            )
        result = {
            "schema_version": EXPLANATION_SCHEMA,
            "producer_version": VERSION,
            "algorithm_source_sha256": file_hash(Path(__file__)),
            "synthetic": bool(
                getattr(
                    self,
                    "explanation_synthetic",
                    self.encoder is not None
                    or self.model_info.get("test_encoder", False),
                )
            ),
            "source_input_sha256": getattr(self, "input_sha256", None),
            "source_run_manifest_sha256": getattr(
                self, "source_run_manifest_sha256", None
            ),
            "options": opts,
            "period": {
                "start": str(start.date()) if start is not None else None,
                "end": str(end.date()) if end is not None else None,
                "date_basis": basis,
                "timezone": "Asia/Jakarta",
            },
            "scope": {
                "sentence_occurrences": len(frame),
                "represented_comments": len(comments),
                "dated_comments": len(cohort),
                "undated_comments": undated,
                "estimated_dates": estimated,
                "topics": len(explanations),
            },
            "denominator": "distinct represented comment/reply source rows per month, including comments whose sentences are noise; topic memberships may overlap",
            "trend_method": "Theil-Sen median pairwise slope + Kendall tau + last complete window comparison; heuristic descriptive labels",
            "partial_months_excluded": [
                key for key, complete in full.items() if not complete
            ],
            "population_inference": False,
            "topics": explanations,
        }
        self.explanations = result
        return result

    @classmethod
    def explain_artifacts(
        cls, data, occurrences=None, *, date_start=None, date_end=None, options=None
    ):
        """Explain already verified artifacts without encoding or fitting again."""
        # Explanation depends on saved data/options, not the historical fitting
        # backend. Legacy ONNX/UMAP/HDBSCAN runs remain readable without retraining.
        stored = data["summary"].get("config", {})
        instance = cls({"explanation": stored.get("explanation", {})})
        instance.topics = data["topics"]
        instance.sentences = data["sentences"]
        instance.occurrences = data["occurrences"]
        instance.embeddings = data.get("embeddings")
        instance.input_sha256 = data["summary"].get("source_sha256")
        instance.explanation_synthetic = bool(
            data["summary"].get("synthetic", False)
            or data["summary"].get("model", {}).get("test_encoder", False)
        )
        instance.source_run_manifest_sha256 = data.get("artifact_manifest_sha256")
        return instance.explain_topics(
            occurrences, date_start=date_start, date_end=date_end, options=options
        )

    @staticmethod
    def save_explanations(result, output_dir, *, source_run=None):
        """Publish a new versioned sidecar; never modify the original run."""
        if result.get("schema_version") != EXPLANATION_SCHEMA:
            raise ValueError("Unsupported explanation schema")
        output = Path(output_dir)
        output = (output if output.is_absolute() else ROOT / output).resolve()
        if output.exists():
            raise FileExistsError("Explanation output exists; use a new directory")
        if source_run is not None and output.is_relative_to(Path(source_run).resolve()):
            raise ValueError("Keep explanation output outside the immutable source run")
        if source_run is not None and result.get(
            "source_run_manifest_sha256"
        ) != file_hash(Path(source_run) / "manifest.json"):
            raise ValueError("Explanation is bound to a different source run")
        output.parent.mkdir(parents=True, exist_ok=True)
        stage = Path(tempfile.mkdtemp(prefix=".topic-explain-", dir=output.parent))
        try:
            json_write(stage / "topic_explanations.json", result)
            json_write(
                stage / "summary.json",
                {
                    "schema_version": EXPLANATION_SCHEMA,
                    "producer_version": VERSION,
                    "synthetic": result.get("synthetic", False),
                    "created_at": datetime.now(timezone.utc).isoformat(),
                    "source_run": str(Path(source_run).resolve())
                    if source_run
                    else None,
                    "source_run_manifest_sha256": result.get(
                        "source_run_manifest_sha256"
                    ),
                    "scope": result["scope"],
                    "directions": dict(
                        Counter(t["trend"]["direction"] for t in result["topics"])
                    ),
                    "partial_months_excluded": result["partial_months_excluded"],
                    "options": result["options"],
                },
            )
            json_write(
                stage / "manifest.json",
                {
                    "schema_version": EXPLANATION_SCHEMA,
                    "files_sha256": {
                        p.name: file_hash(p) for p in stage.iterdir() if p.is_file()
                    },
                },
            )
            stage.chmod(0o755)
            stage.rename(output)
        except BaseException:
            shutil.rmtree(stage, ignore_errors=True)
            raise
        return output

    @staticmethod
    def load_explanations(directory, *, expected_source_manifest_sha256=None):
        directory = Path(directory).resolve()
        manifest = json.loads((directory / "manifest.json").read_text())
        if manifest.get("schema_version") != EXPLANATION_SCHEMA or set(
            manifest.get("files_sha256", {})
        ) != {"topic_explanations.json", "summary.json"}:
            raise ValueError("Unsupported explanation manifest")
        for name, digest in manifest["files_sha256"].items():
            if Path(name).name != name or file_hash(directory / name) != digest:
                raise ValueError("Explanation checksum mismatch: " + name)
        result = json.loads((directory / "topic_explanations.json").read_text())
        summary = json.loads((directory / "summary.json").read_text())
        if result.get("schema_version") != EXPLANATION_SCHEMA or result.get(
            "source_run_manifest_sha256"
        ) != summary.get("source_run_manifest_sha256"):
            raise ValueError("Inconsistent explanation source identity")
        if (
            expected_source_manifest_sha256 is not None
            and result.get("source_run_manifest_sha256")
            != expected_source_manifest_sha256
        ):
            raise ValueError("Explanation belongs to a different source run")
        return result

    def fit_excel(self, path=None, *, embeddings_from=None):
        started = time.monotonic()
        self.prepare_excel(path).encode(
            reuse_run=embeddings_from
        ).compare().describe_topics()
        self.fit_seconds = time.monotonic() - started
        return self

    def save(self, output_dir=None):
        if not self.topics or "cluster_id" not in self.sentences:
            raise ValueError("Fit and describe clusters before saving")
        if not self.coherence["target_met"]:
            raise ValueError("Cannot publish topics below the coherence threshold")
        target = self.config["target_topics"]
        if target is not None and len(self.topics) != target:
            raise ValueError("Cannot publish a run that differs from target_topics")
        baseline = getattr(self, "coherence_baseline", None)
        if baseline and baseline["source_sha256"] != self.input_sha256:
            raise ValueError("Coherence baseline must use the same source workbook")
        if baseline and any(
            baseline[key] != self.coherence[key]
            for key in ("metric", "topn", "window_size", "reference_sha256")
        ):
            raise ValueError("Coherence baseline protocol or reference corpus differs")
        output = Path(output_dir or self.config["output_dir"])
        if not output.is_absolute():
            output = ROOT / output
        if output.exists():
            raise FileExistsError("Use a new NLP output directory")
        output.parent.mkdir(parents=True, exist_ok=True)
        stage = Path(tempfile.mkdtemp(prefix=".nlp-run-", dir=output.parent))
        try:
            self.source_data.to_parquet(stage / "social_data.parquet", index=False)
            self.sentences.to_parquet(stage / "sentences.parquet", index=False)
            self.occurrences.to_parquet(stage / "occurrences.parquet", index=False)
            self.skipped.to_parquet(stage / "skipped.parquet", index=False)
            np.save(stage / "embeddings.npy", self.embeddings, allow_pickle=False)
            np.savez_compressed(
                stage / "umap_embeddings.npz",
                seed_primary=self.reduced_embeddings[0],
                seed_alternate=self.reduced_embeddings[1],
            )
            np.save(stage / "centroids.npy", self.centroids, allow_pickle=False)
            json_write(stage / "topics.json", self.topics)
            json_write(stage / "graph.json", self.graph)
            json_write(stage / "comparison.json", self.comparisons)
            versions = {}
            for name in [
                "numpy",
                "pandas",
                "scikit-learn",
                "sentence-transformers",
                "torch",
                "transformers",
                "tokenizers",
                "gensim",
            ]:
                try:
                    versions[name] = importlib.metadata.version(name)
                except importlib.metadata.PackageNotFoundError:
                    versions[name] = None
            summary = {
                "schema_version": SCHEMA_VERSION,
                "producer_version": VERSION,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "source_file": str(self.input_path),
                "source_sha256": self.input_sha256,
                "preparation": self.preparation,
                "selected": self.selected,
                "model": self.model_info,
                "fit_seconds": getattr(
                    self, "fit_seconds", time.monotonic() - self._started
                ),
                "versions": versions,
                "config": self.config,
                "clusters": len(self.topics),
                "unclustered_unique": int((self.sentences["cluster_id"] < 0).sum()),
                "unclustered_occurrences": int(
                    (self.occurrences["cluster_id"] < 0).sum()
                ),
                "selection_rule": "0.60*max(cosine_silhouette,0)*coverage + 0.25*max(stability_ARI,0) + 0.15*coverage",
                "selection_policy": self.selection_policy,
                "coherence": self.coherence,
                "coherence_baseline": baseline,
                "evaluation_scope": "Internal comparison on this corpus; no labeled semantic ground truth or universal-best claim",
                "stability_scope": "Final topic partitions on two seeded UMAP projections, including configured centroid merging; ARI uses union of assigned points, noise singletons and excludes shared noise. This is not labeled accuracy.",
                "pipeline": "SBERT multilingual -> UMAP -> Euclidean DBSCAN -> optional SBERT centroid topic merging -> c-TF-IDF topic descriptions",
                "reduction": self.reduction_info,
                "topic_reduction": getattr(self, "topic_reduction", None),
                "dates": "Original WIB values/accuracy preserved as source metadata, not clustering features",
                "synthetic": self.model_info.get("test_encoder", False),
            }
            json_write(stage / "summary.json", summary)
            json_write(
                stage / "manifest.json",
                {
                    "schema_version": SCHEMA_VERSION,
                    "files_sha256": {
                        p.name: file_hash(p) for p in stage.iterdir() if p.is_file()
                    },
                },
            )
            stage.chmod(0o755)
            stage.rename(output)
        except BaseException:
            shutil.rmtree(stage, ignore_errors=True)
            raise
        self._event("saved", output_dir=str(output), clusters=len(self.topics))
        return output

    @staticmethod
    def load_run(directory, verify=True):
        directory = Path(directory).resolve()
        manifest_bytes = (directory / "manifest.json").read_bytes()
        manifest = json.loads(manifest_bytes)
        if manifest.get("schema_version") not in {
            SCHEMA_VERSION,
            LEGACY_SCHEMA_VERSION,
        }:
            raise ValueError("Unsupported NLP artifact schema")
        required = {
            "summary.json",
            "topics.json",
            "graph.json",
            "comparison.json",
            "sentences.parquet",
            "occurrences.parquet",
            "skipped.parquet",
            "embeddings.npy",
            "centroids.npy",
        }
        if manifest["schema_version"] == SCHEMA_VERSION:
            required.update({"social_data.parquet", "umap_embeddings.npz"})
        if set(manifest["files_sha256"]) != required:
            raise ValueError("Artifact manifest must cover all required files")
        if verify:
            for name, digest in manifest["files_sha256"].items():
                if Path(name).name != name or file_hash(directory / name) != digest:
                    raise ValueError("NLP artifact checksum mismatch: " + name)
        result = {
            name: json.loads((directory / (name + ".json")).read_text())
            for name in ["summary", "topics", "graph", "comparison"]
        }
        for name in ["sentences", "occurrences", "skipped"]:
            result[name] = pd.read_parquet(directory / (name + ".parquet"))
        if manifest["schema_version"] == SCHEMA_VERSION:
            result["social_data"] = pd.read_parquet(directory / "social_data.parquet")
            with np.load(
                directory / "umap_embeddings.npz", allow_pickle=False
            ) as projections:
                if set(projections.files) != {"seed_primary", "seed_alternate"}:
                    raise ValueError("Unexpected UMAP projection archive")
                result["umap_embeddings"] = [
                    projections["seed_primary"],
                    projections["seed_alternate"],
                ]
            if any(
                len(x) != len(result["sentences"])
                or x.ndim != 2
                or not np.isfinite(x).all()
                for x in result["umap_embeddings"]
            ):
                raise ValueError("UMAP projection shape or values mismatch")
            if (
                len(result["social_data"])
                != result["summary"]["preparation"]["input_rows"]
            ):
                raise ValueError("Source Parquet row count mismatch")
        result["embeddings"] = np.load(directory / "embeddings.npy", allow_pickle=False)
        result["centroids"] = np.load(directory / "centroids.npy", allow_pickle=False)
        if len(result["sentences"]) != len(result["embeddings"]):
            raise ValueError("Sentence/embedding shape mismatch")
        result["artifact_manifest_sha256"] = sha256(manifest_bytes).hexdigest()
        result["artifact_dir"] = str(directory)
        return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=ROOT / "social_nlp.json")
    parser.add_argument("--input", type=Path)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument(
        "--reuse-embeddings-from",
        type=Path,
        help="Reuse checksum-verified SBERT embeddings for identical input/model; refit UMAP and DBSCAN",
    )
    parser.add_argument(
        "--coherence-baseline",
        type=Path,
        help="Measure a previous fitted run with the same c_v protocol before retraining",
    )
    parser.add_argument(
        "--explain-run",
        type=Path,
        help="Explain a saved run without retraining; requires a new --output-dir",
    )
    parser.add_argument("--trend-start", type=date.fromisoformat)
    parser.add_argument("--trend-end", type=date.fromisoformat)
    args = parser.parse_args()
    try:
        if args.reuse_embeddings_from and (args.explain_run or args.prepare_only):
            raise ValueError("Saved embeddings require a training run")
        pipeline = SocialNLPClusterer(
            args.config,
            progress=lambda event: print(
                json.dumps(event, ensure_ascii=False, allow_nan=False), flush=True
            ),
        )
        if args.explain_run:
            if not args.output_dir or args.prepare_only or args.input:
                raise ValueError(
                    "--explain-run requires a new --output-dir and cannot be combined with --input/--prepare-only"
                )
            output = (
                args.output_dir
                if args.output_dir.is_absolute()
                else ROOT / args.output_dir
            )
            if output.exists():
                raise FileExistsError("Explanation output exists; use a new directory")
            data = SocialNLPClusterer.load_run(args.explain_run)
            result = SocialNLPClusterer.explain_artifacts(
                data,
                date_start=args.trend_start,
                date_end=args.trend_end,
                options=pipeline.config["explanation"],
            )
            SocialNLPClusterer.save_explanations(
                result, output, source_run=args.explain_run
            )
            pipeline._event(
                "explained",
                output_dir=str(output),
                topics=len(result["topics"]),
                directions=dict(
                    Counter(t["trend"]["direction"] for t in result["topics"])
                ),
            )
        elif args.prepare_only:
            pipeline.prepare_excel(args.input)
        else:
            output = Path(args.output_dir or pipeline.config["output_dir"])
            output = output if output.is_absolute() else ROOT / output
            if output.exists():
                raise FileExistsError(
                    "Output exists; use --output-dir with a new directory"
                )
            if args.coherence_baseline:
                pipeline.coherence_baseline = pipeline.evaluate_run_coherence(
                    args.coherence_baseline, pipeline.config["coherence"]
                )
                pipeline._event(
                    "coherence_baseline",
                    **{
                        key: value
                        for key, value in pipeline.coherence_baseline.items()
                        if key != "per_topic"
                    },
                )
            pipeline.fit_excel(
                args.input, embeddings_from=args.reuse_embeddings_from
            ).save(args.output_dir)
        return 0
    except (ValueError, OSError, KeyError, TypeError) as error:
        print(
            json.dumps({"status": "ERROR", "error": str(error)}, ensure_ascii=False),
            flush=True,
        )
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
