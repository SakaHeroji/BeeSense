"""Shared export contract for every source script. Python 3.12+.

This is a tabular acquisition projection, not the complete production PRD.
The projection preserves PRD field names; Supabase assigns revision numbers.
"""
from __future__ import annotations

from collections import deque
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from datetime import datetime, timezone
from hashlib import sha256
import hmac
import json
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata
import uuid

CONTRACT_VERSION = "1.0.0"
PRD_REFERENCE = "2.1"
UTC = timezone.utc
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")

def now_iso():
    return datetime.now(UTC).isoformat()

def dumps(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)

def digest(value):
    return sha256(value.encode("utf-8")).hexdigest()

def utc_datetime(value):
    d = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if d.tzinfo is None:
        raise ValueError("Datetime wajib memiliki timezone.")
    return d.astimezone(UTC)

# Each field has one authoritative type and nullability, across every producer.
def fields(spec):
    result = []
    for token in spec.split():
        name, kind = token.split(":")
        result.append({"name":name,"type":kind.rstrip("?"),"nullable":kind.endswith("?")})
    return result + [{"name":"excel_notes","type":"json","nullable":False}]

TABLES = {
    "Contents": fields("contract_version:str project_id:str run_id:str observation_id:str platform:str source_namespace:str object_key:str native_id:str identity_quality:str object_type:str canonical_url:str page_url:str discovered_from:str? title:str? text_source:str text_normalized:str normalizer_version:str text_hash:str text_complete:bool content_revision:int? author_ref:str? publisher_ref:str? published_at:datetime? published_at_display:str? timestamp_quality:str published_interval_start:datetime? published_interval_end:datetime? observed_at:datetime parent_key:str? thread_root_key:str relation_evidence:str? selection_status:str context_only:bool actor_type:str actor_evidence:str? extraction_method:str adapter_version:str quality_flags:json missing_fields:json"),
    "Relations": fields("contract_version:str project_id:str run_id:str relation_id:str subject_key:str predicate:str object_key:str subject_observation_id:str relation_evidence:str resolution_status:str observed_at:datetime"),
    "KeywordHits": fields("contract_version:str project_id:str run_id:str hit_id:str subject_key:str observation_id:str evidence_key:str field_name:str keyword:str matched_variant:str span_start:int span_end:int match_type:str matcher_version:str catalog_version:str"),
    "Metrics": fields("contract_version:str project_id:str run_id:str metric_id:str observation_id:str? platform:str source_namespace:str object_key:str? series_key:str? metric_name:str value:float? value_display:str? unit:str scale_min:float? scale_max:float? period_start:datetime? period_end:datetime? dimensions:json precision:str observed_at:datetime source_url:str quality_flags:json"),
    "Runs": fields("contract_version:str project_id:str run_id:str unit_id:str source:str operation:str target:str status:str reason:str coverage:str contents_written:int comments_seen:int replies_seen:int duplicates_skipped:int invalid_rows:int max_comments:int? max_replies:int? timeout_seconds:int keywords:json source_version:str? warning_codes:json started_at:datetime finished_at:datetime"),
    "Sources": fields("contract_version:str project_id:str run_id:str source:str producer_version:str upstream_version:str? extraction_method:str access_mode:str config:json schema_fingerprint:str observed_at:datetime"),
    "TextChunks": fields("contract_version:str project_id:str run_id:str record_key:str table_name:str field_name:str part_index:int text_part:str encoding:str"),
    "Readme": fields("contract_version:str project_id:str run_id:str item:str value:str"),
}
ENUMS = {
    "Contents.object_type":["post","comment","reply"],
    "Contents.identity_quality":["native_id","canonical_permalink","weak_synthetic"],
    "Contents.timestamp_quality":["exact","relative","date_only","unknown"],
    "Contents.selection_status":["matched","not_matched","context_only","review","excluded"],
    "Contents.actor_type":["unknown","prospective_student","parent","student","alumni","other"],
    "Relations.resolution_status":["resolved_in_export","unresolved_reference"],
    "Metrics.precision":["exact","rounded","unknown"],
}

def schema_document():
    return {"contract_version":CONTRACT_VERSION,"prd_acquisition_reference":PRD_REFERENCE,
            "profile":"tabular_projection_not_full_production_contract",
            "sheet_order":list(TABLES),"tables":TABLES,"enums":ENUMS,
            "time_storage":{"jsonl":"ISO8601 timezone-aware UTC","xlsx":"native Excel datetime UTC without tzinfo"},
            "revision_policy":"content_revision=null until a persistent repository assigns it; observation_id and text_hash bind snapshot evidence"}

def schema_fingerprint():
    return digest(dumps(schema_document()))

def verify_descriptor(path=None):
    path = Path(path) if path else Path(__file__).with_name("contract.json")
    if json.loads(path.read_text(encoding="utf-8")) != schema_document():
        raise ValueError("contract.json tidak sama dengan common module; hentikan ekspor.")

def validate_row(table, row):
    expected = {f["name"] for f in TABLES[table]}
    if set(row) != expected:
        raise ValueError(f"Kolom {table} berbeda: missing={sorted(expected-set(row))}, extra={sorted(set(row)-expected)}")
    for spec in TABLES[table]:
        name, kind = spec["name"], spec["type"]
        v = row[name]
        if v is None:
            if not spec["nullable"]:
                raise ValueError(f"{table}.{name} tidak boleh null")
            continue
        valid = {"str":isinstance(v,str),"int":type(v) is int,
                 "float":type(v) in (int,float) and math.isfinite(v),
                 "bool":type(v) is bool,"json":isinstance(v,(list,dict)),
                 "datetime":isinstance(v,str)}[kind]
        if not valid:
            raise ValueError(f"Tipe {table}.{name} harus {kind}")
        if kind == "datetime":
            if datetime.fromisoformat(v.replace("Z","+00:00")).utcoffset()!=timezone.utc.utcoffset(None):
                raise ValueError(f"{table}.{name} wajib UTC")
            utc_datetime(v)
        if kind == "json":
            dumps(v)
        allowed = ENUMS.get(f"{table}.{name}")
        if allowed and v not in allowed:
            raise ValueError(f"Nilai {table}.{name} tidak dikenal")
    if row["contract_version"] != CONTRACT_VERSION:
        raise ValueError("Versi kontrak berbeda")
    if table == "Contents":
        if row["text_hash"] != digest(row["text_source"]):
            raise ValueError("text_hash tidak cocok")
        if row["timestamp_quality"] != "exact" and row["published_at"] is not None:
            raise ValueError("published_at hanya untuk waktu exact; estimasi tidak boleh dianggap pasti")
        if row["actor_type"]!="unknown" and not row["actor_evidence"]:
            raise ValueError("Peran penulis harus memiliki bukti eksplisit")

class KeywordMatcher:
    """Aho-Corasick; offsets refer to Python Unicode code points in source."""
    VERSION = "aho-casefold-source-offset-1"
    def __init__(self, terms):
        self.terms = tuple(dict.fromkeys(t.strip() for t in terms if t.strip()))
        if len(self.terms)>200 or any(len(t)>200 for t in self.terms):
            raise ValueError("Maksimum 200 keyword, masing-masing 200 karakter")
        self.nodes, self.fail, self.out = [{}], [0], [[]]
        for term in self.terms:
            state=0
            for char in term.casefold():
                if char not in self.nodes[state]:
                    self.nodes[state][char]=len(self.nodes)
                    self.nodes.append({}); self.fail.append(0); self.out.append([])
                state=self.nodes[state][char]
            self.out[state].append(term)
        queue=deque(self.nodes[0].values())
        while queue:
            state=queue.popleft()
            for char,child in self.nodes[state].items():
                queue.append(child); fallback=self.fail[state]
                while fallback and char not in self.nodes[fallback]:
                    fallback=self.fail[fallback]
                self.fail[child]=self.nodes[fallback].get(char,0)
                self.out[child].extend(self.out[self.fail[child]])
        self.catalog_version="literal-"+digest(dumps(self.terms))[:12]

    def find(self, text):
        folded=[]; positions=[]
        for i,char in enumerate(text):
            part=char.casefold(); folded.append(part); positions.extend([i]*len(part))
        value="".join(folded); state=0; seen=set()
        for end,char in enumerate(value):
            while state and char not in self.nodes[state]: state=self.fail[state]
            state=self.nodes[state].get(char,0)
            for keyword in self.out[state]:
                start=end-len(keyword.casefold())+1
                # Do not match a partial casefold expansion, e.g. one half of ß -> ss.
                if (start and positions[start]==positions[start-1]) or (end+1<len(positions) and positions[end]==positions[end+1]): continue
                a,b=positions[start],positions[end]+1
                word=lambda c:c.isalnum() or c=="_"
                if (a and word(text[a-1]) and word(text[a])) or (b<len(text) and word(text[b-1]) and word(text[b])): continue
                key=(keyword,a,b)
                if key not in seen:
                    seen.add(key); yield {"keyword":keyword,"matched_variant":text[a:b],"span_start":a,"span_end":b}

@dataclass(frozen=True, slots=True)
class ContentInput:
    platform: str
    namespace: str
    native_id: str
    object_type: str
    canonical_url: str
    page_url: str
    thread_root_key: str
    text_source: str
    parent_key: str|None=None
    title: str|None=None
    published_at: str|None=None
    published_display: str|None=None
    timestamp_quality: str="unknown"
    text_complete: bool=True
    author_native_id: str|None=None
    publisher_ref: str|None=None
    discovered_from: str|None=None
    observed_at: str=field(default_factory=now_iso)
    quality_flags: tuple[str,...]=()
    extraction_method: str="unofficial_web_client"
    adapter_version: str="unknown"

class RecordFactory:
    def __init__(self, project_id, run_id, matcher, since=None, until=None):
        self.project_id,self.run_id,self.matcher=project_id,run_id,matcher
        self.since=utc_datetime(since) if since else None
        self.until=utc_datetime(until) if until else None
        if self.since and self.until and self.since>=self.until: raise ValueError("since harus lebih awal daripada until")

    @staticmethod
    def key(platform,namespace,native_id):
        return f"{platform}:{namespace}:{native_id}"

    def row(self, table, **values):
        result={f["name"]:None for f in TABLES[table]}
        result.update(contract_version=CONTRACT_VERSION,project_id=self.project_id,run_id=self.run_id,excel_notes={})
        result.update(values)
        for spec in TABLES[table]:
            if spec["type"]=="datetime" and result[spec["name"]] is not None:
                result[spec["name"]]=utc_datetime(result[spec["name"]]).isoformat()
        validate_row(table,result)
        return result

    def content(self, value:ContentInput):
        key=self.key(value.platform,value.namespace,value.native_id)
        obs=str(uuid.uuid5(uuid.UUID(self.run_id),dumps([key,value.text_source,value.title,value.text_complete])))
        hits=[]
        for field_name,text in (("text_source",value.text_source),("title",value.title or "")):
            for hit in self.matcher.find(text):
                hits.append(self.row("KeywordHits",hit_id=digest(dumps([obs,field_name,hit])),subject_key=key,
                    observation_id=obs,evidence_key=key,field_name=field_name,**hit,match_type="direct_literal",
                    matcher_version=self.matcher.VERSION,catalog_version=self.matcher.catalog_version))
        context=value.object_type=="post"
        selection="context_only" if context else "matched" if hits else "not_matched"
        flags=list(value.quality_flags)
        if self.since or self.until:
            if value.published_at:
                dt=utc_datetime(value.published_at)
                if (self.since and dt<self.since) or (self.until and dt>=self.until):
                    selection="context_only" if context else "excluded"; flags.append("outside_requested_window")
            else:
                if not context: selection="review"
                flags.append("window_membership_unknown")
        if not value.text_complete: flags.append("text_incomplete")
        missing=[]
        if not value.text_source: missing.append("text_source")
        if value.timestamp_quality=="unknown": missing.append("published_at")
        author_ref=None
        secret=os.environ.get("AUTHOR_HMAC_KEY")
        if secret and value.author_native_id:
            author_ref=hmac.new(secret.encode(),f"{self.project_id}:{value.platform}:{value.author_native_id}".encode(),"sha256").hexdigest()
        record=self.row("Contents",observation_id=obs,platform=value.platform,source_namespace=value.namespace,
            object_key=key,native_id=value.native_id,identity_quality="native_id",object_type=value.object_type,
            canonical_url=value.canonical_url,page_url=value.page_url,discovered_from=value.discovered_from,
            title=value.title,text_source=value.text_source,text_normalized=unicodedata.normalize("NFC",value.text_source),
            normalizer_version="nfc-1",text_hash=digest(value.text_source),text_complete=value.text_complete,
            content_revision=None,author_ref=author_ref,publisher_ref=value.publisher_ref,published_at=value.published_at,
            published_at_display=value.published_display,timestamp_quality=value.timestamp_quality,
            published_interval_start=None,published_interval_end=None,observed_at=value.observed_at,
            parent_key=value.parent_key,thread_root_key=value.thread_root_key,
            relation_evidence="source_parent_id" if value.parent_key else None,selection_status=selection,
            context_only=context,actor_type="unknown",actor_evidence=None,extraction_method=value.extraction_method,
            adapter_version=value.adapter_version,quality_flags=flags,missing_fields=missing)
        return record,hits

class Dataset:
    """Shared pipeline; adapters cannot choose their own headers or cell formats."""
    def __init__(self, factory, sink):
        self.factory,self.sink=factory,sink
        self.seen=set(); self.object_keys=set(); self.pending_relations=[]
        self.duplicates=0; self.contents=0

    def add_content(self, raw:ContentInput, metrics=()):
        record,hits=self.factory.content(raw)
        identity=record["observation_id"]
        if identity in self.seen:
            self.duplicates+=1; return False
        self.seen.add(identity); self.object_keys.add(record["object_key"]); self.contents+=1
        self.sink.write("Contents",record)
        for h in hits: self.sink.write("KeywordHits",h)
        if record["parent_key"]:
            self.pending_relations.append((record,"reply_to"))
        if record["object_key"]!=record["thread_root_key"]:
            self.pending_relations.append((record,"in_thread"))
        for name,value,unit in metrics:
            if type(value) not in (int,float) or not math.isfinite(value): continue
            self.sink.write("Metrics",self.factory.row("Metrics",metric_id=digest(dumps([identity,name])),
                observation_id=identity,platform=raw.platform,source_namespace=raw.namespace,object_key=record["object_key"],
                series_key=None,metric_name=name,value=value,value_display=None,unit=unit,scale_min=None,scale_max=None,
                period_start=None,period_end=None,dimensions={},precision="unknown",observed_at=raw.observed_at,
                source_url=raw.canonical_url,quality_flags=["library_parsed_count_precision_unknown"]))
        return True

    def finish_relations(self):
        for r,predicate in self.pending_relations:
            target=r["parent_key"] if predicate=="reply_to" else r["thread_root_key"]
            self.sink.write("Relations",self.factory.row("Relations",relation_id=digest(dumps([r["observation_id"],predicate,target])),
                subject_key=r["object_key"],predicate=predicate,object_key=target,subject_observation_id=r["observation_id"],
                relation_evidence="source_parent_id" if predicate=="reply_to" else "source_video_context",
                resolution_status="resolved_in_export" if target in self.object_keys else "unresolved_reference",observed_at=r["observed_at"]))
        self.pending_relations.clear()

class SourceAdapter(ABC):
    @abstractmethod
    def accept(self,kind,raw,url,discovered_from):
        """Normalize a source event; write only through the shared Dataset/Sink."""
        raise NotImplementedError

class BundleSink:
    """Atomic new directory; streaming JSONL + streaming, formula-safe Excel."""
    def __init__(self,path,project_id,run_id):
        verify_descriptor()
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        self.path=Path(path).resolve(); self.project_id,self.run_id=project_id,run_id
        if self.path.exists(): raise FileExistsError("Output directory sudah ada; gunakan nama baru.")
        self.path.parent.mkdir(parents=True,exist_ok=True)
        self.stage=Path(tempfile.mkdtemp(prefix=".source-export-",dir=self.path.parent))
        self.book=Workbook(write_only=True); self.files={}; self.sheets={}; self.counts={t:0 for t in TABLES}
        self.header_font=Font(name="Calibri",size=11,bold=True,color="FFFFFF")
        self.font=Font(name="Calibri",size=11); self.fill=PatternFill("solid",fgColor="163C52")
        self.alignment=Alignment(wrap_text=True,vertical="top"); self.closed=False
        for table,specs in TABLES.items():
            self.files[table]=(self.stage/f"{table}.jsonl").open("w",encoding="utf-8")
            ws=self.book.create_sheet(table); self.sheets[table]=ws
            ws.freeze_panes="A2"; ws.sheet_view.showGridLines=False
            for i,spec in enumerate(specs,1):
                width=62 if spec["name"] in {"text_source","text_normalized","value","reason","config","source_url","canonical_url","page_url","target","quality_flags"} else 25
                ws.column_dimensions[get_column_letter(i)].width=width
            ws.append(self._cells(ws,[s["name"] for s in specs],True))

    def _cells(self,ws,values,header=False):
        from openpyxl.cell import WriteOnlyCell
        result=[]
        for v in values:
            c=WriteOnlyCell(ws,value=v)
            if isinstance(v,str): c.data_type="s"
            c.font=self.header_font if header else self.font; c.alignment=self.alignment
            if header:c.fill=self.fill
            if isinstance(v,datetime):c.number_format="yyyy-mm-dd hh:mm:ss"
            result.append(c)
        return result

    @staticmethod
    def chunks(text):
        current=[]; units=0
        for char in text:
            cost=2 if ord(char)>0xFFFF else 1
            if units+cost>30000:
                yield "".join(current);current=[];units=0
            current.append(char);units+=cost
        yield "".join(current)

    def write(self,table,row):
        validate_row(table,row)
        if self.counts[table]>=1_048_575:raise ValueError(f"Batas baris Excel {table} tercapai")
        notes={}; cells=[]
        key=row.get("observation_id") or row.get("unit_id") or row.get("metric_id") or row.get("hit_id") or f"{table}:{self.counts[table]}"
        for spec in TABLES[table]:
            name=spec["name"];value=row[name]
            if name=="excel_notes":cells.append(dumps(notes));continue
            if spec["type"]=="json":value=dumps(value)
            elif spec["type"]=="datetime" and value:value=utc_datetime(value).replace(tzinfo=None)
            if isinstance(value,str):
                encoding="plain"
                if INVALID_XML.search(value):value=json.dumps(value,ensure_ascii=True);encoding="json_string";notes[name]={"encoding":encoding}
                parts=list(self.chunks(value))
                if len(parts)>1:
                    if table=="TextChunks":raise ValueError("Internal chunking overflow")
                    notes[name]={"encoding":encoding,"parts":len(parts),"record_key":key}
                    for n,part in enumerate(parts,1):
                        self.write("TextChunks",{**{f["name"]:None for f in TABLES["TextChunks"]},"contract_version":CONTRACT_VERSION,
                            "project_id":self.project_id,"run_id":self.run_id,"record_key":key,"table_name":table,
                            "field_name":name,"part_index":n,"text_part":part,"encoding":encoding,"excel_notes":{}})
                value=parts[0]
            cells.append(value)
        stored={**row,"excel_notes":notes}
        self.files[table].write(dumps(stored)+"\n")
        self.sheets[table].append(self._cells(self.sheets[table],cells));self.counts[table]+=1

    def finish(self,metadata):
        from openpyxl.utils import get_column_letter
        for f in self.files.values():f.close()
        for table,ws in self.sheets.items():ws.auto_filter.ref=f"A1:{get_column_letter(len(TABLES[table]))}{self.counts[table]+1}"
        self.book.save(self.stage/"dataset.xlsx")
        (self.stage/"contract.json").write_text(json.dumps(schema_document(),ensure_ascii=False,indent=2),encoding="utf-8")
        def file_hash(path):
            h=sha256()
            with path.open("rb") as stream:
                for chunk in iter(lambda:stream.read(65536),b""):h.update(chunk)
            return h.hexdigest()
        manifest={"contract_version":CONTRACT_VERSION,"prd_reference":PRD_REFERENCE,"schema_fingerprint":schema_fingerprint(),
                  "project_id":self.project_id,"run_id":self.run_id,"created_at":now_iso(),"row_counts":self.counts,
                  "files_sha256":{p.name:file_hash(p) for p in self.stage.iterdir() if p.is_file()},**metadata}
        (self.stage/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
        self.stage.rename(self.path);self.closed=True
        return self.path

    def abort(self):
        for f in self.files.values():
            if not f.closed:f.close()
        shutil.rmtree(self.stage,ignore_errors=True)

def add_readme(sink,factory,extra=()):
    items=[("Kontrak ekspor",CONTRACT_VERSION),("Referensi PRD",PRD_REFERENCE+"; ini proyeksi tabel, bukan implementasi lengkap PRD"),
        ("Header fingerprint",schema_fingerprint()),("Tipe sumber","Semua sumber menggunakan 8 sheet dan tipe kolom identik, termasuk sheet kosong."),
        ("Timestamp","Excel: native UTC tanpa timezone; JSONL: ISO8601 UTC. published_at kosong untuk tanggal relatif/unknown."),
        ("Revision","content_revision null sampai ditetapkan repository persisten. Gunakan observation_id + text_hash untuk evidence snapshot."),
        ("Keyword","Literal casefold + boundary; offset Unicode code point pada field sumber. Tidak menentukan sentimen atau peran penulis."),
        ("Peran penulis","unknown; tidak menebak orangtua/calon mahasiswa dari username atau foto."),
        ("Metrik","Metrik jumlah dan indeks tren berbagi tabel Metrics tetapi dibedakan metric_name, unit, scale dan dimensions; jangan dijumlahkan lintas unit."),
        ("Teks panjang","JSONL menjaga teks penuh. Excel menyimpan sel pertama + seluruh potongan di TextChunks; gabungkan part_index, lalu decode bila encoding=json_string."),
        ("Kelengkapan","LIMIT_REACHED/PARTIAL/UNKNOWN tidak menjamin seluruh komentar. Nol karena gagal berbeda dengan nol komentar terkonfirmasi."),
        ("Data lama","Workbook social-comments-pilot-1.0 terdahulu tidak otomatis kompatibel. Adapter sumber berikutnya wajib memakai common ini; migrasi arsip perlu mapping eksplisit."),
        ("Supabase","Tidak ada koneksi atau transfer DB dalam versi ini. Data kontrak dapat menjadi input staging; adapter repository dan penomoran revisi belum diimplementasikan."),
        ("Sentimen/emosi","Belum dianalisis; tidak ada label otomatis dari crawler."),*extra]
    for item,value in items:sink.write("Readme",factory.row("Readme",item=item,value=str(value)))

def import_pilot(input_path,output_path,project_id,keywords=()):
    """Explicit conversion of the earlier Comments/Posts/Runs pilot workbook.

    This is an import utility, not another source crawler. No network requests.
    It retains historical method/status labels and recomputes literal matches.
    """
    from openpyxl import load_workbook
    path=Path(input_path)
    run_id=str(uuid.uuid4());factory=RecordFactory(project_id,run_id,KeywordMatcher(keywords))
    sink=BundleSink(output_path,project_id,run_id);dataset=Dataset(factory,sink)
    workbook=load_workbook(path,read_only=True,data_only=True)
    namespaces={"youtube":("youtube_video","youtube_comment"),"instagram":("instagram_post_shortcode","instagram_comment"),
                "tiktok":("tiktok_video","tiktok_comment"),"facebook":("facebook_post","facebook_comment")}
    invalid=0;counts={"comment":0,"reply":0};platforms=set();started=now_iso()
    def source_rows(sheet):
        if sheet not in workbook.sheetnames:return
        iterator=workbook[sheet].values;header=next(iterator,None)
        if not header:return
        if len(header)!=len(set(header)):raise ValueError("Kolom legacy duplikat")
        for values in iterator:yield dict(zip(header,values))
    def stamp(value):
        if isinstance(value,datetime):return value.replace(tzinfo=UTC).isoformat() if value.tzinfo is None else value.astimezone(UTC).isoformat()
        if isinstance(value,str):
            try:return utc_datetime(value).isoformat()
            except ValueError:return None
        return None
    try:
        if not {"Comments","Posts","Runs"}.issubset(workbook.sheetnames):raise ValueError("Input bukan workbook pilot Comments/Posts/Runs")
        for sheet in ("Posts","Comments"):
            for r in source_rows(sheet):
                platform=r.get("platform");post_id=r.get("post_id");cid=r.get("comment_id")
                if platform not in namespaces or not isinstance(post_id,str) or not post_id or (sheet=="Comments" and (not isinstance(cid,str) or not cid or cid.startswith("weak:"))):
                    invalid+=1;continue
                source_url=r.get("source_url")
                if not isinstance(source_url,str) or not source_url.startswith("https://"):invalid+=1;continue
                post_ns,comment_ns=namespaces[platform];root=factory.key(platform,post_ns,post_id)
                is_post=sheet=="Posts";parent=r.get("parent_comment_id")
                kind="post" if is_post else "reply" if parent else "comment"
                if r.get("object_type")=="reply" and not parent:kind="reply"
                parent_key=(None if is_post or (kind=="reply" and not parent) else factory.key(platform,comment_ns,str(parent)) if parent else root)
                quality=r.get("timestamp_quality");published=stamp(r.get("published_at_utc"))
                quality="exact" if quality=="source_timestamp" and published else "unknown"
                if quality!="exact":published=None
                flags=["legacy_import","selection_recomputed_on_import"]
                if r.get("quality_flags"):
                    try:
                        old=json.loads(r["quality_flags"])
                        if isinstance(old,list):flags.extend(str(x) for x in old)
                    except (ValueError,TypeError):flags.append("legacy_quality_flags_unparsed")
                if r.get("author_ref"):flags.append("old_project_hmac_not_reused")
                observed=stamp(r.get("collected_at_utc"))
                if not observed:observed=now_iso();flags.append("original_observation_time_missing")
                method=r.get("capture_method") or "legacy_method_unknown"
                raw=ContentInput(platform=platform,namespace=post_ns if is_post else comment_ns,
                    native_id=post_id if is_post else cid,object_type=kind,canonical_url=r.get("source_comment_url") or source_url,
                    page_url=source_url,thread_root_key=root,text_source=r.get("text_source") or "",parent_key=parent_key,
                    published_at=published,published_display=str(r["published_at_display"]) if r.get("published_at_display") else None,
                    timestamp_quality=quality,text_complete=r.get("source_text_complete") is not False,observed_at=observed,
                    quality_flags=tuple(dict.fromkeys(flags)),extraction_method="legacy_import:"+method,adapter_version="pilot-import:1.0.0")
                metric=("comment_count",r.get("displayed_comment_count"),"count") if is_post else ("reaction_count",r.get("reaction_count"),"count")
                dataset.add_content(raw,[metric]);platforms.add(platform)
                if not is_post:counts[kind]+=1
        dataset.finish_relations()
        for n,r in enumerate(source_rows("Runs")):
            sink.write("Runs",factory.row("Runs",unit_id=str(uuid.uuid5(uuid.UUID(run_id),f"legacy-receipt-{n}")),source=r.get("platform") or "legacy",
                operation="imported_historical_receipt",target=r.get("source_url") or "unknown",status=r.get("status") or "UNKNOWN",
                reason="Historical receipt; no new network request. "+str(r.get("reason") or ""),coverage=r.get("coverage") or "UNKNOWN",
                contents_written=0,comments_seen=int(r.get("comments_exported") or 0),replies_seen=int(r.get("replies_exported") or 0),
                duplicates_skipped=0,invalid_rows=0,max_comments=r.get("max_comments"),max_replies=r.get("max_replies"),timeout_seconds=0,
                keywords=[],source_version=r.get("adapter_version"),warning_codes=["historical_receipt_only"],
                started_at=stamp(r.get("started_at_utc")) or started,finished_at=stamp(r.get("finished_at_utc")) or started))
        sink.write("Runs",factory.row("Runs",unit_id=str(uuid.uuid4()),source="legacy_import",operation="import-pilot",target=path.name,
            status="PARTIAL_IMPORT" if invalid else "IMPORTED",reason="explicit_pilot_mapping_no_live_recollection",coverage="INPUT_FILE_ONLY",
            contents_written=dataset.contents,comments_seen=counts["comment"],replies_seen=counts["reply"],duplicates_skipped=dataset.duplicates,
            invalid_rows=invalid,max_comments=None,max_replies=None,timeout_seconds=0,keywords=list(keywords),source_version="pilot-1.0",
            warning_codes=["weak_or_invalid_rows_not_mapped"] if invalid else [],started_at=started,finished_at=now_iso()))
        for platform in sorted(platforms):
            sink.write("Sources",factory.row("Sources",source=platform,producer_version="pilot-import:1.0.0",upstream_version=None,
                extraction_method="legacy_workbook_import",access_mode="local_file_no_network",config={"input_file":path.name},
                schema_fingerprint=schema_fingerprint(),observed_at=now_iso()))
        add_readme(sink,factory,[("Migrasi pilot","Mapping eksplisit Comments/Posts/Runs. Tidak mengubah file asli atau mengklaim data dikumpulkan ulang."),
            ("Tanggal impor","Kolom lama *_utc dianggap UTC sesuai kontrak pilot; quality source_timestamp menjadi exact. Waktu lain tetap unknown."),
            ("Batas mapping","ID lemah/tidak lengkap ditolak dan dihitung invalid_rows. HMAC lama tidak dipakai ulang lintas project. Sentimen tidak ditambahkan.")])
        output=sink.finish({"source":"legacy_import","mode":"local_pilot_import","synthetic":None,"input_file":path.name,"invalid_rows":invalid,
            "limitations":["input_authenticity_not_reverified","persistent_revisions_not_assigned","old_project_author_hmac_not_reused"]})
        return {"output":str(output),"contents":dataset.contents,"comments":counts,"invalid_rows":invalid,"schema_fingerprint":schema_fingerprint()}
    except BaseException:sink.abort();raise
    finally:workbook.close()

if __name__=="__main__":
    import argparse
    parser=argparse.ArgumentParser(description="Common contract utility, not a source crawler.")
    commands=parser.add_subparsers(dest="command",required=True)
    migrate=commands.add_parser("import-pilot",help="Konversi workbook pilot lama ke kontrak bersama")
    migrate.add_argument("--input",required=True,type=Path);migrate.add_argument("--output-dir",required=True,type=Path)
    migrate.add_argument("--project-id",default="university-intelligence");migrate.add_argument("--keyword",action="append",default=[])
    args=parser.parse_args()
    print(json.dumps(import_pilot(args.input,args.output_dir,args.project_id,args.keyword),ensure_ascii=False,indent=2))
