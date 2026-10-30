"""Bounded Instagram acquisition through instagrapi, with the shared Excel contract.

Run as a module from the repository root. Login is an explicit local-terminal
operation; acquisition reuses saved settings and stops on authentication barriers.
"""

from __future__ import annotations

import argparse
import importlib.metadata
import json
import logging
import math
import os
import re
import sys
import tempfile
import time
import uuid
from collections import Counter, deque
from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime, timezone
from hashlib import file_digest
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlsplit

from Crawler.Backend.Reddit import crawl_reddit_single as contract
from Crawler.Backend.social_export import (
    COMPACT_COLUMNS,
    COMPACT_FILE,
    SocialExportMixin,
)

ROOT = Path(__file__).resolve().parents[2]
VERSION = "1.0.0"
METHOD = "instagrapi_authenticated_private_api"
STOP_STATUSES = {
    "LOGIN_REQUIRED",
    "CHALLENGE_REQUIRED",
    "RATE_LIMITED",
    "ACCESS_BLOCKED",
    "DEPENDENCY_MISSING",
    "INTERRUPTED",
    "SESSION_INVALID",
    "CLIENT_UPGRADE_REQUIRED",
}


def upstream_version():
    try:
        return importlib.metadata.version("instagrapi")
    except importlib.metadata.PackageNotFoundError:
        return None


def media_url(value):
    """Accept only Instagram post/reel permalinks; strip tracking parameters."""
    if not isinstance(value, str) or any(ord(c) < 33 for c in value):
        raise ValueError("Expected an Instagram media URL")
    parsed = urlsplit(value)
    if (
        parsed.scheme != "https"
        or parsed.netloc not in {"instagram.com", "www.instagram.com"}
        or not re.fullmatch(r"/(?:p|reel|tv)/[A-Za-z0-9_-]{1,40}/?", parsed.path)
    ):
        raise ValueError("Expected https://www.instagram.com/p/<shortcode>/")
    return "https://www.instagram.com/p/" + parsed.path.strip("/").split("/")[1] + "/"


def native_pk(value):
    value = str(value)
    if not re.fullmatch(r"[1-9][0-9]{0,29}", value):
        raise ValueError("Invalid Instagram native ID")
    return value


def media_pk_from_url(url):
    code = media_url(url).strip("/").split("/")[-1]
    alphabet = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789-_"
    number = 0
    for char in code:
        number = number * 64 + alphabet.index(char)
    return native_pk(number)


def object_key(native, *, comment=False):
    return (
        f"instagram:instagram_{'comment' if comment else 'media'}:{native_pk(native)}"
    )


def error_receipt(error):
    """Classify by exception type/status; never write exception text or sessions."""
    names = {c.__name__ for c in type(error).__mro__}
    response = getattr(error, "response", None)
    status_code = getattr(response, "status_code", None)
    status = "SOURCE_ERROR"
    if isinstance(error, ImportError):
        status = "DEPENDENCY_MISSING"
    elif "ClientNeedsUpgradeError" in names:
        status = "CLIENT_UPGRADE_REQUIRED"
    elif isinstance(error, FileNotFoundError) or names & {
        "LoginRequired",
        "ClientLoginRequired",
        "BadPassword",
        "TwoFactorRequired",
    }:
        status = "LOGIN_REQUIRED"
    elif names & {
        "ChallengeRequired",
        "ChallengeError",
        "ChallengeSelfieCaptcha",
        "ChallengeRedirection",
        "ChallengeUnknownStep",
    }:
        status = "CHALLENGE_REQUIRED"
    elif status_code == 429 or names & {
        "ClientThrottledError",
        "PleaseWaitFewMinutes",
        "RateLimitError",
    }:
        status = "RATE_LIMITED"
    elif status_code == 403 or names & {
        "FeedbackRequired",
        "ClientForbiddenError",
        "SentryBlock",
    }:
        status = "ACCESS_BLOCKED"
    elif isinstance(error, SessionInvalid):
        status = "SESSION_INVALID"
    elif isinstance(error, KeyboardInterrupt):
        status = "INTERRUPTED"
    headers = getattr(response, "headers", {}) or {}
    retry = headers.get("Retry-After")
    retry = int(retry) if isinstance(retry, str) and retry.isdigit() else None
    return dict(
        status=status,
        reason="instagram_" + status.lower(),
        warnings=[type(error).__name__],
        source_version=upstream_version(),
        retry_after_seconds=retry,
    )


class SessionInvalid(ValueError):
    """The supplied settings do not contain an authenticated mobile session."""


def _raise_exception(client, error):
    raise error


def _manual_verification_required(*args, **kwargs):
    # Do not invoke instagrapi's automatic email/SMS/password challenge resolver.
    from instagrapi.exceptions import ChallengeRequired

    raise ChallengeRequired("Complete verification manually")


def create_client():
    from instagrapi import Client

    client = Client(session_retry_total=0, public_request_retries_count=0)
    client.handle_exception = _raise_exception
    client.challenge_code_handler = _manual_verification_required
    client.change_password_handler = _manual_verification_required
    # Upstream request/error logs can contain identifiers, response bodies or codes.
    quiet = logging.Logger("beesense.instagram.silent")
    quiet.addHandler(logging.NullHandler())
    quiet.propagate = False
    client.logger = client.request_logger = client.private_request_logger = quiet
    return client


def load_client(path, request_delay_seconds=2, *, require_authenticated=True):
    path = Path(path)
    if path.stat().st_size > 1_000_000:
        raise SessionInvalid("Session settings exceed 1 MB")
    try:
        settings = json.loads(path.read_text(encoding="utf-8"))
    except (ValueError, UnicodeError) as error:
        raise SessionInvalid("Invalid session settings") from error
    if not isinstance(settings, dict):
        raise SessionInvalid("Session settings must be an object")
    client = create_client()
    # Do not restore stale retry policies from a saved settings file.
    settings.update(session_retry_total=0, public_request_retries_count=0)
    client.set_settings(settings)
    client.request_timeout = request_delay_seconds  # pacing, not HTTP timeout
    if require_authenticated and not client.user_id:
        raise SessionInvalid("No authenticated user in settings")
    return client


def save_session(client, path):
    """Atomic owner-only settings; passwords are never part of this file."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    fd, name = tempfile.mkstemp(prefix=".session-", dir=path.parent)
    try:
        os.fchmod(fd, 0o600)
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(client.get_settings(), stream)
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def login_with_saved_password(client, username, password, code=""):
    """Validate restored state; allow one explicit LoginRequired recovery only."""
    try:
        accepted = client.login(username, password, verification_code=code)
    except Exception as error:
        # Keep this module importable without the optional live dependency.
        if "LoginRequired" not in {c.__name__ for c in type(error).__mro__}:
            raise
        # relogin() does not forward verification_code in instagrapi 3.2.2.
        accepted = (
            client.login(username, password, relogin=True, verification_code=code)
            if code
            else client.relogin()
        )
    if not accepted or not client.user_id:
        raise SessionInvalid("Login was not accepted")


def login_local(path):
    """Interactive credentials stay in the user's terminal, never CLI arguments."""
    from getpass import getpass

    if not sys.stdin.isatty():
        raise SessionInvalid("Run --login in your own interactive terminal")
    client = (
        load_client(path, require_authenticated=False)
        if Path(path).is_file()
        else create_client()
    )
    username = input("Instagram username: ").strip()
    password = getpass("Instagram password (hidden): ")
    code = getpass("2FA code if enabled (hidden; otherwise Enter): ").strip()
    try:
        login_with_saved_password(client, username, password, code)
        save_session(client, path)
    finally:
        password = code = ""


def login_from_file(auth_file, session_file, *, login_method="auto"):
    if login_method not in {"auto", "password", "sessionid"}:
        raise ValueError("Unknown Instagram login method")
    path = Path(auth_file)
    if path.stat().st_size > 1_000_000:
        raise SessionInvalid("Auth file exceeds 1 MB")
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (ValueError, UnicodeError) as error:
        raise SessionInvalid("Auth file must contain valid JSON") from error
    entries = data if isinstance(data, list) else [data]
    entries = [
        e
        for e in entries
        if isinstance(e, dict) and str(e.get("sosmed", "")).casefold() == "instagram"
    ]
    if len(entries) != 1:
        raise SessionInvalid("Exactly one Instagram auth entry is required")
    auth = entries[0]
    sessionid = auth.get("sessionid") if login_method != "password" else None
    if login_method == "sessionid" and sessionid is None:
        raise SessionInvalid("sessionid is required for cookie login")
    if sessionid is not None:
        if (
            not isinstance(sessionid, str)
            or not sessionid.strip()
            or len(sessionid) > 4096
            or any(c.isspace() for c in sessionid)
        ):
            raise SessionInvalid("sessionid must be a non-empty local cookie value")
        # Cookies are opaque: pass the exact browser value to the upstream API.
        if any(ord(c) < 33 or ord(c) > 126 for c in sessionid):
            raise SessionInvalid("sessionid must contain printable ASCII only")
        client = (
            load_client(session_file, require_authenticated=False)
            if Path(session_file).is_file()
            else create_client()
        )
        # Retain the device but remove old in-memory authorization. Upstream
        # login_by_sessionid initializes headers before its first profile lookup;
        # a loaded Bearer header could otherwise refer to the previous cookie.
        client.set_settings(
            {**client.get_settings(), "cookies": {}, "authorization_data": {}}
        )
        if not client.login_by_sessionid(sessionid) or not client.user_id:
            raise SessionInvalid("Instagram did not accept the browser session")
        # Profile recovery can fall back to public GraphQL; require a private read
        # before treating a browser cookie as an accepted acquisition session.
        client.account_info()
        save_session(client, session_file)
        return
    if any(
        not isinstance(auth.get(k), str) or not auth[k].strip()
        for k in ("username", "password")
    ):
        raise SessionInvalid("Instagram username/password are required locally")
    code = auth.get("verification_code", "")
    if not isinstance(code, str):
        raise SessionInvalid("verification_code must be a string")
    client = (
        load_client(session_file, require_authenticated=False)
        if Path(session_file).is_file()
        else create_client()
    )
    login_with_saved_password(client, auth["username"], auth["password"], code)
    save_session(client, session_file)


@dataclass(frozen=True)
class InstagramSettings:
    session_file: str
    request_delay_seconds: float = 2
    search_results: int = 50
    max_comments: int = 100
    max_replies: int = 30
    max_replies_per_thread: int = 10
    search_tab: str = "recent"
    timeout_seconds: int = 90


def _value(item, key, default=None):
    return (
        item.get(key, default)
        if isinstance(item, dict)
        else getattr(item, key, default)
    )


def _timestamp(value):
    if isinstance(value, datetime):
        # instagrapi's extractor declares UTC; reject naive timestamps rather than guess.
        if value.tzinfo is not None and value.utcoffset() is not None:
            return value.astimezone(timezone.utc).isoformat()
    elif isinstance(value, str):
        try:
            return contract.utc_datetime(value).isoformat()
        except (ValueError, TypeError):
            pass
    return None


def normalize_item(item, url, *, media=None, parent=None):
    """Map explicit native IDs/timestamps; do not infer authors or text coverage."""
    url = media_url(url)
    is_media = media is None
    media = item if is_media else media
    root_id = native_pk(_value(media, "pk"))
    if root_id != media_pk_from_url(url):
        raise ValueError("Media native ID does not match the shortcode")
    code = _value(media, "code")
    if media_url("https://www.instagram.com/p/" + str(code) + "/") != url:
        raise ValueError("Media shortcode does not match the requested URL")
    native = native_pk(_value(item, "pk"))
    explicit_parent = _value(item, "replied_to_comment_id")
    if not is_media and explicit_parent:
        explicit_parent = native_pk(explicit_parent)
        if parent is not None and explicit_parent != native_pk(parent):
            raise ValueError("Reply parent mismatch")
        parent = explicit_parent
    if parent is not None:
        parent = native_pk(parent)
        if parent == native:
            raise ValueError("Comment cannot reply to itself")
    text = _value(item, "caption_text" if is_media else "text", "")
    if not isinstance(text, str):
        raise ValueError("Source text must be a string")
    stamp = _timestamp(_value(item, "taken_at" if is_media else "created_at_utc"))
    root = object_key(root_id)
    user = _value(item, "user")
    author = _value(user, "pk") if user is not None else None
    metrics = []
    for key in ("like_count", "comment_count") if is_media else ("like_count",):
        number = _value(item, key)
        if type(number) in (int, float) and math.isfinite(number) and number >= 0:
            metrics.append((key, number, "count"))
    return dict(
        native_id=native,
        object_type="post" if is_media else "reply" if parent else "comment",
        canonical_url=url if is_media else url + "?comment_id=" + native,
        page_url=url,
        thread_root_key=root,
        parent_key=None
        if is_media
        else object_key(parent, comment=True)
        if parent
        else root,
        text_source=text,
        title=(
            str(_value(item, "title", "") or "").strip() or text.split("\n")[0][:240]
        )
        if is_media
        else None,
        published_at=stamp,
        timestamp_quality="exact" if stamp else "unknown",
        text_complete=False,
        author_native_id=str(author) if author else None,
        observed_at=contract.now_iso(),
        quality_flags=[
            "library_text_completeness_unverified",
            "instagram_comments_best_effort",
        ],
        metrics=metrics,
    )


class InstagramAdapter:
    def __init__(self, factory, dataset, version=None, *, method=METHOD):
        self.factory, self.dataset = factory, dataset
        self.version = version or "unknown"
        self.method = method

    def accept(self, kind, raw, url, discovered_from=None):
        if kind != "record":
            return
        url = media_url(url)
        comment = raw["object_type"] != "post"
        if media_url(raw["page_url"]) != url or raw["thread_root_key"].split(":")[
            :2
        ] != ["instagram", "instagram_media"]:
            raise ValueError("Instagram identity/context mismatch")
        native = native_pk(raw["native_id"])
        if raw["thread_root_key"] != object_key(media_pk_from_url(url)):
            raise ValueError("Instagram thread ID/shortcode mismatch")
        if not comment and raw["thread_root_key"] != object_key(native):
            raise ValueError("Instagram root identity mismatch")
        expected_url = url + "?comment_id=" + native if comment else url
        if raw["canonical_url"] != expected_url:
            raise ValueError("Instagram canonical identity mismatch")
        if (
            raw["object_type"] == "comment"
            and raw["parent_key"] != raw["thread_root_key"]
        ):
            raise ValueError("Instagram comment parent mismatch")
        if raw["object_type"] == "reply" and (
            not re.fullmatch(
                r"instagram:instagram_comment:[1-9][0-9]{0,29}", raw["parent_key"] or ""
            )
            or raw["parent_key"] == object_key(native, comment=True)
        ):
            raise ValueError("Instagram reply parent mismatch")
        values = {k: v for k, v in raw.items() if k != "metrics"}
        values["quality_flags"] = tuple(values["quality_flags"])
        value = contract.ContentInput(
            platform="instagram",
            namespace="instagram_comment" if comment else "instagram_media",
            discovered_from=discovered_from,
            extraction_method=self.method,
            adapter_version=f"instagram-adapter:{VERSION};{self.method}:{self.version}",
            **values,
        )
        self.dataset.add_content(value, raw.get("metrics", ()))


def collect(operation, target, settings, emit, *, client=None):
    if operation in {"login", "login_password", "login_sessionid"}:
        login_from_file(
            target,
            settings["session_file"],
            login_method=operation.removeprefix("login_")
            if operation != "login"
            else "auto",
        )
        return dict(
            status="SESSION_SAVED",
            reason="local_auth_file_login",
            warnings=[],
            source_version=upstream_version(),
        )
    client = client or load_client(
        settings["session_file"], settings["request_delay_seconds"]
    )
    if operation == "check_session":
        client.account_info()  # authenticated read, no content acquisition
    elif operation == "search":
        query = json.loads(target)
        medias, cursor = client.hashtag_medias_v1_chunk(
            query["hashtag"],
            max_amount=settings["search_results"],
            tab_key=settings["search_tab"],
            max_id=query.get("cursor"),
        )
        for item in medias[: settings["search_results"]]:
            try:
                url = media_url(
                    "https://www.instagram.com/p/" + str(_value(item, "code")) + "/"
                )
                native_pk(_value(item, "pk"))
            except ValueError:
                continue
            emit(
                "candidate",
                dict(
                    url=url,
                    discovered_from="https://www.instagram.com/explore/tags/"
                    + query["hashtag"]
                    + "/",
                ),
            )
        return dict(
            status="FINISHED_BEST_EFFORT",
            reason="hashtag_page",
            next_cursor=cursor or None,
            warnings=[],
            source_version=upstream_version(),
        )
    elif operation == "media":
        url = media_url(target)
        code = url.strip("/").split("/")[-1]
        pk = native_pk(client.media_pk_from_code(code))
        media = client.media_info_v1(pk)
        if native_pk(_value(media, "pk")) != pk:
            raise ValueError("Returned media ID does not match requested shortcode")
        emit("record", normalize_item(media, url))
        media_id = str(_value(media, "id") or pk)
        seen, parents, cursor, cursors = set(), [], None, set()
        # Calls may return more than requested. Export caps apply to unique IDs.
        while len(parents) < settings["max_comments"]:
            items, next_cursor = client.media_comments_chunk(
                media_id,
                max_amount=settings["max_comments"] - len(parents),
                min_id=cursor,
            )
            for item in items:
                native = native_pk(_value(item, "pk"))
                if native in seen:
                    continue
                seen.add(native)
                emit("record", normalize_item(item, url, media=media))
                parents.append(native)
                if len(parents) >= settings["max_comments"]:
                    break
            if not next_cursor or next_cursor in cursors or not items:
                break
            cursors.add(next_cursor)
            cursor = next_cursor
        replies = 0
        for parent in parents:
            if (
                replies >= settings["max_replies"]
                or not settings["max_replies_per_thread"]
            ):
                break
            cursor, cursors, count = None, set(), 0
            cap = min(
                settings["max_replies_per_thread"], settings["max_replies"] - replies
            )
            while count < cap:
                items, next_cursor = client.media_comment_replies_chunk(
                    media_id, parent, max_amount=cap - count, min_id=cursor
                )
                for item in items:
                    native = native_pk(_value(item, "pk"))
                    if native in seen:
                        continue
                    seen.add(native)
                    emit(
                        "record", normalize_item(item, url, media=media, parent=parent)
                    )
                    count += 1
                    replies += 1
                    if count >= cap:
                        break
                if not next_cursor or next_cursor in cursors or not items:
                    break
                cursors.add(next_cursor)
                cursor = next_cursor
    else:
        raise ValueError("Unknown Instagram operation")
    return dict(
        status="FINISHED_BEST_EFFORT",
        reason=operation + "_completed",
        warnings=["comments_not_exhaustive"],
        source_version=upstream_version(),
    )


def instagram_worker(operation, target, settings, pipe):
    if os.name == "posix":
        os.setsid()
    try:
        # Third-party output may contain identifiers, response bodies or codes.
        with (
            open(os.devnull, "w") as quiet,
            redirect_stdout(quiet),
            redirect_stderr(quiet),
        ):
            result = collect(
                operation, target, settings, lambda kind, row: pipe.send((kind, row))
            )
    except (Exception, KeyboardInterrupt) as error:
        result = error_receipt(error)
    try:
        pipe.send(("done", result))
    finally:
        pipe.close()


def run_instagram_worker(
    operation, target, settings, seconds, emit, *, worker_entry=instagram_worker
):
    return contract.run_worker(
        operation, target, settings, seconds, emit, worker_entry=worker_entry
    )


def _integer(value, name, low, high):
    if type(value) is not int or not low <= value <= high:
        raise ValueError(f"{name} must be an integer in {low}..{high}")


def _configuration(value):
    if isinstance(value, (str, Path)):
        path = Path(value)
        if path.stat().st_size > 1_000_000:
            raise ValueError("config exceeds 1 MB")
        value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("config must be an object")
    cfg = deepcopy(value)
    if set(cfg) - {
        "project_id",
        "since",
        "until",
        "include_keywords",
        "exclude_keywords",
        "include_mode",
        "delay",
        "timeout_seconds",
        "total_seconds",
        "exclude_existing_datasets",
        "platforms",
        "instagrapi",
        "backend",
        "browser",
    }:
        raise ValueError("Unknown Instagram config field")
    if not isinstance(cfg.get("project_id"), str) or not re.fullmatch(
        r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}", str(cfg.get("project_id", ""))
    ):
        raise ValueError("project_id must be a slug")
    for key in ("since", "until"):
        cfg.setdefault(key, None)
        if cfg[key] is not None:
            cfg[key] = contract.utc_datetime(cfg[key]).isoformat()
    if cfg["since"] and cfg["until"] and cfg["since"] >= cfg["until"]:
        raise ValueError("since must precede until")
    for key in ("include_keywords", "exclude_keywords"):
        items = cfg.setdefault(key, [])
        if (
            not isinstance(items, list)
            or len(items) > 500
            or any(
                not isinstance(s, str)
                or not s.strip()
                or len(s) > 300
                or any(ord(c) < 32 for c in s)
                for s in items
            )
        ):
            raise ValueError("Invalid " + key)
        cfg[key] = list(dict.fromkeys(s.strip() for s in items))
    cfg.setdefault("include_mode", "any")
    if cfg["include_mode"] not in {"any", "all"}:
        raise ValueError("include_mode must be any or all")
    _integer(cfg.setdefault("total_seconds", 1800), "total_seconds", 1, 86400)
    _integer(cfg.setdefault("timeout_seconds", 90), "timeout_seconds", 5, 600)
    delay = cfg.setdefault("delay", {"mode": "fixed", "seconds": 2})
    if (
        not isinstance(delay, dict)
        or set(delay) != {"mode", "seconds"}
        or delay["mode"] != "fixed"
        or type(delay["seconds"]) not in (int, float)
        or not math.isfinite(delay["seconds"])
        or not 0 <= delay["seconds"] <= 300
    ):
        raise ValueError("Instagram supports a fixed delay of 0..300 seconds")
    cfg.setdefault("backend", "instagrapi")
    if cfg["backend"] not in {"instagrapi", "browser"}:
        raise ValueError("backend must be instagrapi or browser")
    if cfg["backend"] == "browser":
        from Crawler.Backend.instagram_browser import browser_configuration

        cfg["browser"] = browser_configuration(cfg.get("browser", {}))
    elif "browser" in cfg:
        raise ValueError("browser settings require backend=browser")
    auth = cfg.setdefault("instagrapi", {})
    if not isinstance(auth, dict) or set(auth) - {
        "session_file",
        "request_delay_seconds",
    }:
        raise ValueError(
            "instagrapi only accepts session_file/request_delay_seconds, never credentials"
        )
    auth.setdefault("session_file", "output/instagram/session.json")
    path = auth["session_file"]
    if not isinstance(path, str) or not path.strip():
        raise ValueError("session_file must be a local path")
    auth["session_file"] = str((ROOT / path).resolve())
    delay = auth.setdefault("request_delay_seconds", 2)
    if (
        type(delay) not in (int, float)
        or not math.isfinite(delay)
        or not 1 <= delay <= 60
    ):
        raise ValueError("request_delay_seconds must be 1..60")
    platforms = cfg.get("platforms")
    if (
        not isinstance(platforms, dict)
        or set(platforms) != {"instagram"}
        or not isinstance(platforms["instagram"], dict)
    ):
        raise ValueError("Standalone trial requires only platforms.instagram")
    p = platforms["instagram"]
    defaults = dict(
        enabled=True,
        urls=[],
        hashtags=cfg["include_keywords"][:10],
        min_posts=1000,
        min_comments_per_post=10,
        max_posts=4000,
        max_candidates=6000,
        max_comments=100,
        max_replies=30,
        max_replies_per_thread=10,
        search_results=50,
        max_search_pages=150,
        search_tab="recent",
        rate_limit_cooldown_seconds=360,
        rate_limit_max_retries=0,
    )
    if set(p) - set(defaults):
        raise ValueError("Unknown Instagram platform setting")
    p = {**defaults, **p}
    if p["enabled"] is not True:
        raise ValueError("Enable Instagram for the standalone trial")
    for key, low, high in [
        ("min_posts", 1, 10000),
        ("max_posts", 1, 10000),
        ("max_candidates", 1, 20000),
        ("min_comments_per_post", 0, 2000),
        ("max_comments", 1, 2000),
        ("max_replies", 0, 1000),
        ("max_replies_per_thread", 0, 100),
        ("search_results", 1, 50),
        ("max_search_pages", 1, 500),
        ("rate_limit_cooldown_seconds", 1, 10800),
        ("rate_limit_max_retries", 0, 10),
    ]:
        _integer(p[key], key, low, high)
    if (
        p["min_posts"] > min(p["max_posts"], p["max_candidates"])
        or p["min_comments_per_post"] > p["max_comments"]
    ):
        raise ValueError("Target exceeds acquisition caps")
    if p["search_tab"] not in {"recent", "top", "clips"}:
        raise ValueError("search_tab must be recent, top or clips")
    if not isinstance(p["urls"], list) or len(p["urls"]) > 20000:
        raise ValueError("Invalid URL list")
    p["urls"] = list(dict.fromkeys(media_url(u) for u in p["urls"]))
    if (
        not isinstance(p["hashtags"], list)
        or len(p["hashtags"]) > 100
        or any(
            not isinstance(t, str) or not re.fullmatch(r"[\w]{1,100}", t, re.UNICODE)
            for t in p["hashtags"]
        )
    ):
        raise ValueError("Use hashtag names without # or spaces")
    p["hashtags"] = list(dict.fromkeys(p["hashtags"]))
    if not p["urls"] and not p["hashtags"]:
        raise ValueError("Provide Instagram URLs or hashtags")
    cfg["platforms"] = {"instagram": p}
    paths = cfg.setdefault("exclude_existing_datasets", [])
    if (
        not isinstance(paths, list)
        or len(paths) > 100
        or any(not isinstance(s, str) or not s.strip() for s in paths)
    ):
        raise ValueError("Invalid baseline dataset paths")
    return cfg


class InstagramCrawler(SocialExportMixin):
    """Standalone trial. The production SocialCrawler backend stays unchanged."""

    def __init__(
        self,
        config=None,
        *,
        worker=None,
        progress=None,
        sleep=time.sleep,
        monotonic=time.monotonic,
    ):
        self.config = _configuration(
            ROOT / "config/instagram_crawler.json" if config is None else config
        )
        self.method = METHOD
        self.backend_version = upstream_version()
        self.access_mode = "saved_authenticated_mobile_session;read_only_acquisition"
        self._browser_worker = None
        if self.config["backend"] == "browser":
            from Crawler.Backend.instagram_browser import (
                BROWSER_METHOD,
                CLI_VERSION,
                InstagramBrowserWorker,
            )

            self.method = BROWSER_METHOD
            self.backend_version = CLI_VERSION
            self.access_mode = "manual_login_persistent_chrome;read_only_acquisition"
            if worker is None:
                self._browser_worker = InstagramBrowserWorker(self.config["browser"])
        self.worker = worker or self._browser_worker or run_instagram_worker
        self.synthetic = worker is not None
        self.progress = progress or (lambda event: None)
        self.sleep, self.clock = sleep, monotonic
        self.include = contract.KeywordMatcher(self.config["include_keywords"])
        self.exclude = contract.KeywordMatcher(self.config["exclude_keywords"])
        self.settings = InstagramSettings(
            **self.config["instagrapi"],
            **{
                k: self.config["platforms"]["instagram"][k]
                for k in (
                    "search_results",
                    "max_comments",
                    "max_replies",
                    "max_replies_per_thread",
                    "search_tab",
                )
            },
        )

    def plan(self):
        return dict(
            config=deepcopy(self.config),
            acquisition_backend=self.method,
            upstream_version=self.backend_version,
            local_session_present=(
                Path(self.config["browser"]["profile_dir"]).is_dir()
                if self.config["backend"] == "browser"
                else Path(self.settings.session_file).is_file()
            ),
            live_access_verified=False,
            trial_passed=False,
            platform_scope="instagram_only",
            date_scope="root_publication",
            until_exclusive=True,
            search_scope="hashtag discovery; dates and keywords enforced locally",
            comments_complete=False,
            primary_output=dict(
                file=COMPACT_FILE, sheets=["Data"], columns=list(COMPACT_COLUMNS)
            ),
            integrated_into_social_crawler=False,
        )

    def _baseline(self):
        from openpyxl import load_workbook

        urls, keys, sources = set(), set(), []
        for supplied in self.config["exclude_existing_datasets"]:
            path = (ROOT / supplied).resolve()
            if path.is_dir():
                path /= "dataset.xlsx"
            with path.open("rb") as stream:
                checksum = file_digest(stream, "sha256").hexdigest()
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = book["Contents"].iter_rows(values_only=True)
                headers = next(rows)
                if not {"platform", "object_type", "canonical_url", "native_id"} <= set(
                    headers
                ):
                    raise ValueError("Baseline requires the audit Contents sheet")
                for cells in rows:
                    row = dict(zip(headers, cells))
                    if row["platform"] == "instagram" and row["object_type"] == "post":
                        if native_pk(row["native_id"]) != media_pk_from_url(
                            row["canonical_url"]
                        ):
                            raise ValueError("Baseline native ID/shortcode mismatch")
                        urls.add(media_url(row["canonical_url"]))
                        keys.add(object_key(row["native_id"]))
            finally:
                book.close()
            sources.append(dict(path=str(path), sha256=checksum))
        return urls, keys, sources

    def check_session(self):
        return self.worker(
            "check_session",
            "instagram_account",
            self.settings,
            self.config["timeout_seconds"],
            lambda *args: None,
        )

    def run(self, output_dir):
        excluded_urls, excluded_keys, baselines = self._baseline()
        output = Path(output_dir).resolve()
        run_id = str(uuid.uuid4())
        factory = contract.RecordFactory(
            self.config["project_id"], run_id, self.include
        )
        sink = contract.BundleSink(output, self.config["project_id"], run_id)
        p = self.config["platforms"]["instagram"]
        began = self.clock()
        deadline = began + self.config["total_seconds"]
        counts = dict(
            posts_attempted=0,
            posts_exported=0,
            qualified_posts=0,
            comments=0,
            replies=0,
            invalid_rows=0,
            duplicate_posts_skipped=0,
            existing_posts_skipped=0,
            rate_limit_events=0,
            rate_limit_retries=0,
            search_pages=0,
        )
        seen_urls, exported, all_keys, relations = set(), set(), set(), []
        queue, searches = deque(), deque((t, None) for t in p["hashtags"])
        search_cursors, filtered = set(), Counter()
        previous = None
        source_errors = 0
        stop = "CANDIDATES_EXHAUSTED"
        interrupted = False
        progress_path = output.with_name(output.name + ".progress.json")

        def event(stage, **extra):
            payload = dict(
                stage=stage,
                observed_at=contract.now_iso(),
                platform="instagram",
                elapsed_seconds=round(self.clock() - began, 2),
                remaining_seconds=round(max(0, deadline - self.clock()), 2),
                **counts,
                **extra,
            )
            temp = progress_path.with_suffix(".tmp")
            temp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            temp.replace(progress_path)
            self.progress(payload)

        def enqueue(url, origin=None):
            url = media_url(url)
            if url in seen_urls:
                counts["duplicate_posts_skipped"] += 1
                return
            if len(seen_urls) >= p["max_candidates"]:
                return
            seen_urls.add(url)
            if url in excluded_urls:
                counts["existing_posts_skipped"] += 1
            else:
                queue.append((url, origin))

        for url in p["urls"]:
            enqueue(url)
        try:
            event("started")
            while self.clock() < deadline:
                if counts["qualified_posts"] >= p["min_posts"]:
                    stop = "TARGET_MET"
                    break
                if counts["posts_attempted"] >= p["max_posts"]:
                    stop = "POST_LIMIT"
                    break
                if queue:
                    operation, (target, origin) = "media", queue.popleft()
                    counts["posts_attempted"] += 1
                elif (
                    searches
                    and counts["search_pages"] < p["max_search_pages"]
                    and len(seen_urls) < p["max_candidates"]
                ):
                    operation, origin = "search", None
                    tag, cursor = searches.popleft()
                    target = json.dumps(dict(hashtag=tag, cursor=cursor))
                    counts["search_pages"] += 1
                else:
                    break
                events = []
                if previous is not None:
                    wait = max(
                        0, previous + self.config["delay"]["seconds"] - self.clock()
                    )
                    self.sleep(min(wait, max(0, deadline - self.clock())))
                if self.clock() >= deadline:
                    stop = "TIME_BUDGET_EXHAUSTED"
                    break
                started = contract.now_iso()
                try:
                    result = self.worker(
                        operation,
                        target,
                        self.settings,
                        min(self.config["timeout_seconds"], deadline - self.clock()),
                        lambda kind, row: events.append((kind, row)),
                    )
                except (Exception, KeyboardInterrupt) as error:
                    result = error_receipt(error)
                previous = self.clock()
                source_errors = (
                    source_errors + 1
                    if result["status"] in {"SOURCE_ERROR", "WORKER_FAILED", "TIMEOUT"}
                    else 0
                )
                written = 0
                invalid = 0
                if operation == "search":
                    for kind, row in events:
                        if kind == "candidate":
                            try:
                                enqueue(row["url"], row.get("discovered_from"))
                            except (ValueError, KeyError, TypeError):
                                invalid += 1
                    next_cursor = result.get("next_cursor")
                    if next_cursor and (tag, next_cursor) not in search_cursors:
                        search_cursors.add((tag, next_cursor))
                        searches.append((tag, next_cursor))
                else:
                    rows = {key: [] for key in contract.TABLES}
                    memory = SimpleNamespace(
                        write=lambda table, row: rows[table].append(row)
                    )
                    dataset = contract.Dataset(factory, memory)
                    adapter = InstagramAdapter(
                        factory,
                        dataset,
                        result.get("source_version"),
                        method=self.method,
                    )
                    for kind, raw in events:
                        try:
                            adapter.accept(kind, raw, target, origin)
                        except (ValueError, KeyError, TypeError, OverflowError):
                            invalid += 1
                    dataset.finish_relations()
                    selected, reason = self._select(rows)
                    if selected and selected[0]["object_key"] in excluded_keys:
                        selected, reason = [], "root_seen_in_existing_dataset"
                        counts["existing_posts_skipped"] += 1
                    elif selected and selected[0]["object_key"] in exported:
                        selected, reason = [], "duplicate_root_in_run"
                        counts["duplicate_posts_skipped"] += 1
                    if selected:
                        if sink.counts["Contents"] + len(selected) > 500000:
                            selected, reason = [], "output_row_budget"
                            result = {**result, "status": "OUTPUT_ROW_BUDGET"}
                    if selected:
                        exported.add(selected[0]["object_key"])
                        qualified = len(selected) - 1 >= p["min_comments_per_post"]
                        counts["posts_exported"] += 1
                        counts["qualified_posts"] += int(qualified)
                        ids = {r["observation_id"] for r in selected}
                        keys = {r["object_key"] for r in selected}
                        all_keys.update(keys)
                        for row in selected:
                            row["quality_flags"].append(
                                "social_minimum_comments_met"
                                if qualified
                                else "social_minimum_comments_shortfall"
                            )
                            sink.write("Contents", row)
                            counts["comments"] += row["object_type"] == "comment"
                            counts["replies"] += row["object_type"] == "reply"
                        for table in ("Metrics", "KeywordHits"):
                            for row in rows[table]:
                                if row.get("observation_id") in ids:
                                    sink.write(table, row)
                        relations.extend(
                            r for r in rows["Relations"] if r["subject_key"] in keys
                        )
                        written = len(selected)
                    else:
                        filtered[reason] += 1
                counts["invalid_rows"] += invalid
                sink.write(
                    "Runs",
                    factory.row(
                        "Runs",
                        unit_id=str(uuid.uuid4()),
                        source="instagram",
                        operation=operation,
                        target=("https://www.instagram.com/explore/tags/" + tag + "/")
                        if operation == "search"
                        else target,
                        status=result["status"],
                        reason=result.get("reason", "source_result"),
                        coverage="PARTIAL_NON_EXHAUSTIVE",
                        contents_written=written,
                        comments_seen=sum(
                            k == "record" and r.get("object_type") == "comment"
                            for k, r in events
                        ),
                        replies_seen=sum(
                            k == "record" and r.get("object_type") == "reply"
                            for k, r in events
                        ),
                        duplicates_skipped=0,
                        invalid_rows=invalid,
                        max_comments=p["max_comments"],
                        max_replies=p["max_replies"],
                        timeout_seconds=self.config["timeout_seconds"],
                        keywords=self.config["include_keywords"],
                        source_version=result.get("source_version"),
                        warning_codes=result.get("warnings", []),
                        started_at=started,
                        finished_at=contract.now_iso(),
                    ),
                )
                event(
                    "operation_finished", operation=operation, status=result["status"]
                )
                if result["status"] == "RATE_LIMITED":
                    counts["rate_limit_events"] += 1
                    wait = max(
                        p["rate_limit_cooldown_seconds"],
                        result.get("retry_after_seconds") or 0,
                    )
                    if (
                        not written
                        and counts["rate_limit_retries"] < p["rate_limit_max_retries"]
                        and wait < deadline - self.clock()
                    ):
                        counts["rate_limit_retries"] += 1
                        event("rate_limit_cooldown", cooldown_seconds=wait)
                        remaining = wait
                        while remaining > 0:
                            tick = min(30, remaining, max(0, deadline - self.clock()))
                            if tick <= 0:
                                break
                            self.sleep(tick)
                            remaining -= tick
                        if operation == "media":
                            queue.appendleft((target, origin))
                        else:
                            searches.appendleft((tag, cursor))
                        continue
                if result["status"] in STOP_STATUSES:
                    stop = result["status"]
                    interrupted = stop == "INTERRUPTED"
                    break
                if result["status"] == "OUTPUT_ROW_BUDGET" or source_errors >= 3:
                    stop = (
                        "OUTPUT_ROW_BUDGET"
                        if result["status"] == "OUTPUT_ROW_BUDGET"
                        else "CONSECUTIVE_SOURCE_ERRORS"
                    )
                    break
            else:
                stop = "TIME_BUDGET_EXHAUSTED"
        except KeyboardInterrupt:
            interrupted, stop = True, "INTERRUPTED"
        except BaseException:
            sink.abort()
            raise
        try:
            for row in relations:
                row["resolution_status"] = (
                    "resolved_in_export"
                    if row["object_key"] in all_keys
                    else "unresolved_reference"
                )
                sink.write("Relations", row)
            met = counts["qualified_posts"] >= p["min_posts"]
            live = not self.synthetic and counts["posts_exported"] > 0
            trial_passed = (
                live
                and counts["comments"] + counts["replies"] > 0
                and not counts["invalid_rows"]
                and stop not in STOP_STATUSES
                and stop not in {"CONSECUTIVE_SOURCE_ERRORS", "OUTPUT_ROW_BUDGET"}
            )
            summary = dict(
                status="INTERRUPTED"
                if interrupted
                else "TARGET_MET"
                if met and stop not in STOP_STATUSES
                else "PARTIAL",
                stop_reason=stop,
                targets_met=met,
                live_access_verified=live,
                trial_passed=trial_passed,
                integrated_into_social_crawler=False,
                posts_shortfall=max(0, p["min_posts"] - counts["qualified_posts"]),
                filtered_roots=dict(filtered),
                **counts,
                run_id=run_id,
                output_dir=str(output),
                primary_file=str(output / COMPACT_FILE),
                elapsed_seconds=round(self.clock() - began, 2),
                configured_seconds=self.config["total_seconds"],
                comments_complete=False,
                synthetic=self.synthetic,
                deduplication_sources=baselines,
            )
            sink.write(
                "Sources",
                factory.row(
                    "Sources",
                    source="instagram",
                    producer_version="instagram-crawler:" + VERSION,
                    upstream_version=self.backend_version,
                    extraction_method=self.method,
                    access_mode=self.access_mode,
                    config={**deepcopy(self.config), "summary": summary},
                    schema_fingerprint=contract.schema_fingerprint(),
                    observed_at=contract.now_iso(),
                ),
            )
            for stream in sink.files.values():
                stream.flush()
            summary["compact"] = self._write_compact(
                sink.stage, sink.stage / COMPACT_FILE
            )
            contract.add_readme(
                sink,
                factory,
                [
                    (
                        "InstagramCrawler",
                        "Caption, komentar dan balasan memakai kontrak bersama. Backend: "
                        + self.method
                        + ". Post media/album dihitung satu root; file gambar/video tidak diunduh.",
                    ),
                    (
                        "Cakupan",
                        "Hashtag tidak menjamin hasil lengkap/kronologis. Tanggal root since inklusif/until eksklusif; waktu tidak diketahui tetap unknown. Peran penulis tidak ditebak.",
                    ),
                    (
                        "Autentikasi",
                        "Sesi disimpan di luar Git. LOGIN/CHALLENGE menghentikan crawl tanpa solver atau pergantian identitas. 429 mengikuti cooldown terkonfigurasi, default tanpa retry.",
                    ),
                    ("Ringkasan", json.dumps(summary, ensure_ascii=False)),
                ],
            )
            # Include summary in manifest checksums as well as both Excel exports.
            (sink.stage / "summary.json").write_text(
                json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            sink.finish(
                dict(
                    source="instagram",
                    mode="instagram_crawler",
                    producer_version=VERSION,
                    synthetic=self.synthetic,
                    trial_passed=trial_passed,
                    comments_complete=False,
                    primary_file=COMPACT_FILE,
                )
            )
            event(
                "finished",
                status=summary["status"],
                stop_reason=stop,
                output_dir=str(output),
            )
            return summary
        except BaseException:
            if not sink.closed:
                sink.abort()
            raise
        finally:
            if self._browser_worker is not None:
                self._browser_worker.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--config", type=Path, default=ROOT / "config/instagram_crawler.json"
    )
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument(
        "--auth-file",
        type=Path,
        help="Ignored local credentials file; used only with --login",
    )
    parser.add_argument(
        "--login-method",
        choices=("auto", "password", "sessionid"),
        default="auto",
        help="Select local-file login; password restores settings and validates via login()",
    )
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--dry-run", action="store_true")
    group.add_argument("--login", action="store_true")
    group.add_argument("--check-session", action="store_true")
    args = parser.parse_args()
    if args.auth_file and not args.login:
        parser.error("--auth-file requires --login")
    if args.login_method != "auto" and not (args.login and args.auth_file):
        parser.error("--login-method requires --login and --auth-file")
    crawler = InstagramCrawler(
        args.config,
        progress=lambda event: print(json.dumps(event, ensure_ascii=False), flush=True),
    )
    if args.login and crawler.config["backend"] == "browser":
        parser.error(
            "Browser login is manual; open the configured persistent Chrome profile"
        )
    if args.login:
        try:
            if args.auth_file:
                result = run_instagram_worker(
                    "login"
                    if args.login_method == "auto"
                    else "login_" + args.login_method,
                    str(args.auth_file.resolve()),
                    crawler.settings,
                    crawler.config["timeout_seconds"],
                    lambda *args: None,
                )
            else:
                login_local(crawler.settings.session_file)
                result = dict(
                    status="SESSION_SAVED", session_file=crawler.settings.session_file
                )
        except (Exception, KeyboardInterrupt) as error:
            result = error_receipt(error)
    elif args.dry_run:
        result = crawler.plan()
    elif args.check_session:
        result = crawler.check_session()
    else:
        output = args.output_dir or ROOT / "outputs" / (
            "instagram_"
            + crawler.config["backend"]
            + "_"
            + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        )
        result = crawler.run(output)
    print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
    return (
        0
        if args.dry_run
        or result.get("status")
        in {"TARGET_MET", "SESSION_SAVED", "FINISHED_BEST_EFFORT"}
        else 2
    )


if __name__ == "__main__":
    raise SystemExit(main())
