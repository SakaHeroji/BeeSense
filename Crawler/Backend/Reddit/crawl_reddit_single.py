#!/usr/bin/env python3
"""Crawler Reddit satu file — Friday / University Social Intelligence.

Python 3.12+. Instalasi: python -m pip install openpyxl==3.1.5
Tidak memerlukan PRAW, API key, cookies, browser, file common atau config.
Mengambil HALAMAN HTML publik saja; tidak memakai .json, OAuth atau Data API.

CONTOH:
  python crawl_reddit_single.py run --query "biaya kuliah BINUS" --subreddit indonesia --subreddit karir --max-posts 3 --max-comments 100 --output-dir hasil_reddit
  python crawl_reddit_single.py run --url "https://www.reddit.com/r/karir/comments/1x4sj4l/alternatif_binus/" --output-dir hasil_thread
  python crawl_reddit_single.py import-html --input thread.html --url "https://www.reddit.com/r/karir/comments/1x4sj4l/alternatif_binus/" --output-dir hasil_html
  python crawl_reddit_single.py self-test
  python crawl_reddit_single.py doctor
  python crawl_reddit_single.py contract --json

OUTPUT:
  Folder BARU: dataset.xlsx, delapan tabel JSONL, contract.json, manifest.json.
  Sheet dan kolom IDENTIK dengan crawl_youtube_single.py kontrak 1.0.0:
  Contents, Relations, KeywordHits, Metrics, Runs, Sources, TextChunks, Readme.
  contract.json adalah output; bukan berkas yang harus disiapkan pengguna.

BATAS PENTING:
  Reddit dapat mengirim challenge/blokir bahkan dengan HTTP 200. Kondisi itu
  dicatat di Runs; script tidak menyelesaikan challenge atau melewati login.
  Hanya komentar yang benar-benar ada dalam HTML respons yang dikumpulkan.
  Tidak mengeksekusi JavaScript, menggulir, memperluas "more replies", melakukan
  retry blokir, atau mengambil isi di dalam template inert. HTML dapat parsial.
  Pencarian menemukan POST; bukan pencarian seluruh komentar Reddit.
  --max-comments mencakup komentar utama + seluruh balasan per post.
  Teks adalah teks hasil render HTML, bukan Markdown mentah asli.
  Waktu tanpa timezone/relatif tidak diubah menjadi published_at exact.
  Peran penulis unknown. Tidak ada klasifikasi sentimen/emosi atau upload DB.
  import-html tidak mengakses jaringan; asal HTML tidak diverifikasi independen.
  Filter --since/--until menandai hasil; tidak menjamin coverage semua periode.
  --sort untuk komentar; --search-sort untuk pencarian post. Keduanya best effort.

DESAIN:
  SourceAdapter (Adapter), Web/LocalHtmlSource (Strategy), RecordFactory
  (Factory), Dataset + BundleSink (pipeline). Aho-Corasick dikompilasi sekali.
  HTTP serial berjarak, batas byte/node, subprocess dengan deadline, ekspor
  streaming dan atomic publish. Tidak ada dependency modul proyek lain.

REFERENSI STRUKTUR HTML (dapat berubah):
  https://github.com/relliv/reddit-post-markdown-chrome-extension/blob/main/src/content/extract.ts
  https://github.com/reddit-archive/reddit/tree/master/r2/r2/templates
"""

from __future__ import annotations

from collections import deque
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from datetime import datetime, timezone
from hashlib import sha256
import hmac
import json
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata
import uuid

CONTRACT_VERSION = "1.0.0"
PRD_REFERENCE = "2.1"
UTC = timezone.utc
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")

def now_iso():
    return datetime.now(UTC).isoformat()

def dumps(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)

def digest(value):
    return sha256(value.encode("utf-8")).hexdigest()

def utc_datetime(value):
    d = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if d.tzinfo is None:
        raise ValueError("Datetime wajib memiliki timezone.")
    return d.astimezone(UTC)

# Each field has one authoritative type and nullability, across every producer.
def fields(spec):
    result = []
    for token in spec.split():
        name, kind = token.split(":")
        result.append({"name":name,"type":kind.rstrip("?"),"nullable":kind.endswith("?")})
    return result + [{"name":"excel_notes","type":"json","nullable":False}]

TABLES = {
    "Contents": fields("contract_version:str project_id:str run_id:str observation_id:str platform:str source_namespace:str object_key:str native_id:str identity_quality:str object_type:str canonical_url:str page_url:str discovered_from:str? title:str? text_source:str text_normalized:str normalizer_version:str text_hash:str text_complete:bool content_revision:int? author_ref:str? publisher_ref:str? published_at:datetime? published_at_display:str? timestamp_quality:str published_interval_start:datetime? published_interval_end:datetime? observed_at:datetime parent_key:str? thread_root_key:str relation_evidence:str? selection_status:str context_only:bool actor_type:str actor_evidence:str? extraction_method:str adapter_version:str quality_flags:json missing_fields:json"),
    "Relations": fields("contract_version:str project_id:str run_id:str relation_id:str subject_key:str predicate:str object_key:str subject_observation_id:str relation_evidence:str resolution_status:str observed_at:datetime"),
    "KeywordHits": fields("contract_version:str project_id:str run_id:str hit_id:str subject_key:str observation_id:str evidence_key:str field_name:str keyword:str matched_variant:str span_start:int span_end:int match_type:str matcher_version:str catalog_version:str"),
    "Metrics": fields("contract_version:str project_id:str run_id:str metric_id:str observation_id:str? platform:str source_namespace:str object_key:str? series_key:str? metric_name:str value:float? value_display:str? unit:str scale_min:float? scale_max:float? period_start:datetime? period_end:datetime? dimensions:json precision:str observed_at:datetime source_url:str quality_flags:json"),
    "Runs": fields("contract_version:str project_id:str run_id:str unit_id:str source:str operation:str target:str status:str reason:str coverage:str contents_written:int comments_seen:int replies_seen:int duplicates_skipped:int invalid_rows:int max_comments:int? max_replies:int? timeout_seconds:int keywords:json source_version:str? warning_codes:json started_at:datetime finished_at:datetime"),
    "Sources": fields("contract_version:str project_id:str run_id:str source:str producer_version:str upstream_version:str? extraction_method:str access_mode:str config:json schema_fingerprint:str observed_at:datetime"),
    "TextChunks": fields("contract_version:str project_id:str run_id:str record_key:str table_name:str field_name:str part_index:int text_part:str encoding:str"),
    "Readme": fields("contract_version:str project_id:str run_id:str item:str value:str"),
}
ENUMS = {
    "Contents.object_type":["post","comment","reply"],
    "Contents.identity_quality":["native_id","canonical_permalink","weak_synthetic"],
    "Contents.timestamp_quality":["exact","relative","date_only","unknown"],
    "Contents.selection_status":["matched","not_matched","context_only","review","excluded"],
    "Contents.actor_type":["unknown","prospective_student","parent","student","alumni","other"],
    "Relations.resolution_status":["resolved_in_export","unresolved_reference"],
    "Metrics.precision":["exact","rounded","unknown"],
}

def schema_document():
    return {"contract_version":CONTRACT_VERSION,"prd_acquisition_reference":PRD_REFERENCE,
            "profile":"tabular_projection_not_full_production_contract",
            "sheet_order":list(TABLES),"tables":TABLES,"enums":ENUMS,
            "time_storage":{"jsonl":"ISO8601 timezone-aware UTC","xlsx":"native Excel datetime UTC without tzinfo"},
            "revision_policy":"content_revision=null until a persistent repository assigns it; observation_id and text_hash bind snapshot evidence"}

def schema_fingerprint():
    return digest(dumps(schema_document()))

EXPECTED_SCHEMA_FINGERPRINT = "c6dcff73113f3aba31c2329296cd72b1c24ce6e15db29e22a9e35fa662fc74de"

def verify_descriptor(path=None):
    """Validate the embedded contract; an external descriptor is optional."""
    if schema_fingerprint() != EXPECTED_SCHEMA_FINGERPRINT:
        raise ValueError("Kontrak tertanam berubah; hentikan ekspor.")
    if path is not None and json.loads(Path(path).read_text(encoding="utf-8")) != schema_document():
        raise ValueError("Descriptor eksternal berbeda dari kontrak tertanam.")

def validate_row(table, row):
    expected = {f["name"] for f in TABLES[table]}
    if set(row) != expected:
        raise ValueError(f"Kolom {table} berbeda: missing={sorted(expected-set(row))}, extra={sorted(set(row)-expected)}")
    for spec in TABLES[table]:
        name, kind = spec["name"], spec["type"]
        v = row[name]
        if v is None:
            if not spec["nullable"]:
                raise ValueError(f"{table}.{name} tidak boleh null")
            continue
        valid = {"str":isinstance(v,str),"int":type(v) is int,
                 "float":type(v) in (int,float) and math.isfinite(v),
                 "bool":type(v) is bool,"json":isinstance(v,(list,dict)),
                 "datetime":isinstance(v,str)}[kind]
        if not valid:
            raise ValueError(f"Tipe {table}.{name} harus {kind}")
        if kind == "datetime":
            if datetime.fromisoformat(v.replace("Z","+00:00")).utcoffset()!=timezone.utc.utcoffset(None):
                raise ValueError(f"{table}.{name} wajib UTC")
            utc_datetime(v)
        if kind == "json":
            dumps(v)
        allowed = ENUMS.get(f"{table}.{name}")
        if allowed and v not in allowed:
            raise ValueError(f"Nilai {table}.{name} tidak dikenal")
    if row["contract_version"] != CONTRACT_VERSION:
        raise ValueError("Versi kontrak berbeda")
    if table == "Contents":
        if row["text_hash"] != digest(row["text_source"]):
            raise ValueError("text_hash tidak cocok")
        if row["timestamp_quality"] != "exact" and row["published_at"] is not None:
            raise ValueError("published_at hanya untuk waktu exact; estimasi tidak boleh dianggap pasti")
        if row["actor_type"]!="unknown" and not row["actor_evidence"]:
            raise ValueError("Peran penulis harus memiliki bukti eksplisit")

class KeywordMatcher:
    """Aho-Corasick; offsets refer to Python Unicode code points in source."""
    VERSION = "aho-casefold-source-offset-1"
    def __init__(self, terms):
        self.terms = tuple(dict.fromkeys(t.strip() for t in terms if t.strip()))
        if len(self.terms)>200 or any(len(t)>200 for t in self.terms):
            raise ValueError("Maksimum 200 keyword, masing-masing 200 karakter")
        self.nodes, self.fail, self.out = [{}], [0], [[]]
        for term in self.terms:
            state=0
            for char in term.casefold():
                if char not in self.nodes[state]:
                    self.nodes[state][char]=len(self.nodes)
                    self.nodes.append({}); self.fail.append(0); self.out.append([])
                state=self.nodes[state][char]
            self.out[state].append(term)
        queue=deque(self.nodes[0].values())
        while queue:
            state=queue.popleft()
            for char,child in self.nodes[state].items():
                queue.append(child); fallback=self.fail[state]
                while fallback and char not in self.nodes[fallback]:
                    fallback=self.fail[fallback]
                self.fail[child]=self.nodes[fallback].get(char,0)
                self.out[child].extend(self.out[self.fail[child]])
        self.catalog_version="literal-"+digest(dumps(self.terms))[:12]

    def find(self, text):
        folded=[]; positions=[]
        for i,char in enumerate(text):
            part=char.casefold(); folded.append(part); positions.extend([i]*len(part))
        value="".join(folded); state=0; seen=set()
        for end,char in enumerate(value):
            while state and char not in self.nodes[state]: state=self.fail[state]
            state=self.nodes[state].get(char,0)
            for keyword in self.out[state]:
                start=end-len(keyword.casefold())+1
                # Do not match a partial casefold expansion, e.g. one half of ß -> ss.
                if (start and positions[start]==positions[start-1]) or (end+1<len(positions) and positions[end]==positions[end+1]): continue
                a,b=positions[start],positions[end]+1
                word=lambda c:c.isalnum() or c=="_"
                if (a and word(text[a-1]) and word(text[a])) or (b<len(text) and word(text[b-1]) and word(text[b])): continue
                key=(keyword,a,b)
                if key not in seen:
                    seen.add(key); yield {"keyword":keyword,"matched_variant":text[a:b],"span_start":a,"span_end":b}

@dataclass(frozen=True, slots=True)
class ContentInput:
    platform: str
    namespace: str
    native_id: str
    object_type: str
    canonical_url: str
    page_url: str
    thread_root_key: str
    text_source: str
    parent_key: str|None=None
    title: str|None=None
    published_at: str|None=None
    published_display: str|None=None
    timestamp_quality: str="unknown"
    text_complete: bool=True
    author_native_id: str|None=None
    publisher_ref: str|None=None
    discovered_from: str|None=None
    observed_at: str=field(default_factory=now_iso)
    quality_flags: tuple[str,...]=()
    extraction_method: str="unofficial_web_client"
    adapter_version: str="unknown"

class RecordFactory:
    def __init__(self, project_id, run_id, matcher, since=None, until=None):
        self.project_id,self.run_id,self.matcher=project_id,run_id,matcher
        self.since=utc_datetime(since) if since else None
        self.until=utc_datetime(until) if until else None
        if self.since and self.until and self.since>=self.until: raise ValueError("since harus lebih awal daripada until")

    @staticmethod
    def key(platform,namespace,native_id):
        return f"{platform}:{namespace}:{native_id}"

    def row(self, table, **values):
        result={f["name"]:None for f in TABLES[table]}
        result.update(contract_version=CONTRACT_VERSION,project_id=self.project_id,run_id=self.run_id,excel_notes={})
        result.update(values)
        for spec in TABLES[table]:
            if spec["type"]=="datetime" and result[spec["name"]] is not None:
                result[spec["name"]]=utc_datetime(result[spec["name"]]).isoformat()
        validate_row(table,result)
        return result

    def content(self, value:ContentInput):
        key=self.key(value.platform,value.namespace,value.native_id)
        obs=str(uuid.uuid5(uuid.UUID(self.run_id),dumps([key,value.text_source,value.title,value.text_complete])))
        hits=[]
        for field_name,text in (("text_source",value.text_source),("title",value.title or "")):
            for hit in self.matcher.find(text):
                hits.append(self.row("KeywordHits",hit_id=digest(dumps([obs,field_name,hit])),subject_key=key,
                    observation_id=obs,evidence_key=key,field_name=field_name,**hit,match_type="direct_literal",
                    matcher_version=self.matcher.VERSION,catalog_version=self.matcher.catalog_version))
        context=value.object_type=="post"
        selection="context_only" if context else "matched" if hits else "not_matched"
        flags=list(value.quality_flags)
        if self.since or self.until:
            if value.published_at:
                dt=utc_datetime(value.published_at)
                if (self.since and dt<self.since) or (self.until and dt>=self.until):
                    selection="context_only" if context else "excluded"; flags.append("outside_requested_window")
            else:
                if not context: selection="review"
                flags.append("window_membership_unknown")
        if not value.text_complete: flags.append("text_incomplete")
        missing=[]
        if not value.text_source: missing.append("text_source")
        if value.timestamp_quality=="unknown": missing.append("published_at")
        author_ref=None
        secret=os.environ.get("AUTHOR_HMAC_KEY")
        if secret and value.author_native_id:
            author_ref=hmac.new(secret.encode(),f"{self.project_id}:{value.platform}:{value.author_native_id}".encode(),"sha256").hexdigest()
        record=self.row("Contents",observation_id=obs,platform=value.platform,source_namespace=value.namespace,
            object_key=key,native_id=value.native_id,identity_quality="native_id",object_type=value.object_type,
            canonical_url=value.canonical_url,page_url=value.page_url,discovered_from=value.discovered_from,
            title=value.title,text_source=value.text_source,text_normalized=unicodedata.normalize("NFC",value.text_source),
            normalizer_version="nfc-1",text_hash=digest(value.text_source),text_complete=value.text_complete,
            content_revision=None,author_ref=author_ref,publisher_ref=value.publisher_ref,published_at=value.published_at,
            published_at_display=value.published_display,timestamp_quality=value.timestamp_quality,
            published_interval_start=None,published_interval_end=None,observed_at=value.observed_at,
            parent_key=value.parent_key,thread_root_key=value.thread_root_key,
            relation_evidence="source_parent_id" if value.parent_key else None,selection_status=selection,
            context_only=context,actor_type="unknown",actor_evidence=None,extraction_method=value.extraction_method,
            adapter_version=value.adapter_version,quality_flags=flags,missing_fields=missing)
        return record,hits

class Dataset:
    """Shared pipeline; adapters cannot choose their own headers or cell formats."""
    def __init__(self, factory, sink):
        self.factory,self.sink=factory,sink
        self.seen=set(); self.object_keys=set(); self.pending_relations=[]
        self.duplicates=0; self.contents=0; self.last_record=None

    def add_content(self, raw:ContentInput, metrics=()):
        record,hits=self.factory.content(raw)
        identity=record["observation_id"]
        if identity in self.seen:
            self.duplicates+=1; return False
        self.seen.add(identity); self.object_keys.add(record["object_key"]); self.contents+=1; self.last_record=record
        self.sink.write("Contents",record)
        for h in hits: self.sink.write("KeywordHits",h)
        if record["parent_key"]:
            self.pending_relations.append((record,"reply_to"))
        if record["object_key"]!=record["thread_root_key"]:
            self.pending_relations.append((record,"in_thread"))
        for name,value,unit in metrics:
            if type(value) not in (int,float) or not math.isfinite(value): continue
            self.sink.write("Metrics",self.factory.row("Metrics",metric_id=digest(dumps([identity,name])),
                observation_id=identity,platform=raw.platform,source_namespace=raw.namespace,object_key=record["object_key"],
                series_key=None,metric_name=name,value=value,value_display=None,unit=unit,scale_min=None,scale_max=None,
                period_start=None,period_end=None,dimensions={},precision="unknown",observed_at=raw.observed_at,
                source_url=raw.canonical_url,quality_flags=["library_parsed_count_precision_unknown"]))
        return True

    def finish_relations(self):
        for r,predicate in self.pending_relations:
            target=r["parent_key"] if predicate=="reply_to" else r["thread_root_key"]
            self.sink.write("Relations",self.factory.row("Relations",relation_id=digest(dumps([r["observation_id"],predicate,target])),
                subject_key=r["object_key"],predicate=predicate,object_key=target,subject_observation_id=r["observation_id"],
                relation_evidence=r["relation_evidence"] if predicate=="reply_to" else "source_thread_context",
                resolution_status="resolved_in_export" if target in self.object_keys else "unresolved_reference",observed_at=r["observed_at"]))
        self.pending_relations.clear()

class SourceAdapter(ABC):
    @abstractmethod
    def accept(self,kind,raw,url,discovered_from):
        """Normalize a source event; write only through the shared Dataset/Sink."""
        raise NotImplementedError

class BundleSink:
    """Atomic new directory; streaming JSONL + streaming, formula-safe Excel."""
    def __init__(self,path,project_id,run_id):
        verify_descriptor()
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        self.path=Path(path).resolve(); self.project_id,self.run_id=project_id,run_id
        if self.path.exists(): raise FileExistsError("Output directory sudah ada; gunakan nama baru.")
        self.path.parent.mkdir(parents=True,exist_ok=True)
        self.stage=Path(tempfile.mkdtemp(prefix=".source-export-",dir=self.path.parent))
        self.book=Workbook(write_only=True); self.files={}; self.sheets={}; self.counts={t:0 for t in TABLES}
        self.header_font=Font(name="Calibri",size=11,bold=True,color="FFFFFF")
        self.font=Font(name="Calibri",size=11); self.fill=PatternFill("solid",fgColor="163C52")
        self.alignment=Alignment(wrap_text=True,vertical="top"); self.closed=False
        for table,specs in TABLES.items():
            self.files[table]=(self.stage/f"{table}.jsonl").open("w",encoding="utf-8")
            ws=self.book.create_sheet(table); self.sheets[table]=ws
            ws.freeze_panes="A2"; ws.sheet_view.showGridLines=False
            for i,spec in enumerate(specs,1):
                width=62 if spec["name"] in {"text_source","text_normalized","value","reason","config","source_url","canonical_url","page_url","target","quality_flags"} else 25
                ws.column_dimensions[get_column_letter(i)].width=width
            ws.append(self._cells(ws,[s["name"] for s in specs],True))

    def _cells(self,ws,values,header=False):
        from openpyxl.cell import WriteOnlyCell
        result=[]
        for v in values:
            c=WriteOnlyCell(ws,value=v)
            if isinstance(v,str): c.data_type="s"
            c.font=self.header_font if header else self.font; c.alignment=self.alignment
            if header:c.fill=self.fill
            if isinstance(v,datetime):c.number_format="yyyy-mm-dd hh:mm:ss"
            result.append(c)
        return result

    @staticmethod
    def chunks(text):
        current=[]; units=0
        for char in text:
            cost=2 if ord(char)>0xFFFF else 1
            if units+cost>30000:
                yield "".join(current);current=[];units=0
            current.append(char);units+=cost
        yield "".join(current)

    def write(self,table,row):
        validate_row(table,row)
        if self.counts[table]>=1_048_575:raise ValueError(f"Batas baris Excel {table} tercapai")
        notes={}; cells=[]
        key=row.get("observation_id") or row.get("unit_id") or row.get("metric_id") or row.get("hit_id") or f"{table}:{self.counts[table]}"
        for spec in TABLES[table]:
            name=spec["name"];value=row[name]
            if name=="excel_notes":cells.append(dumps(notes));continue
            if spec["type"]=="json":value=dumps(value)
            elif spec["type"]=="datetime" and value:value=utc_datetime(value).replace(tzinfo=None)
            if isinstance(value,str):
                encoding="plain"
                if INVALID_XML.search(value):value=json.dumps(value,ensure_ascii=True);encoding="json_string";notes[name]={"encoding":encoding}
                parts=list(self.chunks(value))
                if len(parts)>1:
                    if table=="TextChunks":raise ValueError("Internal chunking overflow")
                    notes[name]={"encoding":encoding,"parts":len(parts),"record_key":key}
                    for n,part in enumerate(parts,1):
                        self.write("TextChunks",{**{f["name"]:None for f in TABLES["TextChunks"]},"contract_version":CONTRACT_VERSION,
                            "project_id":self.project_id,"run_id":self.run_id,"record_key":key,"table_name":table,
                            "field_name":name,"part_index":n,"text_part":part,"encoding":encoding,"excel_notes":{}})
                value=parts[0]
            cells.append(value)
        stored={**row,"excel_notes":notes}
        self.files[table].write(dumps(stored)+"\n")
        self.sheets[table].append(self._cells(self.sheets[table],cells));self.counts[table]+=1

    def finish(self,metadata):
        from openpyxl.utils import get_column_letter
        for f in self.files.values():f.close()
        for table,ws in self.sheets.items():ws.auto_filter.ref=f"A1:{get_column_letter(len(TABLES[table]))}{self.counts[table]+1}"
        self.book.save(self.stage/"dataset.xlsx")
        (self.stage/"contract.json").write_text(json.dumps(schema_document(),ensure_ascii=False,indent=2),encoding="utf-8")
        def file_hash(path):
            h=sha256()
            with path.open("rb") as stream:
                for chunk in iter(lambda:stream.read(65536),b""):h.update(chunk)
            return h.hexdigest()
        manifest={"contract_version":CONTRACT_VERSION,"prd_reference":PRD_REFERENCE,"schema_fingerprint":schema_fingerprint(),
                  "project_id":self.project_id,"run_id":self.run_id,"created_at":now_iso(),"row_counts":self.counts,
                  "files_sha256":{p.name:file_hash(p) for p in self.stage.iterdir() if p.is_file()},**metadata}
        (self.stage/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
        self.stage.rename(self.path);self.closed=True
        return self.path

    def abort(self):
        for f in self.files.values():
            if not f.closed:f.close()
        shutil.rmtree(self.stage,ignore_errors=True)

def add_readme(sink,factory,extra=()):
    items=[("Kontrak ekspor",CONTRACT_VERSION),("Referensi PRD",PRD_REFERENCE+"; ini proyeksi tabel, bukan implementasi lengkap PRD"),
        ("Header fingerprint",schema_fingerprint()),("Tipe sumber","Semua sumber menggunakan 8 sheet dan tipe kolom identik, termasuk sheet kosong."),
        ("Timestamp","Excel: native UTC tanpa timezone; JSONL: ISO8601 UTC. published_at kosong untuk tanggal relatif/unknown."),
        ("Revision","content_revision null sampai ditetapkan repository persisten. Gunakan observation_id + text_hash untuk evidence snapshot."),
        ("Keyword","Literal casefold + boundary; offset Unicode code point pada field sumber. Tidak menentukan sentimen atau peran penulis."),
        ("Peran penulis","unknown; tidak menebak orangtua/calon mahasiswa dari username atau foto."),
        ("Metrik","Metrik jumlah dan indeks tren berbagi tabel Metrics tetapi dibedakan metric_name, unit, scale dan dimensions; jangan dijumlahkan lintas unit."),
        ("Teks panjang","JSONL menjaga teks penuh. Excel menyimpan sel pertama + seluruh potongan di TextChunks; gabungkan part_index, lalu decode bila encoding=json_string."),
        ("Kelengkapan","LIMIT_REACHED/PARTIAL/UNKNOWN tidak menjamin seluruh komentar. Nol karena gagal berbeda dengan nol komentar terkonfirmasi."),
        ("Impor HTML","import-html membaca HTML lokal saja; bukan bukti bahwa scraping live berhasil. Waktu capture awal boleh tetap unknown."),
        ("Supabase","Tidak ada koneksi atau transfer DB dalam versi ini. Data kontrak dapat menjadi input staging; adapter repository dan penomoran revisi belum diimplementasikan."),
        ("Sentimen/emosi","Belum dianalisis; tidak ada label otomatis dari crawler."),*extra]
    for item,value in items:sink.write("Readme",factory.row("Readme",item=item,value=str(value)))


# Reddit-specific acquisition. The shared table descriptor above is unchanged.
import argparse
from dataclasses import asdict, replace
from html.parser import HTMLParser
import importlib.metadata
import multiprocessing as mp
import signal
import socket
import ssl
import sys
import time
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urljoin, urlsplit
from urllib.request import Request, build_opener, HTTPSHandler, HTTPRedirectHandler
import unittest

SOURCE = "reddit"
POST_NAMESPACE = "reddit_post"
COMMENT_NAMESPACE = "reddit_comment"
VERSION = "0.1.0-single-file"
PARSER_VERSION = "reddit-html:1;python-html.parser:" + sys.version.split()[0]
USER_AGENT = "UniversitySocialResearch/0.1 (public HTML reader; no API authentication)"
BASE = "https://www.reddit.com"
FULLNAME = re.compile(r"t([13])_([a-z0-9]{3,16})\Z")

class SourceFailure(Exception):
    def __init__(self, code, reason):
        super().__init__(reason)
        self.code, self.reason = code, reason

def safe_web_url(value):
    """Only public HTML pages on the canonical host; never auth/API routes."""
    if not isinstance(value, str) or len(value)>2048 or any(ord(c)<33 for c in value):
        raise ValueError("URL Reddit tidak valid")
    parsed=urlsplit(value)
    if parsed.scheme!="https" or parsed.hostname not in {"www.reddit.com","reddit.com"} or parsed.username or parsed.password or parsed.port is not None:
        raise ValueError("Gunakan URL HTTPS reddit.com tanpa credentials/port")
    if re.search(r"(?:\.json|\.rss)(?:/|$)",parsed.path,re.I) or re.match(r"/(?:api|svc|oauth|account|login|register)(?:/|$)",parsed.path,re.I):
        raise ValueError("Hanya halaman HTML publik, bukan API/login")
    return parsed

def thread_id(value):
    parsed=safe_web_url(value)
    match=re.fullmatch(r"/(?:r/[A-Za-z0-9_]{2,32}/)?comments/([a-z0-9]{3,16})(?:/[A-Za-z0-9_%\-]{0,240})?(?:/[a-z0-9]{3,16})?/?",parsed.path)
    if not match:raise ValueError("Gunakan permalink thread /r/subreddit/comments/ID/judul/ atau /comments/ID/")
    return match.group(1)

def thread_url(value):
    return BASE+"/comments/"+thread_id(value)+"/"

def native_fullname(value, kind=None):
    match=FULLNAME.fullmatch(str(value or ""))
    if not match or (kind and match.group(1)!=str(kind)):return None
    return match.group(0)

def object_key(fullname):
    name=native_fullname(fullname)
    if not name:raise ValueError("Fullname sumber tidak valid")
    return RecordFactory.key(SOURCE,POST_NAMESPACE if name.startswith("t3_") else COMMENT_NAMESPACE,name[3:])

@dataclass(frozen=True, slots=True)
class Settings:
    project_id:str="university-intelligence"
    keywords:tuple[str,...]=("BINUS","universitas","kampus","biaya kuliah","beasiswa","jurusan")
    subreddits:tuple[str,...]=("indonesia","karir","indotech")
    max_posts:int=3
    search_results:int=5
    max_comments:int=100
    max_replies:int=30
    max_replies_per_thread:int=5
    timeout_seconds:int=45
    total_seconds:int=300
    request_delay:float=2.0
    max_html_bytes:int=8_000_000
    sort:str="new"
    search_sort:str="new"
    since:str|None=None
    until:str|None=None
    captured_at:str|None=None

    def validate(self):
        limits={"max_posts":(1,20),"search_results":(1,50),"max_comments":(1,2000),
            "max_replies":(0,2000),"max_replies_per_thread":(0,100),"timeout_seconds":(5,300),
            "total_seconds":(10,1800),"max_html_bytes":(10000,20_000_000)}
        for name,(low,high) in limits.items():
            v=getattr(self,name)
            if type(v) is not int or not low<=v<=high:raise ValueError(f"{name} wajib integer {low}..{high}")
        if type(self.request_delay) not in (int,float) or not math.isfinite(self.request_delay) or not 1<=self.request_delay<=30:
            raise ValueError("request_delay wajib 1..30 detik")
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}",self.project_id):raise ValueError("project_id wajib slug 1..80 karakter")
        if not 1<=len(self.subreddits)<=10 or any(not re.fullmatch(r"[A-Za-z0-9_]{2,32}",s) for s in self.subreddits):
            raise ValueError("Maksimum 10 nama subreddit, tanpa awalan r/")
        if self.sort not in {"new","top","old","best"} or self.search_sort not in {"new","relevance","top"}:
            raise ValueError("Sort tidak dikenal")
        if not isinstance(self.keywords,(list,tuple)) or len(self.keywords)>200 or any(not isinstance(t,str) or not t.strip() or len(t)>200 for t in self.keywords):
            raise ValueError("Maksimum 200 keyword nonkosong, masing-masing 200 karakter")
        RecordFactory(self.project_id,str(uuid.uuid4()),KeywordMatcher([]),self.since,self.until)
        if self.captured_at:utc_datetime(self.captured_at)
        return self

@dataclass(slots=True)
class HtmlPage:
    html:str
    url:str
    observed_at:str
    status:int|None
    byte_count:int
    html_hash:str
    snapshot:bool=False
    original_capture_known:bool=True

class HtmlSource(ABC):
    @abstractmethod
    def fetch(self,url):raise NotImplementedError

class RestrictedRedirect(HTTPRedirectHandler):
    def redirect_request(self,req,fp,code,msg,headers,newurl):
        joined=urljoin(req.full_url,newurl)
        if re.match(r"https://(?:www\.)?reddit\.com/(?:login|account|register)(?:/|\?|$)",joined,re.I):
            raise SourceFailure("LOGIN_REQUIRED","redirect_to_login")
        try:safe_web_url(joined)
        except ValueError:raise SourceFailure("REDIRECT_REJECTED","redirect_outside_public_html_scope") from None
        return super().redirect_request(req,fp,code,msg,headers,joined)

class PublicHtmlSource(HtmlSource):
    def __init__(self,settings):
        self.settings=settings
        self.opener=build_opener(RestrictedRedirect(),HTTPSHandler(context=ssl.create_default_context()))

    def fetch(self,url):
        safe_web_url(url)
        req=Request(url,headers={"User-Agent":USER_AGENT,"Accept":"text/html","Accept-Encoding":"identity"})
        try:
            with self.opener.open(req,timeout=min(20,self.settings.timeout_seconds)) as response:
                safe_web_url(response.url)
                if response.status!=200:raise SourceFailure("HTTP_ERROR","unexpected_http_status")
                if response.headers.get_content_type() not in {"text/html","application/xhtml+xml"}:
                    raise SourceFailure("NON_HTML_RESPONSE","not_a_public_html_document")
                if response.headers.get("Content-Encoding", "identity").lower()!="identity":
                    raise SourceFailure("UNSUPPORTED_ENCODING","server_ignored_identity_encoding")
                data=response.read(self.settings.max_html_bytes+1)
                if len(data)>self.settings.max_html_bytes:raise SourceFailure("RESPONSE_TOO_LARGE","html_byte_budget_exceeded")
                charset=response.headers.get_content_charset() or "utf-8"
                try:html=data.decode(charset,errors="strict")
                except (LookupError,UnicodeError):raise SourceFailure("DECODE_ERROR","html_charset_or_byte_sequence_invalid") from None
                return HtmlPage(html,response.url,now_iso(),response.status,len(data),sha256(data).hexdigest())
        except HTTPError as exc:
            code={401:"LOGIN_REQUIRED",403:"ACCESS_BLOCKED",404:"NOT_FOUND",429:"RATE_LIMITED"}.get(exc.code,"UPSTREAM_ERROR" if exc.code>=500 else "HTTP_ERROR")
            exc.close()
            raise SourceFailure(code,"http_"+str(exc.code)) from None
        except (ssl.SSLError, ssl.CertificateError):raise SourceFailure("TLS_ERROR","tls_verification_failed") from None
        except URLError as exc:
            code="TLS_ERROR" if isinstance(exc.reason,ssl.SSLError) else "NETWORK_TIMEOUT" if isinstance(exc.reason,(TimeoutError,socket.timeout)) else "NETWORK_ERROR"
            raise SourceFailure(code,"request_failed") from None
        except (TimeoutError,socket.timeout):raise SourceFailure("NETWORK_TIMEOUT","socket_timeout") from None

class LocalHtmlSource(HtmlSource):
    def __init__(self,path,settings):self.path,self.settings=Path(path),settings
    def fetch(self,url):
        safe_web_url(url)
        with self.path.open("rb") as stream:data=stream.read(self.settings.max_html_bytes+1)
        if len(data)>self.settings.max_html_bytes:raise SourceFailure("RESPONSE_TOO_LARGE","html_byte_budget_exceeded")
        try:html=data.decode("utf-8-sig",errors="strict")
        except UnicodeError:raise SourceFailure("DECODE_ERROR","saved_html_must_be_utf8") from None
        stamp=utc_datetime(self.settings.captured_at).isoformat() if self.settings.captured_at else now_iso()
        return HtmlPage(html,url,stamp,None,len(data),sha256(data).hexdigest(),True,bool(self.settings.captured_at))

class Node:
    __slots__=("tag","attrs","parent","children")
    def __init__(self,tag,attrs=(),parent=None):
        self.tag,self.attrs,self.parent,self.children=tag,dict(attrs),parent,[]
    def has_class(self,name):return name in self.attrs.get("class","").split()

INERT={"script","style","template","noscript","svg","button"}
VOID={"area","base","br","col","embed","hr","img","input","link","meta","param","source","track","wbr"}
BLOCKS={"p","div","li","ul","ol","blockquote","pre","h1","h2","h3","h4","table","tr","section"}

def content_node(node):
    return node.tag in {"shreddit-post","shreddit-comment"} or (node.has_class("thing") and (node.has_class("link") or node.has_class("comment")))

def walk(node,*,own=False):
    pending=list(reversed(node.children))
    while pending:
        item=pending.pop()
        if not isinstance(item,Node):continue
        if own and content_node(item):continue
        yield item
        if item.tag not in INERT:pending.extend(reversed(item.children))

def first(node,predicate):return next((n for n in walk(node,own=True) if predicate(n)),None)

def rendered_text(node):
    if node is None:return ""
    pieces=[];pending=[(node,False)]
    while pending:
        item,closing=pending.pop()
        if isinstance(item,str):pieces.append(item);continue
        if closing:
            if item.tag in BLOCKS:pieces.append("\n")
            continue
        if item.tag in INERT or "hidden" in item.attrs or item.attrs.get("aria-hidden")=="true" or content_node(item) and item is not node:continue
        style=re.sub(r"\s+","",item.attrs.get("style","").lower())
        if "display:none" in style or "visibility:hidden" in style:continue
        if item.tag=="br":pieces.append("\n");continue
        if item.tag=="img":pieces.append(item.attrs.get("alt",""));continue
        if item.tag in BLOCKS:pieces.append("\n")
        pending.append((item,True));pending.extend((c,False) for c in reversed(item.children))
    text="".join(pieces)
    text=re.sub(r"[^\S\n]+"," ",text)
    return re.sub(r"\n{3,}","\n\n","\n".join(line.strip() for line in text.splitlines())).strip()

class BoundedDom(HTMLParser):
    """Small serialised-DOM reader, not a browser or full HTML5 repair engine."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root=Node("document");self.stack=[self.root];self.nodes=0
    def handle_starttag(self,tag,attrs):
        self.nodes+=1
        if self.nodes>150000 or len(self.stack)>512:raise SourceFailure("HTML_COMPLEXITY_LIMIT","node_or_depth_budget_exceeded")
        node=Node(tag,attrs,self.stack[-1]);self.stack[-1].children.append(node)
        if tag not in VOID:self.stack.append(node)
    def handle_startendtag(self,tag,attrs):
        self.handle_starttag(tag,attrs)
        if tag not in VOID:self.handle_endtag(tag)
    def handle_endtag(self,tag):
        for i in range(len(self.stack)-1,0,-1):
            if self.stack[i].tag==tag:
                del self.stack[i:];break
    def handle_data(self,data):
        if self.stack[-1].children and isinstance(self.stack[-1].children[-1],str):self.stack[-1].children[-1]+=data
        else:self.stack[-1].children.append(data)

def assert_accessible(root):
    nodes=list(walk(root))
    if any(n.tag=="input" and (n.attrs.get("name") or "").lower() in {"js_challenge","jsc_token"} for n in nodes):
        raise SourceFailure("CHALLENGE_REQUIRED","challenge_form_detected_even_if_http_200")
    # Actual content can quote an error. Only classify plain error pages by text.
    if any(content_node(n) for n in nodes):return
    if any(n.tag=="input" and (n.attrs.get("name") or "").lower() in {"cf-turnstile-response","g-recaptcha-response","h-captcha-response"} for n in nodes):
        raise SourceFailure("CHALLENGE_REQUIRED","captcha_form_without_post_content")
    text=rendered_text(root).casefold()
    if any(x in text for x in ("blocked by network security","request has been blocked","you've been blocked","you’ve been blocked")):
        raise SourceFailure("ACCESS_BLOCKED","network_security_page")
    if any(x in text for x in ("verify you are human","checking your browser","enable javascript and cookies","complete the captcha")):
        raise SourceFailure("CHALLENGE_REQUIRED","browser_challenge_page")
    if any(x in text for x in ("log in to continue","login to continue","you must be logged in","log in to view")):
        raise SourceFailure("LOGIN_REQUIRED","login_wall")

def source_time(node,is_post):
    value=node.attrs.get("created-timestamp" if is_post else "created")
    if not value:
        clock=first(node,lambda n:n.tag=="time" and "edited" not in n.attrs.get("class","") or n.tag=="faceplate-timeago")
        if clock:value=clock.attrs.get("datetime") or clock.attrs.get("ts") or rendered_text(clock)
    if not value:return None,None,"unknown"
    try:
        stamp=utc_datetime(value).isoformat()
        return stamp,value,"exact"
    except (ValueError,TypeError,OverflowError):
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}",value):return None,value,"date_only"
        relative=bool(re.search(r"\b(?:ago|lalu|yesterday|kemarin)\b",value,re.I))
        return None,value,"relative" if relative else "unknown"

def parsed_number(value):
    value=str(value or "").strip()
    return int(value) if re.fullmatch(r"-?\d{1,12}",value) else None

class RedditHtmlParser:
    def __init__(self,page):
        self.page=page;dom=BoundedDom();dom.feed(page.html);dom.close();self.root=dom.root
        assert_accessible(self.root)
        self.nodes=[n for n in walk(self.root) if content_node(n)]

    @staticmethod
    def fullname(node):
        for field in ("thingid","id","data-fullname"):
            name=native_fullname(node.attrs.get(field))
            if name:return name
        return None

    def discover(self,limit):
        seen=set();result=[]
        # Links in recognized search cards, not arbitrary sidebar navigation.
        cards=[n for n in walk(self.root) if n.tag in {"shreddit-post","search-telemetry-tracker"} or n.has_class("search-result-link") or n.has_class("link") and n.has_class("thing")]
        for card in cards:
            urls=[card.attrs.get("permalink"),card.attrs.get("content-href")]
            urls += [n.attrs.get("href") for n in walk(card,own=True) if n.tag=="a"]
            for value in urls:
                if not value:continue
                try:url=thread_url(urljoin(self.page.url,value))
                except ValueError:continue
                if url in seen:continue
                seen.add(url);result.append({"url":url,"discovered_from":self.page.url})
                break
            if len(result)>=limit:break
        return result

    def own_body(self,node,is_post):
        slot="text-body" if is_post else "comment"
        holder=first(node,lambda n:n.attrs.get("slot")==slot)
        if holder:
            return holder if holder.has_class("md") else first(holder,lambda n:n.has_class("md")) or holder
        return first(node,lambda n:n.has_class("md"))

    def parent(self,node,root):
        for field in ("parentid","data-parent-fullname","data-parent-id"):
            name=native_fullname(node.attrs.get(field))
            if name:return name,"source_parent_id"
        ancestor=node.parent
        while ancestor:
            name=self.fullname(ancestor)
            if name and name.startswith("t1_"):return name,"dom_nesting"
            ancestor=ancestor.parent
        if node.attrs.get("depth")=="0":return root,"source_depth_and_thread_context"
        ancestor=node.parent
        while ancestor:
            if ancestor.has_class("commentarea"):return root,"dom_nesting"
            ancestor=ancestor.parent
        return None,None

    def thread(self,expected_id):
        root="t3_"+expected_id
        post=next((n for n in self.nodes if self.fullname(n)==root),None)
        if post is None:raise SourceFailure("LAYOUT_UNSUPPORTED","expected_post_id_not_found_in_html")
        records=[];invalid=0
        for node in [post,*[n for n in self.nodes if n.tag=="shreddit-comment" or n.has_class("comment")]]:
            name=self.fullname(node);is_post=node is post
            if not name or not name.startswith("t3_" if is_post else "t1_"):
                invalid+=1;continue
            link=node.attrs.get("linkid") or node.attrs.get("data-link-id")
            if not is_post and link and link!=root:
                invalid+=1;continue
            body=self.own_body(node,is_post);text=rendered_text(body)
            title=node.attrs.get("post-title") if is_post else None
            if is_post and not title:
                title_node=first(node,lambda n:n.tag=="h1" or n.tag=="a" and n.has_class("title"))
                title=rendered_text(title_node) or None
            if not text and not title:
                invalid+=1;continue
            flags=["rendered_html_text_not_original_markdown","static_html_non_exhaustive"]
            if self.page.snapshot:
                flags.append("saved_html_import_not_independently_verified")
                if not self.page.original_capture_known:flags.append("snapshot_original_capture_time_unknown")
            parent,evidence=(None,None) if is_post else self.parent(node,root)
            if parent and parent.startswith("t3_") and parent!=root:
                invalid+=1;continue
            if not is_post and not parent:flags.append("parent_unknown")
            if evidence=="dom_nesting":flags.append("parent_from_dom_nesting")
            if evidence=="source_depth_and_thread_context":flags.append("parent_from_depth_zero")
            deleted=text.strip().casefold() in {"[deleted]","[removed]"}
            if deleted:flags.append("deleted_or_removed_placeholder")
            wrappers=[];ancestor=body
            while ancestor is not None and ancestor is not node:
                wrappers.append(ancestor);ancestor=ancestor.parent
            truncated=any("line-clamp" in n.attrs.get("class","") or n.attrs.get("content-truncated")=="true" for n in wrappers+list(walk(body))) if body else True
            complete=body is not None and not truncated and not deleted
            if truncated:flags.append("visible_body_may_be_truncated_or_missing")
            published,display,quality=source_time(node,is_post)
            score=parsed_number(node.attrs.get("score") or node.attrs.get("data-score"))
            if score is None:
                score_node=first(node,lambda n:n.has_class("score") and n.attrs.get("title") is not None)
                if score_node:score=parsed_number(score_node.attrs.get("title"))
            count=parsed_number(node.attrs.get("comment-count")) if is_post else None
            author=node.attrs.get("author-id") or node.attrs.get("data-author-fullname")
            if not re.fullmatch(r"t2_[a-z0-9]{3,20}",author or ""):author=None
            kind="post" if is_post else "reply" if parent and parent.startswith("t1_") or not parent and (parsed_number(node.attrs.get("depth")) or 0)>0 else "comment"
            records.append({"fullname":name,"kind":kind,"title":title,"text":text,"parent":parent,
                "parent_evidence":evidence,"root":root,"published_at":published,"published_display":display,
                "timestamp_quality":quality,"text_complete":complete,"flags":flags,"author_id":author,
                "score":score,"comment_count":count,"subreddit":post.attrs.get("subreddit-name"),
                "publisher_ref":post.attrs.get("subreddit-id"),"observed_at":self.page.observed_at,"page_url":self.page.url,
                "extraction_method":"saved_html_import" if self.page.snapshot else "public_html_scrape"})
        return records,invalid

def bounded_events(records,settings):
    """Whole-post comment budget includes replies; nested replies share thread cap."""
    by_id={r["fullname"]:r for r in records};seen=set();threads={};total=0;replies=0;limited=False
    for record in records:
        if record["fullname"] in seen:continue
        seen.add(record["fullname"])
        if record["kind"]=="post":yield "record",record;continue
        if total>=settings.max_comments:limited=True;continue
        if record["kind"]=="reply":
            parent=record["parent"];branch=parent;visited=set()
            while parent in by_id and parent not in visited:
                visited.add(parent);candidate=by_id[parent]
                if candidate["kind"]!="reply":branch=parent;break
                parent=candidate["parent"]
                if parent:branch=parent
            if replies>=settings.max_replies or threads.get(branch,0)>=settings.max_replies_per_thread:
                limited=True;continue
            replies+=1;threads[branch]=threads.get(branch,0)+1
        total+=1;yield "record",record
    yield "budget",{"limited":limited,"comments":total-replies,"replies":replies}

def worker(operation,target,settings_dict,pipe):
    if os.name=="posix":os.setsid()
    try:
        settings=Settings(**settings_dict)
        source=LocalHtmlSource(target["input"],settings) if operation=="import_html" else PublicHtmlSource(settings)
        url=target["url"] if isinstance(target,dict) else target
        page=source.fetch(url)
        pipe.send(("evidence",{"url":page.url,"http_status":page.status,"html_bytes":page.byte_count,
            "html_sha256":page.html_hash,"observed_at":page.observed_at,"snapshot":page.snapshot,
            "original_capture_known":page.original_capture_known}))
        parsed=RedditHtmlParser(page)
        if operation=="search":
            candidates=parsed.discover(settings.search_results)
            for candidate in candidates:pipe.send(("candidate",candidate))
            result={"status":"FINISHED_BEST_EFFORT" if candidates else "NO_RESULTS_OR_LAYOUT_UNKNOWN",
                "reason":"bounded_html_post_search" if candidates else "no_recognized_search_cards_not_proof_of_no_results",
                "warnings":["post_search_not_global_comment_search","one_html_page_only"],"invalid":0}
        else:
            records,invalid=parsed.thread(thread_id(url));budget=None
            for kind,payload in bounded_events(records,settings):
                if kind=="record":pipe.send((kind,payload))
                else:budget=payload
            count=budget["comments"]+budget["replies"]
            status="LIMIT_REACHED" if budget["limited"] else "FINISHED_BEST_EFFORT" if count else "NO_COMMENTS_EXTRACTED"
            result={"status":status,"reason":"visible_html_only_not_entire_thread" if count else "post_context_found_but_no_static_comment_text",
                "warnings":["static_html_only","javascript_and_more_replies_not_loaded"],"invalid":invalid}
        result["source_version"]=PARSER_VERSION
        pipe.send(("done",result))
    except SourceFailure as exc:
        pipe.send(("done",{"status":exc.code,"reason":exc.reason,"warnings":[],"invalid":0,"source_version":PARSER_VERSION}))
    except (FileNotFoundError,PermissionError):
        pipe.send(("done",{"status":"INPUT_FILE_ERROR","reason":"local_html_file_unavailable","warnings":[],"invalid":0,"source_version":PARSER_VERSION}))
    except Exception:
        pipe.send(("done",{"status":"EXTRACTION_ERROR","reason":"unexpected_parser_or_transport_error","warnings":[],"invalid":0,"source_version":PARSER_VERSION}))
    finally:pipe.close()

def run_worker(operation,target,settings,seconds,on_event,*,worker_entry=worker):
    context=mp.get_context("spawn");parent,child=context.Pipe(duplex=False)
    proc=context.Process(target=worker_entry,args=(operation,target,asdict(settings),child));proc.start();child.close()
    deadline=time.monotonic()+seconds;result=None
    try:
        while time.monotonic()<deadline:
            if parent.poll(min(0.2,max(0,deadline-time.monotonic()))):
                try:kind,payload=parent.recv()
                except EOFError:break
                if kind=="done":result=payload;break
                on_event(kind,payload)
            elif not proc.is_alive():
                if parent.poll():continue
                break
        return result or {"status":"TIMEOUT" if proc.is_alive() else "WORKER_FAILED",
            "reason":"hard_deadline_partial_events_preserved" if proc.is_alive() else "worker_exited_without_receipt",
            "warnings":[],"invalid":0,"source_version":PARSER_VERSION}
    finally:
        parent.close();proc.join(timeout=0.3)
        if proc.is_alive():
            if os.name=="posix":
                try:
                    if os.getpgid(proc.pid)==proc.pid:os.killpg(proc.pid,signal.SIGTERM)
                    else:proc.terminate()
                except ProcessLookupError:pass
            else:proc.terminate()
            proc.join(timeout=0.5)
            if proc.is_alive():proc.kill();proc.join(timeout=0.5)

class RedditRecordFactory(RecordFactory):
    def content(self,value):
        row,hits=super().content(value)
        if "parent_from_dom_nesting" in value.quality_flags:row["relation_evidence"]="dom_nesting"
        if "parent_from_depth_zero" in value.quality_flags:row["relation_evidence"]="source_depth_and_thread_context"
        return row,hits

class RedditAdapter(SourceAdapter):
    def __init__(self,factory,dataset):self.factory,self.dataset=factory,dataset
    def accept(self,kind,raw,url,discovered_from):
        if kind!="record":return
        native=native_fullname(raw["fullname"])
        root=native_fullname(raw["root"],3)
        if not native or not root or root[3:]!=thread_id(url):raise ValueError("Identitas post/komentar tidak cocok dengan thread")
        parent=native_fullname(raw.get("parent"))
        canonical=BASE+"/comments/"+root[3:]+"/"+("_/"+native[3:]+"/" if native.startswith("t1_") else "")
        publisher=raw.get("publisher_ref")
        if not re.fullmatch(r"t5_[a-z0-9]{3,20}",publisher or ""):publisher=None
        value=ContentInput(platform=SOURCE,namespace=POST_NAMESPACE if native.startswith("t3_") else COMMENT_NAMESPACE,
            native_id=native[3:],object_type=raw["kind"],canonical_url=canonical,page_url=raw.get("page_url") or url,thread_root_key=object_key(root),
            text_source=raw["text"],title=raw.get("title"),parent_key=object_key(parent) if parent else None,
            published_at=raw["published_at"],published_display=raw["published_display"],timestamp_quality=raw["timestamp_quality"],
            text_complete=raw["text_complete"],author_native_id=raw["author_id"],publisher_ref="reddit:subreddit:"+publisher if publisher else None,
            discovered_from=discovered_from,observed_at=raw["observed_at"],quality_flags=tuple(raw["flags"]),
            extraction_method=raw["extraction_method"],adapter_version="reddit-adapter:"+VERSION+";"+PARSER_VERSION)
        if not self.dataset.add_content(value):return
        observation=self.dataset.last_record
        for name,number,unit in (("reddit_score",raw["score"],"score"),("comment_count",raw["comment_count"],"count")):
            if number is None:continue
            self.dataset.sink.write("Metrics",self.factory.row("Metrics",metric_id=digest(dumps([observation["observation_id"],name])),
                observation_id=observation["observation_id"],platform=SOURCE,source_namespace=value.namespace,object_key=observation["object_key"],
                series_key=None,metric_name=name,value=number,value_display=str(number),unit=unit,scale_min=None,scale_max=None,
                period_start=None,period_end=None,dimensions={"subreddit":raw.get("subreddit")},precision="unknown",observed_at=value.observed_at,
                source_url=value.canonical_url,quality_flags=["displayed_score_not_upvote_count" if name=="reddit_score" else "displayed_count_not_coverage_guarantee"]))

def load_settings(args):
    defaults=Settings();values={}
    for name in defaults.__dataclass_fields__:
        supplied=getattr(args,name,None)
        if supplied is not None:values[name]=supplied
    if getattr(args,"keyword",None):values["keywords"]=tuple(args.keyword)
    if getattr(args,"subreddit",None):values["subreddits"]=tuple(dict.fromkeys(args.subreddit))
    settings=replace(defaults,**values).validate()
    urls=list(getattr(args,"url",None) or [])
    if getattr(args,"urls_file",None):
        if args.urls_file.stat().st_size>1_000_000:raise ValueError("urls-file maksimum 1 MB")
        urls.extend(line.strip() for line in args.urls_file.read_text(encoding="utf-8").splitlines() if line.strip() and not line.lstrip().startswith("#"))
    if len(urls)>100:raise ValueError("Maksimum 100 URL input")
    urls=list(dict.fromkeys(thread_url(u) for u in urls))
    queries=list(dict.fromkeys(getattr(args,"query",None) or []))
    if len(queries)>10 or any(not isinstance(q,str) or not q.strip() or len(q)>300 for q in queries):raise ValueError("Maksimum 10 query, masing-masing 1..300 karakter")
    if not urls and not queries:raise ValueError("Berikan --url atau --query")
    if args.command=="import-html" and (len(urls)!=1 or not args.input.is_file()):raise ValueError("import-html memerlukan satu URL dan file HTML yang ada")
    return settings,urls,queries

def run(args):
    verify_descriptor();settings,urls,queries=load_settings(args)
    run_id=str(uuid.uuid4());factory=RedditRecordFactory(settings.project_id,run_id,KeywordMatcher(settings.keywords),settings.since,settings.until)
    output=args.output_dir or Path("output")/("reddit_"+datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")+"_"+run_id[:8])
    sink=BundleSink(output,settings.project_id,run_id);dataset=Dataset(factory,sink);adapter=RedditAdapter(factory,dataset)
    deadline=time.monotonic()+settings.total_seconds;failures=0;interrupted=False;barrier=None;evidence=[];last_request=None
    offline=args.command=="import-html"

    def execute(operation,target,callback,skip=None):
        nonlocal failures,barrier,last_request
        started=now_iso();before=dataset.contents;dupes=dataset.duplicates;counts={"comment":0,"reply":0,"invalid":0}
        def receive(kind,payload):
            if kind=="evidence":evidence.append(payload);return
            if kind=="record" and payload.get("kind") in {"comment","reply"}:counts[payload["kind"]]+=1
            try:callback(kind,payload)
            except (ValueError,TypeError,OverflowError,KeyError):counts["invalid"]+=1
        if skip or barrier:
            result={"status":skip or "SKIPPED_ACCESS_BARRIER","reason":"post_budget_full" if skip else "previous_"+barrier,
                "warnings":[],"invalid":0,"source_version":PARSER_VERSION}
        elif time.monotonic()>=deadline:
            result={"status":"SKIPPED_GLOBAL_DEADLINE","reason":"global_time_budget_exhausted","warnings":[],"invalid":0,"source_version":PARSER_VERSION}
        else:
            if not offline and last_request is not None:
                delay=min(max(0,settings.request_delay-(time.monotonic()-last_request)),max(0,deadline-time.monotonic()))
                if delay:time.sleep(delay)
            remaining=deadline-time.monotonic()
            if remaining<=0:
                result={"status":"SKIPPED_GLOBAL_DEADLINE","reason":"global_time_budget_exhausted","warnings":[],"invalid":0,"source_version":PARSER_VERSION}
            else:
                print(operation+": "+(target["url"] if isinstance(target,dict) else target),flush=True)
                try:result=run_worker(operation,target,settings,min(settings.timeout_seconds,remaining),receive)
                except KeyboardInterrupt:result={"status":"INTERRUPTED","reason":"user_interrupted_partial_events_preserved","warnings":[],"invalid":0,"source_version":PARSER_VERSION}
                last_request=time.monotonic()
        counts["invalid"]+=result.get("invalid",0)
        if result["status"] in {"ACCESS_BLOCKED","CHALLENGE_REQUIRED","RATE_LIMITED","LOGIN_REQUIRED"}:barrier=result["status"]
        if counts["invalid"]:result={**result,"status":"PARTIAL_SCHEMA_ERROR","reason":"some_source_nodes_or_rows_rejected","warnings":[*result["warnings"],"invalid_source_rows"]}
        status=result["status"]
        if status not in {"FINISHED_BEST_EFFORT","LIMIT_REACHED","SKIPPED_POST_LIMIT"}:failures+=1
        row=factory.row("Runs",unit_id=str(uuid.uuid4()),source=SOURCE,operation=operation,target=target["url"] if isinstance(target,dict) else target,
            status=status,reason=result["reason"],coverage="NOT_RUN" if status.startswith("SKIPPED_") else "PARTIAL_HTML_VISIBLE_ONLY" if dataset.contents>before or status=="FINISHED_BEST_EFFORT" else "UNKNOWN",
            contents_written=dataset.contents-before,comments_seen=counts["comment"],replies_seen=counts["reply"],duplicates_skipped=dataset.duplicates-dupes,
            invalid_rows=counts["invalid"],max_comments=settings.max_comments if operation!="search" else None,
            max_replies=settings.max_replies if operation!="search" else None,timeout_seconds=settings.timeout_seconds,keywords=list(settings.keywords),
            source_version=PARSER_VERSION,warning_codes=result["warnings"],started_at=started,finished_at=now_iso())
        sink.write("Runs",row);print(f"  {status}; {row['contents_written']} content rows",flush=True)
        if status=="INTERRUPTED":raise KeyboardInterrupt

    try:
        targets={u:None for u in urls[:settings.max_posts]}
        for query in queries:
            for subreddit in settings.subreddits:
                search_url=BASE+"/r/"+subreddit+"/search/?"+urlencode({"q":query,"restrict_sr":"on","type":"link","sort":settings.search_sort,"t":"all"})
                def candidate(kind,payload):
                    if kind=="candidate" and len(targets)<settings.max_posts:targets.setdefault(thread_url(payload["url"]),payload["discovered_from"])
                execute("search",search_url,candidate,"SKIPPED_POST_LIMIT" if len(targets)>=settings.max_posts else None)
        for url in urls[settings.max_posts:]:execute("thread",url,lambda *_:None,"SKIPPED_POST_LIMIT")
        for url,origin in targets.items():
            target={"url":url,"input":str(args.input.resolve())} if offline else url+"?"+urlencode({"sort":settings.sort})
            execute("import_html" if offline else "thread",target,lambda kind,payload,u=url,o=origin:adapter.accept(kind,payload,u,o))
    except KeyboardInterrupt:interrupted=True;failures+=1
    except BaseException:sink.abort();raise
    try:
        dataset.finish_relations()
        sink.write("Sources",factory.row("Sources",source=SOURCE,producer_version=VERSION,upstream_version=PARSER_VERSION,
            extraction_method="saved_html_import" if offline else "public_html_scrape",access_mode="local_file_no_network" if offline else "guest_no_cookies",
            config={**asdict(settings),"urls":urls,"queries":queries,"html_evidence":evidence},schema_fingerprint=schema_fingerprint(),observed_at=now_iso()))
        add_readme(sink,factory,[("Reddit","HTML publik saja, tanpa .json/OAuth/API key/cookies. Template inert dan konten hasil JS tidak diekstrak."),
            ("Cakupan","Satu halaman HTML per pencarian/thread. Tidak mengklik more replies. Hasil bukan seluruh komentar atau representasi semua pengguna Reddit."),
            ("Teks sumber","text_source adalah teks hasil render HTML, bukan Markdown mentah. Offset keyword mengacu pada teks yang tersimpan ini."),
            ("Waktu","Hanya atribut datetime ISO dengan timezone masuk exact. Tanggal/teks relatif tetap null pada published_at."),
            ("Relasi","t3_ menandai post, t1_ komentar; ID namespace berbeda. Parent dari atribut, DOM nesting, atau depth=0 dibedakan evidence-nya. Parent yang tidak terambil tetap unresolved."),
            ("Score","reddit_score bukan jumlah upvote atau like. score dapat negatif; precision unknown."),
            ("Akses","HTTP 200 tidak membuktikan data tersedia. Barrier menghentikan request berikutnya; Runs membedakan challenge, akses ditolak, layout tak dikenal dan tidak ada komentar terekstrak."),
            ("Mode", "Impor HTML lokal; konten tidak diverifikasi independen. observed_at=captured-at jika diberikan, selain itu waktu impor dengan flag capture awal unknown." if offline else "Pengambilan HTML publik langsung; tidak ada data sintetis yang ditambahkan."),
            ("Hasil run",f"contents={dataset.contents}, duplicates={dataset.duplicates}, failed_units={failures}, interrupted={interrupted}")])
        sink.finish({"source":SOURCE,"producer_version":VERSION,"mode":"saved_html_import" if offline else "live_guest_html",
            "synthetic":None if offline else False,"failed_units":failures,"interrupted":interrupted,
            "limitations":["static_html_visible_only","search_posts_not_comments","no_js_or_more_replies","unknown_total_coverage","rendered_text_not_markdown","no_actor_inference","no_sentiment_labels","no_supabase_transfer"]})
    except BaseException:sink.abort();raise
    print("Output: "+str(sink.path),flush=True)
    return 130 if interrupted else 2 if failures else 0

class MemorySink:
    def __init__(self):self.rows={t:[] for t in TABLES}
    def write(self,t,row):validate_row(t,row);self.rows[t].append(row)


class ContractTests(unittest.TestCase):
    def setUp(self):
            self.factory=RecordFactory("test-project",str(uuid.uuid4()),KeywordMatcher(["BINUS","biaya kuliah","STRASSE"]))
            self.sink=MemorySink();self.data=Dataset(self.factory,self.sink)
            self.raw=ContentInput(platform="reddit",namespace="reddit_comment",native_id="fixture-1",object_type="comment",
                canonical_url="https://example.org/fixture#1",page_url="https://example.org/fixture",thread_root_key="reddit:reddit_post:fixture",
                text_source="BINUS biaya kuliah berapa?",parent_key="reddit:reddit_post:fixture",quality_flags=("synthetic_fixture",))

    def test_01_frozen_header_fingerprint(self):
            verify_descriptor()
            self.assertEqual(schema_fingerprint(),"c6dcff73113f3aba31c2329296cd72b1c24ce6e15db29e22a9e35fa662fc74de")
            self.assertEqual(list(TABLES),["Contents","Relations","KeywordHits","Metrics","Runs","Sources","TextChunks","Readme"])

    def test_02_cross_source_identical_columns_and_types(self):
            a,_=self.factory.content(self.raw)
            b,_=self.factory.content(replace(self.raw,platform="youtube",namespace="youtube_comment"))
            self.assertEqual(list(a),list(b))
            for field in TABLES["Contents"]:
                av,bv=a[field["name"]],b[field["name"]]
                self.assertEqual(type(av),type(bv))
            self.assertIsNone(a["content_revision"])

    def test_03_unknown_columns_rejected(self):
            a,_=self.factory.content(self.raw)
            with self.assertRaises(ValueError):validate_row("Contents",{**a,"source_only_column":1})
            del a["text_hash"]
            with self.assertRaises(ValueError):validate_row("Contents",a)

    def test_04_unicode_offsets_boundary_and_casefold(self):
            text="😀 BINUS bukan BINUSIAN; Straße biaya kuliah"
            hits=list(self.factory.matcher.find(text))
            self.assertEqual(len(hits),3)
            for h in hits:self.assertEqual(text[h["span_start"]:h["span_end"]],h["matched_variant"])
            self.assertEqual(list(KeywordMatcher(["s"]).find("ß")),[])

    def test_05_identical_text_different_ids_is_not_dedup(self):
            self.assertTrue(self.data.add_content(self.raw))
            self.assertFalse(self.data.add_content(self.raw))
            self.assertTrue(self.data.add_content(replace(self.raw,native_id="fixture-2")))
            self.assertEqual(self.data.contents,2);self.assertEqual(self.data.duplicates,1)

    def test_06_changed_text_keeps_distinct_observation(self):
            self.data.add_content(self.raw)
            self.data.add_content(replace(self.raw,text_source="BINUS tidak mahal"))
            records=self.sink.rows["Contents"]
            self.assertEqual(records[0]["object_key"],records[1]["object_key"])
            self.assertNotEqual(records[0]["observation_id"],records[1]["observation_id"])

    def test_07_child_before_parent_and_unresolved(self):
            parent_key="reddit:reddit_comment:parent"
            self.data.add_content(replace(self.raw,object_type="reply",parent_key=parent_key))
            self.data.add_content(replace(self.raw,native_id="parent",parent_key=None))
            self.data.finish_relations()
            relations=self.sink.rows["Relations"]
            reply=next(r for r in relations if r["predicate"]=="reply_to")
            self.assertEqual(reply["resolution_status"],"resolved_in_export")
            self.assertTrue(any(r["resolution_status"]=="unresolved_reference" for r in relations))

    def test_09_no_zero_fill_missing_metrics(self):
            self.data.add_content(self.raw,[("likes",None,"count"),("likes",False,"count"),("likes",0,"count")])
            self.assertEqual(len(self.sink.rows["Metrics"]),1)
            self.assertEqual(self.sink.rows["Metrics"][0]["value"],0)

    def test_10_exact_time_normalized_utc_and_role_evidence(self):
            row,_=self.factory.content(replace(self.raw,published_at="2026-10-25T20:00:00+07:00",timestamp_quality="exact"))
            self.assertEqual(row["published_at"],"2026-10-25T13:00:00+00:00")
            with self.assertRaises(ValueError):validate_row("Contents",{**row,"actor_type":"parent"})
            with self.assertRaises(ValueError):validate_row("Contents",{**row,"timestamp_quality":"relative"})

    def test_13_excel_literal_ids_and_lossless_text(self):
            from openpyxl import load_workbook
            with tempfile.TemporaryDirectory() as td:
                sink=BundleSink(Path(td)/"export",self.factory.project_id,self.factory.run_id)
                ds=Dataset(self.factory,sink)
                huge="=FORMULA 😀\x01"*5000
                ds.add_content(replace(self.raw,native_id="12345678901234567890",text_source=huge))
                ds.add_content(replace(self.raw,native_id="fixture-formula",text_source="=HYPERLINK(\"https://example.org\")"))
                ds.finish_relations();sink.finish({"synthetic":True,"mode":"unit_test"})
                w=load_workbook(sink.path/"dataset.xlsx",read_only=True,data_only=False)
                rows=list(w["Contents"].values);headers=rows[0];rec=dict(zip(headers,rows[1]));formula=dict(zip(headers,rows[2]))
                self.assertEqual(rec["native_id"],"12345678901234567890")
                self.assertEqual(formula["text_source"],"=HYPERLINK(\"https://example.org\")")
                self.assertFalse(any(c.data_type=="f" for ws in w for row in ws for c in row))
                chunkrows=list(w["TextChunks"].values);chunks=[dict(zip(chunkrows[0],r)) for r in chunkrows[1:]]
                parts=sorted((c for c in chunks if c["field_name"]=="text_source"),key=lambda c:c["part_index"])
                self.assertEqual(json.loads("".join(c["text_part"] for c in parts)),huge)
                raw=json.loads((sink.path/"Contents.jsonl").read_text(encoding="utf-8").splitlines()[0])
                self.assertEqual(raw["text_source"],huge)
                for name in TABLES:self.assertEqual(list(next(w[name].values)),[s["name"] for s in TABLES[name]])
                w.close()

    def test_16_trend_index_uses_metrics_without_fake_comment(self):
            row=self.factory.row("Metrics",metric_id="fixture-metric",observation_id=None,platform="google_trends",source_namespace="trends_index",
                object_key=None,series_key="fixture-series",metric_name="search_interest_index",value=75,value_display="75",unit="index",
                scale_min=0,scale_max=100,period_start="2026-10-01T00:00:00+00:00",period_end="2026-10-08T00:00:00+00:00",
                dimensions={"query":"BINUS","geo":"ID"},precision="rounded",observed_at="2026-10-25T00:00:00+00:00",
                source_url="https://example.org/synthetic",quality_flags=["synthetic_fixture"])
            self.sink.write("Metrics",row)
            self.assertEqual(len(self.sink.rows["Contents"]),0)
            self.assertEqual(self.sink.rows["Metrics"][0]["unit"],"index")


# Synthetic regression fixtures only. None is injected into a real run.
FIXTURE_URL=BASE+"/comments/pst123/"
SYNTHETIC_HTML='''<!doctype html><html><body><main>
<shreddit-post id="t3_pst123" post-title="Biaya kuliah BINUS?" post-type="text"
  created-timestamp="2026-10-26T15:00:00+0700" score="12" comment-count="9"
  author="not_exported" author-id="t2_usr123" subreddit-name="indonesia" subreddit-id="t5_sub123">
  <h1 slot="title">Biaya kuliah BINUS?</h1>
  <div slot="text-body"><div class="md"><p>Bandingkan <strong>BINUS</strong> dan kampus lain.</p><p>Bagaimana beasiswa?</p></div></div>
</shreddit-post>
<shreddit-comment-tree post-id="t3_pst123">
  <shreddit-comment thingid="t1_rpl123" parentid="t1_com123" depth="1" author="not_exported">
    <faceplate-timeago ts="1 hour ago"></faceplate-timeago>
    <div slot="comment" class="md"><p>Setuju, cek beasiswa.</p></div>
  </shreddit-comment>
  <shreddit-comment thingid="t1_com123" parentid="t3_pst123" depth="0" created="2026-10-26T08:10:00+00:00" score="-2">
    <div slot="comment"><div class="md"><p>😀 Biaya <b>BINUS</b> perlu dibandingkan.</p></div></div>
    <button>Reply share vote</button>
  </shreddit-comment>
  <shreddit-comment thingid="t1_com456" parentid="t3_pst123" depth="0" created="2026-10-26">
    <div slot="comment" class="md"><p>[removed]</p></div>
  </shreddit-comment>
</shreddit-comment-tree>
<template><shreddit-comment thingid="t1_notreal"><div slot="comment">Template placeholder</div></shreddit-comment></template>
</main></body></html>'''

OLD_SYNTHETIC_HTML='''<html><body>
<div class="thing link self" data-fullname="t3_pst123">
  <div class="entry"><p class="title"><a class="title">Pilihan kampus</a></p><time datetime="2026-10-26T12:00:00Z">kemarin</time>
  <div class="usertext-body"><div class="md"><p>Pertanyaan biaya kuliah</p></div></div></div>
</div>
<div class="commentarea"><div class="sitetable">
<div class="thing comment" data-fullname="t1_aaa123"><div class="entry">
  <time datetime="2026-10-26T12:01:00Z">kemarin</time><time class="edited-timestamp" datetime="2026-10-27T00:00:00Z">edited</time>
  <div class="usertext-body"><div class="md"><p>Orangtua perlu rincian biaya.</p></div></div></div>
  <div class="child"><div class="thing comment" data-fullname="t1_bbb123"><div class="entry">
    <div class="md"><p>Balasan terpisah.</p></div></div></div></div>
</div></div></div></body></html>'''

def fixture_page(html=SYNTHETIC_HTML):
    return HtmlPage(html,FIXTURE_URL,"2026-10-27T00:00:00+00:00",200,len(html.encode()),digest(html),True,False)

def timeout_fixture(operation,target,settings,pipe):
    if os.name=="posix":os.setsid()
    pipe.send(("record",{"fixture":"SYNTHETIC"}));time.sleep(10)

class RedditTests(unittest.TestCase):
    def test_01_shreddit_extracts_body_not_buttons_or_templates(self):
        records,invalid=RedditHtmlParser(fixture_page()).thread("pst123")
        self.assertEqual(invalid,0);self.assertEqual(len(records),4)
        self.assertEqual(records[0]["title"],"Biaya kuliah BINUS?")
        self.assertEqual(records[0]["text"],"Bandingkan BINUS dan kampus lain.\n\nBagaimana beasiswa?")
        self.assertNotIn("Reply share",records[2]["text"])
        self.assertEqual(records[1]["kind"],"reply")

    def test_02_old_nested_comments_do_not_merge_text_or_guess_role(self):
        records,invalid=RedditHtmlParser(fixture_page(OLD_SYNTHETIC_HTML)).thread("pst123")
        self.assertEqual(invalid,0)
        self.assertEqual(records[1]["text"],"Orangtua perlu rincian biaya.")
        self.assertEqual(records[2]["text"],"Balasan terpisah.")
        self.assertEqual(records[2]["parent"],"t1_aaa123")
        self.assertEqual(records[2]["parent_evidence"],"dom_nesting")
        self.assertEqual(records[1]["parent"],"t3_pst123")
        self.assertEqual(records[1]["published_at"],"2026-10-26T12:01:00+00:00")

    def test_03_relative_date_only_and_exact_are_distinct(self):
        records,_=RedditHtmlParser(fixture_page()).thread("pst123")
        self.assertEqual(records[0]["published_at"],"2026-10-26T08:00:00+00:00")
        self.assertEqual(records[1]["timestamp_quality"],"relative")
        self.assertIsNone(records[1]["published_at"])
        self.assertEqual(records[3]["timestamp_quality"],"date_only")
        self.assertIsNone(records[3]["published_at"])
        page=fixture_page(SYNTHETIC_HTML.replace("2026-10-26T08:10:00+00:00","2026-10-26T08:10:00"))
        naive=RedditHtmlParser(page).thread("pst123")[0][2]
        self.assertIsNone(naive["published_at"])

    def test_04_deleted_and_truncated_are_not_full_original_text(self):
        records,_=RedditHtmlParser(fixture_page()).thread("pst123")
        self.assertFalse(records[3]["text_complete"])
        self.assertIn("deleted_or_removed_placeholder",records[3]["flags"])
        changed=SYNTHETIC_HTML.replace('slot="text-body"','slot="text-body" class="line-clamp-3"')
        self.assertFalse(RedditHtmlParser(fixture_page(changed)).thread("pst123")[0][0]["text_complete"])

    def test_05_http_200_javascript_challenge_not_empty_comments(self):
        html='<html><script>document.forms[0].requestSubmit()</script><form><input name="js_challenge" value="1"><input name="solution"></form></html>'
        with self.assertRaises(SourceFailure) as error:RedditHtmlParser(fixture_page(html))
        self.assertEqual(error.exception.code,"CHALLENGE_REQUIRED")

    def test_06_block_login_and_legitimate_quoted_error(self):
        for html,code in (("<main>You've been blocked by network security.</main>","ACCESS_BLOCKED"),("<main>Log in to continue</main>","LOGIN_REQUIRED")):
            with self.assertRaises(SourceFailure) as error:RedditHtmlParser(fixture_page(html))
            self.assertEqual(error.exception.code,code)
        html=SYNTHETIC_HTML.replace('Setuju, cek beasiswa.',"You've been blocked by network security.")
        self.assertEqual(len(RedditHtmlParser(fixture_page(html)).thread("pst123")[0]),4)

    def test_07_wrong_thread_is_rejected_not_relabelled(self):
        with self.assertRaises(SourceFailure):RedditHtmlParser(fixture_page()).thread("xxx999")

    def test_08_urls_reject_credentials_api_and_external_hosts(self):
        for value in ("http://www.reddit.com/comments/pst123/","https://reddit.com.evil.test/comments/pst123/",
                      "https://u:secret@www.reddit.com/comments/pst123/","https://www.reddit.com:443/comments/pst123/",
                      "https://www.reddit.com/comments/pst123.json","https://www.reddit.com/user/a/","https://www.reddit.com/api/morechildren"):
            with self.assertRaises(ValueError):thread_url(value)
        self.assertEqual(thread_url("https://reddit.com/r/karir/comments/pst123/judul/com123/?context=3"),FIXTURE_URL)
        self.assertEqual(thread_url(FIXTURE_URL+"?tracking=1"),FIXTURE_URL)

    def test_09_search_uses_recognized_cards_not_sidebar_links(self):
        html='<aside><a href="/comments/sidebar1/">Sidebar</a></aside><search-telemetry-tracker><a href="/r/karir/comments/pst123/judul/">BINUS</a></search-telemetry-tracker><shreddit-post id="t3_pst999" permalink="/r/karir/comments/pst999/other/"></shreddit-post>'
        found=RedditHtmlParser(fixture_page(html)).discover(1)
        self.assertEqual([r["url"] for r in found],[FIXTURE_URL])

    def test_10_reply_cap_includes_nested_replies_and_preserves_roots(self):
        records,_=RedditHtmlParser(fixture_page()).thread("pst123")
        deep={**records[1],"fullname":"t1_deep123","parent":"t1_rpl123"}
        events=list(bounded_events([*records,deep],Settings(max_comments=3,max_replies=10,max_replies_per_thread=1)))
        self.assertEqual(events[-1],("budget",{"limited":True,"comments":2,"replies":1}))
        names=[row["fullname"] for kind,row in events if kind=="record"]
        self.assertNotIn("t1_deep123",names)
        zero=list(bounded_events(records,Settings(max_replies=0)))
        self.assertEqual(zero[-1][1]["replies"],0);self.assertEqual(zero[-1][1]["comments"],2)

    def test_11_missing_parent_remains_unknown(self):
        html=SYNTHETIC_HTML.replace('parentid="t1_com123" ','')
        raw=RedditHtmlParser(fixture_page(html)).thread("pst123")[0][1]
        self.assertIsNone(raw["parent"]);self.assertEqual(raw["kind"],"reply")
        self.assertIn("parent_unknown",raw["flags"])

    def test_12_relations_child_before_parent_and_score_semantics(self):
        factory=RedditRecordFactory("test",str(uuid.uuid4()),KeywordMatcher(["BINUS"]))
        sink=MemorySink();data=Dataset(factory,sink);adapter=RedditAdapter(factory,data)
        for raw in RedditHtmlParser(fixture_page()).thread("pst123")[0]:adapter.accept("record",raw,FIXTURE_URL,None)
        data.finish_relations()
        relation=next(r for r in sink.rows["Relations"] if r["subject_key"].endswith(":rpl123") and r["predicate"]=="reply_to")
        self.assertEqual(relation["object_key"],"reddit:reddit_comment:com123")
        self.assertEqual(relation["resolution_status"],"resolved_in_export")
        score=next(r for r in sink.rows["Metrics"] if r["object_key"].endswith(":com123"))
        self.assertEqual(score["value"],-2);self.assertEqual(score["unit"],"score")
        self.assertTrue(all(r["actor_type"]=="unknown" for r in sink.rows["Contents"]))
        self.assertTrue(all(r["content_revision"] is None for r in sink.rows["Contents"]))
        self.assertIsNone(sink.rows["Contents"][1]["author_ref"])

    def test_13_dom_parent_provenance_is_not_source_parent_id(self):
        factory=RedditRecordFactory("test",str(uuid.uuid4()),KeywordMatcher([]));sink=MemorySink();data=Dataset(factory,sink)
        adapter=RedditAdapter(factory,data)
        for raw in RedditHtmlParser(fixture_page(OLD_SYNTHETIC_HTML)).thread("pst123")[0]:adapter.accept("record",raw,FIXTURE_URL,None)
        data.finish_relations()
        evidence=[r["relation_evidence"] for r in sink.rows["Relations"] if r["predicate"]=="reply_to"]
        self.assertEqual(evidence,["dom_nesting","dom_nesting"])

    def test_14_settings_bounds_and_no_zero_fill(self):
        for kwargs in ({"max_comments":0},{"max_posts":True},{"request_delay":0},{"subreddits":("../oops",)}):
            with self.assertRaises(ValueError):Settings(**kwargs).validate()
        self.assertIsNone(parsed_number("1.2k"));self.assertEqual(parsed_number("0"),0)

    def test_15_redirect_checks_before_following(self):
        handler=RestrictedRedirect();req=Request(FIXTURE_URL)
        for url,code in (("https://evil.test/","REDIRECT_REJECTED"),(BASE+"/login/","LOGIN_REQUIRED")):
            with self.assertRaises(SourceFailure) as error:handler.redirect_request(req,None,302,"",{},url)
            self.assertEqual(error.exception.code,code)

    def test_16_hard_timeout_keeps_already_received_event(self):
        events=[]
        result=run_worker("thread","synthetic",Settings(),1.5,lambda kind,row:events.append((kind,row)),worker_entry=timeout_fixture)
        self.assertEqual(result["status"],"TIMEOUT")
        self.assertEqual(events,[("record",{"fixture":"SYNTHETIC"})])

    def test_17_import_html_produces_eight_sheets_and_unknown_capture_flag(self):
        import contextlib,io
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            p=Path(td);(p/"input.html").write_text(SYNTHETIC_HTML,encoding="utf-8")
            args=parser().parse_args(["import-html","--input",str(p/"input.html"),"--url",FIXTURE_URL,"--output-dir",str(p/"out")])
            with contextlib.redirect_stdout(io.StringIO()):code=run(args)
            self.assertEqual(code,0)
            book=load_workbook(p/"out/dataset.xlsx",read_only=True)
            self.assertEqual(book.sheetnames,list(TABLES))
            for table in TABLES:self.assertEqual(list(next(book[table].values)),[f["name"] for f in TABLES[table]])
            book.close()
            rows=[json.loads(x) for x in (p/"out/Contents.jsonl").read_text().splitlines()]
            self.assertEqual(len(rows),4)
            self.assertTrue(all("snapshot_original_capture_time_unknown" in r["quality_flags"] for r in rows))
            source=json.loads((p/"out/Sources.jsonl").read_text())
            self.assertEqual(source["access_mode"],"local_file_no_network")
            self.assertIsNone(json.loads((p/"out/manifest.json").read_text())["synthetic"])

    def test_18_challenge_snapshot_has_receipt_and_zero_content_without_success(self):
        import contextlib,io
        with tempfile.TemporaryDirectory() as td:
            p=Path(td);(p/"input.html").write_text('<form><input name="jsc_token"></form>')
            args=parser().parse_args(["import-html","--input",str(p/"input.html"),"--url",FIXTURE_URL,"--output-dir",str(p/"out")])
            with contextlib.redirect_stdout(io.StringIO()):code=run(args)
            self.assertEqual(code,2)
            self.assertEqual((p/"out/Contents.jsonl").read_text(),"")
            receipt=json.loads((p/"out/Runs.jsonl").read_text())
            self.assertEqual(receipt["status"],"CHALLENGE_REQUIRED")
            self.assertEqual(receipt["coverage"],"UNKNOWN")

def run_self_tests():
    try:import openpyxl
    except ImportError:
        print("Install: python -m pip install openpyxl==3.1.5",file=sys.stderr);return 2
    suite=unittest.TestSuite()
    for cls in (ContractTests,RedditTests):suite.addTests(unittest.defaultTestLoader.loadTestsFromTestCase(cls))
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    return 0 if result.wasSuccessful() else 2

def parser():
    p=argparse.ArgumentParser(description="Reddit satu file, scraping HTML publik; Excel identik kontrak YouTube 1.0.0.")
    p.add_argument("--version",action="version",version=VERSION)
    subs=p.add_subparsers(dest="command",required=True)
    subs.add_parser("guide",help="Panduan dan batas pengambilan")
    con=subs.add_parser("contract",help="Validasi kontrak tertanam")
    con.add_argument("--json",action="store_true")
    subs.add_parser("doctor",help="Cek Python/dependency/kontrak tanpa jaringan")
    subs.add_parser("self-test",help="Regression test offline; fixture sintetis hanya di folder temporer")
    live=subs.add_parser("run",help="Pencarian post/thread publik; hanya HTML respons awal")
    snap=subs.add_parser("import-html",help="Konversi HTML tersimpan tanpa jaringan")
    for command in (live,snap):
        command.add_argument("--output-dir",type=Path,help="Folder baru, menolak overwrite")
        command.add_argument("--url",action="append",help="Permalink thread; dapat diulang pada mode run")
        command.add_argument("--keyword",action="append",help="Keyword literal; menggantikan daftar default")
        command.add_argument("--project-id")
        for name in ("max_comments","max_replies","max_replies_per_thread","timeout_seconds","total_seconds","max_html_bytes"):
            command.add_argument("--"+name.replace("_","-"),type=int)
        command.add_argument("--since",help="Awal inklusif ISO8601 timezone-aware")
        command.add_argument("--until",help="Akhir eksklusif ISO8601 timezone-aware")
    live.add_argument("--query",action="append",help="Mencari post, bukan global seluruh komentar")
    live.add_argument("--subreddit",action="append",help="Nama tanpa r/, default: indonesia, karir, indotech")
    live.add_argument("--urls-file",type=Path)
    live.add_argument("--max-posts",type=int,help="Batas post seluruh run, default 3")
    live.add_argument("--search-results",type=int,help="Kandidat maksimum per halaman pencarian, default 5")
    live.add_argument("--request-delay",type=float,help="Jarak antarrequest 1..30 detik, default 2")
    live.add_argument("--sort",choices=["new","top","old","best"],help="Urutan komentar, default new; best effort server")
    live.add_argument("--search-sort",choices=["new","relevance","top"],help="Urutan pencarian post, default new")
    snap.add_argument("--input",type=Path,required=True,help="HTML UTF-8, bukan PDF/MHTML; halaman yang berisi komentar")
    snap.add_argument("--captured-at",help="Opsional: waktu capture asli ISO8601 dengan timezone, dinyatakan pengguna")
    return p

def main():
    if len(sys.argv)==1:parser().print_help();return 0
    args=parser().parse_args()
    try:
        if args.command=="guide":print(__doc__);return 0
        if args.command=="contract":
            verify_descriptor();print(json.dumps(schema_document(),ensure_ascii=False,indent=2) if args.json else schema_fingerprint());return 0
        if args.command=="doctor":
            print("Python:",sys.version.split()[0]);verify_descriptor()
            try:print("openpyxl:",importlib.metadata.version("openpyxl"))
            except importlib.metadata.PackageNotFoundError:print("openpyxl: MISSING; install openpyxl==3.1.5");return 2
            print("Contract:",schema_fingerprint());print("Doctor tidak menguji akses Reddit live.");return 0
        if args.command=="self-test":verify_descriptor();return run_self_tests()
        return run(args)
    except ImportError:
        print("Dependency belum tersedia. Install: python -m pip install openpyxl==3.1.5",file=sys.stderr);return 2
    except (ValueError,FileNotFoundError,FileExistsError,PermissionError,OSError) as exc:
        print("Input/export error: "+str(exc),file=sys.stderr);return 2

if __name__=="__main__":
    mp.freeze_support()
    raise SystemExit(main())
