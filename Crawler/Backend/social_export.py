"""Shared root selection and formula-safe Excel projection for social crawlers."""

import calendar
import json
import re
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from Crawler.Backend.Reddit import crawl_reddit_single as reddit

COMPACT_FILE = "social_data.xlsx"
COMPACT_TIMEZONE = ZoneInfo("Asia/Jakarta")
COMPACT_COLUMNS = (
    "platform",
    "jenis",
    "judul_post",
    "tanggal_post_wib",
    "akurasi_tanggal_post",
    "tanggal_konten_wib",
    "akurasi_tanggal_konten",
    "teks",
    "url",
    "teks_dipotong",
)


class SocialExportMixin:
    """Requires config, include and exclude matchers; performs no acquisition."""

    @staticmethod
    def _calendar_date(row):
        """Return a local date for display; relative source dates stay estimates."""
        quality = row.get("timestamp_quality")
        if isinstance(row.get("published_at"), str) and row["published_at"]:
            try:
                stamp = reddit.utc_datetime(row["published_at"]).astimezone(
                    COMPACT_TIMEZONE
                )
                if quality == "exact":
                    return stamp.replace(tzinfo=None), "pasti"
                return (
                    stamp.date(),
                    "perkiraan" if quality == "relative" else "tanggal_saja",
                )
            except (TypeError, ValueError, OverflowError):
                pass
        display = str(row.get("published_at_display") or "").strip()
        if quality == "date_only":
            for fmt in ("%Y%m%d", "%Y-%m-%d"):
                try:
                    return datetime.strptime(display, fmt).date(), "tanggal_saja"
                except ValueError:
                    pass
        text = (
            re.sub(r"\s*\((?:edited|diedit)\)\s*$", "", display, flags=re.I)
            .casefold()
            .strip()
        )
        match = re.fullmatch(
            r"(\d+)\s+(seconds?|minutes?|hours?|days?|weeks?|months?|years?|"
            r"detik|menit|jam|hari|minggu|bulan|tahun)\s+(?:ago|(?:yang\s+)?lalu)",
            text,
        )
        relative = {
            "today": 0,
            "hari ini": 0,
            "just now": 0,
            "baru saja": 0,
            "yesterday": 1,
            "kemarin": 1,
        }
        if not match and text not in relative:
            return None, "tidak_diketahui"
        if not isinstance(row.get("observed_at"), str):
            return None, "tidak_diketahui"
        try:
            anchor = reddit.utc_datetime(row["observed_at"]).astimezone(
                COMPACT_TIMEZONE
            )
            if not match:
                return (anchor - timedelta(days=relative[text])).date(), "perkiraan"
            count = int(match[1])
            unit = match[2].rstrip("s")
            seconds = {
                "second": 1,
                "detik": 1,
                "minute": 60,
                "menit": 60,
                "hour": 3600,
                "jam": 3600,
                "day": 86400,
                "hari": 86400,
                "week": 604800,
                "minggu": 604800,
            }
            if unit in seconds:
                return (
                    anchor - timedelta(seconds=count * seconds[unit])
                ).date(), "perkiraan"
            months = count * (12 if unit in {"year", "tahun"} else 1)
            index = anchor.year * 12 + anchor.month - 1 - months
            year, month = divmod(index, 12)
            month += 1
            return date(
                year, month, min(anchor.day, calendar.monthrange(year, month)[1])
            ), "perkiraan"
        except (KeyError, TypeError, ValueError, OverflowError):
            return None, "tidak_diketahui"

    def _write_compact(self, directory, output):
        """Project the full audit rows into one sheet, applying current filters."""
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        contents = Path(directory) / "Contents.jsonl"
        roots = {}
        all_roots = 0

        def records():
            with contents.open(encoding="utf-8") as stream:
                for line in stream:
                    row = json.loads(line)
                    reddit.validate_row("Contents", row)
                    yield row

        for row in records():
            if row["object_type"] != "post":
                continue
            all_roots += 1
            if not self.config["platforms"].get(row["platform"], {}).get("enabled"):
                continue
            selected, _ = self._select({"Contents": [row]})
            if selected:
                roots[row["object_key"]] = {
                    "title": row["title"],
                    "date": self._calendar_date(row),
                }
        book = Workbook(write_only=True)
        sheet = book.create_sheet("Data")
        sheet.freeze_panes = "D2"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_format.defaultRowHeight = 45
        sheet.row_dimensions[1].height = 28
        for i, name in enumerate(COMPACT_COLUMNS, 1):
            sheet.column_dimensions[get_column_letter(i)].width = (
                70 if name == "teks" else 48 if name in {"judul_post", "url"} else 24
            )
        font = Font(name="Calibri", size=11)
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="163C52")
        alignment = Alignment(wrap_text=True, vertical="top")

        def cells(values, header=False):
            result = []
            for value in values:
                if isinstance(value, str) and reddit.INVALID_XML.search(value):
                    value = json.dumps(value, ensure_ascii=True)
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = header_font if header else font
                cell.alignment = alignment
                if header:
                    cell.fill = fill
                elif isinstance(value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
                result.append(cell)
            return result

        sheet.append(cells(COMPACT_COLUMNS, True))
        counts = Counter()
        dates = Counter()
        written = skipped_comments = truncated = 0
        for row in records():
            root = roots.get(row["thread_root_key"])
            if not root:
                continue
            if row["object_type"] != "post" and any(
                self.exclude.find(row["text_source"])
            ):
                skipped_comments += 1
                continue
            stamp, accuracy = self._calendar_date(row)
            text = row["text_source"]
            if reddit.INVALID_XML.search(text):
                text = json.dumps(text, ensure_ascii=True)
            first = next(reddit.BundleSink.chunks(text))
            cut = first != text
            if cut:
                first += (
                    "\n[Teks lengkap tersedia di Contents.jsonl pada bundle sumber.]"
                )
            kind = (
                "video"
                if row["platform"] == "youtube" and row["object_type"] == "post"
                else row["object_type"]
            )
            sheet.append(
                cells(
                    [
                        row["platform"],
                        kind,
                        root["title"],
                        root["date"][0],
                        root["date"][1],
                        stamp,
                        accuracy,
                        first,
                        row["canonical_url"],
                        cut,
                    ]
                )
            )
            written += 1
            truncated += cut
            dates[accuracy] += 1
            counts[row["platform"] + ":" + kind] += 1
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(COMPACT_COLUMNS))}{written + 1}"
        )
        book.save(output)
        return dict(
            file=COMPACT_FILE,
            sheets=["Data"],
            columns=list(COMPACT_COLUMNS),
            timezone=str(COMPACT_TIMEZONE),
            rows=written,
            counts=dict(counts),
            date_accuracy_counts=dict(dates),
            truncated_texts=truncated,
            filtered_roots=all_roots - len(roots),
            filtered_comments=skipped_comments,
            date_estimate_anchor="source_observed_at",
            relative_date_precision="calendar_date_not_exact_timestamp",
        )

    def _select(self, rows):
        roots = [row for row in rows["Contents"] if row["object_type"] == "post"]
        if len(roots) != 1:
            return [], "root_not_verified"
        root = roots[0]
        if self.config["since"] or self.config["until"]:
            if not root["published_at"]:
                return [], "root_publication_date_unknown"
            stamp = reddit.utc_datetime(root["published_at"])
            if (
                self.config["since"]
                and stamp < reddit.utc_datetime(self.config["since"])
                or self.config["until"]
                and stamp >= reddit.utc_datetime(self.config["until"])
            ):
                return [], "root_outside_date_window"
        fields = [root.get("title") or "", root["text_source"]]
        if any(any(self.exclude.find(text)) for text in fields):
            return [], "root_excluded_keyword"
        matched = {hit["keyword"] for text in fields for hit in self.include.find(text)}
        keywords = set(self.config["include_keywords"])
        if keywords and (
            not matched
            if self.config["include_mode"] == "any"
            else not keywords <= matched
        ):
            return [], "root_missing_include_keyword"
        selected = [root]
        seen = {root["object_key"]}
        for row in rows["Contents"]:
            if (
                row is root
                or row["object_key"] in seen
                or row["thread_root_key"] != root["object_key"]
                or not row["text_source"].strip()
            ):
                continue
            if any(self.exclude.find(row["text_source"])):
                continue
            selected.append(row)
            seen.add(row["object_key"])
        return selected, "root_selected"
