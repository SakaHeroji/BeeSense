#!/usr/bin/env python3
"""Check exported evidence, relationships, workbook values and checksums offline."""
import argparse
from collections import Counter
from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import load_workbook
import crawl_reddit_single as common


def verify(directory, config):
    manifest = json.loads((directory / 'manifest.json').read_text())
    assert json.loads((directory / 'contract.json').read_text()) == common.schema_document()
    assert manifest['schema_fingerprint'] == common.schema_fingerprint()
    for name, expected in manifest['files_sha256'].items():
        assert sha256((directory / name).read_bytes()).hexdigest() == expected, name
    rows = {}
    for table in common.TABLES:
        rows[table] = [json.loads(line) for line in (directory / (table + '.jsonl')).read_text().splitlines()]
        assert len(rows[table]) == manifest['row_counts'][table], table
        for row in rows[table]:
            common.validate_row(table, row)

    contents = rows['Contents']
    by_key = {row['object_key']: row for row in contents}
    assert len(by_key) == len(contents), 'duplicate native content identity'
    posts = [row for row in contents if row['object_type'] == 'post']
    assert len(posts) == config['max_posts']
    assert len({row['native_id'] for row in posts}) == len(posts)
    since, until = common.utc_datetime(config['since']), common.utc_datetime(config['until'])
    for row in contents:
        if row['published_at']:
            assert since <= common.utc_datetime(row['published_at']) < until
        assert row['identity_quality'] == 'native_id'
        assert row['thread_root_key'] in by_key
        assert by_key[row['thread_root_key']]['object_type'] == 'post'
        assert row['text_complete'] is False
        if row['object_type'] == 'post':
            assert common.thread_id(row['canonical_url']) == row['native_id']
            assert common.safe_web_url(row['canonical_url']).path.split('/')[2].lower() in config['subreddits']
        else:
            assert '/comment/' + row['native_id'] + '/' in row['canonical_url']
            assert common.thread_id(row['canonical_url']) == by_key[row['thread_root_key']]['native_id']
            if row['parent_key'] in by_key:
                assert by_key[row['parent_key']]['thread_root_key'] == row['thread_root_key']
        if row['adapter_version'] == 'public-web-excerpt:1':
            assert len(((row['title'] or '') + ' ' + row['text_source']).split()) <= 200
            assert row['object_type'] == 'post'
            assert 'comments_not_acquired_by_web_reader' in row['quality_flags']
            assert 'bounded_browser_dom' not in row['quality_flags']
        if row['adapter_version'] == 'search-metadata-only':
            assert row['text_source'] == ''
            assert 'body_not_acquired' in row['quality_flags']
    for row in rows['Relations']:
        assert row['subject_key'] in by_key
        resolved = row['object_key'] in by_key
        assert (row['resolution_status'] == 'resolved_in_export') == resolved
    by_observation = {row['observation_id']: row for row in contents}
    for hit in rows['KeywordHits']:
        row = by_observation[hit['observation_id']]
        field = row[hit['field_name']]
        assert field[hit['span_start']:hit['span_end']] == hit['matched_variant']

    chunk_groups = {}
    for chunk in rows['TextChunks']:
        key = (chunk['table_name'], chunk['record_key'], chunk['field_name'])
        chunk_groups.setdefault(key, []).append(chunk)
    for table, records in rows.items():
        for record in records:
            for field, note in record['excel_notes'].items():
                if 'parts' not in note:
                    continue
                chunks = sorted(chunk_groups[(table, note['record_key'], field)], key=lambda item: item['part_index'])
                assert [item['part_index'] for item in chunks] == list(range(1, note['parts'] + 1))
                value = ''.join(item['text_part'] for item in chunks)
                if note['encoding'] == 'json_string':
                    value = json.loads(value)
                expected = record[field]
                if next(spec['type'] for spec in common.TABLES[table] if spec['name'] == field) == 'json':
                    expected = common.dumps(expected)
                assert value == expected, 'long text reconstruction differs from JSONL'

    book = load_workbook(directory / 'dataset.xlsx', read_only=True, data_only=False)
    assert book.sheetnames == list(common.TABLES)
    checked = 0
    for table, specs in common.TABLES.items():
        iterator = iter(book[table])
        assert [cell.value for cell in next(iterator)] == [spec['name'] for spec in specs]
        for record in rows[table]:
            cells = next(iterator)
            for cell, spec in zip(cells, specs):
                value = record[spec['name']]
                if spec['type'] == 'json':
                    value = common.dumps(value)
                elif spec['type'] == 'datetime' and value:
                    value = common.utc_datetime(value).replace(tzinfo=None)
                if isinstance(value, str):
                    if common.INVALID_XML.search(value):
                        value = json.dumps(value, ensure_ascii=True)
                    value = next(common.BundleSink.chunks(value))
                if isinstance(value, datetime):
                    assert isinstance(cell.value, datetime) and abs((cell.value - value).total_seconds()) <= .001, (table, spec['name'])
                else:
                    assert cell.value == value or value == '' and cell.value is None, (table, spec['name'])
                assert cell.data_type != 'f', 'source text became an Excel formula'
                checked += 1
        assert next(iterator, None) is None, 'extra workbook rows'
    book.close()
    result = {'status': 'PASS', 'unique_posts': len(posts), 'counts': dict(Counter(row['object_type'] for row in contents)),
              'post_source_versions': dict(Counter(row['adapter_version'] for row in posts)),
              'row_counts': manifest['row_counts'], 'checksum_files': len(manifest['files_sha256']),
              'excel_cells_compared': checked, 'all_publication_dates_in_window': True,
              'unresolved_relations': sum(row['resolution_status'] == 'unresolved_reference' for row in rows['Relations']),
              'rss_used': manifest['rss_used'], 'comments_complete': False,
              'metadata_target_reached': manifest['metadata_target_reached'],
              'post_section_target_reached': manifest['post_section_target_reached']}
    (directory / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    parser.add_argument('--config', type=Path, default=Path(__file__).with_name('config_reddit_1000.json'))
    args = parser.parse_args()
    print(json.dumps(verify(args.directory, json.loads(args.config.read_text())), ensure_ascii=False, indent=2))
