#!/usr/bin/env python3
"""SocialCrawler: bounded YouTube/Reddit acquisition with one export contract.

Publication dates and include keywords select root posts/videos. Exclude keywords
also remove comments. A minimum counts qualifying roots with enough unique
comments/replies; shortages, unknown dates and access barriers remain explicit.
"""
from __future__ import annotations

import argparse
import calendar
from collections import Counter, deque
from copy import deepcopy
from dataclasses import replace
from datetime import date, datetime, timedelta, timezone
from hashlib import file_digest
import json
import math
from pathlib import Path
import random
import re
import shutil
import tempfile
import time
from types import SimpleNamespace
import uuid
from urllib.parse import urlencode, urlsplit
from zoneinfo import ZoneInfo

from Youtube import crawl_youtube as youtube
from Reddit import crawl_reddit_single as reddit

VERSION = '1.1.0'
ROOT = Path(__file__).resolve().parent
COMPACT_FILE = 'social_data.xlsx'
COMPACT_TIMEZONE = ZoneInfo('Asia/Jakarta')
COMPACT_COLUMNS = ('platform','jenis','judul_post','tanggal_post_wib','akurasi_tanggal_post',
    'tanggal_konten_wib','akurasi_tanggal_konten',
    'teks','url','teks_dipotong')
BARRIERS = {'CAPTCHA_REQUIRED','CHALLENGE_REQUIRED','LOGIN_OR_CHALLENGE','LOGIN_REQUIRED',
            'ACCESS_BLOCKED','ACCESS_BARRIER','CLIENT_BROWSER_BLOCK','RATE_LIMITED','TLS_ERROR'}
STOP_STATUSES = BARRIERS | {'DEPENDENCY_MISSING','INTERRUPTED'}
DEFAULT_PLATFORM = dict(enabled=True, urls=[], queries=[], min_posts=1, min_comments_per_post=1,
    max_posts=20, max_candidates=100, max_comments=100, max_replies=30, max_replies_per_thread=5,
    search_results=10, sort='new')
DEFAULT_CONFIG = dict(project_id='social-intelligence', since=None, until=None,
    include_keywords=[], exclude_keywords=[], include_mode='any',
    delay=dict(mode='fixed', seconds=2.0), timeout_seconds=120, total_seconds=1800,
    platforms={'youtube':{}, 'reddit':{}})


def _number(value, name, low, high, integer=False):
    if type(value) not in ((int,) if integer else (int,float)) or not math.isfinite(value) or not low <= value <= high:
        raise ValueError(f'{name} must be {"integer" if integer else "number"} in {low}..{high}')


def _keywords(value, name):
    if not isinstance(value, (list,tuple)) or len(value)>200 or any(not isinstance(x,str) or not x.strip() or len(x)>200 or any(ord(c)<32 for c in x) for x in value):
        raise ValueError(name + ' must contain at most 200 nonempty strings, each at most 200 characters')
    return list(dict.fromkeys(x.strip() for x in value))


def _config(value):
    cfg = deepcopy(DEFAULT_CONFIG)
    if isinstance(value, (str,Path)):
        path = Path(value)
        if path.stat().st_size > 1_000_000:
            raise ValueError('config exceeds 1 MB')
        value = json.loads(path.read_text())
    if not isinstance(value, dict) or set(value)-set(DEFAULT_CONFIG):
        raise ValueError('config must be an object with known SocialCrawler fields')
    cfg.update(deepcopy(value))
    if not isinstance(cfg['project_id'],str) or not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_-]{0,79}',cfg['project_id']):
        raise ValueError('project_id must be a slug of 1..80 characters')
    for key in ['since','until']:
        if cfg[key] is not None:
            if not isinstance(cfg[key],str) or not cfg[key].strip():
                raise ValueError(key+' must be a timezone-aware ISO datetime or null')
            cfg[key] = reddit.utc_datetime(cfg[key]).isoformat()
    if cfg['since'] and cfg['until'] and reddit.utc_datetime(cfg['since']) >= reddit.utc_datetime(cfg['until']):
        raise ValueError('since must be before until; until is exclusive')
    for key in ['include_keywords','exclude_keywords']:
        cfg[key] = _keywords(cfg[key],key)
    if cfg['include_mode'] not in {'any','all'}:
        raise ValueError('include_mode must be any or all')
    _number(cfg['timeout_seconds'],'timeout_seconds',5,600,True)
    _number(cfg['total_seconds'],'total_seconds',1,86400,True)
    delay = cfg['delay']
    if not isinstance(delay,dict) or delay.get('mode') not in {'fixed','random'}:
        raise ValueError('delay.mode must be fixed or random')
    fields = {'mode','seconds'} if delay['mode']=='fixed' else {'mode','min_seconds','max_seconds'}
    if set(delay) != fields:
        raise ValueError('delay fields do not match selected mode')
    for key in fields-{'mode'}:
        _number(delay[key],'delay.'+key,0,300)
    if delay['mode']=='random' and delay['min_seconds']>delay['max_seconds']:
        raise ValueError('random delay min_seconds must not exceed max_seconds')
    platforms = cfg['platforms']
    if not isinstance(platforms,dict) or not platforms or set(platforms)-{'youtube','reddit'}:
        raise ValueError('platforms must contain youtube and/or reddit')
    result = {}
    extra = {'youtube':{'cert_store'},'reddit':{'subreddits','search_sort','backend','import_checkpoint_dir','session','cli','artifact_dir'}}
    for name, settings in platforms.items():
        if not isinstance(settings,dict) or set(settings)-(set(DEFAULT_PLATFORM)|extra[name]):
            raise ValueError('unknown settings for '+name)
        p = {**deepcopy(DEFAULT_PLATFORM),**settings}
        if type(p['enabled']) is not bool:
            raise ValueError(name+'.enabled must be boolean')
        for field,low,high in [('min_posts',0,10000),('max_posts',1,10000),('max_candidates',1,20000),
                               ('min_comments_per_post',0,2000),('max_comments',1,2000),('max_replies',0,1000),
                               ('max_replies_per_thread',0,100),('search_results',1,500 if name=='youtube' else 50)]:
            _number(p[field],name+'.'+field,low,high,True)
        if p['min_posts']>p['max_posts'] or p['min_posts']>p['max_candidates']:
            raise ValueError(name+': min_posts must not exceed max_posts or max_candidates')
        if p['min_comments_per_post']>p['max_comments']:
            raise ValueError(name+': minimum comments exceeds acquisition cap')
        for key in ['urls','queries']:
            if not isinstance(p[key],list) or any(not isinstance(x,str) for x in p[key]):
                raise ValueError(name+'.'+key+' must be a list of strings')
        if len(p['urls'])>20000 or len(p['queries'])>100 or any(len(q)>300 or any(ord(c)<32 for c in q) for q in p['queries']):
            raise ValueError(name+': URL/query budget exceeded')
        p['urls'] = list(dict.fromkeys(_url(name,u) for u in p['urls']))
        p['queries'] = list(dict.fromkeys(q.strip() for q in p['queries'] if q.strip()))
        if not p['queries'] and not p['urls']:
            p['queries'] = cfg['include_keywords'][:10]
        if name=='youtube':
            p.setdefault('cert_store','system')
            if p['sort'] not in {'new','top'} or p['cert_store'] not in {'system','bundled'}:
                raise ValueError('invalid YouTube sort/cert_store')
        else:
            p.setdefault('backend','playwright')
            p.setdefault('subreddits',['indonesia','karir','indotech'])
            p.setdefault('search_sort','new')
            p.setdefault('import_checkpoint_dir',None)
            p.setdefault('session','beesense-reddit')
            p.setdefault('cli',None)
            p.setdefault('artifact_dir',str(ROOT/'output/playwright/reddit'))
            if p['backend'] not in {'html','playwright','checkpoint'}:
                raise ValueError('Reddit backend must be html, playwright or checkpoint')
            if not isinstance(p['subreddits'],list) or not 1<=len(p['subreddits'])<=10 or any(not isinstance(s,str) or not re.fullmatch(r'[A-Za-z0-9_]{2,32}',s) for s in p['subreddits']):
                raise ValueError('invalid subreddit allowlist')
            p['subreddits'] = list(dict.fromkeys(s.lower() for s in p['subreddits']))
            if p['sort'] not in {'new','top','old','best'} or p['search_sort'] not in {'new','relevance','top'}:
                raise ValueError('invalid Reddit sort')
            if p['backend']=='checkpoint' and not p['import_checkpoint_dir']:
                raise ValueError('checkpoint backend requires import_checkpoint_dir')
            if not isinstance(p['session'],str) or not re.fullmatch(r'[A-Za-z0-9_-]+',p['session']):
                raise ValueError('invalid Playwright session')
            for key in ['import_checkpoint_dir','artifact_dir','cli']:
                if p[key] is not None and not isinstance(p[key],str):
                    raise ValueError(name+'.'+key+' must be a path string or null')
        result[name] = p
    if not any(p['enabled'] for p in result.values()):
        raise ValueError('enable at least one platform')
    cfg['platforms'] = result
    return cfg


def _url(platform, value):
    return youtube.video_url(value) if platform=='youtube' else reddit.thread_url(value)


class SocialCrawler:
    """Compose the original adapters and workers; one JSON configuration and bundle.

    Worker injection is intended for deterministic offline verification. A worker
    has the same interface as the existing run_worker functions. Date boundaries
    require explicit timezones; since is inclusive and until is exclusive.
    """
    def __init__(self, config=None, *, workers=None, sleep=time.sleep, monotonic=time.monotonic, rng=None):
        self.config = _config(ROOT/'social_crawler.json' if config is None else config)
        self.workers = {'youtube':youtube.run_worker,'reddit':reddit.run_worker,**(workers or {})}
        self.sleep, self.clock, self.rng = sleep, monotonic, rng or random.Random()
        self.include = reddit.KeywordMatcher(self.config['include_keywords'])
        self.exclude = reddit.KeywordMatcher(self.config['exclude_keywords'])
        if youtube.schema_fingerprint()!=reddit.schema_fingerprint():
            raise ValueError('YouTube/Reddit table contracts differ')
        self._last_finished = None
        self._browser = None
        self._synthetic = bool(workers)

    def plan(self):
        """Validate and return the effective plan without browser/network access."""
        return {'config':deepcopy(self.config),'date_scope':'root_publication',
                'comment_scope':'all_visible_context_except_excluded_keywords',
                'minimum_scope':'per_platform_roots_with_minimum_unique_comments_including_replies',
                'delay_scope':'between_search_or_content_operations_not_each_HTTP_subrequest',
                'until_exclusive':True,'live_access_verified':False,
                'primary_output':dict(file=COMPACT_FILE,sheets=['Data'],timezone=str(COMPACT_TIMEZONE),
                                      columns=list(COMPACT_COLUMNS),relative_dates='estimated_calendar_date')}

    @staticmethod
    def _calendar_date(row):
        """Return a local date for display; relative source dates stay estimates."""
        quality=row.get('timestamp_quality')
        if isinstance(row.get('published_at'),str) and row['published_at']:
            try:
                stamp=reddit.utc_datetime(row['published_at']).astimezone(COMPACT_TIMEZONE)
                if quality=='exact':return stamp.replace(tzinfo=None),'pasti'
                return stamp.date(),'perkiraan' if quality=='relative' else 'tanggal_saja'
            except (TypeError,ValueError,OverflowError):
                pass
        display=str(row.get('published_at_display') or '').strip()
        if quality=='date_only':
            for fmt in ('%Y%m%d','%Y-%m-%d'):
                try:return datetime.strptime(display,fmt).date(),'tanggal_saja'
                except ValueError:pass
        text=re.sub(r'\s*\((?:edited|diedit)\)\s*$', '',display,flags=re.I).casefold().strip()
        match=re.fullmatch(r'(\d+)\s+(seconds?|minutes?|hours?|days?|weeks?|months?|years?|'
                          r'detik|menit|jam|hari|minggu|bulan|tahun)\s+(?:ago|(?:yang\s+)?lalu)',text)
        relative={'today':0,'hari ini':0,'just now':0,'baru saja':0,'yesterday':1,'kemarin':1}
        if not match and text not in relative:return None,'tidak_diketahui'
        if not isinstance(row.get('observed_at'),str):return None,'tidak_diketahui'
        try:
            anchor=reddit.utc_datetime(row['observed_at']).astimezone(COMPACT_TIMEZONE)
            if not match:return (anchor-timedelta(days=relative[text])).date(),'perkiraan'
            count=int(match[1]);unit=match[2].rstrip('s')
            seconds={'second':1,'detik':1,'minute':60,'menit':60,'hour':3600,'jam':3600,
                     'day':86400,'hari':86400,'week':604800,'minggu':604800}
            if unit in seconds:
                return (anchor-timedelta(seconds=count*seconds[unit])).date(),'perkiraan'
            months=count*(12 if unit in {'year','tahun'} else 1)
            index=anchor.year*12+anchor.month-1-months
            year,month=divmod(index,12);month+=1
            return date(year,month,min(anchor.day,calendar.monthrange(year,month)[1])),'perkiraan'
        except (KeyError,TypeError,ValueError,OverflowError):
            return None,'tidak_diketahui'

    def _write_compact(self, directory, output):
        """Project the full audit rows into one sheet, applying current filters."""
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment,Font,PatternFill
        from openpyxl.utils import get_column_letter
        contents=Path(directory)/'Contents.jsonl'
        roots={};all_roots=0
        def records():
            with contents.open(encoding='utf-8') as stream:
                for line in stream:
                    row=json.loads(line);reddit.validate_row('Contents',row)
                    yield row
        for row in records():
            if row['object_type']!='post':continue
            all_roots+=1
            if not self.config['platforms'].get(row['platform'],{}).get('enabled'):continue
            selected,_=self._select({'Contents':[row]})
            if selected:
                roots[row['object_key']]={'title':row['title'],
                                          'date':self._calendar_date(row)}
        book=Workbook(write_only=True);sheet=book.create_sheet('Data')
        sheet.freeze_panes='D2';sheet.sheet_view.showGridLines=False
        sheet.sheet_format.defaultRowHeight=45;sheet.row_dimensions[1].height=28
        for i,name in enumerate(COMPACT_COLUMNS,1):
            sheet.column_dimensions[get_column_letter(i)].width=70 if name=='teks' else 48 if name in {'judul_post','url'} else 24
        font=Font(name='Calibri',size=11);header_font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
        fill=PatternFill('solid',fgColor='163C52');alignment=Alignment(wrap_text=True,vertical='top')
        def cells(values,header=False):
            result=[]
            for value in values:
                if isinstance(value,str) and reddit.INVALID_XML.search(value):value=json.dumps(value,ensure_ascii=True)
                cell=WriteOnlyCell(sheet,value=value)
                if isinstance(value,str):cell.data_type='s'
                cell.font=header_font if header else font;cell.alignment=alignment
                if header:cell.fill=fill
                elif isinstance(value,datetime):cell.number_format='yyyy-mm-dd hh:mm:ss'
                elif isinstance(value,date):cell.number_format='yyyy-mm-dd'
                result.append(cell)
            return result
        sheet.append(cells(COMPACT_COLUMNS,True))
        counts=Counter();dates=Counter();written=skipped_comments=truncated=0
        for row in records():
            root=roots.get(row['thread_root_key'])
            if not root:continue
            if row['object_type']!='post' and any(self.exclude.find(row['text_source'])):
                skipped_comments+=1;continue
            stamp,accuracy=self._calendar_date(row)
            text=row['text_source']
            if reddit.INVALID_XML.search(text):text=json.dumps(text,ensure_ascii=True)
            first=next(reddit.BundleSink.chunks(text));cut=first!=text
            if cut:first+='\n[Teks lengkap tersedia di Contents.jsonl pada bundle sumber.]'
            kind='video' if row['platform']=='youtube' and row['object_type']=='post' else row['object_type']
            sheet.append(cells([row['platform'],kind,root['title'],
                root['date'][0],root['date'][1],stamp,accuracy,first,row['canonical_url'],cut]))
            written+=1;truncated+=cut;dates[accuracy]+=1;counts[row['platform']+':'+kind]+=1
        sheet.auto_filter.ref=f'A1:{get_column_letter(len(COMPACT_COLUMNS))}{written+1}'
        book.save(output)
        return dict(file=COMPACT_FILE,sheets=['Data'],columns=list(COMPACT_COLUMNS),timezone=str(COMPACT_TIMEZONE),
            rows=written,counts=dict(counts),date_accuracy_counts=dict(dates),truncated_texts=truncated,
            filtered_roots=all_roots-len(roots),filtered_comments=skipped_comments,
            date_estimate_anchor='source_observed_at',relative_date_precision='calendar_date_not_exact_timestamp')

    def export_compact(self, bundle_dir, output_dir):
        """Re-export saved Contents without network or modifying the source bundle."""
        source=Path(bundle_dir).resolve();output=Path(output_dir).resolve()
        if output.exists():raise FileExistsError('Output directory sudah ada; gunakan nama baru.')
        contents=source/'Contents.jsonl'
        def file_hash(path):
            with path.open('rb') as stream:return file_digest(stream,'sha256').hexdigest()
        digest=file_hash(contents)
        source_manifest=json.loads((source/'manifest.json').read_text())
        if source_manifest['files_sha256']['Contents.jsonl']!=digest:
            raise ValueError('Checksum Contents.jsonl sumber tidak cocok')
        if source_manifest['schema_fingerprint']!=reddit.schema_fingerprint():
            raise ValueError('Kontrak bundle sumber tidak cocok')
        output.parent.mkdir(parents=True,exist_ok=True)
        stage=Path(tempfile.mkdtemp(prefix='.social-compact-',dir=output.parent))
        try:
            compact=self._write_compact(source,stage/COMPACT_FILE)
            report=dict(status='EXPORTED',mode='compact_reexport',producer_version=VERSION,output_dir=str(output),
                primary_file=str(output/COMPACT_FILE),compact=compact,source_bundle=str(source),
                source_run_id=source_manifest['run_id'],source_contents_sha256=digest,
                source_synthetic=source_manifest.get('synthetic',False),source_observations_refreshed=False,
                selection={k:deepcopy(self.config[k]) for k in ['since','until','include_keywords','exclude_keywords','include_mode']})
            (stage/'summary.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
            (stage/'manifest.json').write_text(json.dumps({**report,'created_at':reddit.now_iso(),
                'files_sha256':{p.name:file_hash(p) for p in stage.iterdir() if p.is_file()}},ensure_ascii=False,indent=2)+'\n')
            stage.rename(output)
            return report
        except BaseException:
            shutil.rmtree(stage,ignore_errors=True)
            raise

    def _delay(self):
        d=self.config['delay']
        return d['seconds'] if d['mode']=='fixed' else self.rng.uniform(d['min_seconds'],d['max_seconds'])

    def _pace(self, deadline):
        if self._last_finished is not None:
            wait=max(0,self._delay()-(self.clock()-self._last_finished))
            if self.clock()+wait>=deadline:
                return False
            while wait>0:
                part=min(30,wait);self.sleep(part);wait-=part
        return self.clock()<deadline

    def _legacy_settings(self, name, p):
        shared=dict(project_id=self.config['project_id'],keywords=tuple(self.config['include_keywords']),
                    max_comments=p['max_comments'],max_replies=p['max_replies'],
                    max_replies_per_thread=p['max_replies_per_thread'],search_results=min(10,p['search_results']) if name=='youtube' else p['search_results'],
                    request_delay=2,sort=p['sort'])
        if name=='youtube':
            settings=youtube.Settings(**shared,max_videos=1,timeout_seconds=max(10,min(600,self.config['timeout_seconds'])),
                                      cert_store=p['cert_store'])
        else:
            settings=reddit.Settings(**shared,max_posts=1,timeout_seconds=min(300,self.config['timeout_seconds']),
                                     subreddits=tuple(p['subreddits']),search_sort=p['search_sort'])
        settings.validate()
        if name=='youtube':
            # This coordinator owns and validates the larger discovery budget;
            # the original CLI retains its existing 10-result search bound.
            settings=replace(settings,search_results=p['search_results'])
        return settings

    def _checkpoint_items(self, p):
        if not p.get('import_checkpoint_dir'):
            return []
        directory=Path(p['import_checkpoint_dir']).expanduser().resolve()
        candidates=json.loads((directory/'candidates.json').read_text())
        items=[]
        for candidate in candidates:
            if candidate.get('subreddit','').lower() not in p['subreddits']:
                continue
            try:
                url=_url('reddit',candidate['canonical_url'])
                if reddit.thread_id(url)!=candidate['native_id']:
                    continue
            except (KeyError,ValueError):
                continue
            path=directory/'raw'/(candidate['native_id']+'.json')
            raw=json.loads(path.read_text()) if path.exists() else None
            if raw and raw.get('extractor_version') not in {'browser-dom:2','playwright-dom:1'}:
                raw=None
            if p['backend']!='checkpoint' or raw:
                items.append((url,candidate.get('discovered_from'),raw))
            if len(items)>=p['max_candidates']:
                break
        # Existing DOM evidence is reused before starting new browser requests.
        return sorted(items,key=lambda item:item[2] is None)

    def _dom_events(self, raw, url, p):
        from Reddit import crawl_reddit_playwright as browser
        post=raw['post'];native=reddit.thread_id(url)
        candidate=dict(native_id=native,title=post.get('title'),subreddit=post.get('subreddit'),canonical_url=url)
        validation=dict(subreddits=p['subreddits'],max_comments=p['max_comments'],
                        since='0001-01-01T00:00:00+00:00',until='9999-12-31T23:59:59+00:00')
        if str(candidate['subreddit'] or '').lower() not in p['subreddits']:
            raise ValueError('source subreddit outside allowlist')
        post,comments=browser.validate_capture(candidate,raw,validation)
        stamp=raw.get('captured_at') or reddit.now_iso()
        base=dict(root='t3_'+native,observed_at=stamp,page_url=url,
                  extraction_method=raw.get('method','standalone_playwright_cli'),author_id=None,publisher_ref=None,
                  subreddit=post['subreddit'],text_complete=False,
                  flags=['bounded_browser_dom','more_replies_not_expanded','parent_from_dom_nesting',
                         'imported_dom_snapshot_not_independently_verified'] if raw.get('extractor_version') else
                        ['bounded_browser_dom','more_replies_not_expanded','parent_from_dom_nesting'])
        records=[{**base,'fullname':'t3_'+native,'kind':'post','title':post['title'],'text':post.get('text',''),
                  'parent':None,'parent_evidence':None,'published_at':post['published_at'],'published_display':None,
                  'timestamp_quality':'exact','score':reddit.parsed_number(post.get('score_display')),
                  'comment_count':reddit.parsed_number(post.get('comment_count_display'))}]
        for item in comments:
            native_comment=re.search(r'/comment/([a-z0-9]+)/',item['url'])[1]
            parent=re.search(r'/comment/([a-z0-9]+)/',item.get('parent_url') or '')
            try:
                published=reddit.utc_datetime(item['published_at']).isoformat() if item.get('published_at') else None
            except ValueError:
                published=None
            records.append({**base,'fullname':'t1_'+native_comment,'kind':'reply' if parent else 'comment',
                'title':None,'text':item['text'],'parent':'t1_'+parent[1] if parent else 't3_'+native,
                'parent_evidence':'dom_nesting','published_at':published,'published_display':item.get('published_display'),
                'timestamp_quality':'exact' if published else 'unknown','score':None,'comment_count':None})
        return list(reddit.bounded_events(records,SimpleNamespace(**p)))

    def _browser_call(self, operation, target, p, seconds):
        from Reddit import crawl_reddit_playwright as browser
        self._browser_deadline=self.clock()+seconds
        if self._browser is None:
            self._browser=browser.PlaywrightCLI(browser.find_cli(p.get('cli')),p['session'],Path(p['artifact_dir']).resolve(),
                                               timeout=min(90,self.config['timeout_seconds']+15))
            cli=self._browser
            original_call=cli.call
            def bounded_call(*args,**kwargs):
                remaining=self._browser_deadline-self.clock()
                if remaining<=0:raise TimeoutError('browser_operation_deadline')
                cli.timeout=min(90,self.config['timeout_seconds']+15,remaining)
                return original_call(*args,**kwargs)
            cli.call=bounded_call
            try:
                current=self._browser.inspect()
            except RuntimeError as error:
                if 'is not open' not in str(error):
                    raise
                self._browser.call('open','about:blank','--browser=chrome','--headed',
                                   '--profile='+str(self._browser.artifacts/'profile'))
                current=self._browser.inspect()
            status,wait=browser.barrier_status(current)
            if status:
                self._browser=None
                return [],dict(status=status,reason='existing_browser_access_barrier',warnings=[],source_version='playwright-dom:1',retry_after_seconds=wait)
        if operation=='search':
            script='''async page => {
              const response=await page.goto(URL_VALUE,{waitUntil:'domcontentloaded',timeout:TIMEOUT_VALUE});
              const headings=await page.getByRole('heading',{level:2}).count();
              const candidates=await page.getByRole('heading',{level:2}).evaluateAll(elements=>elements.map(h=>({
                url:h.querySelector('a')?.href,discovered_from:location.href})).filter(x=>x.url));
              return {http_status:response?.status(),retry_after:response?.headers()['retry-after'],post:!!headings,
                page_title:await page.title(),barrier_text:headings ? (await page.locator('[role=banner],header').allTextContents()).join('\\n') : await page.locator('body').innerText(),candidates};
            }'''.replace('URL_VALUE',json.dumps(target)).replace('TIMEOUT_VALUE',str(min(45000,self.config['timeout_seconds']*1000)))
            file=self._browser.artifacts/'social-search.js';file.write_text(script)
            capture=self._browser.call('run-code','--filename',file,decode=True)
            self._browser.call('snapshot','--filename',self._browser.artifacts/'social-search.yml')
            status,wait=browser.barrier_status(capture)
            if status:
                return [],dict(status=status,reason='browser_search_access_barrier',warnings=[],source_version='playwright-dom:1',retry_after_seconds=wait)
            return [('candidate',value) for value in capture['candidates'][:p['search_results']]],dict(status='FINISHED_BEST_EFFORT',reason='bounded_browser_search',warnings=[],source_version='playwright-dom:1')
        capture=self._browser.capture(dict(native_id=reddit.thread_id(target),canonical_url=target),
                                      dict(max_comments=p['max_comments'],timeout_seconds=self.config['timeout_seconds']))
        status,wait=browser.barrier_status(capture)
        if status:
            return [],dict(status=status,reason='browser_thread_access_barrier',warnings=[],source_version='playwright-dom:1',retry_after_seconds=wait)
        raw={**capture,'captured_at':reddit.now_iso(),'method':'standalone_playwright_cli'}
        return self._dom_events(raw,target,p),dict(status='FINISHED_BEST_EFFORT',reason='rendered_dom_not_complete_thread',warnings=['more_replies_not_expanded'],source_version='playwright-dom:1')

    def _execute(self, name, operation, target, settings, p, deadline, cached=None):
        if cached is not None:
            try:
                return self._dom_events(cached,target,p),dict(status='FINISHED_BEST_EFFORT',reason='reused_dom_checkpoint',warnings=['snapshot_not_refreshed'],source_version=cached['extractor_version'])
            except (ValueError,TypeError,KeyError) as error:
                return [],dict(status='SOURCE_ERROR',reason='checkpoint_record_unverified',warnings=[type(error).__name__],source_version=None)
        events=[]
        try:
            if not self._pace(deadline):
                return [],dict(status='TIME_BUDGET_EXHAUSTED',reason='global_deadline',warnings=[],source_version=None)
            if name=='reddit' and p['backend']=='playwright' and self.workers[name] is reddit.run_worker:
                return self._browser_call(operation,target,p,min(settings.timeout_seconds,deadline-self.clock()))
            result=self.workers[name](operation,target,settings,min(settings.timeout_seconds,deadline-self.clock()),
                                      lambda kind,payload:events.append((kind,payload)))
            return events,result
        except KeyboardInterrupt:
            return events,dict(status='INTERRUPTED',reason='interrupted_partial_events_preserved',warnings=[],source_version=None)
        except Exception as error:
            # Chrome can reject navigation before Playwright returns a response.
            # Inspect the existing error page, without making another request.
            if name=='reddit' and p['backend']=='playwright' and self._browser is not None:
                from Reddit import crawl_reddit_playwright as browser
                try:
                    current=self._browser.inspect()
                    status,wait=browser.barrier_status(current)
                    if status:
                        return events,dict(status=status,reason='browser_error_page_access_barrier',
                            warnings=[type(error).__name__],source_version='playwright-dom:1',retry_after_seconds=wait)
                except Exception:
                    pass
            return events,dict(status='SOURCE_ERROR',reason='source_call_failed',warnings=[type(error).__name__],source_version=None)
        finally:
            self._last_finished=self.clock()

    def _normalize(self, name, events, url, origin, factory, version):
        rows={key:[] for key in reddit.TABLES}
        sink=SimpleNamespace(write=lambda table,row:rows[table].append(row))
        dataset=reddit.Dataset(factory,sink)
        adapter=youtube.YouTubeAdapter(factory,dataset,version or 'unknown') if name=='youtube' else reddit.RedditAdapter(factory,dataset)
        invalid=0
        for kind,payload in events:
            if name=='reddit' and kind!='record' or name=='youtube' and kind not in {'video','comment'}:
                continue
            try:
                adapter.accept(kind,payload,url,origin)
            except (ValueError,TypeError,KeyError,OverflowError):
                invalid+=1
        dataset.finish_relations()
        return rows,invalid

    def _select(self, rows):
        roots=[row for row in rows['Contents'] if row['object_type']=='post']
        if len(roots)!=1:
            return [],'root_not_verified'
        root=roots[0]
        if self.config['since'] or self.config['until']:
            if not root['published_at']:
                return [],'root_publication_date_unknown'
            stamp=reddit.utc_datetime(root['published_at'])
            if self.config['since'] and stamp<reddit.utc_datetime(self.config['since']) or self.config['until'] and stamp>=reddit.utc_datetime(self.config['until']):
                return [],'root_outside_date_window'
        fields=[root.get('title') or '',root['text_source']]
        if any(any(self.exclude.find(text)) for text in fields):
            return [],'root_excluded_keyword'
        matched={hit['keyword'] for text in fields for hit in self.include.find(text)}
        keywords=set(self.config['include_keywords'])
        if keywords and (not matched if self.config['include_mode']=='any' else not keywords<=matched):
            return [],'root_missing_include_keyword'
        selected=[root];seen={root['object_key']}
        for row in rows['Contents']:
            if row is root or row['object_key'] in seen or row['thread_root_key']!=root['object_key'] or not row['text_source'].strip():
                continue
            if any(self.exclude.find(row['text_source'])):
                continue
            selected.append(row);seen.add(row['object_key'])
        return selected,'root_selected'

    def run(self, output_dir=None):
        """Return an explicit target/shortfall report and write one atomic bundle."""
        reddit.verify_descriptor()
        self._last_finished=None;self._browser=None
        run_id=str(uuid.uuid4())
        factory=reddit.RedditRecordFactory(self.config['project_id'],run_id,self.include)
        output=Path(output_dir) if output_dir else ROOT/'outputs'/('social_'+datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')+'_'+run_id[:8])
        sink=reddit.BundleSink(output,self.config['project_id'],run_id)
        deadline=self.clock()+self.config['total_seconds']
        summaries={};object_keys=set();relations=[];interrupted=False;row_count=0
        try:
            for name,p in self.config['platforms'].items():
                summary=dict(enabled=p['enabled'],min_posts=p['min_posts'],min_comments_per_post=p['min_comments_per_post'],
                    posts_attempted=0,posts_exported=0,qualified_posts=0,comments=0,replies=0,filtered_roots={},
                    insufficient_comment_posts=0,invalid_rows=0,target_met=False,stop_reason='DISABLED')
                summaries[name]=summary
                if not p['enabled']:
                    continue
                settings=self._legacy_settings(name,p)
                queue=deque();known=set();filtered=Counter()
                def enqueue(url,origin=None,cached=None):
                    try:url=_url(name,url)
                    except ValueError:return
                    if url not in known and len(known)<p['max_candidates']:
                        known.add(url);queue.append((url,origin,cached))
                if name=='reddit':
                    for url,origin,cached in self._checkpoint_items(p):enqueue(url,origin,cached)
                for url in p['urls']:enqueue(url)
                searches=(('search',query) for query in p['queries']) if name=='youtube' else (
                    ('search',reddit.BASE+'/r/'+sub+'/search/?'+urlencode(dict(q=query,restrict_sr='on',type='link',sort=p['search_sort'],t='all')))
                    for query in p['queries'] for sub in p['subreddits'])
                summary['stop_reason']='CANDIDATES_EXHAUSTED'
                source_errors=0;versions=set()
                while summary['qualified_posts']<p['min_posts']:
                    if interrupted:
                        summary['stop_reason']='INTERRUPTED';break
                    if self.clock()>=deadline:
                        summary['stop_reason']='TIME_BUDGET_EXHAUSTED';break
                    if summary['posts_attempted']>=p['max_posts']:
                        summary['stop_reason']='MAX_POSTS_REACHED';break
                    if row_count>=500000:
                        summary['stop_reason']='OUTPUT_ROW_BUDGET';break
                    search=not queue
                    if search:
                        if name=='reddit' and p['backend']=='checkpoint' or len(known)>=p['max_candidates']:
                            break
                        try:operation,target=next(searches)
                        except StopIteration:break
                        origin,cached=None,None
                    else:
                        target,origin,cached=queue.popleft()
                        operation='video' if name=='youtube' else 'thread'
                        summary['posts_attempted']+=1
                        if name=='reddit' and p['backend']=='checkpoint' and cached is None:
                            continue
                    started=reddit.now_iso()
                    events,result=self._execute(name,operation,target,settings,p,deadline,cached)
                    if result.get('source_version'):versions.add(result['source_version'])
                    source_errors=source_errors+1 if result['status']=='SOURCE_ERROR' else 0
                    written=comments=replies=invalid=seen_comments=seen_replies=0
                    if search:
                        for kind,payload in events:
                            if kind=='candidate' and isinstance(payload,dict):enqueue(payload.get('url',''),payload.get('discovered_from'))
                    else:
                        rows,invalid=self._normalize(name,events,target,origin,factory,result.get('source_version'))
                        seen_comments=sum(row['object_type']=='comment' for row in rows['Contents'])
                        seen_replies=sum(row['object_type']=='reply' for row in rows['Contents'])
                        selected,reason=self._select(rows)
                        if not selected:filtered[reason]+=1
                        if row_count+len(selected)>500000:
                            selected=[];result={**result,'status':'OUTPUT_ROW_BUDGET','reason':'bounded_export_size'}
                        if selected:
                            count=len(selected)-1
                            qualified=count>=p['min_comments_per_post']
                            summary['qualified_posts']+=int(qualified)
                            summary['insufficient_comment_posts']+=int(not qualified)
                            summary['posts_exported']+=1
                            selected_ids={row['observation_id'] for row in selected}
                            keys={row['object_key'] for row in selected}
                            for row in selected:
                                row['quality_flags'] += ['social_crawler_root_date_scope','social_minimum_comments_met' if qualified else 'social_minimum_comments_shortfall']
                                sink.write('Contents',row);object_keys.add(row['object_key'])
                                comments+=row['object_type']=='comment';replies+=row['object_type']=='reply'
                            for table in ['KeywordHits','Metrics']:
                                for row in rows[table]:
                                    if row.get('observation_id') in selected_ids:sink.write(table,row)
                            relations.extend(row for row in rows['Relations'] if row['subject_key'] in keys)
                            written=len(selected);row_count+=written
                            summary['comments']+=comments;summary['replies']+=replies
                    summary['invalid_rows']+=invalid
                    warnings=list(result.get('warnings',[]))
                    if invalid:warnings.append('invalid_source_rows')
                    if result.get('retry_after_seconds') is not None:warnings.append('retry_after_seconds:'+str(result['retry_after_seconds']))
                    sink.write('Runs',factory.row('Runs',unit_id=str(uuid.uuid4()),source=name,operation=operation,target=target,
                        status=result['status'],reason=result.get('reason','source_result'),coverage='PARTIAL_NON_EXHAUSTIVE',contents_written=written,
                        comments_seen=seen_comments,replies_seen=seen_replies,duplicates_skipped=0,invalid_rows=invalid,
                        max_comments=None if search else p['max_comments'],max_replies=None if search else p['max_replies'],
                        timeout_seconds=self.config['timeout_seconds'],keywords=self.config['include_keywords'],source_version=result.get('source_version'),
                        warning_codes=warnings,started_at=started,finished_at=reddit.now_iso()))
                    if result['status'] in STOP_STATUSES or result['status'] in {'TIME_BUDGET_EXHAUSTED','OUTPUT_ROW_BUDGET'}:
                        summary['stop_reason']=result['status'];interrupted=result['status']=='INTERRUPTED';break
                    if source_errors>=3:
                        summary['stop_reason']='CONSECUTIVE_SOURCE_ERRORS';break
                summary['filtered_roots']=dict(filtered)
                summary['target_met']=summary['qualified_posts']>=p['min_posts']
                summary['posts_shortfall']=max(0,p['min_posts']-summary['qualified_posts'])
                if summary['target_met'] and summary['stop_reason'] not in STOP_STATUSES:
                    summary['stop_reason']='TARGET_MET'
                sink.write('Sources',factory.row('Sources',source=name,producer_version='social-crawler:'+VERSION,
                    upstream_version=';'.join(sorted(versions)) or None,
                    extraction_method='yt_dlp_unofficial_web' if name=='youtube' else p['backend'],
                    access_mode='public_guest_or_local_dom_checkpoint',config={**deepcopy(self.config),'platform_summary':summary},
                    schema_fingerprint=reddit.schema_fingerprint(),observed_at=reddit.now_iso()))
            for row in relations:
                row['resolution_status']='resolved_in_export' if row['object_key'] in object_keys else 'unresolved_reference'
                sink.write('Relations',row)
            enabled=[s for s in summaries.values() if s['enabled']]
            met=all(s['target_met'] for s in enabled)
            report=dict(status='INTERRUPTED' if interrupted else 'TARGET_MET' if met and not any(s['stop_reason'] in STOP_STATUSES for s in enabled) else 'PARTIAL',
                        targets_met=met,platforms=summaries,output_dir=str(sink.path),comments_complete=False,
                        date_scope='root_publication',delay_scope='between_operations',run_id=run_id)
            for stream in sink.files.values():stream.flush()
            compact=self._write_compact(sink.stage,sink.stage/COMPACT_FILE)
            report.update(primary_file=str(sink.path/COMPACT_FILE),compact=compact)
            reddit.add_readme(sink,factory,[('SocialCrawler','YouTube dan Reddit dalam satu kontrak delapan sheet.'),
                ('Minimum','Minimum dihitung per platform: root lolos tanggal/keyword dan memiliki minimum komentar unik termasuk balasan. Kekurangan tidak direkayasa.'),
                ('Tanggal','since inklusif; until eksklusif; wajib timezone. Tanggal berlaku untuk root. Root bertanggal tidak diketahui dikeluarkan jika rentang aktif. Komentar merupakan konteks, termasuk waktu relatif/tidak diketahui.'),
                ('Keyword','Include any/all pada judul dan teks root; exclude mengalahkan include dan juga membuang komentar. Komentar lain dipertahankan sebagai konteks.'),
                ('Jeda','Tetap/acak antar operasi pencarian atau post/video. Satu operasi bisa memicu banyak request upstream. Jeda tidak menjamin bebas rate limit/CAPTCHA.'),
                ('Cakupan','Komentar publik yang tersedia; more replies Reddit tidak diperluas. Snapshot checkpoint dapat lebih lama dari waktu run.'),
                ('Excel ringkas',COMPACT_FILE+' berisi satu sheet Data. Tanggal kalender WIB; pasti/perkiraan/tanggal_saja/tidak_diketahui dibedakan. Tanggal relatif dihitung dari observed_at sumber dan tidak mengubah published_at exact dalam audit.'),
                ('Target',json.dumps(report,ensure_ascii=False))])
            sink.finish(dict(source='social',mode='social_crawler',producer_version=VERSION,synthetic=self._synthetic,
                             interrupted=interrupted,targets_met=met,platform_summaries=summaries,comments_complete=False,
                             primary_file=COMPACT_FILE,compact=compact))
            (sink.path/'summary.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
            return report
        except BaseException:
            if not sink.closed:sink.abort()
            raise


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config',type=Path,default=ROOT/'social_crawler.json')
    parser.add_argument('--output-dir',type=Path)
    parser.add_argument('--dry-run',action='store_true')
    parser.add_argument('--from-bundle',type=Path,help='Ekspor ringkas bundle tersimpan tanpa crawling jaringan')
    args=parser.parse_args()
    try:
        crawler=SocialCrawler(args.config)
        if args.from_bundle:
            if args.dry_run or args.output_dir is None:
                raise ValueError('--from-bundle memerlukan --output-dir baru dan tidak memakai --dry-run')
            result=crawler.export_compact(args.from_bundle,args.output_dir)
        else:
            result=crawler.plan() if args.dry_run else crawler.run(args.output_dir)
        print(json.dumps(result,ensure_ascii=False,indent=2))
        return 0 if args.dry_run or result['status'] in {'TARGET_MET','EXPORTED'} else 130 if result['status']=='INTERRUPTED' else 2
    except (ValueError,OSError,KeyError,TypeError) as error:
        print(json.dumps({'status':'CONFIG_OR_RUNTIME_ERROR','error_type':type(error).__name__,'reason':str(error)},ensure_ascii=False))
        return 2


if __name__=='__main__':
    import multiprocessing
    multiprocessing.freeze_support()
    raise SystemExit(main())
