"""Merge compact crawl workbooks, rejecting repeated native content identities."""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import tempfile
from collections import Counter
from hashlib import file_digest, sha256
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

COLUMNS = (
    "platform",
    "jenis",
    "judul_post",
    "tanggal_post_wib",
    "akurasi_tanggal_post",
    "tanggal_konten_wib",
    "akurasi_tanggal_konten",
    "teks",
    "url",
    "teks_dipotong",
)


def sha(path):
    with Path(path).open("rb") as stream:
        return file_digest(stream, "sha256").hexdigest()


def identity(row):
    """Return global native object ID plus its root; text never defines identity."""
    platform, kind = row["platform"], row["jenis"]
    url = urlsplit(row["url"])
    if platform == "youtube" and url.hostname in {"youtube.com", "www.youtube.com"}:
        params = parse_qs(url.query)
        root = params.get("v", [""])[0]
        native = root if kind == "video" else params.get("lc", [""])[0]
        root_kind = kind == "video"
    elif platform == "reddit" and url.hostname in {"reddit.com", "www.reddit.com"}:
        match = re.fullmatch(
            r"/comments/([a-z0-9]+)/(?:(?:_|[^/]+)/([a-z0-9]+)/)?", url.path
        )
        if not match:
            raise ValueError(f"Unsupported Reddit native URL: {row['url']}")
        root, comment = match.groups()
        root_kind = kind == "post"
        native = root if root_kind else comment
    elif (
        platform == "instagram"
        and url.scheme == "https"
        and url.netloc in {"instagram.com", "www.instagram.com"}
    ):
        match = re.fullmatch(
            r"/(?:[\w.]+/)?(?:p|reel|tv)/([A-Za-z0-9_-]{1,40})/(?:c/([1-9][0-9]{0,29})/?)?",
            url.path,
        )
        if not match:
            raise ValueError("Unsupported Instagram native URL")
        code, path_comment = match.groups()
        alphabet = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789-_"
        root_number = 0
        for char in code:
            root_number = root_number * 64 + alphabet.index(char)
        if root_number <= 0:
            raise ValueError("Invalid Instagram media ID")
        root = str(root_number)
        query = parse_qs(url.query).get("comment_id", [])
        if len(query) > 1 or query and path_comment and query[0] != path_comment:
            raise ValueError("Conflicting Instagram comment identity")
        comment = path_comment or (query[0] if query else None)
        root_kind = kind == "post"
        if root_kind and comment:
            raise ValueError("Instagram root cannot have a comment identity")
        native = root if root_kind else comment
        if native is None or not re.fullmatch(r"[1-9][0-9]{0,29}", native):
            raise ValueError("Missing Instagram native ID")
    else:
        raise ValueError(f"Unsupported platform or URL: {platform}")
    if kind not in (
        {"video", "comment", "reply"}
        if platform == "youtube"
        else {"post", "comment", "reply"}
    ):
        raise ValueError(f"Unsupported content kind: {kind}")
    if not native or not root:
        raise ValueError("Missing native content ID")
    return (platform, "post" if root_kind else "comment", native), (platform, root)


def merge(sources, output, provenance, *, duplicate_policy="error"):
    if duplicate_policy == "prefer_latest":
        return merge_unique(sources, output, provenance)
    if duplicate_policy != "error":
        raise ValueError("Unsupported duplicate policy")
    sources = [Path(p).resolve() for p in sources]
    output, provenance = Path(output).resolve(), Path(provenance).resolve()
    if output.exists() or provenance.exists():
        raise ValueError("Destination already exists; use a new versioned path")
    if len(set(sources)) != len(sources):
        raise ValueError("Source repeated")
    output.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix="." + output.name + "-", dir=output.parent))
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Data")
    sheet.append(COLUMNS)
    seen, roots, comment_roots = set(), set(), set()
    counts, lineage = Counter(), []
    total = 0
    try:
        for source in sources:
            before = sha(source)
            original = load_workbook(source, read_only=True, data_only=False)
            per_source = Counter()
            first = total + 2
            try:
                rows = original["Data"].iter_rows()
                if tuple(cell.value for cell in next(rows)) != COLUMNS:
                    raise ValueError(f"Unexpected columns: {source}")
                for cells in rows:
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError("Formula found in source data")
                    values = tuple(cell.value for cell in cells)
                    row = dict(zip(COLUMNS, values))
                    key, root = identity(row)
                    if key in seen:
                        raise ValueError(f"Duplicate native content: {key}")
                    seen.add(key)
                    if key[1] == "post":
                        roots.add(root)
                    else:
                        comment_roots.add(root)
                        if row["teks_dipotong"]:
                            raise ValueError(
                                "Truncated comment cannot become training input"
                            )
                    safe = []
                    for value in values:
                        cell = WriteOnlyCell(sheet, value=value)
                        if isinstance(value, str):
                            cell.data_type = "s"
                        safe.append(cell)
                    sheet.append(safe)
                    total += 1
                    counts[(row["platform"], row["jenis"])] += 1
                    per_source[(row["platform"], row["jenis"])] += 1
            finally:
                original.close()
            if sha(source) != before:
                raise ValueError(f"Source changed during merge: {source}")
            lineage.append(
                {
                    "path": str(source),
                    "sha256": before,
                    "first_excel_row": first,
                    "last_excel_row": total + 1,
                    "rows": sum(per_source.values()),
                    "counts": {
                        f"{p}/{k}": n for (p, k), n in sorted(per_source.items())
                    },
                }
            )
        if comment_roots - roots:
            raise ValueError("Orphan comments in merged input")
        target = stage / "social_data.xlsx"
        book.save(target)
        # Read back native identities and exact row count before publishing.
        check = load_workbook(target, read_only=True, data_only=False)
        rows = check["Data"].values
        assert tuple(next(rows)) == COLUMNS
        verified = {identity(dict(zip(COLUMNS, values)))[0] for values in rows}
        check.close()
        assert verified == seen and len(verified) == total
        report = {
            "status": "PASS",
            "output": str(output / target.name),
            "rows": total,
            "roots": len(roots),
            "comments": total - len(roots),
            "duplicate_native_ids": 0,
            "counts": {f"{p}/{k}": n for (p, k), n in sorted(counts.items())},
            "source_workbooks": lineage,
            "output_sha256": sha(target),
            "row_order": "source order, then original worksheet order",
            "source_observations_refreshed": False,
            "synthetic": False,
        }
        os.replace(stage, output)
        provenance.parent.mkdir(parents=True, exist_ok=True)
        provenance.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
        return report
    finally:
        if not sheet.closed:
            sheet.close()
        book.close()
        if stage.exists():
            shutil.rmtree(stage)


def merge_unique(sources, output, provenance):
    """Keep the latest supplied observation per native ID, with full row lineage.

    Source order is explicit and must be chronological. Identical text on different
    native IDs is preserved. Conflicting root identities are rejected; corrected
    comment/reply classifications use the latest observation and are audited.
    """
    sources = [Path(p).resolve() for p in sources]
    output, provenance = Path(output).resolve(), Path(provenance).resolve()
    if not sources or len(set(sources)) != len(sources):
        raise ValueError("Provide distinct source workbooks")
    if output.exists() or provenance.exists():
        raise ValueError("Destination already exists; use a new versioned path")
    output.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix="." + output.name + "-", dir=output.parent))
    winners, lineage, duplicates = {}, [], []
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Data")
    sheet.append(COLUMNS)
    try:
        for source_index, source in enumerate(sources):
            before = sha(source)
            original = load_workbook(source, read_only=True, data_only=False)
            count = 0
            try:
                rows = original["Data"].iter_rows()
                if tuple(cell.value for cell in next(rows)) != COLUMNS:
                    raise ValueError(f"Unexpected columns: {source}")
                for row_number, cells in enumerate(rows, 2):
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError("Formula found in source data")
                    values = tuple(cell.value for cell in cells)
                    row = dict(zip(COLUMNS, values))
                    key, root = identity(row)
                    if key[1] == "comment" and row["teks_dipotong"]:
                        raise ValueError(
                            "Truncated comment cannot become training input"
                        )
                    digest = sha256(
                        json.dumps(values, default=str, ensure_ascii=False).encode()
                    ).hexdigest()
                    previous = winners.get(key)
                    if previous:
                        if previous["root"] != root:
                            raise ValueError("Conflicting native root identity")
                        duplicates.append(
                            {
                                "native_key": key,
                                "discarded_source": str(sources[previous["source"]]),
                                "discarded_row": previous["row"],
                                "retained_source": str(source),
                                "retained_row": row_number,
                                "same_record": previous["sha256"] == digest,
                                "kind_changed": previous["kind"] != row["jenis"],
                            }
                        )
                    winners[key] = {
                        "values": values,
                        "source": source_index,
                        "row": row_number,
                        "root": root,
                        "kind": row["jenis"],
                        "sha256": digest,
                    }
                    count += 1
            finally:
                original.close()
            if sha(source) != before:
                raise ValueError("Source changed during merge")
            lineage.append({"path": str(source), "sha256": before, "input_rows": count})
        roots = {row["root"] for key, row in winners.items() if key[1] == "post"}
        if {row["root"] for key, row in winners.items() if key[1] == "comment"} - roots:
            raise ValueError("Orphan comments in merged input")
        counts, retained = Counter(), Counter()
        with (stage / "record_lineage.jsonl").open("w", encoding="utf-8") as audit:
            for excel_row, (key, row) in enumerate(winners.items(), 2):
                safe = []
                for value in row["values"]:
                    cell = WriteOnlyCell(sheet, value=value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                    safe.append(cell)
                sheet.append(safe)
                counts[(key[0], row["kind"])] += 1
                retained[row["source"]] += 1
                audit.write(
                    json.dumps(
                        {
                            "excel_row": excel_row,
                            "native_key": key,
                            "source": str(sources[row["source"]]),
                            "source_row": row["row"],
                            "record_sha256": row["sha256"],
                        },
                        ensure_ascii=False,
                    )
                    + "\n"
                )
        target = stage / "social_data.xlsx"
        book.save(target)
        check = load_workbook(target, read_only=True, data_only=False)
        try:
            rows = check["Data"].values
            assert tuple(next(rows)) == COLUMNS
            ids = [identity(dict(zip(COLUMNS, row)))[0] for row in rows]
            assert len(ids) == len(set(ids)) == len(winners)
            assert set(ids) == set(winners)
        finally:
            check.close()
        for index, source in enumerate(lineage):
            source["retained_rows"] = retained[index]
        with (stage / "duplicate_records.jsonl").open("w", encoding="utf-8") as audit:
            for row in duplicates:
                audit.write(json.dumps(row, ensure_ascii=False) + "\n")
        report = {
            "status": "PASS",
            "output": str(output / target.name),
            "rows": len(winners),
            "roots": len(roots),
            "comments": len(winners) - len(roots),
            "duplicate_native_ids": 0,
            "duplicates_removed": len(duplicates),
            "duplicate_policy": "prefer_latest_supplied_observation",
            "counts": {f"{p}/{k}": n for (p, k), n in sorted(counts.items())},
            "source_workbooks": lineage,
            "output_sha256": sha(target),
            "row_order": "first native appearance; latest supplied record retained",
            "source_observations_refreshed": False,
            "synthetic": False,
            "files_sha256": {p.name: sha(p) for p in stage.iterdir() if p.is_file()},
        }
        (stage / "merge_report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        )
        os.replace(stage, output)
        provenance.parent.mkdir(parents=True, exist_ok=True)
        provenance.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
        return report
    finally:
        if not sheet.closed:
            sheet.close()
        book.close()
        if stage.exists():
            shutil.rmtree(stage)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", action="append", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--provenance", type=Path, required=True)
    parser.add_argument(
        "--duplicate-policy", choices=["error", "prefer_latest"], default="error"
    )
    args = parser.parse_args()
    print(
        json.dumps(
            merge(
                args.source,
                args.output_dir,
                args.provenance,
                duplicate_policy=args.duplicate_policy,
            ),
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
