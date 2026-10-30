"""Prepare an original-sentence export and seal it into a new NLP snapshot."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from MachineLearning.social_nlp import (
    ROOT,
    TOPIC_SENTENCES_EXTENSION,
    TOPIC_SENTENCES_FILE,
    SocialNLPClusterer,
    file_hash,
    json_write,
)

HEADERS = ("topik", "kalimat yang berada dalam cluster topik")
HEADER_ROW = 5


def export_payload(data):
    """Keep original variants; remove repeated occurrences of the same text."""
    labels = {
        row["cluster_id"]: f"T{row['cluster_id'] + 1:02d} · {row['label']}"
        for row in data["topics"]
    }
    assigned = data["occurrences"].query("cluster_id >= 0")
    if set(assigned.cluster_id.unique()) != set(labels):
        raise ValueError("Export requires every fitted topic to have sentences")
    distinct = assigned.drop_duplicates(["cluster_id", "sentence"]).sort_values(
        ["cluster_id", "source_row", "span_start"], kind="stable"
    )
    rows = [
        [labels[int(row.cluster_id)], row.sentence] for row in distinct.itertuples()
    ]
    if any(not isinstance(text, str) or not text for _, text in rows):
        raise ValueError("Export requires nonempty original sentences")
    counts = distinct.cluster_id.value_counts()
    return {
        "headers": list(HEADERS),
        "rows": rows,
        "topics": [
            {
                "label": labels[row["cluster_id"]],
                "coherence_cv": row["coherence_cv"],
                "sentences": int(counts[row["cluster_id"]]),
            }
            for row in sorted(data["topics"], key=lambda row: row["cluster_id"])
        ],
        "macro_coherence": data["summary"]["coherence"]["macro_mean"],
        "min_coherence": data["summary"]["coherence"]["min_score"],
        "noise_occurrences": int((data["occurrences"].cluster_id < 0).sum()),
        "source_sha256": data["summary"]["source_sha256"],
        "source_url": "https://github.com/SakaHeroji/BeeSense/blob/main/"
        + data["summary"]["config"]["input_file"],
    }


def seal(run, destination):
    """Verify every Excel membership before publishing the complete snapshot."""
    from openpyxl import load_workbook

    run, destination = Path(run), Path(destination)
    if destination.exists():
        raise FileExistsError("Use a new destination for the fitted snapshot")
    data = SocialNLPClusterer.load_run(run)
    payload = export_payload(data)
    workbook = run / TOPIC_SENTENCES_FILE
    original = load_workbook(workbook, read_only=True, data_only=False)
    try:
        sheet = original["Kalimat"]
        rows = sheet.iter_rows(min_row=HEADER_ROW)
        if tuple(cell.value for cell in next(rows)) != HEADERS:
            raise ValueError("Topic sentence Excel headers differ")
        actual = []
        for row in rows:
            if len(row) != 2 or any(cell.data_type in {"f", "e"} for cell in row):
                raise ValueError("Sentence table must contain two literal text columns")
            actual.append([cell.value for cell in row])
        if actual != payload["rows"]:
            raise ValueError("Excel sentences differ from original cluster membership")
    finally:
        original.close()
    summary = data["summary"]
    summary["topic_sentences_export"] = {
        "file": TOPIC_SENTENCES_FILE,
        "rows": len(payload["rows"]),
        "topics": len(payload["topics"]),
        "source_sha256": summary["source_sha256"],
        "noise_excluded": True,
        "deduplication": "Exact original sentence within each final cluster; normalized variants retained",
    }
    json_write(run / "summary.json", summary)
    manifest = json.loads((run / "manifest.json").read_text())
    manifest.setdefault("extensions", []).append(TOPIC_SENTENCES_EXTENSION)
    manifest["files_sha256"].update(
        {
            "summary.json": file_hash(run / "summary.json"),
            TOPIC_SENTENCES_FILE: file_hash(workbook),
        }
    )
    json_write(run / "manifest.json", manifest)
    SocialNLPClusterer.load_run(run)
    destination.parent.mkdir(parents=True, exist_ok=True)
    run.rename(destination)
    return summary["topic_sentences_export"]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=["prepare", "seal"])
    parser.add_argument("run", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if args.action == "prepare":
        data = SocialNLPClusterer.load_run(args.run)
        source = ROOT / data["summary"]["config"]["input_file"]
        if file_hash(source) != data["summary"]["source_sha256"]:
            raise ValueError("Source workbook differs from the fitted snapshot")
        json_write(args.output, export_payload(data))
    else:
        print(json.dumps(seal(args.run, args.output), ensure_ascii=False))


if __name__ == "__main__":
    main()
