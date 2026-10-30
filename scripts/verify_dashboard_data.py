"""Verify published source workbooks and the dashboard artifact bundle."""

from __future__ import annotations

import argparse
import json
import math
import subprocess
from pathlib import Path

from MachineLearning.social_nlp import ROOT, SocialNLPClusterer, file_hash


def verify_native_dataset(path):
    """Verify record identity independently of sentence-text deduplication."""
    from openpyxl import load_workbook

    from scripts.merge_social_datasets import COLUMNS, identity

    book = load_workbook(path, read_only=True, data_only=False)
    seen, roots, parents = set(), set(), set()
    try:
        rows = book["Data"].iter_rows()
        if tuple(cell.value for cell in next(rows)) != COLUMNS:
            raise ValueError(
                "Published dataset columns differ from the crawler contract"
            )
        for cells in rows:
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError("Formula found in published dataset")
            row = dict(zip(COLUMNS, (cell.value for cell in cells)))
            key, root = identity(row)
            if key in seen:
                raise ValueError("Duplicate native record in published dataset")
            seen.add(key)
            (roots if key[1] == "post" else parents).add(root)
    finally:
        book.close()
    if parents - roots:
        raise ValueError("Orphan comment in published dataset")
    return {"rows": len(seen), "roots": len(roots), "duplicate_native_ids": 0}


def published_source_files(source):
    """A merge report binds its small, explicit audit sidecars to the input."""
    allowed = {source}
    report_path = source.with_name("merge_report.json")
    if report_path.exists():
        report = json.loads(report_path.read_text())
        if (
            report.get("output_sha256") != file_hash(source)
            or report.get("duplicate_native_ids") != 0
        ):
            raise ValueError("Published merge report does not match the dataset")
        expected = {
            "social_data.xlsx",
            "record_lineage.jsonl",
            "duplicate_records.jsonl",
        }
        if set(report.get("files_sha256", {})) != expected:
            raise ValueError("Unexpected published merge audit files")
        for name, digest in report["files_sha256"].items():
            path = source.with_name(name)
            if not path.is_file() or file_hash(path) != digest:
                raise ValueError("Published merge audit checksum mismatch")
            allowed.add(path)
        allowed.add(report_path)
    return allowed


def verify(check_git=False):
    config = json.loads((ROOT / "config/social_nlp.json").read_text())
    run = ROOT / config["output_dir"]
    source = ROOT / config["input_file"]
    audit = ROOT / "outputs/social_20261028_01/dataset.xlsx"
    baseline = audit.with_name("social_data.xlsx")
    data = SocialNLPClusterer.load_run(run)
    if file_hash(source) != data["summary"]["source_sha256"]:
        raise ValueError("Excel source checksum differs from the fitted run")
    native_records = verify_native_dataset(source)
    target = config.get("target_topics")
    if target is not None and len(data["topics"]) != target:
        raise ValueError("Active run differs from configured target_topics")
    minimum = config.get("coherence", {}).get("min_score")
    if minimum is not None:
        score = data["summary"].get("coherence", {}).get("macro_mean")
        if score is None or not math.isfinite(score) or score < minimum:
            raise ValueError("Active run fails the configured coherence threshold")
    if (
        config.get("gemini", {}).get("enabled")
        and len(data.get("topic_narratives", [])) != 20
    ):
        raise ValueError("Active run is missing its 20 Gemini topic descriptions")
    manifest = json.loads((run / "manifest.json").read_text())
    runs = {run: manifest}
    sources = {source, baseline}
    # Versioned NLP snapshots are retained across retraining. Crawler output and
    # unfinished staging directories are separate from published NLP bundles.
    for path in (ROOT / "outputs").glob("nlp_*/manifest.json"):
        if path.parent == run:
            continue
        descriptor = json.loads(path.read_text())
        if descriptor.get("schema_version") in {
            "beesense-sentence-clusters:1",
            "beesense-sentence-clusters:2",
        }:
            archived = SocialNLPClusterer.load_run(path.parent)
            archived_source = Path(archived["summary"].get("source_file", ""))
            if archived_source.parts[-3:] == baseline.parts[-3:] and (
                file_hash(baseline) != archived["summary"]["source_sha256"]
            ):
                raise ValueError(
                    "Preserved baseline checksum differs from its fitted run"
                )
            runs[path.parent] = descriptor
            # Source paths in old summaries can belong to another clone. Retain
            # versioned inputs using their repository outputs/<run>/<file> suffix.
            parts = archived_source.parts
            if (
                len(parts) >= 3
                and parts[-3] == "outputs"
                and parts[-1] == "social_data.xlsx"
            ):
                preserved = ROOT.joinpath(*parts[-3:])
                if (
                    not preserved.is_file()
                    or file_hash(preserved) != archived["summary"]["source_sha256"]
                ):
                    raise ValueError("Preserved training source checksum mismatch")
                sources.add(preserved)
    allowed = {
        str((directory / name).relative_to(ROOT))
        for directory, descriptor in runs.items()
        for name in [*descriptor["files_sha256"], "manifest.json"]
    } | {
        str(path.relative_to(ROOT))
        for path in {
            audit,
            *[file for item in sources for file in published_source_files(item)],
        }
    }
    actual = {
        str(path.relative_to(ROOT))
        for directory in {audit.parent, *[item.parent for item in sources], *runs}
        for path in directory.rglob("*")
        if path.is_file()
        and path.name != ".DS_Store"
        and not (path.name.startswith("~$") and path.suffix.lower() == ".xlsx")
    }
    if actual != allowed:
        raise ValueError(
            "Published source workbooks and NLP snapshots contain missing or unexpected files"
        )
    if check_git:
        tracked = set(
            subprocess.check_output(
                ["git", "ls-files", "--", "outputs"], cwd=ROOT, text=True
            ).splitlines()
        )
        if tracked != allowed:
            raise ValueError("Git dataset allowlist differs from the published bundle")
    return {
        "status": "PASS",
        "dataset": config["input_file"],
        "run": config["output_dir"],
        "published_data_files": len(allowed),
        "archived_runs": [str(path.relative_to(ROOT)) for path in runs if path != run],
        "input_sha256": data["summary"]["source_sha256"],
        "audit_dataset": str(audit.relative_to(ROOT)),
        "audit_sha256": file_hash(audit),
        "manifest_sha256": data["artifact_manifest_sha256"],
        "topics": len(data["topics"]),
        "sentence_occurrences": len(data["occurrences"]),
        "unique_sentences": len(data["sentences"]),
        "native_records": native_records,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check-git", action="store_true")
    args = parser.parse_args()
    try:
        print(json.dumps(verify(args.check_git), ensure_ascii=False, indent=2))
    except (OSError, ValueError, KeyError, subprocess.CalledProcessError) as error:
        print(json.dumps({"status": "FAIL", "error": str(error)}))
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
