import json
import random
import tempfile
import unittest
from copy import deepcopy
from datetime import date, datetime
from functools import partial
from hashlib import sha256
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from social_crawler import COMPACT_COLUMNS, SocialCrawler, reddit, youtube

YT1 = "AAAAAAAAAAA"
YT2 = "BBBBBBBBBBB"
YT3 = "CCCCCCCCCCC"
START = "2026-06-01T00:00:00+00:00"
END = "2026-07-01T00:00:00+00:00"


def receipt(status="FINISHED_BEST_EFFORT"):
    return dict(
        status=status,
        reason="synthetic_fixture",
        warnings=[],
        source_version="fixture:1",
    )


def yt_events(vid=YT1, title="Kuliah kampus", count=2, timestamp=START):
    root = dict(
        id=vid,
        title=title,
        description="Informasi beasiswa",
        timestamp=datetime.fromisoformat(timestamp).timestamp() if timestamp else None,
        upload_date="20260601",
    )
    return [("video", root)] + [
        (
            "comment",
            dict(
                id="c" + str(i),
                text="Komentar tanpa keyword " + str(i),
                parent="root",
                _time_text="1 month ago",
            ),
        )
        for i in range(count)
    ]


def rd_events(native="abc", count=2):
    base = dict(
        root="t3_" + native,
        observed_at=START,
        page_url="https://www.reddit.com/comments/" + native + "/",
        extraction_method="synthetic_fixture",
        author_id=None,
        publisher_ref=None,
        subreddit="indonesia",
        text_complete=False,
        flags=[],
        published_at=START,
        published_display=None,
        timestamp_quality="exact",
        score=None,
        comment_count=None,
        parent_evidence=None,
    )
    rows = [
        {
            **base,
            "fullname": "t3_" + native,
            "kind": "post",
            "title": "Kuliah kampus",
            "text": "Isi root",
            "parent": None,
        }
    ]
    rows += [
        {
            **base,
            "fullname": "t1_xyz" + str(i),
            "kind": "comment",
            "title": None,
            "text": "Komentar " + str(i),
            "parent": "t3_" + native,
        }
        for i in range(count)
    ]
    return [("record", row) for row in rows]


def spawn_youtube_fixture(operation, target, settings, pipe):
    for event in yt_events():
        pipe.send(event)
    pipe.send(("done", receipt()))
    pipe.close()


class Clock:
    def __init__(self):
        self.now = 0
        self.sleeps = []

    def clock(self):
        return self.now

    def sleep(self, seconds):
        self.sleeps.append(seconds)
        self.now += seconds


class SocialCrawlerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.clock = Clock()
        self.calls = []
        self.config = dict(
            project_id="social-test",
            since=START,
            until=END,
            include_keywords=["kuliah"],
            exclude_keywords=[],
            delay=dict(mode="fixed", seconds=0),
            total_seconds=100,
            platforms={
                "youtube": dict(
                    urls=["https://youtu.be/" + YT1],
                    min_posts=1,
                    min_comments_per_post=2,
                    max_posts=5,
                )
            },
        )

    def crawler(self, worker, config=None, rng=None):
        def capture(operation, target, settings, seconds, on_event):
            self.calls.append((operation, target))
            events, result = worker(operation, target)
            for kind, payload in events:
                on_event(kind, payload)
            return result

        return SocialCrawler(
            config or self.config,
            workers={"youtube": capture, "reddit": capture},
            sleep=self.clock.sleep,
            monotonic=self.clock.clock,
            rng=rng,
        )

    def read(self, table):
        return [
            json.loads(line)
            for line in (self.directory / "bundle" / (table + ".jsonl"))
            .read_text()
            .splitlines()
        ]

    def test_combined_bundle_uses_original_contract_and_separate_platform_minimums(
        self,
    ):
        cfg = deepcopy(self.config)
        cfg["platforms"]["reddit"] = dict(
            backend="html",
            urls=["https://www.reddit.com/comments/abc/"],
            min_posts=1,
            min_comments_per_post=2,
        )
        crawler = self.crawler(
            lambda op, target: (
                yt_events() if op == "video" else rd_events(),
                receipt(),
            ),
            cfg,
        )
        result = crawler.run(self.directory / "bundle")
        self.assertEqual(result["status"], "TARGET_MET")
        self.assertEqual(
            {name: s["qualified_posts"] for name, s in result["platforms"].items()},
            {"youtube": 1, "reddit": 1},
        )
        manifest = json.loads((self.directory / "bundle/manifest.json").read_text())
        self.assertTrue(manifest["synthetic"])
        for name, digest in manifest["files_sha256"].items():
            self.assertEqual(
                sha256((self.directory / "bundle" / name).read_bytes()).hexdigest(),
                digest,
            )
        book = load_workbook(self.directory / "bundle/dataset.xlsx", read_only=True)
        self.assertEqual(book.sheetnames, list(reddit.TABLES))
        for table in reddit.TABLES:
            rows = self.read(table)
            for row in rows:
                reddit.validate_row(table, row)
            self.assertEqual(
                next(book[table].values),
                tuple(field["name"] for field in reddit.TABLES[table]),
            )
        book.close()
        self.assertEqual(len(self.read("Contents")), 6)
        self.assertEqual(len(self.read("Relations")), 8)

    def test_date_window_is_inclusive_start_exclusive_end_and_unknown_is_excluded(self):
        self.config["platforms"]["youtube"]["urls"] = [
            "https://youtu.be/" + vid for vid in [YT2, YT3, YT1]
        ]

        def worker(op, url):
            vid = url.split("v=")[1]
            return yt_events(
                vid, timestamp=END if vid == YT2 else None if vid == YT3 else START
            ), receipt()

        result = self.crawler(worker).run(self.directory / "bundle")
        self.assertEqual(
            result["platforms"]["youtube"]["filtered_roots"],
            {"root_outside_date_window": 1, "root_publication_date_unknown": 1},
        )
        self.assertEqual(
            [
                row["native_id"]
                for row in self.read("Contents")
                if row["object_type"] == "post"
            ],
            [YT1],
        )

    def test_include_all_across_title_and_description_and_exclude_wins(self):
        self.config.update(
            include_keywords=["kuliah", "beasiswa"],
            include_mode="all",
            exclude_keywords=["gacor"],
        )
        self.config["platforms"]["youtube"]["urls"] = [
            "https://youtu.be/" + YT2,
            "https://youtu.be/" + YT1,
        ]

        def worker(op, url):
            vid = url.split("v=")[1]
            return yt_events(
                vid, title="KULIAH gacor" if vid == YT2 else "KULIAH"
            ), receipt()

        result = self.crawler(worker).run(self.directory / "bundle")
        self.assertTrue(result["targets_met"])
        self.assertEqual(
            result["platforms"]["youtube"]["filtered_roots"],
            {"root_excluded_keyword": 1},
        )

    def test_comment_exclusion_preserves_child_with_unresolved_parent_and_context_comments(
        self,
    ):
        self.config["exclude_keywords"] = ["gacor"]
        events = yt_events(count=0) + [
            ("comment", dict(id="parent", text="gacor", parent="root")),
            ("comment", dict(id="child", text="Balasan biasa", parent="parent")),
            ("comment", dict(id="other", text="Komentar konteks", parent="root")),
        ]
        result = self.crawler(lambda *args: (events, receipt())).run(
            self.directory / "bundle"
        )
        self.assertTrue(result["targets_met"])
        self.assertEqual(len(self.read("Contents")), 3)
        self.assertEqual(
            sum(
                row["resolution_status"] == "unresolved_reference"
                for row in self.read("Relations")
            ),
            1,
        )
        self.assertEqual(
            sum(
                row["comments_seen"] + row["replies_seen"] for row in self.read("Runs")
            ),
            3,
        )

    def test_minimum_counts_only_roots_with_enough_unique_comments(self):
        self.config["platforms"]["youtube"]["urls"] = [
            "https://youtu.be/" + YT2,
            "https://youtu.be/" + YT1,
        ]

        def worker(op, url):
            vid = url.split("v=")[1]
            events = yt_events(vid, count=0 if vid == YT2 else 2)
            return events + [events[-1]], receipt()

        result = self.crawler(worker).run(self.directory / "bundle")["platforms"][
            "youtube"
        ]
        self.assertEqual(
            (
                result["posts_exported"],
                result["qualified_posts"],
                result["insufficient_comment_posts"],
            ),
            (2, 1, 1),
        )
        self.assertEqual(result["comments"], 2)

    def test_minimum_shortage_is_reported_at_hard_post_cap(self):
        self.config["platforms"]["youtube"].update(min_posts=2, max_posts=2)
        result = self.crawler(lambda *args: (yt_events(count=0), receipt())).run(
            self.directory / "bundle"
        )
        self.assertFalse(result["targets_met"])
        self.assertEqual(result["platforms"]["youtube"]["posts_shortfall"], 2)
        self.assertEqual(result["status"], "PARTIAL")

    def test_discovery_refills_after_a_filtered_candidate_without_repeating_urls(self):
        p = self.config["platforms"]["youtube"]
        p.update(urls=[], queries=["first", "second"])

        def worker(op, target):
            if op == "search":
                return [
                    (
                        "candidate",
                        {
                            "url": "https://youtu.be/"
                            + (YT2 if target == "first" else YT1),
                            "discovered_from": "https://www.youtube.com/results",
                        },
                    )
                ], receipt()
            vid = target.split("v=")[1]
            return yt_events(
                vid, title="Other topic" if vid == YT2 else "kuliah"
            ), receipt()

        result = self.crawler(worker).run(self.directory / "bundle")
        self.assertTrue(result["targets_met"])
        self.assertEqual(
            [op for op, target in self.calls], ["search", "video", "search", "video"]
        )

    def test_random_and_fixed_delays_are_between_operations(self):
        for delay in [
            dict(mode="fixed", seconds=2),
            dict(mode="random", min_seconds=2, max_seconds=5),
        ]:
            cfg = deepcopy(self.config)
            cfg["delay"] = delay
            cfg["platforms"]["youtube"]["urls"] = [
                "https://youtu.be/" + YT2,
                "https://youtu.be/" + YT1,
            ]
            clock = Clock()
            crawler = self.crawler(
                lambda op, url: (
                    yt_events(
                        url.split("v=")[1], title="other" if YT2 in url else "kuliah"
                    ),
                    receipt(),
                ),
                cfg,
                random.Random(7),
            )
            crawler.clock = clock.clock
            crawler.sleep = clock.sleep
            crawler.run(
                self.directory / ("fixed" if delay["mode"] == "fixed" else "random")
            )
            self.assertEqual(len(clock.sleeps), 1)
            self.assertTrue(2 <= clock.sleeps[0] <= 5)
            if delay["mode"] == "fixed":
                self.assertEqual(clock.sleeps[0], 2)

    def test_delay_cannot_start_another_request_after_global_budget(self):
        self.config.update(total_seconds=1, delay=dict(mode="fixed", seconds=2))
        self.config["platforms"]["youtube"]["urls"] = [
            "https://youtu.be/" + YT2,
            "https://youtu.be/" + YT1,
        ]
        result = self.crawler(
            lambda op, url: (yt_events(YT2, title="other"), receipt())
        ).run(self.directory / "bundle")
        self.assertEqual(len(self.calls), 1)
        self.assertEqual(
            result["platforms"]["youtube"]["stop_reason"], "TIME_BUDGET_EXHAUSTED"
        )

    def test_barrier_stops_platform_and_keeps_partial_bundle(self):
        self.config["platforms"] = {
            "reddit": dict(
                backend="html",
                urls=[
                    "https://www.reddit.com/comments/abc/",
                    "https://www.reddit.com/comments/def/",
                ],
            )
        }
        result = self.crawler(lambda *args: ([], receipt("CHALLENGE_REQUIRED"))).run(
            self.directory / "bundle"
        )
        self.assertEqual(len(self.calls), 1)
        self.assertFalse(result["targets_met"])
        self.assertEqual(
            result["platforms"]["reddit"]["stop_reason"], "CHALLENGE_REQUIRED"
        )
        self.assertEqual(self.read("Contents"), [])

    def test_partial_metadata_survives_a_worker_timeout(self):
        def worker(*args):
            self.clock.now += 101
            return yt_events(count=0), receipt("TIMEOUT")

        result = self.crawler(worker).run(self.directory / "bundle")
        self.assertEqual(len(self.read("Contents")), 1)
        self.assertEqual(
            result["platforms"]["youtube"]["stop_reason"], "TIME_BUDGET_EXHAUSTED"
        )

    def test_interrupt_during_delay_publishes_already_captured_evidence(self):
        self.config["delay"] = dict(mode="fixed", seconds=2)
        self.config["platforms"]["youtube"].update(
            min_posts=2, urls=["https://youtu.be/" + YT1, "https://youtu.be/" + YT2]
        )
        crawler = self.crawler(
            lambda op, url: (yt_events(url.split("v=")[1]), receipt())
        )

        def interrupt(seconds):
            raise KeyboardInterrupt

        crawler.sleep = interrupt
        result = crawler.run(self.directory / "bundle")
        self.assertEqual(result["status"], "INTERRUPTED")
        self.assertEqual(len(self.calls), 1)
        self.assertEqual(len(self.read("Contents")), 3)

    def test_invalid_native_identity_does_not_become_a_success(self):
        result = self.crawler(lambda *args: (yt_events(YT2), receipt())).run(
            self.directory / "bundle"
        )
        self.assertFalse(result["targets_met"])
        self.assertEqual(result["platforms"]["youtube"]["invalid_rows"], 1)
        self.assertEqual(self.read("Contents"), [])

    def test_source_exception_is_sanitized_and_stops_after_three_calls(self):
        self.config["platforms"]["youtube"]["urls"] = [
            "https://youtu.be/" + vid for vid in [YT1, YT2, YT3]
        ]

        def worker(*args):
            raise RuntimeError("PRIVATE_ERROR_TOKEN")

        result = self.crawler(worker).run(self.directory / "bundle")
        self.assertEqual(
            result["platforms"]["youtube"]["stop_reason"], "CONSECUTIVE_SOURCE_ERRORS"
        )
        self.assertNotIn(
            "PRIVATE_ERROR_TOKEN", (self.directory / "bundle/Runs.jsonl").read_text()
        )

    def test_browser_navigation_error_checks_current_page_and_stops_on_429(self):
        self.config["platforms"] = {
            "reddit": dict(
                backend="playwright",
                urls=[
                    "https://www.reddit.com/comments/abc/",
                    "https://www.reddit.com/comments/def/",
                ],
            )
        }
        crawler = SocialCrawler(
            self.config, sleep=self.clock.sleep, monotonic=self.clock.clock
        )
        inspections = []

        def inspect():
            inspections.append(True)
            return dict(
                page_title="www.reddit.com", post=False, barrier_text="HTTP ERROR 429"
            )

        def fail(*args):
            crawler._browser = SimpleNamespace(inspect=inspect)
            raise RuntimeError("PRIVATE_BROWSER_TOKEN")

        with patch.object(crawler, "_browser_call", side_effect=fail) as call:
            result = crawler.run(self.directory / "bundle")
        self.assertEqual(call.call_count, 1)
        self.assertEqual(len(inspections), 1)
        self.assertEqual(result["platforms"]["reddit"]["stop_reason"], "RATE_LIMITED")
        self.assertEqual(
            self.read("Runs")[0]["reason"], "browser_error_page_access_barrier"
        )
        self.assertNotIn(
            "PRIVATE_BROWSER_TOKEN", (self.directory / "bundle/Runs.jsonl").read_text()
        )

    def test_checkpoint_dom_import_does_not_make_network_calls(self):
        support = self.directory / "checkpoint"
        (support / "raw").mkdir(parents=True)
        candidate = dict(
            native_id="abc",
            subreddit="indonesia",
            canonical_url="https://www.reddit.com/r/indonesia/comments/abc/kuliah/",
        )
        (support / "candidates.json").write_text(json.dumps([candidate]))
        raw = dict(
            extractor_version="browser-dom:2",
            captured_at=START,
            post=dict(
                native_id="abc",
                title="Kuliah",
                text="Isi",
                published_at=START,
                permalink="/r/indonesia/comments/abc/kuliah/",
                subreddit="Indonesia",
            ),
            comments=[
                dict(
                    url=f"https://www.reddit.com/r/indonesia/comments/abc/comment/xyz{i}/",
                    text="Komentar",
                    parent_url=None,
                    published_at=START,
                )
                for i in range(2)
            ],
        )
        (support / "raw/abc.json").write_text(json.dumps(raw))
        self.config["platforms"] = {
            "reddit": dict(
                backend="checkpoint",
                import_checkpoint_dir=str(support),
                min_comments_per_post=2,
            )
        }
        result = self.crawler(
            lambda *args: (_ for _ in ()).throw(AssertionError("network called"))
        ).run(self.directory / "bundle")
        self.assertTrue(result["targets_met"])
        self.assertEqual(self.calls, [])
        self.assertEqual(len(self.read("Contents")), 3)

    def test_original_youtube_process_worker_boundary_works_as_package_import(self):
        crawler = SocialCrawler(
            self.config,
            workers={
                "youtube": partial(
                    youtube.run_worker, worker_entry=spawn_youtube_fixture
                )
            },
            sleep=self.clock.sleep,
            monotonic=self.clock.clock,
        )
        self.assertTrue(crawler.run(self.directory / "bundle")["targets_met"])

    def test_invalid_config_is_rejected_before_output_or_requests(self):
        cases = [
            {"since": "2026-01-01"},
            {"since": 1},
            {"since": END, "until": START},
            {"total_seconds": True},
            {"delay": {"mode": "random", "min_seconds": 5, "max_seconds": 2}},
            {"delay": {"mode": "fixed", "seconds": float("nan")}},
            {"exclude_keywords": [""]},
            {"platforms": {"youtube": {"min_posts": 3, "max_posts": 1}}},
            {
                "platforms": {
                    "youtube": {"min_posts": 3, "max_posts": 5, "max_candidates": 2}
                }
            },
            {"platforms": {"reddit": {"backend": "unsupported"}}},
            {"platforms": {"reddit": {"subreddits": ["../../bad"]}}},
            {"platforms": {"youtube": {"unknown_option": 1}}},
        ]
        for changed in cases:
            with self.subTest(changed=changed), self.assertRaises(ValueError):
                SocialCrawler({**self.config, **changed})

    def test_historical_posts_are_skipped_before_requests_and_ids_are_canonical(self):
        baseline = self.directory / "baseline"
        self.crawler(lambda *args: (yt_events(), receipt())).run(baseline)
        checksum = sha256((baseline / "dataset.xlsx").read_bytes()).hexdigest()
        cfg = deepcopy(self.config)
        cfg["exclude_existing_datasets"] = [str(baseline)]
        cfg["platforms"]["youtube"]["urls"] = [
            "https://www.youtube.com/shorts/" + YT1,
            "https://youtu.be/" + YT2,
        ]
        self.calls = []
        result = self.crawler(lambda op, url: (yt_events(YT2), receipt()), cfg).run(
            self.directory / "bundle"
        )
        self.assertEqual(len(self.calls), 1)
        self.assertIn(YT2, self.calls[0][1])
        self.assertEqual(result["platforms"]["youtube"]["existing_posts_skipped"], 1)
        self.assertEqual(result["deduplication_sources"][0]["sha256"], checksum)
        self.assertEqual(
            sha256((baseline / "dataset.xlsx").read_bytes()).hexdigest(), checksum
        )
        self.assertEqual(
            [
                r["native_id"]
                for r in self.read("Contents")
                if r["object_type"] == "post"
            ],
            [YT2],
        )

    def test_redirect_to_historical_root_is_not_exported(self):
        baseline = self.directory / "baseline"
        self.crawler(lambda *args: (yt_events(), receipt())).run(baseline)
        cfg = deepcopy(self.config)
        cfg["exclude_existing_datasets"] = [str(baseline)]
        cfg["platforms"]["youtube"]["urls"] = ["https://youtu.be/" + YT2]
        result = self.crawler(lambda *args: (yt_events(), receipt()), cfg).run(
            self.directory / "bundle"
        )
        self.assertEqual(self.read("Contents"), [])
        self.assertEqual(
            result["platforms"]["youtube"]["filtered_roots"], {"root_not_verified": 1}
        )
        self.assertGreater(result["platforms"]["youtube"]["invalid_rows"], 0)

    def test_round_robin_gives_both_platforms_turns_and_shares_deadline(self):
        cfg = deepcopy(self.config)
        cfg.update(platform_execution="round_robin", total_seconds=3)
        cfg["platforms"]["youtube"].update(
            min_posts=2, urls=["https://youtu.be/" + YT1, "https://youtu.be/" + YT2]
        )
        cfg["platforms"]["reddit"] = dict(
            backend="html",
            min_posts=2,
            urls=[
                "https://www.reddit.com/comments/abc/",
                "https://www.reddit.com/comments/def/",
            ],
        )

        def worker(op, url):
            self.clock.now += 1
            return (
                yt_events(YT1 if YT1 in url else YT2)
                if op == "video"
                else rd_events("abc" if "abc" in url else "def")
            ), receipt()

        result = self.crawler(worker, cfg).run(self.directory / "bundle")
        self.assertEqual([op for op, url in self.calls], ["video", "thread", "video"])
        self.assertEqual(result["platforms"]["reddit"]["qualified_posts"], 1)
        self.assertEqual(
            result["platforms"]["reddit"]["stop_reason"], "TIME_BUDGET_EXHAUSTED"
        )
        self.assertEqual(result["elapsed_seconds"], 3)
        progress = json.loads((self.directory / "bundle.progress.json").read_text())
        self.assertEqual(progress["stage"], "finished")
        self.assertEqual(progress["remaining_seconds"], 0)

    def test_reddit_cooldown_leaves_youtube_runnable_and_retries_same_thread(self):
        cfg = deepcopy(self.config)
        cfg.update(platform_execution="round_robin", total_seconds=10)
        cfg["platforms"]["youtube"].update(
            min_posts=2, urls=["https://youtu.be/" + YT1, "https://youtu.be/" + YT2]
        )
        cfg["platforms"]["reddit"] = dict(
            backend="html",
            urls=["https://www.reddit.com/comments/abc/"],
            min_comments_per_post=2,
            rate_limit_cooldown_seconds=4,
            rate_limit_max_retries=1,
        )
        cfg["platforms"] = {
            name: cfg["platforms"][name] for name in ["reddit", "youtube"]
        }
        attempts = []

        def worker(op, target):
            attempts.append((op, target, self.clock.now))
            self.clock.now += 1
            if len(attempts) == 1:
                return [], {**receipt("RATE_LIMITED"), "retry_after_seconds": 2}
            if op == "thread":
                return rd_events(), receipt()
            return yt_events(YT1 if YT1 in target else YT2), receipt()

        result = self.crawler(worker, cfg).run(self.directory / "bundle")
        self.assertEqual(result["status"], "TARGET_MET")
        self.assertEqual(
            [a[0] for a in attempts], ["thread", "video", "video", "thread"]
        )
        self.assertEqual(attempts[0][1], attempts[-1][1])
        self.assertEqual([a[2] for a in attempts], [0, 1, 2, 5])
        self.assertEqual(result["platforms"]["reddit"]["rate_limit_retries"], 1)
        self.assertEqual(sum(self.clock.sleeps), 2)
        self.assertEqual(
            len([r for r in self.read("Contents") if r["object_type"] == "post"]), 3
        )

    def test_reddit_cooldown_honors_longer_server_wait_and_retry_cap(self):
        cfg = deepcopy(self.config)
        cfg["total_seconds"] = 30
        cfg["platforms"] = {
            "reddit": dict(
                backend="html",
                urls=["https://www.reddit.com/comments/abc/"],
                rate_limit_cooldown_seconds=4,
                rate_limit_max_retries=2,
            )
        }
        attempts = []

        def worker(op, target):
            attempts.append(self.clock.now)
            return [], {**receipt("RATE_LIMITED"), "retry_after_seconds": 6.5}

        result = self.crawler(worker, cfg).run(self.directory / "bundle")
        self.assertEqual(attempts, [0, 7, 14])
        self.assertEqual(result["platforms"]["reddit"]["rate_limit_events"], 3)
        self.assertEqual(result["platforms"]["reddit"]["rate_limit_retries"], 2)
        self.assertEqual(result["platforms"]["reddit"]["stop_reason"], "RATE_LIMITED")
        self.assertEqual(len(self.read("Runs")), 3)

    def test_reddit_cooldown_never_waits_or_retries_beyond_global_deadline(self):
        cfg = deepcopy(self.config)
        cfg["total_seconds"] = 5
        cfg["platforms"] = {
            "reddit": dict(
                backend="html",
                urls=["https://www.reddit.com/comments/abc/"],
                rate_limit_cooldown_seconds=6,
                rate_limit_max_retries=3,
            )
        }
        result = self.crawler(lambda *args: ([], receipt("RATE_LIMITED")), cfg).run(
            self.directory / "bundle"
        )
        self.assertEqual(len(self.calls), 1)
        self.assertEqual(self.clock.sleeps, [])
        self.assertEqual(result["platforms"]["reddit"]["stop_reason"], "RATE_LIMITED")

    def test_reddit_cooldown_retries_failed_search_without_losing_query(self):
        cfg = deepcopy(self.config)
        cfg["platforms"] = {
            "reddit": dict(
                backend="html",
                queries=["kuliah"],
                rate_limit_cooldown_seconds=3,
                rate_limit_max_retries=1,
            )
        }

        def worker(op, target):
            if len(self.calls) == 1:
                return [], receipt("RATE_LIMITED")
            if op == "search":
                return [
                    ("candidate", {"url": "https://www.reddit.com/comments/abc/"})
                ], receipt()
            return rd_events(), receipt()

        result = self.crawler(worker, cfg).run(self.directory / "bundle")
        self.assertEqual(result["status"], "TARGET_MET")
        self.assertEqual([op for op, _ in self.calls], ["search", "search", "thread"])
        self.assertEqual(self.calls[0], self.calls[1])
        self.assertEqual(sum(self.clock.sleeps), 3)

    def test_reddit_cooldown_rejects_invalid_config_values(self):
        for key, values in {
            "rate_limit_cooldown_seconds": [0, 10801, True, 2.5],
            "rate_limit_max_retries": [-1, 11, True, 1.5],
        }.items():
            for value in values:
                with self.subTest(key=key, value=value):
                    cfg = deepcopy(self.config)
                    cfg["platforms"]["reddit"] = {key: value}
                    with self.assertRaises(ValueError):
                        SocialCrawler(cfg)

    def test_unrepresentable_server_cooldown_preserves_partial_bundle(self):
        cfg = deepcopy(self.config)
        cfg["platforms"] = {
            "reddit": dict(
                backend="html",
                urls=["https://www.reddit.com/comments/abc/"],
                rate_limit_max_retries=3,
            )
        }
        result = self.crawler(
            lambda *args: (
                [],
                {**receipt("RATE_LIMITED"), "retry_after_seconds": 10**500},
            ),
            cfg,
        ).run(self.directory / "bundle")
        self.assertEqual(result["status"], "PARTIAL")
        self.assertEqual(len(self.calls), 1)
        self.assertIsNone(result["platforms"]["reddit"]["rate_limit_cooldown_until"])
        self.assertTrue((self.directory / "bundle/manifest.json").exists())

    def test_manual_captcha_wait_only_observes_and_obeys_global_deadline(self):
        crawler = self.crawler(lambda *args: ([], receipt()))
        crawler._browser_deadline = 100
        crawler._run_deadline = 6
        state = dict(
            page_title="Reddit",
            barrier_text="Verify that you are human",
            post=False,
            page_url="https://www.reddit.com/",
        )
        calls = []

        def inspect():
            calls.append(self.clock.now)
            return state

        crawler._browser = SimpleNamespace(inspect=inspect)
        returned = crawler._wait_for_captcha(
            state, dict(captcha_wait_seconds=600, captcha_poll_seconds=5)
        )
        self.assertEqual(returned, state)
        self.assertEqual(self.clock.now, 6)
        self.assertEqual(calls, [5])
        self.clock.now = 0
        crawler._browser = SimpleNamespace(
            inspect=lambda: dict(post=True, page_title="Reddit", barrier_text="")
        )
        returned = crawler._wait_for_captcha(
            state, dict(captcha_wait_seconds=600, captcha_poll_seconds=2)
        )
        self.assertTrue(returned["post"])
        self.assertEqual(self.clock.now, 2)

    def test_default_uses_one_json_and_plan_does_not_access_sources(self):
        cfg = deepcopy(self.config)
        cfg["platforms"]["reddit"] = {"enabled": False}
        (self.directory / "social_crawler.json").write_text(json.dumps(cfg))
        with patch("social_crawler.ROOT", self.directory):
            plan = SocialCrawler().plan()
        self.assertEqual(set(plan["config"]["platforms"]), {"youtube", "reddit"})
        self.assertFalse(plan["live_access_verified"])

    def test_expanded_youtube_search_budget_reaches_worker_without_old_cli_cap(self):
        self.config["platforms"]["youtube"].update(
            urls=[], queries=["kuliah"], search_results=500
        )
        observed = []

        def worker(operation, target, settings, seconds, on_event):
            observed.append(settings.search_results)
            events = (
                [("candidate", dict(url="https://youtu.be/" + YT1))]
                if operation == "search"
                else yt_events()
            )
            for kind, payload in events:
                on_event(kind, payload)
            return receipt()

        result = SocialCrawler(self.config, workers={"youtube": worker}).run(
            self.directory / "bundle"
        )
        self.assertTrue(result["targets_met"])
        self.assertEqual(observed, [500, 500])

    def test_candidate_cap_and_post_cap_are_independent_when_minimum_fits(self):
        self.config["platforms"]["youtube"].update(max_posts=500, max_candidates=100)
        result = self.crawler(lambda *args: (yt_events(), receipt())).run(
            self.directory / "bundle"
        )
        self.assertTrue(result["targets_met"])

    def test_single_sheet_compact_export_repeats_root_date_and_keeps_audit_dates(self):
        cfg = deepcopy(self.config)
        cfg["platforms"]["reddit"] = dict(
            backend="html",
            urls=["https://www.reddit.com/comments/abc/"],
            min_comments_per_post=2,
        )
        crawler = self.crawler(
            lambda op, target: (
                yt_events() if op == "video" else rd_events(),
                receipt(),
            ),
            cfg,
        )
        report = crawler.run(self.directory / "bundle")
        self.assertEqual(report["compact"]["rows"], 6)
        self.assertEqual(
            report["compact"]["date_accuracy_counts"], {"pasti": 4, "perkiraan": 2}
        )
        book = load_workbook(report["primary_file"], read_only=True)
        self.assertEqual(book.sheetnames, ["Data"])
        rows = list(book["Data"].values)
        self.assertEqual(rows[0], COMPACT_COLUMNS)
        records = [dict(zip(COMPACT_COLUMNS, row)) for row in rows[1:]]
        for row in records:
            self.assertEqual(row["tanggal_post_wib"], datetime(2026, 6, 1, 7))
            self.assertEqual(row["akurasi_tanggal_post"], "pasti")
            self.assertTrue(row["judul_post"])
        for row in records:
            if row["platform"] == "youtube" and row["jenis"] == "comment":
                self.assertIsInstance(row["tanggal_konten_wib"], datetime)
                self.assertEqual(row["akurasi_tanggal_konten"], "perkiraan")
        self.assertTrue(
            all(
                row["published_at"] is None
                for row in self.read("Contents")
                if row["platform"] == "youtube" and row["object_type"] == "comment"
            )
        )
        book.close()

    def test_relative_calendar_dates_use_source_capture_timezone_and_edit_suffix(self):
        base = dict(
            published_at=None,
            timestamp_quality="relative",
            observed_at="2026-10-27T18:30:00Z",
        )
        for display, expected in [
            ("10 days ago", date(2026, 10, 18)),
            ("10 hari yang lalu", date(2026, 10, 18)),
            ("3 weeks ago (edited)", date(2026, 10, 7)),
            ("kemarin", date(2026, 10, 27)),
            ("2 hours ago", date(2026, 10, 27)),
            ("2 tahun lalu", date(2024, 10, 28)),
        ]:
            with self.subTest(display=display):
                self.assertEqual(
                    SocialCrawler._calendar_date(
                        {**base, "published_at_display": display}
                    ),
                    (expected, "perkiraan"),
                )

    def test_calendar_month_boundary_date_only_and_unrecognized_dates(self):
        row = dict(
            published_at=None,
            timestamp_quality="relative",
            observed_at="2024-03-31T00:00:00Z",
            published_at_display="1 month ago",
        )
        self.assertEqual(
            SocialCrawler._calendar_date(row), (date(2024, 2, 29), "perkiraan")
        )
        self.assertEqual(
            SocialCrawler._calendar_date(
                {
                    **row,
                    "timestamp_quality": "date_only",
                    "published_at_display": "20260601",
                }
            ),
            (date(2026, 6, 1), "tanggal_saja"),
        )
        self.assertEqual(
            SocialCrawler._calendar_date(
                {**row, "published_at_display": "a while ago"}
            ),
            (None, "tidak_diketahui"),
        )
        self.assertEqual(
            SocialCrawler._calendar_date({**row, "observed_at": None}),
            (None, "tidak_diketahui"),
        )
        self.assertEqual(
            SocialCrawler._calendar_date(
                {**row, "published_at_display": "99999999999999999 years ago"}
            ),
            (None, "tidak_diketahui"),
        )

    def test_compact_reexport_refilters_saved_data_and_keeps_source_unchanged(self):
        events = yt_events(count=0) + [
            (
                "comment",
                dict(
                    id="political",
                    text="DPR soal kampus",
                    parent="root",
                    _time_text="10 days ago",
                ),
            ),
            (
                "comment",
                dict(
                    id="education",
                    text="Info kuliah",
                    parent="root",
                    _time_text="10 days ago",
                ),
            ),
        ]
        self.crawler(lambda *args: (events, receipt())).run(self.directory / "bundle")
        source = self.directory / "bundle"
        before = {
            name: sha256((source / name).read_bytes()).hexdigest()
            for name in ["Contents.jsonl", "manifest.json", "social_data.xlsx"]
        }
        cfg = deepcopy(self.config)
        cfg["exclude_keywords"] = ["dpr"]
        crawler = SocialCrawler(
            cfg,
            workers={
                "youtube": lambda *args: (_ for _ in ()).throw(
                    AssertionError("network called")
                )
            },
        )
        report = crawler.export_compact(source, self.directory / "view")
        self.assertEqual(report["compact"]["rows"], 2)
        self.assertEqual(report["compact"]["filtered_comments"], 1)
        self.assertFalse(report["source_observations_refreshed"])
        self.assertTrue(report["source_synthetic"])
        for name, digest in before.items():
            self.assertEqual(sha256((source / name).read_bytes()).hexdigest(), digest)
        with self.assertRaises(FileExistsError):
            crawler.export_compact(source, self.directory / "view")
        (source / "Contents.jsonl").write_text(
            (source / "Contents.jsonl").read_text() + "\n"
        )
        with self.assertRaisesRegex(ValueError, "Checksum"):
            crawler.export_compact(source, self.directory / "corrupt")
        self.assertFalse((self.directory / "corrupt").exists())

    def test_topic_exclusions_remove_roots_and_comments_without_excluding_normal_diplomas(
        self,
    ):
        terms = [
            "mbg",
            "makan bergizi gratis",
            "kopdes",
            "koperasi desa merah putih",
            "ijazah jokowi",
            "ijazah gibran",
            "dpr",
        ]
        cfg = deepcopy(self.config)
        cfg["exclude_keywords"] = terms
        cfg["platforms"]["youtube"]["urls"] = [
            "https://youtu.be/" + YT2,
            "https://youtu.be/" + YT1,
        ]
        for term in terms:
            with self.subTest(term=term):

                def worker(op, url):
                    vid = url.split("v=")[1]
                    events = yt_events(
                        vid,
                        title="Kuliah " + term.upper()
                        if vid == YT2
                        else "Kuliah dan ijazah S1",
                    )
                    events.append(
                        (
                            "comment",
                            dict(
                                id="filtered",
                                text="Bahas " + term.upper(),
                                parent="root",
                            ),
                        )
                    )
                    return events, receipt()

                crawler = self.crawler(worker, cfg)
                report = crawler.run(
                    self.directory / ("exclude_" + term.replace(" ", "_"))
                )
                self.assertEqual(
                    report["platforms"]["youtube"]["filtered_roots"],
                    {"root_excluded_keyword": 1},
                )
                self.assertEqual(report["platforms"]["youtube"]["comments"], 2)
                self.assertEqual(report["compact"]["rows"], 3)

    def test_compact_long_unicode_text_is_marked_and_formula_text_stays_literal(self):
        events = yt_events(
            title='=HYPERLINK("https://example.test","kuliah")', count=0
        ) + [
            ("comment", dict(id="long", text="😀" * 20000, parent="root")),
            ("comment", dict(id="formula", text="=1+1", parent="root")),
        ]
        report = self.crawler(lambda *args: (events, receipt())).run(
            self.directory / "bundle"
        )
        self.assertEqual(report["compact"]["truncated_texts"], 1)
        self.assertEqual(
            len(
                next(
                    row for row in self.read("Contents") if row["native_id"] == "long"
                )["text_source"]
            ),
            20000,
        )
        book = load_workbook(report["primary_file"], read_only=True)
        rows = list(book["Data"])
        for row in rows:
            self.assertTrue(all(cell.data_type != "f" for cell in row))
        self.assertTrue(rows[2][-1].value)
        text_index = COMPACT_COLUMNS.index("teks")
        self.assertIn("Contents.jsonl", rows[2][text_index].value)
        self.assertEqual(rows[3][text_index].value, "=1+1")
        book.close()


if __name__ == "__main__":
    unittest.main()
