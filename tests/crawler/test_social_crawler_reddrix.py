"""Exercise the default Reddrix route with YouTube and the shared export pipeline."""

import asyncio
import json
import tempfile
import unittest
from copy import deepcopy
from dataclasses import asdict
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from test_reddit_crawler import CONFIG, URL, FakeReddrix, records
from test_social_crawler import YT1, Clock, receipt, yt_events

from Crawler.Backend.social_crawler import (
    REDDRIX_METHOD,
    SocialCrawler,
    collect_reddrix,
    reddit,
    reddrix_error_receipt,
    reddrix_worker,
    run_reddrix_worker,
)


class FixtureReddrix(FakeReddrix):
    items = records()


def spawn_reddrix_fixture(operation, target, settings, pipe):
    # Patch only the external library boundary in the spawned test process.
    with patch(
        "Crawler.Backend.social_crawler.load_reddrix", return_value=FixtureReddrix
    ):
        reddrix_worker(operation, target, settings, pipe, display_mode="headless")


class SocialCrawlerReddrixTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        self.clock = Clock()
        self.config = deepcopy(CONFIG)
        self.config["platform_execution"] = "round_robin"
        self.config["platforms"]["youtube"] = {
            "urls": ["https://youtu.be/" + YT1],
            "min_comments_per_post": 2,
        }
        self.config["reddrix"] = {"display_mode": "headless"}

    @staticmethod
    def youtube_worker(operation, target, settings, seconds, emit):
        for kind, row in yt_events():
            emit(kind, row)
        return receipt()

    def runner(self, operation, target, settings, seconds, emit, *, worker_entry):
        self.assertIs(worker_entry.func, reddrix_worker)
        self.assertEqual(worker_entry.keywords, {"display_mode": "headless"})
        result = asyncio.run(
            collect_reddrix(
                operation, target, asdict(settings), emit, factory=FixtureReddrix
            )
        )
        return {**result, "source_version": "fixture:reddrix"}

    def crawler(self, config=None):
        return SocialCrawler(
            self.config if config is None else config,
            workers={"youtube": self.youtube_worker},
            sleep=self.clock.sleep,
            monotonic=self.clock.clock,
        )

    def test_default_config_selects_reddrix_without_loading_browser(self):
        with patch(
            "Crawler.Backend.social_crawler.load_reddrix",
            side_effect=AssertionError("network"),
        ):
            crawler = SocialCrawler()
            plan = crawler.plan()
        self.assertEqual(plan["config"]["platforms"]["reddit"]["backend"], "reddrix")
        self.assertEqual(plan["config"]["total_seconds"], 10800)
        self.assertEqual(set(plan["acquisition_backends"]), {"youtube", "reddit"})
        self.assertEqual(plan["acquisition_backends"]["reddit"], REDDRIX_METHOD)
        self.assertIs(crawler.workers["reddit"].func, run_reddrix_worker)
        self.assertFalse(crawler._synthetic)

    def test_mixed_default_workers_publish_one_bundle_with_correct_provenance(self):
        with (
            patch(
                "Crawler.Backend.social_crawler.reddit.run_worker",
                side_effect=self.runner,
            ) as run,
            patch.object(
                SocialCrawler, "_browser_call", side_effect=AssertionError("legacy")
            ),
        ):
            result = self.crawler().run(self.directory / "mixed")
        self.assertEqual(run.call_count, 1)
        self.assertEqual(result["status"], "TARGET_MET")
        for summary in result["platforms"].values():
            self.assertEqual(summary["qualified_posts"], 1)
        output = Path(result["output_dir"])
        rows = [
            json.loads(line)
            for line in (output / "Contents.jsonl").read_text().splitlines()
        ]
        self.assertEqual(len(rows), 6)
        rd = [row for row in rows if row["platform"] == "reddit"]
        self.assertEqual({row["extraction_method"] for row in rd}, {REDDRIX_METHOD})
        self.assertTrue(
            all("reddrix:fixture:reddrix" in row["adapter_version"] for row in rd)
        )
        self.assertEqual(sum(row["object_type"] == "reply" for row in rd), 1)
        sources = {
            row["source"]: row
            for line in (output / "Sources.jsonl").read_text().splitlines()
            for row in [json.loads(line)]
        }
        self.assertEqual(
            sources["reddit"]["access_mode"], "reddrix_public_web;display=headless"
        )
        self.assertEqual(
            sources["youtube"]["extraction_method"], "yt_dlp_unofficial_web"
        )
        book = load_workbook(result["primary_file"], read_only=True)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["Data"])
        self.assertEqual(sum(1 for _ in book["Data"].iter_rows(values_only=True)), 7)
        self.assertTrue(json.loads((output / "manifest.json").read_text())["synthetic"])

    def test_mixed_baseline_skips_both_platforms_before_worker_requests(self):
        with patch(
            "Crawler.Backend.social_crawler.reddit.run_worker", side_effect=self.runner
        ):
            first = self.crawler().run(self.directory / "baseline")
        config = deepcopy(self.config)
        config["exclude_existing_datasets"] = [first["output_dir"]]
        crawler = self.crawler(config)
        crawler.workers["youtube"] = Mock(side_effect=AssertionError("duplicate"))
        with patch(
            "Crawler.Backend.social_crawler.reddit.run_worker",
            side_effect=AssertionError("duplicate"),
        ):
            result = crawler.run(self.directory / "deduplicated")
        for summary in result["platforms"].values():
            self.assertEqual(summary["existing_posts_skipped"], 1)
            self.assertEqual(summary["posts_exported"], 0)

    def test_missing_reddrix_exports_partial_and_keeps_youtube_running(self):
        failure = reddrix_error_receipt(ImportError("missing optional browser"))
        with (
            patch(
                "Crawler.Backend.social_crawler.reddit.run_worker", return_value=failure
            ),
            patch.object(
                SocialCrawler, "_browser_call", side_effect=AssertionError("fallback")
            ),
        ):
            result = self.crawler().run(self.directory / "blocked")
        self.assertEqual(result["status"], "PARTIAL")
        self.assertEqual(
            result["platforms"]["reddit"]["stop_reason"], "DEPENDENCY_MISSING"
        )
        self.assertEqual(result["platforms"]["youtube"]["qualified_posts"], 1)

    def test_reddrix_retry_uses_coordinator_cooldown_and_same_worker(self):
        self.config["platforms"]["reddit"].update(
            rate_limit_cooldown_seconds=4, rate_limit_max_retries=1
        )
        attempts = []

        def throttled(*args, **kwargs):
            attempts.append(self.clock.now)
            if len(attempts) == 1:
                error = RuntimeError("verified 429 fixture")
                error.status_code, error.headers = 429, {"Retry-After": "6"}
                return reddrix_error_receipt(error)
            return self.runner(*args, **kwargs)

        with patch(
            "Crawler.Backend.social_crawler.reddit.run_worker", side_effect=throttled
        ):
            result = self.crawler().run(self.directory / "retry")
        self.assertEqual(result["status"], "TARGET_MET")
        self.assertEqual(attempts, [0, 6])
        self.assertEqual(result["platforms"]["reddit"]["rate_limit_retries"], 1)

    def test_invalid_display_and_checkpoint_mix_fail_before_acquisition(self):
        for options in [None, {"display_mode": False}, {"unknown": "auto"}]:
            config = deepcopy(self.config)
            config["reddrix"] = options
            with self.subTest(options=options), self.assertRaises(ValueError):
                SocialCrawler(config)
        config = deepcopy(self.config)
        config["platforms"]["reddit"]["import_checkpoint_dir"] = "not-read"
        with self.assertRaisesRegex(ValueError, "does not import DOM checkpoints"):
            SocialCrawler(config)

    def test_spawned_worker_imports_shared_implementation_without_live_access(self):
        events = []
        settings = self.crawler()._legacy_settings(
            "reddit", self.crawler().config["platforms"]["reddit"]
        )
        result = reddit.run_worker(
            "thread",
            URL,
            settings,
            10,
            lambda kind, payload: events.append((kind, payload)),
            worker_entry=spawn_reddrix_fixture,
        )
        self.assertEqual(result["status"], "FINISHED_BEST_EFFORT")
        self.assertEqual(len(events), 3)
        self.assertTrue(
            all(row["extraction_method"] == REDDRIX_METHOD for _, row in events)
        )


if __name__ == "__main__":
    unittest.main()
