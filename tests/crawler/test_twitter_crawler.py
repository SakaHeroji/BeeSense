"""Offline contracts for Twitter/X acquisition; no live credentials required."""

import asyncio
import json
import os
import tempfile
import time
import unittest
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from Crawler.Backend import twitter_crawler as x


def tweet(ident="100", text="Kuliah kampus", parent=None, comments=0):
    return dict(
        tweet_id=ident,
        text=text,
        timestamp="Fri Oct 30 10:00:00 +0000 2026",
        tweet_url=f"https://x.com/example/status/{ident}",
        in_reply_to_tweet_id=parent,
        comments=comments,
        likes=2,
    )


def persistent_worker(options, pipe):
    count = 0
    while pipe.recv() is not None:
        count += 1
        pipe.send(
            ("done", dict(status="FINISHED_BEST_EFFORT", count=count, pid=os.getpid()))
        )


def hanging_worker(options, pipe):
    pipe.recv()
    pipe.send(("record", tweet()))
    time.sleep(20)


class TwitterCrawlerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name)
        self.cfg = json.loads((x.ROOT / "config/twitter_crawler.json").read_text())
        self.cfg["delay"]["seconds"] = 0
        self.cfg["platforms"]["twitter"].update(
            queries=["kuliah"], min_posts=1, max_posts=2, max_searches=1
        )
        self.cfg["scweet"].update(
            auth_file=str(self.path / "auth.json"), state_dir=str(self.path / "state")
        )

    def test_plan_never_reads_auth_or_contacts_network(self):
        with patch.object(x, "auth_material", side_effect=AssertionError("auth read")):
            plan = x.TwitterCrawler(self.cfg).plan()
        self.assertEqual(plan["config"]["total_seconds"], 1800)
        self.assertFalse(plan["live_access_verified"])
        self.assertFalse(plan["integrated_into_social_crawler"])
        self.assertEqual(plan["primary_output"]["columns"], list(x.COMPACT_COLUMNS))

    def test_validation_rejects_embedded_secrets_and_other_platforms(self):
        for change in (
            {"auth_token": "secret"},
            {"scweet": {"sessionid": "secret"}},
            {"platforms": {"instagram": {}}},
            {"total_seconds": True},
            {"since": "2027-01-01T00:00:00Z"},
        ):
            with self.subTest(change=list(change)), self.assertRaises(ValueError):
                x.configuration({**self.cfg, **change})

    def test_auth_alias_and_local_secret_errors_are_not_values(self):
        path = self.path / "auth.json"
        path.write_text(
            json.dumps(
                [
                    dict(sosmed="instagram", sessionid="other"),
                    dict(
                        sosmed="twitter", username="@example", sessionid="local-secret"
                    ),
                ]
            )
        )
        self.assertEqual(
            x.auth_material(path), ("example", {"auth_token": "local-secret"})
        )
        path.write_text(
            json.dumps(
                [dict(sosmed="twitter", username="example", sessionid="bad\nsecret")]
            )
        )
        with self.assertRaises(ValueError) as caught:
            x.auth_material(path)
        self.assertNotIn("bad", str(caught.exception))

    def test_native_identity_parentage_and_unknown_dates(self):
        root = x.normalize_item(tweet())
        direct = x.normalize_item(tweet("101", parent="100"), root_id="100")
        nested = x.normalize_item(tweet("102", parent="101"), root_id="100")
        self.assertEqual(root.published_at, "2026-10-30T10:00:00+00:00")
        self.assertEqual(direct.object_type, "comment")
        self.assertEqual(nested.object_type, "reply")
        self.assertEqual(nested.parent_key, x.object_key("101"))
        self.assertEqual(nested.thread_root_key, x.object_key("100"))
        self.assertFalse(root.text_complete)
        with self.assertRaises(ValueError):
            x.normalize_item(tweet("101"), root_id="100")
        bad = tweet()
        bad["timestamp"] = "recently"
        self.assertEqual(x.normalize_item(bad).timestamp_quality, "unknown")
        bad["tweet_url"] = "https://x.com/example/status/999"
        with self.assertRaises(ValueError):
            x.normalize_item(bad)

    def test_url_canonicalization_rejects_foreign_hosts_and_credentials(self):
        self.assertEqual(
            x.post_url("https://twitter.com/example/status/100?s=20"),
            "https://x.com/i/web/status/100",
        )
        for url in (
            "https://x.com.evil.org/example/status/100",
            "http://x.com/example/status/100",
            "https://secret@x.com/example/status/100",
            "https://x.com/example/status/not-id",
        ):
            with self.subTest(url=url), self.assertRaises(ValueError):
                x.post_url(url)

    def test_shared_excel_dedup_formula_safety_parent_relations_and_manifest(self):
        def worker(operation, target, seconds, emit):
            if operation == "search":
                for row in [tweet(text="=kuliah kampus", comments=3)] * 2 + [
                    tweet("110", "politik kuliah"),
                    tweet("111", parent="109"),
                ]:
                    emit(row)
            else:
                for row in [
                    tweet("101", "Biaya kuliah?", "100"),
                    tweet("102", "Beasiswa", "101"),
                    tweet("101", "edit duplicate", "100"),
                ]:
                    emit(row)
            return dict(status="FINISHED_BEST_EFFORT", reason="fixture", warnings=[])

        output = self.path / "export"
        summary = x.TwitterCrawler(self.cfg, worker=worker).run(output)
        self.assertEqual(
            (summary["posts_exported"], summary["comments"], summary["replies"]),
            (1, 1, 1),
        )
        self.assertEqual(summary["duplicates_skipped"], 2)
        self.assertFalse(summary["live_access_verified"])
        self.assertFalse(summary["trial_passed"])
        compact = load_workbook(output / x.COMPACT_FILE)
        self.assertEqual(compact.sheetnames, ["Data"])
        self.assertEqual(tuple(c.value for c in compact["Data"][1]), x.COMPACT_COLUMNS)
        self.assertEqual(compact["Data"]["H2"].data_type, "s")
        self.assertEqual(compact["Data"].max_row, 4)
        compact.close()
        audit = load_workbook(output / "dataset.xlsx", read_only=True)
        self.assertEqual(audit.sheetnames, list(x.contract.TABLES))
        audit.close()
        rows = [
            json.loads(line)
            for line in (output / "Contents.jsonl").read_text().splitlines()
        ]
        self.assertEqual(len({r["native_id"] for r in rows}), 3)
        relations = [
            json.loads(line)
            for line in (output / "Relations.jsonl").read_text().splitlines()
        ]
        self.assertTrue(
            all(r["resolution_status"] == "resolved_in_export" for r in relations)
        )
        manifest = json.loads((output / "manifest.json").read_text())
        for name, digest in manifest["files_sha256"].items():
            self.assertEqual(x.sha256((output / name).read_bytes()).hexdigest(), digest)
        again = deepcopy(self.cfg)
        again["exclude_existing_datasets"] = [str(output)]
        skipped = x.TwitterCrawler(again, worker=worker).run(self.path / "baseline")
        self.assertEqual(skipped["posts_exported"], 0)
        self.assertGreater(skipped["existing_records_skipped"], 0)

    def test_barrier_keeps_partial_records_and_stops_remaining_operations(self):
        worker = Mock(
            side_effect=lambda operation, target, seconds, emit: (
                emit(tweet(comments=1))
                or dict(
                    status="RATE_LIMITED",
                    reason="fixture",
                    warnings=[],
                    http_status=429,
                )
            )
        )
        summary = x.TwitterCrawler(self.cfg, worker=worker).run(self.path / "limited")
        self.assertEqual(worker.call_count, 1)
        self.assertEqual(summary["stop_reason"], "RATE_LIMITED")
        self.assertEqual(summary["posts_exported"], 1)
        self.assertFalse(summary["trial_passed"])
        self.assertTrue((self.path / "limited/social_data.xlsx").exists())

    def test_time_budget_caps_operations_and_preserves_records(self):
        now = [0.0]
        self.cfg["total_seconds"] = 2

        def worker(operation, target, seconds, emit):
            self.assertEqual(seconds, 2)
            emit(tweet())
            now[0] += seconds
            return dict(status="FINISHED_BEST_EFFORT", reason="fixture", warnings=[])

        summary = x.TwitterCrawler(
            self.cfg, worker=worker, monotonic=lambda: now[0]
        ).run(self.path / "budget")
        self.assertEqual(summary["stop_reason"], "TIME_BUDGET_EXHAUSTED")
        self.assertEqual(summary["posts_exported"], 1)

    def test_library_cooldown_waits_then_retries_same_operation_within_budget(self):
        now, calls = [0.0], []

        def worker(operation, target, seconds, emit):
            calls.append((operation, target))
            if len(calls) == 1:
                return dict(
                    status="COOLDOWN",
                    reason="scweet_account_cooldown",
                    retry_after_seconds=3,
                    warnings=[],
                )
            emit(tweet())
            return dict(status="FINISHED_BEST_EFFORT", reason="fixture", warnings=[])

        summary = x.TwitterCrawler(
            self.cfg,
            worker=worker,
            monotonic=lambda: now[0],
            sleep=lambda seconds: now.__setitem__(0, now[0] + seconds),
        ).run(self.path / "cooldown")
        self.assertEqual(calls[0], calls[1])
        self.assertEqual(summary["cooldown_waits"], 1)
        self.assertEqual(summary["cooldown_seconds"], 3)
        self.assertEqual(summary["posts_exported"], 1)
        self.assertEqual(now[0], 3)

    def test_interrupt_during_delay_finalizes_partial_excel(self):
        self.cfg["delay"]["seconds"] = 15

        def worker(operation, target, seconds, emit):
            emit(tweet(comments=1))
            return dict(status="FINISHED_BEST_EFFORT", reason="fixture", warnings=[])

        crawler = x.TwitterCrawler(
            self.cfg, worker=worker, sleep=Mock(side_effect=KeyboardInterrupt)
        )
        summary = crawler.run(self.path / "interrupt")
        self.assertEqual(summary["stop_reason"], "INTERRUPTED")
        self.assertTrue(summary["interrupted"])
        self.assertEqual(summary["posts_exported"], 1)
        self.assertTrue((self.path / "interrupt/social_data.xlsx").exists())

    def test_pool_diagnostics_expose_wait_only_and_keep_daily_cap_terminal(self):
        client = Mock()
        client._accounts_repo.eligibility_diagnostics.return_value = dict(
            blocked_counts={"cooldown": 1},
            blocked_samples=[dict(available_til=time.time() + 60, username="hidden")],
        )
        result = x.account_cooldown(client, dict(status="ACCOUNT_UNAVAILABLE"))
        self.assertEqual(result["status"], "COOLDOWN")
        self.assertNotIn("hidden", json.dumps(result))
        client._accounts_repo.eligibility_diagnostics.return_value["blocked_counts"][
            "daily_limit"
        ] = 1
        self.assertEqual(
            x.account_cooldown(client, dict(status="ACCOUNT_UNAVAILABLE"))["status"],
            "ACCOUNT_UNAVAILABLE",
        )

    def test_persistent_subprocess_reuses_one_pid(self):
        worker = x.ScweetWorker({}, target=persistent_worker)
        self.addCleanup(worker.close)
        first = worker("search", {}, 5, Mock())
        second = worker("search", {}, 5, Mock())
        self.assertEqual(first["pid"], second["pid"])
        self.assertEqual(second["count"], 2)

    def test_deadline_kills_owned_worker_and_retains_streamed_rows(self):
        worker = x.ScweetWorker({}, target=hanging_worker)
        self.addCleanup(worker.close)
        rows = []
        result = worker("search", {}, 1, rows.append)
        self.assertEqual(result["status"], "TIMEOUT")
        self.assertEqual(len(rows), 1)
        self.assertIsNone(worker.process)

    def test_429_retry_headers_and_http200_challenge(self):
        response = SimpleNamespace(
            status_code=429,
            url="https://x.com/i/api/graphql/test",
            headers={"Retry-After": "360"},
        )
        barrier = x.response_barrier(response)
        self.assertEqual(
            (barrier.status, barrier.retry_after_seconds), ("RATE_LIMITED", 360)
        )
        response.status_code = 200
        response.headers = {"content-type": "application/json"}
        response.json = lambda: {"errors": [{"code": 326, "message": "sensitive"}]}
        self.assertEqual(x.response_barrier(response).status, "CHALLENGE_REQUIRED")
        self.assertNotIn(
            "sensitive", json.dumps(x.error_receipt(x.response_barrier(response)))
        )

    def test_saved_state_reuses_account_without_resetting_counters(self):
        options = x.configuration(self.cfg)["scweet"]
        directory = Path(options["state_dir"])
        directory.mkdir()
        Path(options["auth_file"]).write_text(
            json.dumps(
                [dict(sosmed="twitter", username="example", sessionid="test-token")]
            )
        )
        (directory / "state.db").touch()
        x.private_json(
            directory / "identity.json",
            dict(username="example", token_sha256=x.sha256(b"test-token").hexdigest()),
        )
        client = Mock()
        client._accounts_repo.eligibility_diagnostics.return_value = {"eligible": 1}
        with patch("Scweet.Scweet", return_value=client) as constructor:
            self.assertIs(x.create_client(options), client)
        self.assertFalse(constructor.call_args.kwargs["provision"])
        self.assertIsNone(constructor.call_args.kwargs["cookies"])
        x.private_json(
            directory / "identity.json", dict(username="other", token_sha256="other")
        )
        with patch("Scweet.Scweet") as constructor, self.assertRaises(ValueError):
            x.create_client(options)
        constructor.assert_not_called()

    def test_saved_429_and_login_barriers_prevent_new_network_clients(self):
        options = x.configuration(self.cfg)["scweet"]
        directory = Path(options["state_dir"])
        directory.mkdir()
        Path(options["auth_file"]).write_text(
            json.dumps(
                [dict(sosmed="twitter", username="example", sessionid="test-token")]
            )
        )
        for status, not_before in [
            ("RATE_LIMITED", time.time() + 600),
            ("LOGIN_REQUIRED", time.time() - 600),
        ]:
            x.private_json(
                directory / "barrier.json",
                dict(
                    status=status,
                    http_status=429 if status == "RATE_LIMITED" else 401,
                    not_before=not_before,
                ),
            )
            with (
                self.subTest(status=status),
                patch("Scweet.Scweet") as constructor,
                self.assertRaises(x.AccessBarrier),
            ):
                x.create_client(options)
            constructor.assert_not_called()

    def test_transport_stops_all_network_after_first_429_sync_and_async(self):
        from curl_cffi.requests import AsyncSession, Session

        options = x.configuration(self.cfg)["scweet"]
        Path(options["state_dir"]).mkdir()
        response = SimpleNamespace(
            status_code=429, url="https://x.com/test", headers={"retry-after": "1200"}
        )
        original = Mock(return_value=response)
        with (
            patch.object(Session, "request", original),
            x.guarded_transport(options) as state,
        ):
            Session().get("https://x.com/test")
            with self.assertRaises(x.AccessBarrier):
                Session().get("https://x.com/test")
            with self.assertRaises(x.AccessBarrier):
                asyncio.run(AsyncSession().get("https://x.com/test"))
        self.assertEqual(original.call_count, 1)
        self.assertEqual(state["http_status_counts"], {"429": 1})
        saved = json.loads((Path(options["state_dir"]) / "barrier.json").read_text())
        self.assertGreater(saved["not_before"], time.time() + 1190)


if __name__ == "__main__":
    unittest.main()
