"""Offline contract/deadline tests; fixture data is always marked synthetic."""

import json
import os
import tempfile
import time
import unittest
from copy import deepcopy
from dataclasses import asdict
from datetime import datetime, timezone
from hashlib import file_digest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from Crawler.Backend import instagram_crawler as ig
from Crawler.Backend.social_crawler import SocialCrawler

URL = "https://www.instagram.com/p/B/"  # native media PK 1
STAMP = datetime(2026, 10, 1, 10, tzinfo=timezone.utc)
MEDIA = dict(
    pk="1",
    id="1_9",
    code="B",
    taken_at=STAMP,
    caption_text="Kuliah kampus 😊",
    user={"pk": "9"},
    like_count=5,
    comment_count=2,
)
COMMENT = dict(
    pk="21", text="Biaya kuliah?", created_at_utc=STAMP, user={"pk": "10"}, like_count=0
)
REPLY = dict(
    pk="22",
    text="Ada beasiswa",
    created_at_utc=STAMP,
    user={"pk": "9"},
    replied_to_comment_id="21",
)


def records():
    return [
        ig.normalize_item(MEDIA, URL),
        ig.normalize_item(COMMENT, URL, media=MEDIA),
        ig.normalize_item(REPLY, URL, media=MEDIA, parent="21"),
    ]


class Clock:
    now = 0

    def clock(self):
        return self.now

    def sleep(self, seconds):
        self.now += seconds


def spawn_fixture(operation, target, settings, pipe):
    if os.name == "posix":
        os.setsid()
    pipe.send(("record", records()[0]))
    pipe.send(("done", {"status": "FINISHED_BEST_EFFORT", "source_version": "fixture"}))
    pipe.close()


def spawn_hanging(operation, target, settings, pipe):
    if os.name == "posix":
        os.setsid()
    pipe.send(("record", records()[0]))
    time.sleep(10)


class FakeClient:
    def media_pk_from_code(self, code):
        return "1"

    def media_info_v1(self, pk):
        return deepcopy(MEDIA)

    def media_comments_chunk(self, media_id, max_amount, min_id=None):
        return [deepcopy(COMMENT), deepcopy(COMMENT)], "same"

    def media_comment_replies_chunk(self, media_id, parent, max_amount, min_id=None):
        return [deepcopy(REPLY), deepcopy(REPLY)], "same"

    def hashtag_medias_v1_chunk(self, name, **kwargs):
        return [deepcopy(MEDIA)], "next"


class InstagramCrawlerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name)
        self.config = json.loads(
            (ig.ROOT / "config/instagram_crawler.json").read_text()
        )
        self.config["exclude_existing_datasets"] = []
        self.config["platforms"]["instagram"].update(
            urls=[URL], hashtags=[], min_posts=1, min_comments_per_post=2
        )
        self.clock = Clock()

    def worker(self, operation, target, settings, seconds, emit):
        for record in records():
            emit("record", record)
        return {
            "status": "FINISHED_BEST_EFFORT",
            "source_version": "fixture",
            "warnings": [],
        }

    def crawler(self, worker=None):
        return ig.InstagramCrawler(
            self.config,
            worker=worker or self.worker,
            sleep=self.clock.sleep,
            monotonic=self.clock.clock,
        )

    def test_tracked_trial_is_30_minutes_and_main_remains_two_platforms(self):
        crawler = ig.InstagramCrawler()
        plan = crawler.plan()
        self.assertEqual(plan["config"]["total_seconds"], 1800)
        self.assertEqual(set(plan["config"]["platforms"]), {"instagram"})
        self.assertFalse(plan["trial_passed"])
        self.assertFalse(plan["integrated_into_social_crawler"])
        self.assertEqual(
            set(SocialCrawler().config["platforms"]), {"youtube", "reddit"}
        )

    def test_url_canonicalization_and_identity(self):
        self.assertEqual(ig.media_url("https://instagram.com/reel/B/?igsh=abc"), URL)
        self.assertEqual(ig.media_pk_from_url(URL), "1")
        for bad in [
            "http://www.instagram.com/p/B/",
            "https://www.instagram.com.evil/p/B/",
            "https://www.instagram.com@evil/p/B/",
            "https://instagram.com:443/p/B/",
            "https://instagram.com/user/",
            "https://instagram.com/p/B/\n",
        ]:
            with self.subTest(url=bad), self.assertRaises(ValueError):
                ig.media_url(bad)
        with self.assertRaises(ValueError):
            ig.normalize_item({**MEDIA, "pk": "2"}, URL)

    def test_source_identity_parent_and_unknown_time_are_preserved(self):
        rows = records()
        self.assertEqual([r["object_type"] for r in rows], ["post", "comment", "reply"])
        self.assertEqual(rows[2]["parent_key"], ig.object_key("21", comment=True))
        self.assertFalse(rows[0]["text_complete"])
        self.assertEqual(rows[0]["published_at"], "2026-10-01T10:00:00+00:00")
        unknown = ig.normalize_item(
            {**COMMENT, "created_at_utc": STAMP.replace(tzinfo=None)}, URL, media=MEDIA
        )
        self.assertIsNone(unknown["published_at"])
        with self.assertRaises(ValueError):
            ig.normalize_item(REPLY, URL, media=MEDIA, parent="99")

    def test_real_collector_deduplicates_and_bounds_pages_and_replies(self):
        events = []
        settings = asdict(self.crawler().settings)
        result = ig.collect(
            "media",
            URL,
            settings,
            lambda *args: events.append(args),
            client=FakeClient(),
        )
        self.assertEqual(result["status"], "FINISHED_BEST_EFFORT")
        self.assertEqual(
            [r["object_type"] for k, r in events], ["post", "comment", "reply"]
        )
        settings.update(max_comments=1, max_replies=0)
        events.clear()
        ig.collect(
            "media",
            URL,
            settings,
            lambda *args: events.append(args),
            client=FakeClient(),
        )
        self.assertEqual(len(events), 2)

    def test_search_uses_explicit_hashtag_cursor(self):
        client, events = FakeClient(), []
        client.hashtag_medias_v1_chunk = Mock(wraps=client.hashtag_medias_v1_chunk)
        result = ig.collect(
            "search",
            json.dumps(dict(hashtag="kuliah", cursor="before")),
            asdict(self.crawler().settings),
            lambda *args: events.append(args),
            client=client,
        )
        self.assertEqual(result["next_cursor"], "next")
        self.assertEqual(events[0][1]["url"], URL)
        self.assertEqual(
            client.hashtag_medias_v1_chunk.call_args.kwargs["max_id"], "before"
        )

    def test_complete_bundle_matches_social_excel_contract_and_checksums(self):
        result = self.crawler().run(self.path / "trial")
        self.assertEqual(result["status"], "TARGET_MET")
        self.assertTrue(result["synthetic"])
        self.assertFalse(result["trial_passed"])
        self.assertFalse(result["live_access_verified"])
        self.assertEqual(
            (result["posts_exported"], result["comments"], result["replies"]), (1, 1, 1)
        )
        out = Path(result["output_dir"])
        manifest = json.loads((out / "manifest.json").read_text())
        for name, expected in manifest["files_sha256"].items():
            with (out / name).open("rb") as stream:
                self.assertEqual(file_digest(stream, "sha256").hexdigest(), expected)
        book = load_workbook(result["primary_file"], read_only=True)
        try:
            self.assertEqual(book.sheetnames, ["Data"])
            rows = list(book["Data"].values)
            self.assertEqual(rows[0], ig.COMPACT_COLUMNS)
            self.assertEqual(len(rows), 4)
            self.assertEqual(rows[1][5], datetime(2026, 10, 1, 17))
        finally:
            book.close()
        book = load_workbook(out / "dataset.xlsx", read_only=True)
        try:
            self.assertEqual(book.sheetnames, list(ig.contract.TABLES))
        finally:
            book.close()
        rows = [
            json.loads(s) for s in (out / "Contents.jsonl").read_text().splitlines()
        ]
        self.assertEqual({r["extraction_method"] for r in rows}, {ig.METHOD})
        self.assertTrue(
            all(r["actor_type"] == "unknown" and r["author_ref"] is None for r in rows)
        )

    def test_formula_and_long_unicode_text_remain_safe_and_full_in_jsonl(self):
        raw = records()
        raw[0]["text_source"] = "=kuliah " + "😊" * 16000

        def worker(op, target, settings, seconds, emit):
            for row in raw:
                emit("record", row)
            return dict(status="FINISHED_BEST_EFFORT")

        result = self.crawler(worker).run(self.path / "long")
        book = load_workbook(result["primary_file"], read_only=True)
        try:
            row = list(book["Data"].rows)[1]
            self.assertEqual(row[7].data_type, "s")
            self.assertTrue(row[9].value)
        finally:
            book.close()
        stored = json.loads(
            (Path(result["output_dir"]) / "Contents.jsonl").read_text().splitlines()[0]
        )
        self.assertEqual(stored["text_source"], raw[0]["text_source"])

    def test_baseline_dedup_before_request_and_source_remains_unchanged(self):
        first = self.crawler().run(self.path / "baseline")
        path = Path(first["output_dir"]) / "dataset.xlsx"
        original = path.read_bytes()
        self.config["exclude_existing_datasets"] = [str(path)]
        worker = Mock(side_effect=AssertionError("duplicate request"))
        result = self.crawler(worker).run(self.path / "new")
        worker.assert_not_called()
        self.assertEqual(result["existing_posts_skipped"], 1)
        self.assertEqual(result["posts_exported"], 0)
        self.assertEqual(path.read_bytes(), original)

    def test_filtering_dates_keywords_and_minimum_uses_unique_comments(self):
        self.config["until"] = "2026-10-01T10:00:00Z"
        result = self.crawler().run(self.path / "outside")
        self.assertEqual(result["posts_exported"], 0)
        self.assertEqual(result["filtered_roots"], {"root_outside_date_window": 1})
        self.config["until"] = "2026-10-27T00:00:00Z"
        self.config["exclude_keywords"] = ["beasiswa"]
        result = self.crawler().run(self.path / "exclude")
        self.assertEqual(result["posts_exported"], 1)
        self.assertEqual(result["qualified_posts"], 0)
        self.assertEqual(result["replies"], 0)

    def test_authentication_barrier_keeps_partial_records_without_retry(self):
        def worker(op, target, settings, seconds, emit):
            emit("record", records()[0])
            return dict(status="CHALLENGE_REQUIRED", warnings=["ChallengeRequired"])

        worker = Mock(side_effect=worker)
        result = self.crawler(worker).run(self.path / "challenge")
        self.assertEqual(worker.call_count, 1)
        self.assertEqual(result["stop_reason"], "CHALLENGE_REQUIRED")
        self.assertEqual(result["posts_exported"], 1)
        self.assertFalse(result["trial_passed"])

    def test_verified_429_respects_retry_after_and_budget(self):
        self.config["platforms"]["instagram"].update(
            rate_limit_max_retries=1, rate_limit_cooldown_seconds=4
        )
        attempts = []

        def worker(*args):
            attempts.append(self.clock.now)
            if len(attempts) == 1:
                return dict(status="RATE_LIMITED", retry_after_seconds=6)
            return self.worker(*args)

        result = self.crawler(worker).run(self.path / "rate")
        self.assertEqual(attempts, [0, 6])
        self.assertEqual(result["rate_limit_retries"], 1)
        self.assertEqual(result["status"], "TARGET_MET")
        self.config["total_seconds"] = 5
        attempts.clear()
        self.clock.now = 0
        result = self.crawler(worker).run(self.path / "too_late")
        self.assertEqual(len(attempts), 1)
        self.assertEqual(result["stop_reason"], "RATE_LIMITED")

    def test_global_deadline_prevents_late_requests_and_keeps_partial(self):
        self.config["total_seconds"] = 5
        self.config["platforms"]["instagram"].update(
            urls=[URL, "https://instagram.com/p/C/"], min_posts=2
        )

        def worker(*args):
            result = self.worker(*args)
            self.clock.sleep(5)
            return result

        worker = Mock(side_effect=worker)
        result = self.crawler(worker).run(self.path / "deadline")
        self.assertEqual(worker.call_count, 1)
        self.assertEqual(result["posts_exported"], 1)
        self.assertEqual(result["stop_reason"], "TIME_BUDGET_EXHAUSTED")

    def test_repeated_source_errors_stop_after_three_operations(self):
        self.config["platforms"]["instagram"].update(
            urls=[], hashtags=["kuliah", "kampus", "beasiswa", "jurusan"]
        )
        worker = Mock(return_value=dict(status="SOURCE_ERROR"))
        result = self.crawler(worker).run(self.path / "errors")
        self.assertEqual(worker.call_count, 3)
        self.assertEqual(result["stop_reason"], "CONSECUTIVE_SOURCE_ERRORS")

    def test_missing_login_is_reported_without_credentials_or_exception_message(self):
        error = FileNotFoundError("secret/password/cookie")
        result = ig.error_receipt(error)
        self.assertEqual(result["status"], "LOGIN_REQUIRED")
        self.assertNotIn("secret", json.dumps(result))
        error = RuntimeError("not necessarily 429")
        self.assertEqual(ig.error_receipt(error)["status"], "SOURCE_ERROR")
        error.response = SimpleNamespace(status_code=429, headers={"Retry-After": "7"})
        self.assertEqual(ig.error_receipt(error)["retry_after_seconds"], 7)

    def test_saved_settings_are_atomic_owner_only_and_retry_settings_are_reset(self):
        path = self.path / "auth/session.json"
        client = Mock()
        client.get_settings.return_value = {
            "authorization_data": {"ds_user_id": "9"},
            "session_retry_total": 8,
        }
        ig.save_session(client, path)
        if os.name == "posix":
            self.assertEqual(path.stat().st_mode & 0o777, 0o600)
        client.user_id = "9"
        with patch.object(ig, "create_client", return_value=client):
            ig.load_client(path)
        self.assertEqual(
            client.set_settings.call_args.args[0]["session_retry_total"], 0
        )
        self.assertNotIn("sessionid", json.dumps(self.crawler().plan()))
        client.user_id = None
        with (
            patch.object(ig, "create_client", return_value=client),
            self.assertRaises(ig.SessionInvalid),
        ):
            ig.load_client(path)
        with patch.object(ig, "create_client", return_value=client):
            self.assertIs(ig.load_client(path, require_authenticated=False), client)

    def test_credentials_cannot_be_embedded_in_config(self):
        for change in [{"password": "secret"}, {"username": "user"}]:
            cfg = deepcopy(self.config)
            cfg["instagrapi"].update(change)
            with self.assertRaises(ValueError):
                ig.InstagramCrawler(cfg)

    def test_auth_file_login_selects_instagram_and_never_exports_credentials(self):
        path = self.path / "auth.json"
        path.write_text(
            json.dumps(
                [
                    {"sosmed": "reddit", "username": "other", "password": "unused"},
                    {
                        "sosmed": "Instagram",
                        "username": "fixture_user",
                        "password": "fixture_secret",
                    },
                ]
            )
        )
        session = self.path / "session.json"
        client = Mock()
        client.login.return_value = True
        client.get_settings.return_value = {"authorization_data": {"ds_user_id": "9"}}
        with patch.object(ig, "create_client", return_value=client):
            ig.login_from_file(path, session)
        client.login.assert_called_once_with(
            "fixture_user", "fixture_secret", verification_code=""
        )
        self.assertNotIn("fixture_secret", session.read_text())
        path.write_text("{invalid json}")
        with self.assertRaises(ig.SessionInvalid):
            ig.login_from_file(path, session)

    def test_upgrade_gate_is_distinct_from_throttling_and_unknown_source_errors(self):
        error_type = type("ClientNeedsUpgradeError", (Exception,), {})
        result = ig.error_receipt(error_type("secret response body"))
        self.assertEqual(result["status"], "CLIENT_UPGRADE_REQUIRED")
        self.assertIn(result["status"], ig.STOP_STATUSES)
        self.assertNotIn("secret", json.dumps(result))

    def test_local_session_cookie_is_used_without_password_login(self):
        path = self.path / "auth.json"
        path.write_text(
            json.dumps({"sosmed": "instagram", "sessionid": "fixture_cookie"})
        )
        client = Mock()
        client.user_id = "9"
        client.login_by_sessionid.return_value = True
        client.get_settings.return_value = {"authorization_data": {"ds_user_id": "9"}}
        with patch.object(ig, "create_client", return_value=client):
            ig.login_from_file(path, self.path / "session.json")
        client.login_by_sessionid.assert_called_once_with("fixture_cookie")
        client.account_info.assert_called_once_with()
        client.login.assert_not_called()
        client.user_id = None
        with (
            patch.object(ig, "create_client", return_value=client),
            self.assertRaises(ig.SessionInvalid),
        ):
            ig.login_from_file(path, self.path / "rejected.json")

    def test_browser_cookie_is_opaque_and_device_survives_stale_auth_removal(self):
        path = self.path / "auth.json"
        path.write_text(
            json.dumps(
                {
                    "sosmed": "instagram",
                    "sessionid": "123%3Afixture%3Apart%3Avalue%2Bkeep%252F",
                }
            )
        )
        session = self.path / "session.json"
        session.write_text("{}")
        client = Mock()
        client.user_id = "123"
        client.login_by_sessionid.return_value = True
        client.get_settings.return_value = {
            "uuids": {"uuid": "same-device"},
            "authorization_data": {"ds_user_id": "123", "sessionid": "old-cookie"},
            "cookies": {"sessionid": "old-cookie"},
        }
        with (
            patch.object(ig, "load_client", return_value=client) as load,
            patch.object(
                ig, "create_client", side_effect=AssertionError("device rotation")
            ),
        ):
            ig.login_from_file(path, session)
        load.assert_called_once_with(session, require_authenticated=False)
        client.login_by_sessionid.assert_called_once_with(
            "123%3Afixture%3Apart%3Avalue%2Bkeep%252F"
        )
        client.set_settings.assert_called_once_with(
            {
                "uuids": {"uuid": "same-device"},
                "authorization_data": {},
                "cookies": {},
            }
        )
        client.login.assert_not_called()

    def test_password_mode_restores_settings_and_ignores_browser_cookie(self):
        auth = self.path / "auth.json"
        auth.write_text(
            json.dumps(
                {
                    "sosmed": "instagram",
                    "username": "fixture_user",
                    "password": "fixture_secret",
                    "sessionid": "browser_cookie_not_used",
                }
            )
        )
        session = self.path / "session.json"
        session.write_text("{}")
        client = Mock(user_id="9")
        client.login.return_value = True
        client.get_settings.return_value = {
            "uuids": {"uuid": "same-device"},
            "authorization_data": {"ds_user_id": "9"},
        }
        with patch.object(ig, "load_client", return_value=client) as load:
            ig.login_from_file(auth, session, login_method="password")
        load.assert_called_once_with(session, require_authenticated=False)
        client.login.assert_called_once_with(
            "fixture_user", "fixture_secret", verification_code=""
        )
        client.login_by_sessionid.assert_not_called()
        client.relogin.assert_not_called()
        self.assertEqual(
            json.loads(session.read_text())["uuids"]["uuid"], "same-device"
        )
        self.assertNotIn("fixture_secret", session.read_text())
        saved = session.read_bytes()
        upgrade = type("ClientNeedsUpgradeError", (Exception,), {})
        client.login.side_effect = upgrade("private response")
        with (
            patch.object(ig, "load_client", return_value=client),
            self.assertRaises(upgrade),
        ):
            ig.login_from_file(auth, session, login_method="password")
        self.assertEqual(session.read_bytes(), saved)
        with patch.object(ig, "login_from_file") as login:
            ig.collect(
                "login_password", str(auth), {"session_file": str(session)}, Mock()
            )
        login.assert_called_once_with(str(auth), str(session), login_method="password")

    def test_password_recovery_is_one_relogin_and_propagates_failure(self):
        required = type("LoginRequired", (Exception,), {})
        challenge = type("ChallengeRequired", (Exception,), {})
        client = Mock(user_id="9")
        client.login.side_effect = required()
        client.relogin.return_value = True
        ig.login_with_saved_password(client, "fixture", "secret")
        client.login.assert_called_once()
        client.relogin.assert_called_once_with()
        client.reset_mock()
        client.relogin.side_effect = challenge("not exported")
        with self.assertRaises(challenge):
            ig.login_with_saved_password(client, "fixture", "secret")
        client.login.assert_called_once()
        client.relogin.assert_called_once_with()

    def test_password_recovery_retains_explicit_verification_code(self):
        required = type("LoginRequired", (Exception,), {})
        client = Mock(user_id="9")
        client.login.side_effect = [required(), True]
        ig.login_with_saved_password(client, "fixture", "secret", "fixture-code")
        self.assertEqual(client.login.call_count, 2)
        self.assertEqual(
            client.login.call_args.kwargs,
            {"relogin": True, "verification_code": "fixture-code"},
        )
        client.relogin.assert_not_called()

    def test_password_upgrade_challenge_and_rate_limit_are_not_retried(self):
        for name in (
            "ClientNeedsUpgradeError",
            "ChallengeRequired",
            "PleaseWaitFewMinutes",
        ):
            error = type(name, (Exception,), {})
            client = Mock(user_id="9")
            client.login.side_effect = error("private response")
            with self.subTest(name=name), self.assertRaises(error):
                ig.login_with_saved_password(client, "fixture", "secret")
            client.login.assert_called_once()
            client.relogin.assert_not_called()

    def test_invalid_input_fails_before_any_worker(self):
        for change in [
            {"total_seconds": True},
            {"total_seconds": 0},
            {"until": "2026-01-01"},
            {"include_keywords": [""]},
        ]:
            with self.subTest(change=change), self.assertRaises(ValueError):
                ig.InstagramCrawler({**self.config, **change})

    def test_spawned_worker_preserves_events_and_kills_hanging_process(self):
        events = []
        result = ig.run_instagram_worker(
            "media",
            URL,
            self.crawler().settings,
            3,
            lambda *args: events.append(args),
            worker_entry=spawn_fixture,
        )
        self.assertEqual(result["status"], "FINISHED_BEST_EFFORT")
        self.assertEqual(len(events), 1)
        start = time.monotonic()
        result = ig.run_instagram_worker(
            "media",
            URL,
            self.crawler().settings,
            0.8,
            lambda *args: None,
            worker_entry=spawn_hanging,
        )
        self.assertEqual(result["status"], "TIMEOUT")
        self.assertLess(time.monotonic() - start, 4)

    def test_actual_instagram_worker_missing_session_stops_without_network(self):
        settings = ig.InstagramSettings(session_file=str(self.path / "missing.json"))
        result = ig.run_instagram_worker(
            "check_session",
            "instagram_account",
            settings,
            5,
            lambda *args: self.fail("Unexpected data"),
        )
        self.assertEqual(result["status"], "LOGIN_REQUIRED")


if __name__ == "__main__":
    unittest.main()
