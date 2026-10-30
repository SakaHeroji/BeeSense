"""Reddrix field/identity boundaries and the real, shared Excel exporter."""

import asyncio
import json
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from Crawler.Backend.reddit_crawler import (
    RedditCrawler,
    collect,
    error_receipt,
    load_reddrix,
    normalize_item,
    receipt,
    resolve_display,
    run_reddrix_worker,
)
from Crawler.Backend.social_crawler import COMPACT_COLUMNS, SocialCrawler, reddit

URL = "https://www.reddit.com/r/indonesia/comments/abc123/kuliah/"
STAMP = "2026-06-01T07:00:00+00:00"
CONFIG = dict(
    project_id="reddrix-test",
    since="2026-06-01T00:00:00+00:00",
    until="2026-07-01T00:00:00+00:00",
    include_keywords=["kuliah"],
    delay=dict(mode="fixed", seconds=0),
    total_seconds=60,
    platforms={
        "reddit": dict(urls=[URL], subreddits=["indonesia"], min_comments_per_post=2)
    },
)


def item(native="abc123", parent=None):
    post = native == "abc123"
    return dict(
        submissionId=native,
        submissionIdPrefixed=("t3_" if post else "t1_") + native,
        submissionType="post" if post else "comment",
        submissionUrl=URL if post else URL + native + "/",
        parentId=parent,
        parentPostId=None if post else "t3_abc123",
        submittedAt=STAMP,
        title="Kuliah dan kampus" if post else None,
        bodyText="=literal teks 😀",
        subreddit="r/indonesia",
        score=5,
        numComments=2 if post else None,
        depth=0 if parent == "t3_abc123" else 1,
    )


def records():
    return [item(), item("com123", "t3_abc123"), item("rpl123", "t1_com123")]


class FakeReddrix:
    options = None
    items = []

    def __init__(self, **kwargs):
        type(self).options = kwargs

    async def __aenter__(self):
        return self

    async def __aexit__(self, *args):
        pass

    async def get_items(self):
        for row in self.items:
            yield row


class RedditCrawlerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    @staticmethod
    def worker(operation, target, settings, seconds, emit):
        for row in records():
            emit("record", normalize_item(row, URL, STAMP))
        return receipt("FINISHED_BEST_EFFORT", "synthetic_test")

    def test_active_configuration_keeps_every_reddit_rule_and_never_runs_youtube(self):
        raw = json.loads(Path("config/social_crawler.json").read_text())
        before = deepcopy(raw)
        crawler = RedditCrawler(raw)
        original = SocialCrawler(raw)
        self.assertEqual(
            crawler.config["platforms"]["reddit"],
            original.config["platforms"]["reddit"],
        )
        for key in [
            "since",
            "until",
            "include_keywords",
            "exclude_keywords",
            "delay",
            "total_seconds",
            "timeout_seconds",
            "exclude_existing_datasets",
        ]:
            self.assertEqual(crawler.config[key], original.config[key])
        self.assertEqual(set(crawler.config["platforms"]), {"reddit"})
        self.assertEqual(raw, before)
        self.assertFalse(crawler._synthetic)
        self.assertEqual(crawler.plan()["acquisition_backend"], "reddrix")

    def test_trial_config_changes_only_duration_and_platform_scope(self):
        original = SocialCrawler(
            json.loads(Path("config/social_crawler.json").read_text())
        ).config
        trial = RedditCrawler().config
        self.assertEqual(trial["total_seconds"], 1800)
        expected = deepcopy(original)
        expected["total_seconds"] = 1800
        expected["platforms"] = {"reddit": original["platforms"]["reddit"]}
        self.assertEqual(trial["reddrix"], {"display_mode": "auto"})
        self.assertEqual(trial, expected)

    def test_complete_excel_contract_formula_safety_relations_and_provenance(self):
        result = RedditCrawler(CONFIG, workers={"reddit": self.worker}).run(
            self.directory / "bundle"
        )
        self.assertEqual(result["status"], "TARGET_MET")
        self.assertEqual(result["platforms"]["reddit"]["qualified_posts"], 1)
        folder = Path(result["output_dir"])
        with (folder / "Contents.jsonl").open() as stream:
            rows = [json.loads(line) for line in stream]
        self.assertEqual([r["object_type"] for r in rows], ["post", "comment", "reply"])
        self.assertEqual(rows[2]["parent_key"], rows[1]["object_key"])
        self.assertEqual(rows[0]["extraction_method"], "reddrix_unofficial_web")
        self.assertTrue(rows[0]["adapter_version"].startswith("reddrix-adapter:"))
        self.assertIsNone(rows[1]["author_ref"])
        audit = load_workbook(folder / "dataset.xlsx", read_only=True)
        self.addCleanup(audit.close)
        self.assertEqual(audit.sheetnames, list(reddit.TABLES))
        for name, fields in reddit.TABLES.items():
            self.assertEqual(
                list(next(audit[name].values)), [f["name"] for f in fields]
            )
        compact = load_workbook(folder / "social_data.xlsx", read_only=True)
        self.addCleanup(compact.close)
        self.assertEqual(compact.sheetnames, ["Data"])
        self.assertEqual(tuple(next(compact["Data"].values)), COMPACT_COLUMNS)
        self.assertEqual(compact["Data"]["H2"].data_type, "s")
        self.assertEqual(len(list(compact["Data"].values)), 4)
        manifest = json.loads((folder / "manifest.json").read_text())
        from hashlib import sha256

        for name, expected in manifest["files_sha256"].items():
            self.assertEqual(sha256((folder / name).read_bytes()).hexdigest(), expected)
        self.assertTrue(manifest["synthetic"])
        with self.assertRaises(FileExistsError):
            RedditCrawler(CONFIG).run(folder)

    def test_baseline_blocks_native_roots_across_runs(self):
        first = RedditCrawler(CONFIG, workers={"reddit": self.worker}).run(
            self.directory / "old"
        )
        cfg = deepcopy(CONFIG)
        cfg["exclude_existing_datasets"] = [first["output_dir"]]

        def forbidden(*args):
            self.fail("Excluded thread must not be requested")

        result = RedditCrawler(cfg, workers={"reddit": forbidden}).run(
            self.directory / "new"
        )
        self.assertEqual(result["platforms"]["reddit"]["posts_exported"], 0)
        self.assertEqual(result["platforms"]["reddit"]["existing_posts_skipped"], 1)

    def test_barrier_exports_partial_without_calling_it_success(self):
        result = RedditCrawler(
            CONFIG, workers={"reddit": lambda *args: receipt("ACCESS_BLOCKED", "test")}
        ).run(self.directory / "blocked")
        self.assertEqual(result["status"], "PARTIAL")
        self.assertFalse(result["targets_met"])
        self.assertEqual(result["platforms"]["reddit"]["stop_reason"], "ACCESS_BLOCKED")
        self.assertTrue(Path(result["primary_file"]).exists())

    def test_missing_parent_and_timezone_are_preserved_as_unknown(self):
        raw = item("com123")
        raw["submittedAt"] = "yesterday"
        row = normalize_item(raw, URL, STAMP)
        self.assertIsNone(row["parent"])
        self.assertEqual(row["kind"], "reply")
        self.assertIsNone(row["published_at"])
        self.assertFalse(row["text_complete"])

    def test_rejects_cross_thread_and_inconsistent_ids(self):
        for key, value in [
            ("submissionId", "wrong"),
            ("parentPostId", "t3_xyz123"),
            ("submissionUrl", "https://evil.test/"),
            ("bodyText", {}),
            ("parentId", "t1_com123"),
        ]:
            raw = item("com123", "t3_abc123")
            raw[key] = value
            with self.subTest(key=key), self.assertRaises((ValueError, TypeError)):
                normalize_item(raw, URL, STAMP)

    def test_unknown_antibot_status_is_not_guessed_as_http429(self):
        error = type("FatalAntiBotBlockError", (Exception,), {})()
        self.assertEqual(error_receipt(error)["status"], "ACCESS_BLOCKED")
        error = RuntimeError("sensitive cookie must not be exported")
        error.status_code, error.headers = 429, {"Retry-After": "720"}
        result = error_receipt(error)
        self.assertEqual(result["retry_after_seconds"], 720)
        self.assertNotIn("sensitive", json.dumps(result))

    def test_stream_enforces_duplicates_and_reply_budgets(self):
        settings = reddit.Settings(
            subreddits=("indonesia",),
            max_comments=2,
            max_replies=1,
            max_replies_per_thread=1,
        )
        from dataclasses import asdict

        FakeReddrix.items = records() + [
            item("rpl456", "t1_com123"),
            item("com123", "t3_abc123"),
        ]
        events = []
        asyncio.run(
            collect(
                "thread",
                URL,
                asdict(settings),
                lambda k, v: events.append((k, v)),
                factory=FakeReddrix,
            )
        )
        self.assertEqual(len(events), 3)
        self.assertFalse(FakeReddrix.options["stealth"])
        self.assertEqual(FakeReddrix.options["max_results"], 3)

    def test_search_preserves_subreddit_scope_without_fetching_comments(self):
        from dataclasses import asdict

        settings = reddit.Settings(subreddits=("indonesia",), search_results=3)
        FakeReddrix.items = [item(), item()]
        events = []
        asyncio.run(
            collect(
                "search",
                "https://www.reddit.com/r/indonesia/search/?q=kuliah",
                asdict(settings),
                lambda k, v: events.append((k, v)),
                factory=FakeReddrix,
            )
        )
        self.assertEqual(len(events), 1)
        self.assertEqual(FakeReddrix.options["query"], "subreddit:indonesia (kuliah)")
        self.assertFalse(FakeReddrix.options["include_comments"])

    def test_dry_run_does_not_import_or_start_reddrix(self):
        with patch(
            "Crawler.Backend.social_crawler.run_reddrix_worker",
            side_effect=AssertionError("no network"),
        ):
            self.assertFalse(RedditCrawler(CONFIG).plan()["live_access_verified"])

    def test_adaptive_display_matrix(self):
        for platform, xvfb, expected in [
            ("darwin", "/opt/X11/bin/Xvfb", True),
            ("darwin", None, True),
            ("linux", "/usr/bin/Xvfb", "virtual"),
            ("linux", None, True),
            ("win32", None, True),
        ]:
            with self.subTest(platform=platform, xvfb=xvfb):
                display = resolve_display(
                    "auto", platform=platform, finder=lambda _: xvfb
                )
                self.assertEqual(display["headless"], expected)
        self.assertIs(
            resolve_display(
                "headless", platform="linux", finder=lambda _: "/usr/bin/Xvfb"
            )["headless"],
            True,
        )

    def test_explicit_virtual_requires_linux_and_xvfb(self):
        for platform, xvfb in [("darwin", "/opt/X11/bin/Xvfb"), ("linux", None)]:
            with (
                self.subTest(platform=platform),
                self.assertRaisesRegex(ValueError, "requires Linux and Xvfb"),
            ):
                resolve_display("virtual", platform=platform, finder=lambda _: xvfb)
        self.assertEqual(
            resolve_display(
                "virtual", platform="linux", finder=lambda _: "/usr/bin/Xvfb"
            )["headless"],
            "virtual",
        )
        for mode in [False, "invalid", {}]:
            with self.subTest(mode=mode), self.assertRaises(ValueError):
                resolve_display(mode)

    def test_display_config_and_override_reach_plan_and_worker(self):
        cfg = deepcopy(CONFIG)
        cfg["reddrix"] = {"display_mode": "headless"}
        crawler = RedditCrawler(cfg)
        self.assertEqual(crawler.plan()["display"]["effective_mode"], "headless")
        self.assertEqual(crawler.workers["reddit"].keywords["display_mode"], "headless")
        self.assertEqual(
            RedditCrawler(cfg, display_mode="auto").config["reddrix"],
            {"display_mode": "auto"},
        )
        self.assertEqual(cfg["reddrix"], {"display_mode": "headless"})
        cfg["reddrix"]["unknown"] = True
        with self.assertRaises(ValueError):
            RedditCrawler(cfg)

    def test_browser_factory_overrides_initial_and_refreshed_sessions(self):
        from types import SimpleNamespace
        from unittest.mock import Mock

        browser = Mock()
        module = SimpleNamespace(Reddrix=FakeReddrix)
        camoufox_module = SimpleNamespace(AsyncCamoufox=browser)
        modules = {"reddrix.reddrix": module, "camoufox.async_api": camoufox_module}
        with (
            patch(
                "Crawler.Backend.social_crawler.resolve_reddrix_display",
                return_value={"headless": True},
            ),
            patch("importlib.import_module", side_effect=lambda name: modules[name]),
        ):
            self.assertIs(load_reddrix("auto"), FakeReddrix)
        module.AsyncCamoufox(headless="virtual", humanize=True)
        module.AsyncCamoufox(headless=False)
        self.assertEqual(
            [c.kwargs["headless"] for c in browser.call_args_list], [True, True]
        )
        self.assertIs(camoufox_module.AsyncCamoufox, browser)

    def test_display_mode_propagates_through_process_boundary(self):
        with patch(
            "Crawler.Backend.reddit_crawler.reddit.run_worker",
            return_value={"status": "test"},
        ) as run:
            run_reddrix_worker("thread", URL, None, 90, None, display_mode="headless")
        entry = run.call_args.kwargs["worker_entry"]
        self.assertEqual(entry.keywords, {"display_mode": "headless"})


if __name__ == "__main__":
    unittest.main()
