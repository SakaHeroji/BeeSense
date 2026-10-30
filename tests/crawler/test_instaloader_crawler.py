"""Offline contracts for the separate Instaloader experiment."""

import json
import os
import tempfile
import time
import unittest
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace as Obj
from unittest.mock import Mock, patch

import requests
from openpyxl import load_workbook

from Crawler.Backend import instaloader_crawler as il

URL = "https://www.instagram.com/p/B/"
STAMP = datetime(2026, 10, 1, 10)


def fixture_post():
    reply = Obj(
        id=22,
        text="Ada beasiswa",
        created_at_utc=STAMP,
        owner=Obj(userid=9),
        likes_count=0,
    )
    comment = Obj(
        id=21,
        text="Biaya kuliah?",
        created_at_utc=STAMP,
        owner=Obj(userid=10),
        likes_count=1,
        answers=[reply, reply],
    )
    return Obj(
        mediaid=1,
        shortcode="B",
        caption="Kuliah kampus",
        date_utc=STAMP,
        owner_id=9,
        likes=5,
        comments=2,
        get_comments=lambda: iter([comment, comment]),
    )


def persistent_fixture(settings, pipe):
    count = 0
    while pipe.recv() is not None:
        count += 1
        pipe.send(
            ("done", dict(status="FINISHED_BEST_EFFORT", pid=os.getpid(), count=count))
        )
    pipe.close()


def hanging_fixture(settings, pipe):
    pipe.recv()
    pipe.send(("done", dict(status="FINISHED_BEST_EFFORT")))
    pipe.recv()
    pipe.send(("record", {"native_id": "1"}))
    time.sleep(20)


class InstaloaderCrawlerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name)
        self.config = json.loads(
            (il.ig.ROOT / "config/instaloader_crawler.json").read_text()
        )
        self.config["exclude_existing_datasets"] = []
        self.config["delay"]["seconds"] = 0
        self.config["platforms"]["instagram"].update(
            urls=[URL],
            hashtags=[],
            min_posts=1,
            min_comments_per_post=1,
            max_comments=1,
            max_replies=1,
            max_replies_per_thread=1,
        )
        self.config["instaloader"].update(
            auth_file=str(self.path / "auth.json"),
            session_file=str(self.path / "session.json"),
        )

    def settings(self):
        crawler = il.InstaloaderCrawler(self.config)
        self.addCleanup(crawler.close)
        return crawler.settings

    def adapter(self, settings=None):
        module = Obj(Post=Obj(from_shortcode=Mock(return_value=fixture_post())))
        loader = Mock()
        return il.InstaloaderAdapter(
            settings or self.settings(),
            loader=loader,
            module=module,
            comments=lambda post: post.get_comments(),
        )

    def test_normalization_and_caps_preserve_native_parents_and_utc(self):
        adapter, events = self.adapter(), []
        with patch.object(il, "save_session"):
            result = adapter.collect("media", URL, lambda kind, row: events.append(row))
        self.assertEqual(result["status"], "FINISHED_BEST_EFFORT")
        self.assertEqual([row["native_id"] for row in events], ["1", "21", "22"])
        self.assertEqual(events[2]["parent_key"], il.ig.object_key("21", comment=True))
        self.assertEqual(
            events[0]["published_at"], STAMP.replace(tzinfo=timezone.utc).isoformat()
        )
        self.assertTrue(all(row["text_complete"] is False for row in events))

    def test_mismatched_shortcode_rejected_before_emission(self):
        adapter = self.adapter()
        adapter.module.Post.from_shortcode.return_value.mediaid = 2
        emit = Mock()
        with self.assertRaises(ValueError):
            adapter.collect("media", URL, emit)
        emit.assert_not_called()

    def test_search_retains_iterator_without_fetching_one_extra(self):
        adapter = self.adapter(replace(self.settings(), search_results=1))
        fetched = []

        def posts():
            for code in ("B", "C"):
                fetched.append(code)
                yield Obj(shortcode=code)

        hashtag = Obj(get_posts_resumable=lambda: posts())
        adapter.module.Hashtag = Obj(from_name=Mock(return_value=hashtag))
        first = adapter.collect(
            "search", json.dumps(dict(hashtag="kuliah", cursor=None)), Mock()
        )
        self.assertEqual(fetched, ["B"])
        second = adapter.collect(
            "search",
            json.dumps(dict(hashtag="kuliah", cursor=first["next_cursor"])),
            Mock(),
        )
        self.assertEqual(fetched, ["B", "C"])
        self.assertEqual(second["next_cursor"], "2")
        adapter.module.Hashtag.from_name.assert_called_once()

    def test_discovered_post_reuses_metadata_instead_of_shortcode_query(self):
        adapter = self.adapter()
        post = fixture_post()
        adapter.module.Hashtag = Obj(
            from_name=Mock(return_value=Obj(get_posts_resumable=lambda: iter([post])))
        )
        adapter.collect(
            "search", json.dumps(dict(hashtag="kuliah", cursor=None)), Mock()
        )
        emitted = []
        with patch.object(il, "save_session"):
            adapter.collect("media", URL, lambda kind, row: emitted.append(row))
        adapter.module.Post.from_shortcode.assert_not_called()
        self.assertEqual([r["native_id"] for r in emitted], ["1", "21", "22"])

    def test_web_graphql_applies_supplied_patch_only_to_owned_context(self):
        loader = Obj(context=Obj())
        with patch(
            "Crawler.Backend.instaloader_browser_patch.apply_browser_patch"
        ) as apply:
            il.configure_web_graphql(loader)
        apply.assert_called_once_with(loader.context)

    def test_web_comments_pagination_ids_and_previews_do_not_use_mobile(self):
        adapter = self.adapter()
        node = dict(
            pk="21",
            text="Biaya kuliah?",
            created_at=1790848800,
            user=dict(pk="10"),
            comment_like_count=1,
            preview_child_comments=[
                dict(
                    pk="22",
                    text="Ada beasiswa",
                    created_at=1790848801,
                    user=dict(pk="9"),
                )
            ],
        )
        connection = dict(
            edges=[dict(node=node)],
            page_info=dict(has_next_page=True, end_cursor="next"),
        )
        adapter.loader.context.doc_id_graphql_query.return_value = {
            "data": {"xdt_api__v1__media__media_id__comments__connection": connection}
        }
        rows = list(adapter.web_comments(fixture_post()))
        self.assertEqual([r.id for r in rows], ["21"])
        self.assertEqual([r.id for r in rows[0].answers], ["22"])
        adapter.loader.context.get_iphone_json.assert_not_called()
        self.assertEqual(
            adapter.loader.context.doc_id_graphql_query.call_args.args[0],
            "25060748103519434",
        )
        self.assertTrue(
            adapter.loader.context.doc_id_graphql_query.call_args.kwargs["web_api"]
        )

    def test_legacy_graphql_comments_do_not_take_the_mobile_branch(self):
        adapter = self.adapter()
        node = dict(
            id="21",
            text="Biaya kuliah?",
            created_at=1790848800,
            owner=dict(id="10"),
            edge_liked_by=dict(count=3),
            edge_threaded_comments=dict(edges=[]),
        )
        response = dict(
            data=dict(
                shortcode_media=dict(
                    edge_media_to_parent_comment=dict(edges=[dict(node=node)])
                )
            )
        )

        def iterator(context, query_hash, extractor, wrapper, variables, referer):
            self.assertEqual(query_hash, "97b41c52301f77ce508f55e66d17620e")
            self.assertEqual(variables, {"shortcode": "B"})
            return iter(wrapper(e["node"]) for e in extractor(response)["edges"])

        adapter.module.NodeIterator = iterator
        rows = list(adapter.legacy_comments(fixture_post()))
        self.assertEqual([(r.id, r.likes_count) for r in rows], [("21", 3)])
        adapter.loader.context.get_iphone_json.assert_not_called()

    def test_local_session_reuse_and_owner_only_json_persistence(self):
        settings = self.settings()
        auth = Path(settings.auth_file)
        auth.write_text(
            json.dumps(
                [dict(sosmed="instagram", sessionid="local-cookie", password="unused")]
            )
        )
        cookies = il.session_cookies(settings)
        self.assertEqual(cookies, {"csrftoken": "", "sessionid": "local-cookie"})
        loader = Mock()
        loader.save_session.return_value = {
            "sessionid": "renewed-cookie",
            "csrftoken": "csrf",
            "unrelated": "discard",
        }
        il.save_session(loader, settings)
        auth.unlink()
        self.assertEqual(il.session_cookies(settings)["sessionid"], "renewed-cookie")
        self.assertEqual(Path(settings.session_file).stat().st_mode & 0o777, 0o600)
        self.assertNotIn("password", Path(settings.session_file).read_text())
        self.assertNotIn("unrelated", Path(settings.session_file).read_text())
        with self.assertRaises(il.ig.SessionInvalid):
            il.session_cookies(replace(settings, username="another_account"))

    def test_http_429_observed_on_copied_sessions_and_not_retried(self):
        response = requests.Response()
        response.status_code = 429
        response.headers["Retry-After"] = "720"
        response._content = b"sensitive response excluded"
        response._content_consumed = True
        with patch.object(requests.Session, "send", return_value=response) as transport:
            with il.observe_requests(5):
                with self.assertRaises(il.AccessBarrier) as raised:
                    requests.Session().get(URL)
            transport.assert_called_once()
        receipt = il.error_receipt(raised.exception)
        self.assertEqual(receipt["status"], "RATE_LIMITED")
        self.assertEqual(receipt["http_status"], 429)
        self.assertEqual(receipt["retry_after_seconds"], 720)
        self.assertNotIn("sensitive", json.dumps(receipt))

    def test_login_redirect_and_structured_challenge_stop(self):
        for code, headers, body, expected in [
            (302, {"Location": "/accounts/login/?next=/"}, b"", "LOGIN_REQUIRED"),
            (
                200,
                {"Content-Type": "application/json"},
                b'{"message":"challenge_required"}',
                "CHALLENGE_REQUIRED",
            ),
            (
                400,
                {"Content-Type": "application/json"},
                b'{"message":"feedback_required"}',
                "ACCESS_BLOCKED",
            ),
        ]:
            response = requests.Response()
            response.status_code, response._content = code, body
            response._content_consumed = True
            response.headers.update(headers)
            with (
                self.subTest(expected=expected),
                patch.object(requests.Session, "send", return_value=response),
            ):
                with (
                    il.observe_requests(5),
                    self.assertRaises(il.AccessBarrier) as caught,
                ):
                    requests.Session().get(URL)
                self.assertEqual(caught.exception.status, expected)

    def test_upstream_abort_is_terminal_and_diagnostics_exclude_error_text(self):
        abort = type("AbortDownloadException", (Exception,), {})
        result = il.error_receipt(abort("private response not for logs"))
        self.assertEqual(result["status"], "ACCESS_BLOCKED")
        self.assertNotIn("private response", json.dumps(result))

    def test_real_loader_keeps_builtin_throttling_and_disables_retries(self):
        module = Obj(Instaloader=Mock())
        with patch.object(
            il,
            "session_cookies",
            return_value={"sessionid": "fixture", "csrftoken": ""},
        ):
            il.InstaloaderAdapter(self.settings(), module=module)
        options = module.Instaloader.call_args.kwargs
        self.assertTrue(options["sleep"])
        self.assertEqual(options["max_connection_attempts"], 1)
        self.assertFalse(options["iphone_support"])
        module.Instaloader.return_value.login.assert_not_called()

    def test_persistent_worker_retains_process_between_operations(self):
        worker = il.InstaloaderWorker(persistent_fixture)
        self.addCleanup(worker.close)
        first = worker("search", "one", self.settings(), 5, Mock())
        second = worker("media", URL, self.settings(), 5, Mock())
        self.assertEqual(first["pid"], second["pid"])
        self.assertEqual(second["count"], 2)
        worker.close()
        self.assertFalse(worker.process.is_alive())

    def test_rate_limited_context_waits_for_coordinator_without_recreating_session(
        self,
    ):
        from dataclasses import asdict

        pipe = Mock()
        pipe.recv.side_effect = [("media", URL), ("media", URL), None]
        with patch.object(il, "InstaloaderAdapter") as create:
            create.return_value.collect.side_effect = [
                dict(status="RATE_LIMITED"),
                dict(status="FINISHED_BEST_EFFORT"),
            ]
            il.instaloader_worker(asdict(self.settings()), pipe)
        create.assert_called_once()
        self.assertEqual(create.return_value.collect.call_count, 2)
        self.assertEqual(pipe.send.call_count, 2)
        create.return_value.loader.close.assert_called_once()

    def test_hard_deadline_preserves_streamed_events_and_stops_worker(self):
        worker, events = il.InstaloaderWorker(hanging_fixture), []
        self.addCleanup(worker.close)
        # Separate spawn/import time from the deliberate hang. The deadline is
        # still measured in the parent, and streamed rows must survive the kill.
        self.assertEqual(
            worker("check_session", "ready", self.settings(), 10, Mock())["status"],
            "FINISHED_BEST_EFFORT",
        )
        began = time.monotonic()
        result = worker(
            "media", URL, self.settings(), 1.5, lambda k, r: events.append((k, r))
        )
        self.assertEqual(result["status"], "TIMEOUT")
        self.assertEqual(events, [("record", {"native_id": "1"})])
        self.assertLess(time.monotonic() - began, 4)
        self.assertFalse(worker.process.is_alive())
        self.assertEqual(
            worker("media", URL, self.settings(), 1, Mock())["status"], "WORKER_FAILED"
        )

    def test_shared_excel_export_provenance_and_baseline_deduplication(self):
        def worker(operation, target, settings, seconds, emit):
            with patch.object(il, "save_session"):
                return self.adapter(settings).collect(operation, target, emit)

        output = self.path / "first"
        crawler = il.InstaloaderCrawler(self.config, worker=worker)
        result = crawler.run(output)
        self.assertEqual(
            (result["posts_exported"], result["comments"], result["replies"]), (1, 1, 1)
        )
        self.assertTrue(result["synthetic"])
        self.assertFalse(result["live_access_verified"])
        book = load_workbook(output / "social_data.xlsx", read_only=True)
        try:
            self.assertEqual(book.sheetnames, ["Data"])
            self.assertEqual(sum(1 for _ in book["Data"].values), 4)
        finally:
            book.close()
        metadata = (output / "summary.json").read_text()
        self.assertNotIn("local-cookie", metadata)
        self.config["exclude_existing_datasets"] = [str(output)]
        never = Mock()
        second = il.InstaloaderCrawler(self.config, worker=never).run(
            self.path / "second"
        )
        self.assertEqual(second["existing_posts_skipped"], 1)
        never.assert_not_called()

    def test_coordinator_stops_at_first_429(self):
        worker = Mock(
            return_value=dict(
                status="RATE_LIMITED", warnings=[], retry_after_seconds=360
            )
        )
        result = il.InstaloaderCrawler(self.config, worker=worker).run(
            self.path / "rate"
        )
        self.assertEqual(result["stop_reason"], "RATE_LIMITED")
        self.assertEqual(result["rate_limit_retries"], 0)
        worker.assert_called_once()

    def test_configuration_rejects_retry_or_unbounded_requests(self):
        self.config["platforms"]["instagram"]["rate_limit_max_retries"] = 1
        with self.assertRaises(ValueError):
            il.InstaloaderCrawler(self.config)
        self.config["platforms"]["instagram"]["rate_limit_max_retries"] = 0
        self.config["instaloader"]["request_interval_seconds"] = 0
        with self.assertRaises(ValueError):
            il.InstaloaderCrawler(self.config)


if __name__ == "__main__":
    unittest.main()
