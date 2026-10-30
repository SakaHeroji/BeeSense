"""Offline export/worker tests, plus opt-in real Chrome DOM fixtures."""

import json
import os
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
from urllib.parse import quote

from openpyxl import load_workbook

from Crawler.Backend.social_crawler import COMPACT_COLUMNS, SocialCrawler, reddit
from Crawler.Backend.social_crawler_selenium import (
    METHOD,
    PAGE_STATE,
    SEARCH_LINKS,
    SOURCE_VERSION,
    YOUTUBE_COMMENTS,
    YOUTUBE_VIDEO,
    SocialCrawlerSelenium,
    _barrier,
    _create_driver,
    _receipt,
    _reddit_records,
    _SeleniumSource,
    _timestamp,
    selenium_worker,
)

VIDEO_ID = "AAAAAAAAAAA"
VIDEO_URL = "https://www.youtube.com/watch?v=" + VIDEO_ID
STAMP = "2026-06-01T07:00:00+00:00"
CONFIG = dict(
    project_id="selenium-test",
    since="2026-06-01T00:00:00+00:00",
    until="2026-07-01T00:00:00+00:00",
    include_keywords=["kuliah"],
    exclude_keywords=["gacor"],
    delay=dict(mode="fixed", seconds=0),
    total_seconds=100,
    platforms=dict(youtube=dict(urls=[VIDEO_URL], min_comments_per_post=2)),
)


def youtube_events():
    return [
        (
            "video",
            dict(
                id=VIDEO_ID,
                title="Kuliah kampus",
                description="Beasiswa",
                published=STAMP,
                observed_at=STAMP,
            ),
        ),
        (
            "comment",
            dict(
                id="comment1",
                text="=literal formula",
                parent="root",
                _time_text="1 month ago",
                observed_at=STAMP,
            ),
        ),
        (
            "comment",
            dict(
                id="reply1",
                text="Balasan 😀",
                parent="comment1",
                _time_text="yesterday",
                observed_at=STAMP,
            ),
        ),
    ]


def capture():
    return dict(
        page_url="https://www.reddit.com/r/indonesia/comments/abc/kuliah/",
        post=dict(
            native_id="abc",
            title="Kuliah",
            text="Kampus",
            subreddit="indonesia",
            published_at=STAMP,
            permalink="/r/indonesia/comments/abc/kuliah/",
        ),
        comments=[
            dict(
                url="https://www.reddit.com/r/indonesia/comments/abc/comment/com123/",
                text="Komentar",
                published_at=STAMP,
            ),
            dict(
                url="https://www.reddit.com/r/indonesia/comments/abc/comment/rpl123/",
                parent_url="https://www.reddit.com/r/indonesia/comments/abc/comment/com123/",
                text="Balasan",
                published_display="yesterday",
            ),
        ],
    )


class SocialCrawlerSeleniumTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    @staticmethod
    def worker(operation, target, settings, seconds, on_event):
        for kind, row in youtube_events():
            on_event(kind, row)
        return _receipt()

    def test_dry_run_uses_same_config_and_excel_contract_without_browser(self):
        with patch(
            "Crawler.Backend.social_crawler_selenium._create_driver",
            side_effect=AssertionError("no browser"),
        ):
            crawler = SocialCrawlerSelenium(CONFIG)
            plan = crawler.plan()
        self.assertEqual(crawler.config, SocialCrawler(CONFIG).config)
        self.assertEqual(
            plan["primary_output"], SocialCrawler(CONFIG).plan()["primary_output"]
        )
        self.assertEqual(plan["acquisition_backend"], "selenium")
        self.assertFalse(plan["live_access_verified"])

    def test_search_scroll_deduplicates_and_enforces_candidate_cap(self):
        class Driver:
            scrolls = 0

            def get(self, url):
                self.url = url

            def set_page_load_timeout(self, seconds):
                pass

            def set_script_timeout(self, seconds):
                pass

            def execute_script(self, script, *args):
                if script == PAGE_STATE:
                    return dict(url=self.url, content=True)
                if script == SEARCH_LINKS:
                    return [VIDEO_URL, "https://evil.example/watch?v=AAAAAAAAAAA"] + (
                        ["https://youtu.be/BBBBBBBBBBB", "https://youtu.be/CCCCCCCCCCC"]
                        if self.scrolls
                        else []
                    )
                self.scrolls += 1

        cfg = deepcopy(CONFIG)
        cfg["platforms"]["youtube"]["search_results"] = 2
        crawler = SocialCrawlerSelenium(cfg)
        settings = crawler._legacy_settings(
            "youtube", crawler.config["platforms"]["youtube"]
        )
        driver, events = Driver(), []
        source = _SeleniumSource(
            driver, "youtube", settings, events.append, cfg, sleep=lambda _: None
        )
        source.search("kuliah")
        self.assertEqual(driver.scrolls, 1)
        self.assertEqual(len(events), 2)
        self.assertEqual(
            events[1][1]["url"], "https://www.youtube.com/watch?v=BBBBBBBBBBB"
        )

    def test_youtube_acquisition_enforces_total_and_reply_caps(self):
        class Driver:
            def get(self, url):
                self.url = url

            def set_page_load_timeout(self, seconds):
                pass

            def set_script_timeout(self, seconds):
                pass

            def find_elements(self, *args):
                return []

            def execute_script(self, script, *args):
                if script == PAGE_STATE:
                    return dict(url=self.url, content=True)
                if script == YOUTUBE_VIDEO:
                    return youtube_events()[0][1]
                if script == YOUTUBE_COMMENTS:
                    return [event[1] for event in youtube_events()[1:]] + [
                        dict(id="reply2", text="Melebihi batas", parent="comment1")
                    ]

        cfg = deepcopy(CONFIG)
        cfg["platforms"]["youtube"].update(
            max_comments=2, max_replies=1, max_replies_per_thread=1
        )
        crawler = SocialCrawlerSelenium(cfg)
        settings = crawler._legacy_settings(
            "youtube", crawler.config["platforms"]["youtube"]
        )
        events = []
        source = _SeleniumSource(
            Driver(), "youtube", settings, events.append, cfg, sleep=lambda _: None
        )
        source.youtube(VIDEO_URL)
        self.assertEqual([kind for kind, _ in events], ["video", "comment", "comment"])
        self.assertEqual(source.replies, 1)
        self.assertIn("comment_cap_reached", source.warnings)

    def test_both_workbooks_preserve_schema_and_selenium_provenance(self):
        cfg = deepcopy(CONFIG)
        cfg["platforms"]["reddit"] = dict(
            urls=["https://www.reddit.com/comments/abc/"], min_comments_per_post=2
        )

        def reddit_worker(operation, target, settings, seconds, on_event):
            for event in reddit.bounded_events(
                _reddit_records(capture(), target, ["indonesia"]), settings
            ):
                on_event(*event)
            return _receipt()

        crawler = SocialCrawlerSelenium(
            cfg, workers=dict(youtube=self.worker, reddit=reddit_worker)
        )
        report = crawler.run(self.directory / "bundle")
        self.assertEqual(report["status"], "TARGET_MET")
        self.assertEqual(report["compact"]["rows"], 6)
        for filename, sheets in [
            ("social_data.xlsx", ["Data"]),
            ("dataset.xlsx", list(reddit.TABLES)),
        ]:
            book = load_workbook(self.directory / "bundle" / filename, read_only=True)
            self.assertEqual(book.sheetnames, sheets)
            for table in sheets:
                expected = (
                    COMPACT_COLUMNS
                    if table == "Data"
                    else tuple(field["name"] for field in reddit.TABLES[table])
                )
                self.assertEqual(next(book[table].values), expected)
            if filename == "social_data.xlsx":
                cells = list(book["Data"].rows)
                self.assertEqual(cells[2][7].data_type, "s")
                self.assertEqual(cells[2][7].value, "=literal formula")
                self.assertEqual(cells[2][6].value, "perkiraan")
            book.close()
        for table in reddit.TABLES:
            for row in self.rows(table):
                reddit.validate_row(table, row)
        self.assertEqual(
            {row["extraction_method"] for row in self.rows("Contents")}, {METHOD}
        )
        self.assertEqual(
            {row["extraction_method"] for row in self.rows("Sources")}, {METHOD}
        )
        replies = [
            row for row in self.rows("Contents") if row["object_type"] == "reply"
        ]
        self.assertEqual(len(replies), 2)
        self.assertTrue(
            all(row["relation_evidence"] == "dom_nesting" for row in replies)
        )
        manifest = json.loads((self.directory / "bundle/manifest.json").read_text())
        self.assertTrue(manifest["synthetic"])
        for name, digest in manifest["files_sha256"].items():
            from hashlib import sha256

            self.assertEqual(
                sha256((self.directory / "bundle" / name).read_bytes()).hexdigest(),
                digest,
            )

    def rows(self, table):
        return [
            json.loads(line)
            for line in (self.directory / "bundle" / (table + ".jsonl"))
            .read_text()
            .splitlines()
        ]

    def test_barrier_stops_after_first_call_and_still_exports_headers(self):
        calls = []

        def barrier(*args):
            calls.append(args[1])
            return _receipt("CAPTCHA_REQUIRED")

        cfg = deepcopy(CONFIG)
        cfg["platforms"]["youtube"]["urls"].append("https://youtu.be/BBBBBBBBBBB")
        report = SocialCrawlerSelenium(cfg, workers=dict(youtube=barrier)).run(
            self.directory / "bundle"
        )
        self.assertEqual(report["status"], "PARTIAL")
        self.assertEqual(
            report["platforms"]["youtube"]["stop_reason"], "CAPTCHA_REQUIRED"
        )
        self.assertEqual(len(calls), 1)
        self.assertEqual(report["compact"]["rows"], 0)

    def test_timeout_keeps_streamed_evidence_and_source_version(self):
        def timed_out(operation, target, settings, seconds, on_event):
            on_event(*youtube_events()[0])
            return _receipt("TIMEOUT")

        report = SocialCrawlerSelenium(CONFIG, workers=dict(youtube=timed_out)).run(
            self.directory / "bundle"
        )
        self.assertEqual(report["status"], "PARTIAL")
        self.assertEqual(len(self.rows("Contents")), 1)
        self.assertEqual(self.rows("Runs")[0]["source_version"], SOURCE_VERSION)

    def test_existing_output_is_not_overwritten(self):
        crawler = SocialCrawlerSelenium(CONFIG, workers=dict(youtube=self.worker))
        crawler.run(self.directory / "bundle")
        before = (self.directory / "bundle/social_data.xlsx").read_bytes()
        with self.assertRaises(FileExistsError):
            crawler.run(self.directory / "bundle")
        self.assertEqual(
            before, (self.directory / "bundle/social_data.xlsx").read_bytes()
        )

    def test_relative_and_date_only_are_not_promoted_to_exact(self):
        self.assertEqual(_timestamp("2026-06-01"), (None, "2026-06-01", "date_only"))
        self.assertEqual(_timestamp("1 month ago"), (None, "1 month ago", "relative"))
        self.assertEqual(
            _timestamp("2026-06-01T12:00:00"), (None, "2026-06-01T12:00:00", "unknown")
        )
        self.assertEqual(_timestamp(STAMP)[2], "exact")

    def test_captcha_quotes_are_ignored_when_content_is_present(self):
        self.assertIsNone(_barrier(dict(content=True, title="Kuliah", text="")))
        self.assertEqual(
            _barrier(dict(content=False, title="Blocked", text="Verify you are human")),
            "CAPTCHA_REQUIRED",
        )
        self.assertEqual(_barrier(dict(content=True, status=429)), "RATE_LIMITED")
        self.assertEqual(
            _barrier(dict(content=True, challenge=True)), "CHALLENGE_REQUIRED"
        )
        self.assertEqual(
            _barrier(dict(content=False, text="Your connection is not private")),
            "TLS_ERROR",
        )

    def test_reddit_identity_and_subreddit_checks(self):
        value = capture()
        with self.assertRaises(reddit.SourceFailure):
            _reddit_records(
                value, "https://www.reddit.com/comments/def/", ["indonesia"]
            )
        with self.assertRaises(reddit.SourceFailure):
            _reddit_records(value, "https://www.reddit.com/comments/abc/", ["karir"])
        value["comments"][0]["url"] = (
            "https://evil.example/comments/abc/comment/com123/"
        )
        with self.assertRaises(ValueError):
            _reddit_records(
                value, "https://www.reddit.com/comments/abc/", ["indonesia"]
            )

    def test_disabled_checkpoint_and_remote_credentials_are_rejected(self):
        cfg = deepcopy(CONFIG)
        cfg["platforms"]["reddit"] = dict(import_checkpoint_dir="private-checkpoint")
        with self.assertRaisesRegex(ValueError, "remove import_checkpoint_dir"):
            SocialCrawlerSelenium(cfg)
        with self.assertRaisesRegex(ValueError, "without credentials"):
            SocialCrawlerSelenium(CONFIG, remote_url="http://secret@localhost:4444")

    def test_worker_sanitizes_errors_and_closes_browser(self):
        class Driver:
            closed = False

            def quit(self):
                self.closed = True

            def execute_script(self, script):
                return dict(content=False, text="Verify you are human")

        class Pipe:
            closed = False

            def __init__(self):
                self.events = []

            def send(self, event):
                self.events.append(event)

            def close(self):
                self.closed = True

        driver, pipe = Driver(), Pipe()
        with (
            patch("Crawler.Backend.social_crawler_selenium.os.setsid"),
            patch(
                "Crawler.Backend.social_crawler_selenium._create_driver",
                return_value=driver,
            ),
            patch.object(
                _SeleniumSource,
                "search",
                side_effect=RuntimeError("PRIVATE_ERROR_SECRET"),
            ),
        ):
            selenium_worker(
                "search",
                "kuliah",
                dict(timeout_seconds=20),
                pipe,
                platform="youtube",
                browser_options={},
                config=CONFIG,
            )
        self.assertTrue(driver.closed)
        self.assertTrue(pipe.closed)
        self.assertEqual(pipe.events[-1][1]["status"], "CAPTCHA_REQUIRED")
        self.assertNotIn("PRIVATE_ERROR_SECRET", str(pipe.events))


@unittest.skipUnless(
    os.environ.get("BEESENSE_TEST_SELENIUM") == "1", "opt-in local Chrome DOM fixtures"
)
class SeleniumDOMTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.driver = _create_driver(dict(headless=True))

    @classmethod
    def tearDownClass(cls):
        cls.driver.quit()

    def test_youtube_scripts_extract_only_native_ids_and_own_text(self):
        html = f'''<html><head><meta itemprop="datePublished" content="{STAMP}"></head><body>
        <ytd-watch-flexy video-id="{VIDEO_ID}"><h1 class="ytd-watch-metadata">Kuliah kampus</h1>
        <div id="description"><yt-formatted-string>Beasiswa</yt-formatted-string></div>
        <ytd-comment-thread-renderer><ytd-comment-view-model>
          <div id="content-text">Komentar asli</div><div id="published-time-text"><a href="{VIDEO_URL}&lc=comment1">1 month ago</a></div>
        </ytd-comment-view-model><ytd-comment-replies-renderer><ytd-comment-view-model>
          <div id="content-text">Balasan</div><div id="published-time-text"><a href="{VIDEO_URL}&lc=reply1">yesterday</a></div>
        </ytd-comment-view-model></ytd-comment-replies-renderer></ytd-comment-thread-renderer>
        <ytd-comment-view-model><div id="content-text">Tanpa ID, jangan ekspor</div></ytd-comment-view-model>
        </ytd-watch-flexy></body></html>'''
        self.driver.get("data:text/html;charset=utf-8," + quote(html))
        post = self.driver.execute_script(YOUTUBE_VIDEO)
        comments = self.driver.execute_script(YOUTUBE_COMMENTS, 10)
        self.assertEqual(post["id"], VIDEO_ID)
        self.assertEqual(post["published"], STAMP)
        self.assertEqual(len(comments), 2)
        self.assertEqual(comments[0]["text"], "Komentar asli")
        self.assertEqual(comments[1]["parent"], "comment1")
        self.assertTrue(self.driver.execute_script(PAGE_STATE)["content"])


if __name__ == "__main__":
    unittest.main()
