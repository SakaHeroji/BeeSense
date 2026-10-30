import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from scripts.merge_social_datasets import COLUMNS, merge


class MergeTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)

    def source(self, name, rows):
        path = self.root / (name + ".xlsx")
        book = Workbook()
        sheet = book.active
        sheet.title = "Data"
        sheet.append(COLUMNS)
        for platform, kind, url, text in rows:
            values = [
                platform,
                kind,
                "Title",
                None,
                "pasti",
                None,
                "pasti",
                text,
                url,
                False,
            ]
            safe = []
            for value in values:
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                safe.append(cell)
            sheet.append(safe)
        book.save(path)
        return path

    def run_merge(self, sources):
        return merge(sources, self.root / "combined", self.root / "provenance.json")

    def test_merge_preserves_identical_text_distinct_ids_and_formula_like_text(self):
        a = self.source(
            "a",
            [
                (
                    "youtube",
                    "video",
                    "https://www.youtube.com/watch?v=videoA",
                    "Root A",
                ),
                (
                    "youtube",
                    "comment",
                    "https://www.youtube.com/watch?v=videoA&lc=c1",
                    "=same text",
                ),
            ],
        )
        b = self.source(
            "b",
            [
                ("reddit", "post", "https://www.reddit.com/comments/abc/", "Root B"),
                (
                    "reddit",
                    "reply",
                    "https://www.reddit.com/comments/abc/_/c2/",
                    "=same text",
                ),
            ],
        )
        hashes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (a, b)]
        result = self.run_merge([a, b])
        self.assertEqual(
            (result["rows"], result["roots"], result["comments"]), (4, 2, 2)
        )
        self.assertEqual([s["sha256"] for s in result["source_workbooks"]], hashes)
        self.assertEqual(
            [s["first_excel_row"] for s in result["source_workbooks"]], [2, 4]
        )
        self.assertEqual(
            [hashlib.sha256(p.read_bytes()).hexdigest() for p in (a, b)], hashes
        )
        book = load_workbook(result["output"], read_only=True, data_only=False)
        cells = list(book["Data"].iter_rows())
        self.assertEqual(
            (cells[2][7].value, cells[2][7].data_type), ("=same text", "s")
        )
        self.assertEqual(cells[4][7].value, "=same text")
        book.close()

    def test_same_native_id_different_kind_is_rejected_without_publication(self):
        a = self.source(
            "a",
            [
                ("reddit", "post", "https://www.reddit.com/comments/abc/", "root"),
                (
                    "reddit",
                    "comment",
                    "https://www.reddit.com/comments/abc/_/c1/",
                    "text",
                ),
            ],
        )
        b = self.source(
            "b",
            [
                (
                    "reddit",
                    "reply",
                    "https://www.reddit.com/comments/abc/_/c1/",
                    "changed",
                )
            ],
        )
        with self.assertRaisesRegex(ValueError, "Duplicate native content"):
            self.run_merge([a, b])
        self.assertFalse((self.root / "combined").exists())

    def test_orphan_comment_is_rejected(self):
        a = self.source(
            "a",
            [
                (
                    "youtube",
                    "comment",
                    "https://www.youtube.com/watch?v=videoA&lc=c1",
                    "orphan",
                )
            ],
        )
        with self.assertRaisesRegex(ValueError, "Orphan comments"):
            self.run_merge([a])

    def test_instagram_latest_record_wins_without_removing_equal_text_distinct_ids(
        self,
    ):
        import json

        old = self.source(
            "old",
            [
                ("instagram", "post", "https://www.instagram.com/p/B/", "Kuliah"),
                (
                    "instagram",
                    "comment",
                    "https://www.instagram.com/p/B/?comment_id=21",
                    "lama",
                ),
            ],
        )
        new = self.source(
            "new",
            [
                (
                    "instagram",
                    "post",
                    "https://www.instagram.com/reel/B/",
                    "Kuliah baru",
                ),
                (
                    "instagram",
                    "comment",
                    "https://www.instagram.com/p/B/c/21/",
                    "teks sama",
                ),
                (
                    "instagram",
                    "comment",
                    "https://www.instagram.com/p/B/?comment_id=22",
                    "teks sama",
                ),
            ],
        )
        result = merge(
            [old, new],
            self.root / "combined",
            self.root / "provenance.json",
            duplicate_policy="prefer_latest",
        )
        self.assertEqual(result["rows"], 3)
        self.assertEqual(result["duplicates_removed"], 2)
        self.assertEqual(result["duplicate_native_ids"], 0)
        book = load_workbook(result["output"], read_only=True)
        rows = list(book["Data"].values)
        book.close()
        self.assertEqual(
            [row[7] for row in rows[1:]], ["Kuliah baru", "teks sama", "teks sama"]
        )
        lineage = [
            json.loads(row)
            for row in (self.root / "combined/record_lineage.jsonl")
            .read_text()
            .splitlines()
        ]
        self.assertEqual(len(lineage), 3)
        self.assertTrue(all(row["source"] == str(new.resolve()) for row in lineage))

    def test_instagram_comment_identity_conflict_and_fake_host_are_rejected(self):
        from scripts.merge_social_datasets import identity

        for url in [
            "https://www.instagram.com/p/B/c/21/?comment_id=22",
            "https://www.instagram.com.evil/p/B/?comment_id=21",
        ]:
            with self.subTest(url=url), self.assertRaises(ValueError):
                identity({"platform": "instagram", "jenis": "comment", "url": url})

    def test_latest_policy_still_rejects_comment_attached_to_different_root(self):
        old = self.source(
            "old",
            [
                ("instagram", "post", "https://www.instagram.com/p/B/", "Kuliah"),
                (
                    "instagram",
                    "comment",
                    "https://www.instagram.com/p/B/?comment_id=21",
                    "Isi",
                ),
            ],
        )
        new = self.source(
            "new",
            [
                ("instagram", "post", "https://www.instagram.com/p/C/", "Kampus"),
                (
                    "instagram",
                    "comment",
                    "https://www.instagram.com/p/C/?comment_id=21",
                    "Isi",
                ),
            ],
        )
        with self.assertRaisesRegex(ValueError, "Conflicting native root"):
            merge(
                [old, new],
                self.root / "combined",
                self.root / "provenance.json",
                duplicate_policy="prefer_latest",
            )
        self.assertFalse((self.root / "combined").exists())

    def test_latest_policy_audits_updated_comment_reply_classification(self):
        import json

        old = self.source(
            "old",
            [
                ("reddit", "post", "https://www.reddit.com/comments/abc/", "Root"),
                (
                    "reddit",
                    "comment",
                    "https://www.reddit.com/comments/abc/_/c1/",
                    "Isi",
                ),
            ],
        )
        new = self.source(
            "new",
            [("reddit", "reply", "https://www.reddit.com/comments/abc/_/c1/", "Isi")],
        )
        result = merge(
            [old, new],
            self.root / "combined",
            self.root / "provenance.json",
            duplicate_policy="prefer_latest",
        )
        self.assertEqual(result["rows"], 2)
        self.assertEqual(result["counts"]["reddit/reply"], 1)
        audit = json.loads((self.root / "combined/duplicate_records.jsonl").read_text())
        self.assertTrue(audit["kind_changed"])


if __name__ == "__main__":
    unittest.main()
